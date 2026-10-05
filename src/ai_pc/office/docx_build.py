"""The .docx from a resolved plan (docplan.resolve), the way a professional sets a document:

real Word styles (Normal, Title, Heading 1-3, Caption, lists), set once from the theme, so the navigation pane, the
table of contents and later edits all work; headings kept with the text that follows them; tables with a header
row that repeats on every page, numbers right-aligned, column widths from their content, rows that never split;
native charts (charts.py); captions above tables and below figures; a cover and a table of contents for long work;
headers and "Page X of Y" footers (none on the cover); numbered lists that restart; letters, CVs, invoices and
certificates laid out the way they are expected to look.

build(plan, path) -> {"path", "blocks": [{"i", "type", "text"}], "text_width_cm"}
"""

import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ai_pc.office import charts
from ai_pc.office.docplan import fmt, plain, runs

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "centre": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
PAGE = {"A4": (21.0, 29.7), "Letter": (21.59, 27.94)}


# ------------------------------------------------------------------------------------------------ low-level helpers
def _rgb(h):
    return RGBColor.from_string(str(h).lstrip("#").upper())


def _rfonts(rpr, name, cs=None):
    """Set a font on an rPr for every script and drop the theme-font attributes that would override it."""
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), name)
    rf.set(qn("w:cs"), cs or name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def _style_font(style, name, size=None, color=None, bold=None, italic=None, cs=None):
    _rfonts(style.element.get_or_add_rPr(), name, cs)
    f = style.font
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = _rgb(color)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _borders(el_pr, tag, edges):
    """edges: {"top": (size_eighths, color, "single"), ...} on a tblPr / tcPr / pPr."""
    old = el_pr.find(qn(tag))
    if old is not None:
        el_pr.remove(old)
    b = OxmlElement(tag)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        sz, color, val = edges[edge]
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), val if sz else "nil")
        if sz:
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
        b.append(e)
    el_pr.append(b)


def _cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblpr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblpr.append(m)


def _row_flags(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trpr.append(cs)
    if header:
        h = OxmlElement("w:tblHeader")
        trpr.append(h)


def _table_width(table, widths_cm):
    tblpr = table._tbl.tblPr
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    w.set(qn("w:type"), "dxa")
    for old in tblpr.findall(qn("w:tblW")):
        tblpr.remove(old)
    tblpr.append(w)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    grid = table._tbl.tblGrid
    for gc, wc in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(wc * 567)))
    table.autofit = False
    for row in table.rows:
        for cell, wc in zip(row.cells, widths_cm):
            cell.width = Cm(wc)


def _field(paragraph, code, placeholder=""):
    run = paragraph.add_run()
    for kind in ("begin", "instr", "separate", "text", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {code} "
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = placeholder
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    return run


def _rule(paragraph, color, size=6, where="bottom", space=4):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{where}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(size))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    bdr.append(e)
    ppr.append(bdr)


def _rtl(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.append(OxmlElement("w:bidi"))
    for r in paragraph.runs:
        rpr = r._r.get_or_add_rPr()
        if rpr.find(qn("w:rtl")) is None:
            rpr.append(OxmlElement("w:rtl"))


def _add_text(paragraph, text, base=None):
    """Runs from inline **bold** / *italic* marks; base: {"bold", "italic", "size", "color", "font"} for every run."""
    base = base or {}
    lines = str(text).split("\n")
    for k, line in enumerate(lines):
        for t, f in runs(line):
            r = paragraph.add_run(t)
            if f.get("bold") or base.get("bold"):
                r.bold = True
            if f.get("italic") or base.get("italic"):
                r.italic = True
            if f.get("underline") or base.get("underline"):
                r.underline = True
            if base.get("size"):
                r.font.size = Pt(base["size"])
            if base.get("color"):
                r.font.color.rgb = _rgb(base["color"])
            if base.get("font"):
                _rfonts(r._r.get_or_add_rPr(), base["font"])
        if k < len(lines) - 1:
            paragraph.add_run().add_break()
    return paragraph


def _restart_numbering(doc, paragraph, style_name="List Number"):
    """A numbered list that starts again at 1 (python-docx would continue the previous list's numbers)."""
    try:
        st = doc.styles[style_name]
        num_id = st.element.pPr.numPr.numId.val
        numbering = doc.part.numbering_part.element
        abstract = numbering.num_having_numId(num_id).abstractNumId.val
        num = numbering.add_num(abstract)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = 0
        numpr.get_or_add_numId().val = num.numId
        return num.numId
    except Exception:  # noqa: BLE001
        return None


def _use_num(paragraph, num_id, level=0):
    if num_id is None:
        return
    numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_ilvl().val = level
    numpr.get_or_add_numId().val = num_id


# ------------------------------------------------------------------------------------------------ the document
class _Builder:
    def __init__(self, plan):
        self.p = plan
        self.th = plan["th"]
        self.doc = Document()
        w, h = PAGE[plan["page"]["size"]]
        if plan["page"]["orientation"] == "landscape":
            w, h = h, w
        self.text_w = round(w - 2 * plan["page"]["margins"], 2)
        self.rtl = plan["language"] == "ur"
        self.out_blocks = []

    # styles: set once, used everywhere
    def styles(self):
        th, d = self.th, self.doc
        cs = "Segoe UI" if self.rtl else None
        normal = d.styles["Normal"]
        _style_font(normal, th["body"], th["size"], "262626", cs=cs)
        pf = normal.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(th["after"])
        pf.line_spacing = th["line"]
        pf.widow_control = True
        if th["justify"]:
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        k = th.get("h_scale", 1.0)
        for lv, key, before, after in ((1, "h1", 18 * k, 6 * k), (2, "h2", 12 * k, 4 * k), (3, "h3", 10 * k, 3 * k)):
            s = d.styles[f"Heading {lv}"]
            semibold = "semibold" in th["head"].lower()
            _style_font(s, th["head"], th[key], th["head_color"], bold=not semibold, italic=False, cs=cs)
            f = s.paragraph_format
            f.space_before, f.space_after = Pt(before), Pt(after)
            f.keep_with_next = True
            f.alignment = WD_ALIGN_PARAGRAPH.RIGHT if self.rtl else WD_ALIGN_PARAGRAPH.LEFT
            f.line_spacing = 1.0
        t = d.styles["Title"]
        _style_font(t, th["head"], th["title"], th["accent"], bold=False, cs=cs)
        tppr = t.element.get_or_add_pPr()
        for old in tppr.findall(qn("w:pBdr")):
            tppr.remove(old)
        t.paragraph_format.space_after = Pt(4)
        t.paragraph_format.line_spacing = 1.0
        sub = d.styles["Subtitle"]
        _style_font(sub, th["body"], th["size"] + 3, th["muted"], italic=False, cs=cs)
        sub.paragraph_format.space_after = Pt(10)
        cap = d.styles["Caption"]
        _style_font(cap, th["body"], max(8.5, th["size"] - 2), th["muted"], bold=False, italic=False, cs=cs)
        cap.paragraph_format.space_before, cap.paragraph_format.space_after = Pt(4), Pt(10)
        q = d.styles["Quote"]
        _style_font(q, th["body"], th["size"], th["muted"], italic=True, cs=cs)
        for name in ("List Bullet", "List Number", "List Bullet 2", "List Number 2"):
            try:
                ls = d.styles[name]
                ls.paragraph_format.space_after = Pt(2)
                ls.paragraph_format.line_spacing = th["line"]
            except KeyError:
                continue

    def page(self):
        pg = self.p["page"]
        w, h = PAGE[pg["size"]]
        for s in self.doc.sections:
            if pg["orientation"] == "landscape":
                s.orientation = WD_ORIENT.LANDSCAPE
                w, h = max(w, h), min(w, h)
            s.page_width, s.page_height = Cm(w), Cm(h)
            m = Cm(pg["margins"])
            s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = m
            s.header_distance = s.footer_distance = Cm(1.2)

    def header_footer(self):
        th, p = self.th, self.p
        s = self.doc.sections[0]
        if p["cover"]:
            s.different_first_page_header_footer = True
        if p.get("header"):
            hp = s.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _add_text(hp, p["header"], {"size": 8.5, "color": th["muted"]})
            _rule(hp, th["table"]["border"], size=4, where="bottom", space=4)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p["footer"].get("text"):
            _add_text(fp, p["footer"]["text"] + ("   ·   " if p["footer"].get("page_numbers") else ""), {"size": 8.5, "color": th["muted"]})
        if p["footer"].get("page_numbers"):
            for part in ("Page ", ("PAGE", "1"), " of ", ("NUMPAGES", "1")):
                r = _field(fp, *part) if isinstance(part, tuple) else fp.add_run(part)
                r.font.size = Pt(8.5)
                r.font.color.rgb = _rgb(th["muted"])

    def para(self, text="", style=None, align=None, base=None, after=None, before=None, keep=False):
        try:
            par = self.doc.add_paragraph(style=style)
        except KeyError:  # the document has no such style: plain text, set by hand
            par = self.doc.add_paragraph()
            if style and style.startswith("List"):
                par.paragraph_format.left_indent = Cm(0.63 * (2 if style.endswith(" 2") else 1))
                par.paragraph_format.first_line_indent = Cm(-0.63)
                text = ("• " if "Bullet" in style else "") + text
        if text:
            _add_text(par, text, base)
        if align:
            par.alignment = ALIGN.get(align, None)
        if after is not None:
            par.paragraph_format.space_after = Pt(after)
        if before is not None:
            par.paragraph_format.space_before = Pt(before)
        if keep:
            par.paragraph_format.keep_with_next = True
        if self.rtl:
            _rtl(par)
            if not align:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return par

    # front matter -------------------------------------------------------------------------------------------
    def cover(self):
        th, p, m = self.th, self.p, self.p["meta"]
        if p["doctype"] in ("assignment", "thesis", "research"):
            for k in ("university", "org", "institute", "school", "department"):
                if m.get(k):
                    self.para(
                        m[k],
                        align="center",
                        base={"bold": k != "department", "size": 16 if k != "department" else 12, "color": th["head_color"]},
                        after=4,
                    )
            self.para("", after=60)
            self.para(p["title"], style="Title", align="center", after=6)
            if p["subtitle"]:
                self.para(p["subtitle"], style="Subtitle", align="center")
            self.para("", after=50)
            rows = [
                (lbl, m.get(k))
                for lbl, k in (
                    ("Course", "course"),
                    ("Submitted to", "instructor"),
                    ("Submitted by", "author"),
                    ("Roll No.", "roll_no"),
                    ("Class / Section", "class"),
                    ("Date", "date"),
                )
                if m.get(k)
            ]
            if rows:
                self.kv({"items": [[a, b] for a, b in rows]}, centered=True, width=11.0)
        else:
            if m.get("org"):
                self.para(m["org"].upper(), base={"size": 10, "color": th["muted"], "bold": True}, after=0)
            self.para("", after=150)
            title = self.para(p["title"], style="Title", after=6)
            title.paragraph_format.keep_with_next = True
            if p["subtitle"]:
                self.para(p["subtitle"], style="Subtitle")
            rule = self.para("", after=24)
            _rule(rule, th["accent"], size=18, where="top", space=1)
            for lbl, k in (("Prepared by", "author"), ("Prepared for", "client"), ("Date", "date"), ("Version", "version")):
                if m.get(k):
                    self.para(f"**{lbl}:** {m[k]}", base={"color": th["muted"]}, after=2)
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def title_block(self):
        th, p, m = self.th, self.p, self.p["meta"]
        dt_ = p["doctype"]
        if dt_ in ("letter", "application", "invoice", "quotation", "certificate"):
            return
        if dt_ == "cv":
            self.para(p["title"], style="Title", after=2)
            if p["subtitle"]:
                self.para(p["subtitle"], base={"size": th["size"] + 2, "color": th["accent"]}, after=4)
            contact = [m.get(k) for k in ("location", "phone", "email", "linkedin", "website", "github") if m.get(k)]
            if contact:
                c = self.para("   |   ".join(contact), base={"size": th["size"] - 1, "color": th["muted"]}, after=10)
                _rule(c, th["accent"], size=8, where="bottom", space=6)
            return
        if dt_ in ("notice", "memo", "agenda", "minutes"):
            t = self.para(p["title"].upper() if dt_ == "notice" else p["title"], style="Title", align="center", after=4)
            for r in t.runs:
                r.font.size = Pt(th["title"] - 4)
            if p["subtitle"]:
                self.para(p["subtitle"], align="center", base={"color": th["muted"]}, after=6)
            line = [m.get(k) for k in ("org", "date") if m.get(k)]
            if line:
                self.para("   ·   ".join(line), align="center", base={"size": th["size"] - 1, "color": th["muted"]}, after=12)
            return
        if p["cover"]:
            return
        self.para(p["title"], style="Title", after=4)
        if p["subtitle"]:
            self.para(p["subtitle"], style="Subtitle", after=4)
        line = [m.get(k) for k in ("author", "org", "date") if m.get(k)]
        if line:
            par = self.para("   ·   ".join(line), base={"size": th["size"] - 1, "color": th["muted"]}, after=14)
            if th["rule"]:
                _rule(par, th["accent"], size=8, where="bottom", space=6)

    def toc(self):
        th = self.th
        self.para(
            "Contents", base={"size": th["h1"], "color": th["head_color"], "bold": "semibold" not in th["head"].lower(), "font": th["head"]}, after=10
        )
        par = self.doc.add_paragraph()
        _field(par, 'TOC \\o "1-3" \\h \\z \\u', "Right-click to update the table of contents.")
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # blocks ---------------------------------------------------------------------------------------------------
    def heading(self, b):
        h = self.doc.add_heading(level=b["level"])
        _add_text(h, b["text"])
        if b.get("page_break_before"):
            h.paragraph_format.page_break_before = True
        if self.rtl:
            _rtl(h)

    def paragraph(self, b):
        th = self.th
        st = str(b.get("style") or "normal").lower()
        base = {
            "lead": {"size": th["size"] + 1.5, "color": th["muted"]},
            "note": {"italic": True, "size": th["size"] - 1, "color": th["muted"]},
            "small": {"size": th["size"] - 1.5},
        }.get(st, {})
        if b.get("bold"):
            base = {**base, "bold": True}
        self.para(b["text"], align=b.get("align"), base=base)

    def quote(self, b):
        par = self.para(b["text"], style="Quote")
        par.paragraph_format.left_indent = Cm(1.0)
        par.paragraph_format.right_indent = Cm(1.0)
        ppr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        e = OxmlElement("w:left")
        for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "8"), ("w:color", self.th["accent"])):
            e.set(qn(k), v)
        bdr.append(e)
        ppr.append(bdr)
        if b.get("by"):
            self.para(f"{b['by']}", align="right", base={"size": self.th["size"] - 1, "color": self.th["muted"]})

    def bullets(self, b):
        style = "List Number" if b["numbered"] else "List Bullet"
        num_id = None
        items = b["items"]
        for k, it in enumerate(items):
            par = self.para(it["text"], style=style, after=2 if k < len(items) - 1 or it.get("items") else self.th["after"])
            if b["numbered"]:
                num_id = _restart_numbering(self.doc, par) if k == 0 else (_use_num(par, num_id) or num_id)
            subs = it.get("items") or []
            for j, sub in enumerate(subs):
                sp = self.para(sub["text"], style=style + " 2", after=2 if j < len(subs) - 1 or k < len(items) - 1 else self.th["after"])
                if b["numbered"] and num_id is not None:
                    _use_num(sp, num_id, 1)

    def caption(self, b, keep=False):
        if not b.get("label"):
            return
        par = self.doc.add_paragraph(style="Caption")
        r = par.add_run(f"{b['label']}" + (": " if b.get("caption") else ""))
        r.bold = True
        if b.get("caption"):
            _add_text(par, b["caption"])
        if keep:
            par.paragraph_format.keep_with_next = True
        if self.rtl:
            _rtl(par)

    def table(self, b):
        th = self.th
        tt = th["table"]
        cols, rows = b["columns"], b["rows"] + ([b["total_row"]] if b.get("total_row") else [])
        if b.get("label"):
            self.caption(b, keep=True)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(cols))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        # widths from content: the longest words and the typical length of each column
        weight = []
        for j, c in enumerate(cols):
            cells = [plain(c)] + [plain(str(r[j])) for r in rows]
            longest_word = max((len(w) for x in cells for w in x.split()), default=4)
            typical = sorted(len(x) for x in cells)[len(cells) * 3 // 4]
            weight.append(max(4, longest_word * 1.1, min(typical, 40), min(len(plain(c)), 24) * 0.85))  # the header fits on one line where it can
        total_w = self.text_w if sum(weight) > 30 or len(cols) > 3 else min(self.text_w, max(9.0, sum(weight) * 0.32))
        widths = [total_w * w_ / sum(weight) for w_ in weight]
        _table_width(t, widths)
        _cell_margins(t)
        grid = th["name"] == "academic"
        light = tt["border"] if not grid else "000000"
        edges = {
            "top": (8 if not grid else 6, tt["header_fill"] if not grid else "000000", "single"),
            "bottom": (8 if not grid else 6, tt["header_fill"] if not grid else "000000", "single"),
            "insideH": (4, light, "single"),
            "left": (6 if grid else 0, "000000", "single"),
            "right": (6 if grid else 0, "000000", "single"),
            "insideV": (4 if grid else 0, "000000", "single"),
        }
        if th["name"] == "minimal":
            edges["insideH"] = (0, light, "single")
        _borders(t._tbl.tblPr, "w:tblBorders", edges)
        size = tt["size"] - (1.5 if b.get("small") else 0)  # "small": the fixer's answer to a table wider than the page
        numeric = b.get("numeric") or [False] * len(cols)
        for j, c in enumerate(cols):
            cell = t.rows[0].cells[j]
            cell.text = ""
            par = cell.paragraphs[0]
            _add_text(par, c, {"bold": True, "size": size, "color": tt["header_text"]})
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if numeric[j] else WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.keep_with_next = True
            _shade(cell, tt["header_fill"])
        _row_flags(t.rows[0], header=True)
        for i, r in enumerate(rows):
            row = t.rows[i + 1]
            _row_flags(row)
            is_total = bool(b.get("total_row")) and i == len(rows) - 1
            for j, v in enumerate(r):
                cell = row.cells[j]
                par = cell.paragraphs[0]
                txt = v if not isinstance(v, (int, float)) or isinstance(v, bool) else fmt(v, 0 if float(v).is_integer() else 2)
                _add_text(par, str(txt), {"size": size, "bold": is_total})
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if numeric[j] else WD_ALIGN_PARAGRAPH.LEFT
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                if is_total:
                    _borders(cell._tc.get_or_add_tcPr(), "w:tcBorders", {"top": (8, tt["header_fill"] if not grid else "000000", "single")})
                elif tt.get("stripe") and i % 2 == 1:
                    _shade(cell, tt["stripe"])
        self.para("", after=4, base={"size": 4})

    def chart(self, b):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        pie = b["chart"] in ("pie", "doughnut")
        w = min(self.text_w, 13.0 if pie else 15.5)
        charts.add_chart(self.doc, par, b, self.th, width_cm=w, height_cm=7.5 if pie else 8.5)
        self.caption(b)

    def image(self, b):
        from pathlib import Path

        if not Path(str(b["path"])).exists():
            self.para(f"[image not found: {b['path']}]", base={"italic": True, "color": "C00000"})
            return
        w = min(self.text_w, float(b.get("width_cm") or self.text_w * 0.8))
        self.doc.add_picture(str(b["path"]), width=Cm(w))
        pic = self.doc.paragraphs[-1]
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.keep_with_next = True
        self.caption(b)

    def callout(self, b):
        th = self.th
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_width(t, [self.text_w])
        _cell_margins(t, 120, 120, 200, 160)
        cell = t.rows[0].cells[0]
        _shade(cell, th["table"].get("stripe") or "F2F2F2")
        _borders(
            cell._tc.get_or_add_tcPr(),
            "w:tcBorders",
            {"left": (24, th["accent"], "single"), "top": (0, "", ""), "bottom": (0, "", ""), "right": (0, "", "")},
        )
        par = cell.paragraphs[0]
        if b.get("title"):
            _add_text(par, b["title"], {"bold": True, "color": th["head_color"]})
            par.paragraph_format.space_after = Pt(3)
            par = cell.add_paragraph()
        _add_text(par, b.get("text") or "")
        par.paragraph_format.space_after = Pt(0)
        self.para("", after=4, base={"size": 4})

    def kv(self, b, centered=False, width=None):
        th = self.th
        items = b["items"]
        t = self.doc.add_table(rows=len(items), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER if centered else WD_TABLE_ALIGNMENT.LEFT
        w = width or min(self.text_w, 12.0)
        _table_width(t, [w * 0.38, w * 0.62])
        _cell_margins(t, 30, 30, 60, 60)
        _borders(t._tbl.tblPr, "w:tblBorders", {e: (0, "", "") for e in ("top", "left", "bottom", "right", "insideH", "insideV")})
        for i, (k, v) in enumerate(items):
            a, c = t.rows[i].cells
            _add_text(a.paragraphs[0], k, {"bold": True, "color": th["muted"]})
            _add_text(c.paragraphs[0], v)
            for cell in (a, c):
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        self.para("", after=6, base={"size": 4})

    def entry(self, b):
        th = self.th
        par = self.doc.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.tab_stops.add_tab_stop(Cm(self.text_w), WD_TAB_ALIGNMENT.RIGHT)
        _add_text(par, b["title"], {"bold": True, "size": th["size"] + 0.5})
        if b.get("org"):
            _add_text(par, "  ·  " + str(b["org"]), {"color": th["accent"]})
        if b.get("dates"):
            par.add_run("\t")
            _add_text(par, str(b["dates"]), {"size": th["size"] - 1, "color": th["muted"]})
        if b.get("place"):
            self.para(str(b["place"]), base={"italic": True, "size": th["size"] - 1, "color": th["muted"]}, after=2, keep=True)
        if b.get("text"):
            self.para(b["text"], after=2 if b.get("bullets") else th["after"] + 2)
        if b.get("bullets"):
            self.bullets({"items": [{"text": x, "items": []} for x in b["bullets"]], "numbered": False})
        else:
            par.paragraph_format.space_after = Pt(th["after"]) if not b.get("text") and not b.get("place") else par.paragraph_format.space_after

    def invoice_items(self, b):
        th = self.th
        c = b["computed"]
        cur, dec = c["currency"], c["decimals"]
        money = lambda x: f"{cur} {x:,.{dec}f}".strip()
        rows = [
            [str(i + 1), it["desc"], fmt(it["qty"], 0 if float(it["qty"]).is_integer() else 2), money(it["price"]), money(line)]
            for i, (it, line) in enumerate(zip(b["items"], c["lines"]))
        ]
        self.table({"columns": ["#", "Description", "Qty", "Unit price", "Amount"], "rows": rows, "numeric": [False, False, True, True, True]})
        tot = [["Subtotal", money(c["subtotal"])]]
        if c["discount"]:
            tot.append(["Discount", "- " + money(c["discount"])])
        if c["tax"]:
            tot.append([c["tax_label"] or "Tax", money(c["tax"])])
        tot.append(["Total", money(c["total"])])
        t = self.doc.add_table(rows=len(tot), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.RIGHT
        _table_width(t, [4.2, 4.0])
        _cell_margins(t, 40, 40, 90, 90)
        _borders(t._tbl.tblPr, "w:tblBorders", {e: (0, "", "") for e in ("top", "left", "bottom", "right", "insideH", "insideV")})
        for i, (k, v) in enumerate(tot):
            a, cc = t.rows[i].cells
            last = i == len(tot) - 1
            _add_text(a.paragraphs[0], k, {"bold": last, "color": th["table"]["header_text"] if last else th["muted"]})
            _add_text(cc.paragraphs[0], v, {"bold": last, "color": th["table"]["header_text"] if last else None})
            cc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for cell in (a, cc):
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                if last:
                    _shade(cell, th["table"]["header_fill"])
        self.para("", after=8, base={"size": 4})

    def address(self, b):
        self.para("\n".join(b["lines"]), align=b.get("align"), after=10)

    def date(self, b):
        self.para(b["text"], align=b.get("align"), after=10)

    def subject(self, b):
        txt = b["text"] if re.match(r"^(subject|re)\s*:", b["text"], re.I) else f"Subject: {b['text']}"
        self.para(txt, base={"bold": True}, after=10)

    def salutation(self, b):
        self.para(b["text"], after=8)

    def closing(self, b):
        self.para(b.get("text") or "Yours sincerely,", after=30, before=6)
        if b.get("name"):
            self.para(b["name"], base={"bold": True}, after=0)
        if b.get("title"):
            self.para(b["title"], base={"color": self.th["muted"]}, after=0)

    def signature(self, b):
        self.para("", after=24)
        self.para("_" * 28, align=b.get("align") or "left", after=2)
        for k in ("name", "title", "org", "date"):
            if b.get(k):
                self.para(
                    str(b[k]), align=b.get("align") or "left", base={"bold": k == "name", "color": None if k == "name" else self.th["muted"]}, after=0
                )

    def references(self, b):
        for x in b["items"]:
            par = self.para(x, base={"size": self.th["size"] - 0.5}, after=4)
            par.paragraph_format.left_indent = Cm(1.0)
            par.paragraph_format.first_line_indent = Cm(-1.0)

    def page_break(self, b):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def spacer(self, b):
        self.para("", after=float(b.get("height") or 12))

    # special layouts -------------------------------------------------------------------------------------------
    def invoice_head(self):
        th, p, m = self.th, self.p, self.p["meta"]
        t = self.doc.add_table(rows=1, cols=2)
        _table_width(t, [self.text_w * 0.55, self.text_w * 0.45])
        _borders(t._tbl.tblPr, "w:tblBorders", {e: (0, "", "") for e in ("top", "left", "bottom", "right", "insideH", "insideV")})
        left, right = t.rows[0].cells
        _add_text(left.paragraphs[0], m.get("org") or m.get("from") or "", {"bold": True, "size": th["size"] + 4, "color": th["head_color"]})
        for k in ("org_address", "address", "phone", "email", "ntn", "strn"):
            if m.get(k):
                q = left.add_paragraph()
                _add_text(q, (k.upper() + ": " if k in ("ntn", "strn") else "") + m[k], {"size": th["size"] - 1, "color": th["muted"]})
                q.paragraph_format.space_after = Pt(0)
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _add_text(
            rp, (p["title"] or ("QUOTATION" if p["doctype"] == "quotation" else "INVOICE")).upper(), {"size": th["title"] - 2, "color": th["accent"]}
        )
        for lbl, k in (("No.", "number"), ("Date", "date"), ("Due", "due")):
            if m.get(k):
                q = right.add_paragraph()
                q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                _add_text(q, f"**{lbl}** {m[k]}", {"size": th["size"] - 0.5})
                q.paragraph_format.space_after = Pt(0)
        self.para("", after=12, base={"size": 4})
        if m.get("bill_to") or m.get("client"):
            self.para("BILL TO", base={"bold": True, "size": th["size"] - 1.5, "color": th["muted"]}, after=0)
            self.para(m.get("bill_to") or m.get("client"), after=12)

    def certificate(self):
        th, p, m = self.th, self.p, self.p["meta"]
        s = self.doc.sections[0]
        sp = s._sectPr
        pb = OxmlElement("w:pgBorders")
        pb.set(qn("w:offsetFrom"), "page")
        for edge in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{edge}")
            for k, v in (("w:val", "double"), ("w:sz", "18"), ("w:space", "24"), ("w:color", th["accent"])):
                e.set(qn(k), v)
            pb.append(e)
        sp.append(pb)
        self.para("", after=30)
        if m.get("org"):
            self.para(m["org"].upper(), align="center", base={"bold": True, "color": th["muted"], "size": 12}, after=18)
        self.para(p["title"] or "Certificate of Achievement", style="Title", align="center", after=10)
        self.para(p["subtitle"] or "This is to certify that", align="center", base={"italic": True, "size": 13, "color": th["muted"]}, after=12)
        if m.get("recipient"):
            r = self.para(m["recipient"], align="center", base={"size": 30, "color": th["head_color"], "font": th["head"]}, after=12)
            _rule(r, th["accent"], size=6, where="bottom", space=6)

    def run(self, path):
        self.styles()
        self.page()
        self.header_footer()
        dt_ = self.p["doctype"]
        if dt_ == "certificate":
            self.certificate()
        if self.p["cover"]:
            self.cover()
        if self.p["toc"]:
            self.toc()
        if dt_ in ("invoice", "quotation"):
            self.invoice_head()
        self.title_block()
        for i, b in enumerate(self.p["blocks"]):
            if b["type"] == "toc":
                self.toc()
                continue
            fn = getattr(self, b["type"], None)
            if fn:
                fn(b)
                self.out_blocks.append(
                    {
                        "i": i,
                        "type": b["type"],
                        "text": plain(b.get("text") or b.get("title") or b.get("caption") or "")[:120],
                        "label": b.get("label"),
                    }
                )
        cp = self.doc.core_properties
        cp.title = self.p["title"][:250]
        cp.subject = self.p["subtitle"][:250]
        cp.author = self.p["meta"].get("author", "")[:250]
        cp.comments = "Made by the AI PC document agent"
        self.doc.save(str(path))
        return {"path": str(path), "blocks": self.out_blocks, "text_width_cm": self.text_w}


def build(plan, path):
    """Write the .docx for a plan (resolved here if it is not yet). Returns {"path", "blocks", "text_width_cm"}."""
    from ai_pc.office.docplan import resolve

    return _Builder(plan if "th" in plan else resolve(plan)).run(path)


# ------------------------------------------------------------------------------------------------ existing documents
def theme_of(doc, base="corporate"):
    """A theme for an existing document: its own body font and size, heading font and colour; the rest from `base`.
    New parts written into the document then look like the parts already there."""
    from ai_pc.office import themes

    th = themes.get(base)
    th = dict(th)
    th["table"] = dict(th["table"])

    def font_of(style):
        try:
            rpr = style.element.rPr
            if rpr is not None and rpr.rFonts is not None:
                f = rpr.rFonts.get(qn("w:ascii")) or rpr.rFonts.get(qn("w:hAnsi"))
                if f:
                    return f
            return style.font.name
        except Exception:  # noqa: BLE001
            return None

    normal = doc.styles["Normal"]
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    size = normal.font.size.pt if normal.font.size else None
    if size is None and defaults is not None:
        sz = defaults.find(".//" + qn("w:sz"))
        size = int(sz.get(qn("w:val"))) / 2 if sz is not None else None
    th["body"] = font_of(normal) or th["body"]
    th["size"] = size or 11.0
    try:
        h1 = doc.styles["Heading 1"]
        th["head"] = font_of(h1) or th["body"]
        col = h1.font.color.rgb if h1.font.color is not None and h1.font.color.type is not None else None
        if col is not None:
            hx = str(col)
            th["head_color"] = th["accent"] = hx
            th["table"]["header_fill"] = hx
            th["chart"] = [hx] + [c for c in th["chart"] if c != hx]
    except KeyError:
        th["head"] = th["body"]
    return th


def builder_on(doc, th, language="en"):
    """A _Builder writing into an existing document with its own look (no styles or page set up again)."""
    b = _Builder.__new__(_Builder)
    b.doc, b.th = doc, th
    sec = doc.sections[0]
    w = (
        (sec.page_width.cm if sec.page_width else 21.0)
        - (sec.left_margin.cm if sec.left_margin else 2.5)
        - (sec.right_margin.cm if sec.right_margin else 2.5)
    )
    b.text_w = round(w, 2)
    b.rtl = language == "ur"
    b.out_blocks = []
    b.p = {
        "th": th,
        "page": {"size": "A4", "orientation": "portrait", "margins": sec.left_margin.cm if sec.left_margin else 2.5},
        "language": language,
        "doctype": "other",
        "meta": {},
        "title": "",
        "subtitle": "",
        "cover": False,
        "toc": False,
        "header": "",
        "footer": {},
    }
    return b


def insert_blocks(doc, blocks, anchor=None, where="after", th=None, language="en", numbering=(0, 0)):
    """Plan blocks (docplan) written into an existing document: after / before the `anchor` body element, or at the
    end. Returns the new body elements. numbering: tables and figures already in the document (new labels continue)."""
    from ai_pc.office.docplan import resolve

    th = th or theme_of(doc)
    rp = resolve({"doctype": "other", "blocks": blocks, "cover": False, "toc": False, "header": "", "footer": {"page_numbers": False}})
    nt, nf = numbering
    for b in rp["blocks"]:
        if b.get("label", "").startswith("Table "):
            nt += 1
            b["label"] = f"Table {nt}"
        elif b.get("label", "").startswith("Figure "):
            nf += 1
            b["label"] = f"Figure {nf}"
    body = doc.element.body
    before = set(body)  # the elements themselves: ids of lxml proxies are reused once a proxy is freed
    bld = builder_on(doc, th, language)
    for b in rp["blocks"]:
        fn = getattr(bld, b["type"], None)
        if fn:
            fn(b)
    new = [el for el in body if el not in before and el.tag != qn("w:sectPr")]
    if anchor is not None and new:
        if where == "after":
            for el in reversed(new):
                anchor.addnext(el)
        else:
            for el in new:
                anchor.addprevious(el)
    return new
