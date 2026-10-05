"""A conversation about one document: questions, changes, versions (the video chat's recipe, for Word files).

  c = DocChat.start("out/docs/report/report.docx")          a chat about a copy of the file (the original is never changed)
  print(c.say("make the headings dark blue, add page numbers and a table of contents"))
  print(c.say("rewrite the introduction shorter and translate the conclusion into urdu"))
  print(c.say("undo"))  /  "go back to v1"  /  "what changed?"  /  "how many pages?"  /  "export"

A message is split into requests (doc_parse.clauses). Each is a question (answered from the document's map, or by the
cheap model for "what does X say"), a version command (undo, redo, go to vN, versions, compare, export), or a change.
Changes are read by rules (doc_parse); what the rules cannot read goes to the cheap model, which answers in the same
operations (checked against the document before use). All changes of a message make ONE new version: each operation is
applied (docx_ops), checked on the result, the version is rendered by Word in the background (contents page and page
numbers refreshed), and the reply says what was done, what could not be done and why, and what the checks found.
Chats live in out/docs/chats/<id>/ (chat.json, v0.docx, v1.docx, ... and their PDFs).
"""

import json
import re
import shutil
import time
from pathlib import Path

from docx import Document

from ai_pc.core.config import ROOT
from ai_pc.core.util import parse_json
from ai_pc.office import doc_parse as P
from ai_pc.office import docmap as DM
from ai_pc.office import docx_build as DB
from ai_pc.office import docx_ops as OP
from ai_pc.office import render as RN

CHATS = ROOT / "out" / "docs" / "chats"
UNDO = re.compile(
    r"^\s*(?:undo|revert|go back|take (?:that|it) back|cancel (?:that|it|the last (?:change|edit))|scratch that|never ?mind|"
    r"put it back|change it back|undo (?:that|the last (?:change|edit)|it))\b",
    re.I,
)
REDO = re.compile(r"^\s*redo\b|\bundo the undo\b", re.I)
GOTO = re.compile(r"\b(?:go|jump|switch|get|revert|take me) (?:back )?to (?:version|v) ?(\d+)\b|\buse (?:version|v) ?(\d+)\b|^\s*v(\d+)\s*$", re.I)
RESET = re.compile(r"\b(?:start over|from scratch|back to the original|the original (?:version|file|document)|reset (?:it|everything|all))\b", re.I)
VERSIONS = re.compile(
    r"\b(?:show|list|see) (?:me )?(?:the |all )?(?:versions|history|changes)\b|\bhow many versions\b|\bversion history\b|^\s*(?:versions?|history)\s*\??\s*$",
    re.I,
)
HISTORY = re.compile(r"\bwhat (?:did you|have you|you) (?:just )?(?:change|do|did)\b|\bwhat(?:'s| is| has)? (?:changed|different)\b", re.I)
COMPARE = re.compile(r"\bcompare\b|\bdifference between\b|\bdifferent (?:from|than) (?:the )?(?:original|first|v\d)", re.I)
EXPORT = re.compile(
    r"\b(?:export|save (?:it )?as (?:a )?pdf|(?:give|send) me (?:the )?(?:file|pdf|document|docx)|download|make (?:a|the) pdf|pdf (?:version|copy))\b",
    re.I,
)
ACK = re.compile(r"^\s*(?:ok(?:ay)?|cool|nice|great|perfect|awesome|thanks?(?: you)?|thank you|good|looks good|love it|fine|done)\W*$", re.I)
POLITE = re.compile(
    r"^\s*(?:(?:can|could|would|will) (?:you|u) (?:please |pls )?|please |pls |kindly |i (?:want|need|would like) (?:you )?(?:to )?|i'd like (?:you )?to |let's |lets )",
    re.I,
)
QUESTION = re.compile(
    r"^\s*(?:what|what's|which|how many|how much|how long|how big|is there|are there|is it|are the|does|do|did|who|where|when|why|"
    r"tell me|show me|list|summari[sz]e (?:the |this |my )?(?:document|file|report|it)\b|give me a summary)\b",
    re.I,
)
EDIT_START = re.compile(r"^\s*(?:" + P.EDIT_VERB + r")\b", re.I)
# "actually, undo that" / "ok so make it blue" / "no, the other one": the words in front are not the request
FILLER = re.compile(r"^\s*(?:(?:actually|oh|hmm+|wait|sorry|oops|ugh|alright|ok(?:ay)?|so|hey)\b[,!.]*\s+|(?:no|nah|nope)\s*[,!.]+\s*)+(?=\S)", re.I)

OPS_SYSTEM = """You turn a client's request about their Word document into edit operations for a program. Reply with ONE JSON
object: {"ops": [...], "ask": "<a short question back, only if the request cannot be done as asked>", "answer": "<only if it was a question>"}.
Operations (target = a descriptor below; leave out what you do not need):
  {"op": "style", "target": T, "set": {"font", "size" (pt or "+2"), "color", "bold", "italic", "underline", "align": "left|center|right|justify",
   "line_spacing", "space_after", "first_line_indent" (cm), "highlight", "case": "upper|lower|title|sentence"}}
  {"op": "page", "set": {"margins" (cm), "orientation": "portrait|landscape", "size": "A4|Letter|Legal", "columns"}}
  {"op": "page_numbers", "position": "bottom|top", "align": "left|center|right", "format": "Page X of Y|Page X|X", "remove": false}
  {"op": "header"|"footer", "text": "...", "remove": false}     {"op": "toc", "remove": false}
  {"op": "cover", "title": "...", "subtitle": "...", "lines": ["**Label:** value", ...], "remove": false}
  {"op": "theme", "name": "corporate|academic|modern|elegant|minimal|warm", "accent": "<colour>"}
  {"op": "table_style", "target": T, "accent": "<colour>"}     {"op": "structure"}     {"op": "heading_numbers", "remove": false}
  {"op": "list", "target": T, "kind": "bullets|numbered|none"}  {"op": "page_breaks", "level": 1, "remove": false}
  {"op": "replace", "find": "...", "with": "..."}              {"op": "emphasis", "find" | "terms": [...] | "auto": 8, "set": {"bold": true}, "target": T}
  {"op": "delete", "target": T}     {"op": "move", "target": T, "to": {"target": T2, "where": "before|after"}}
  {"op": "insert", "where": "before|after|start|end", "anchor": T, "heading": "...", "about": "<what to write>", "text": "<exact text>", "words": 150}
  {"op": "rewrite", "target": T, "instruction": "...", "keep_count": true|false}     {"op": "translate", "target": T, "language": "..."}
  {"op": "summarize", "target": T, "kind": "executive summary|abstract|key points|conclusion", "where": "start|end"}
  {"op": "table_sort", "target": T, "by": "<column>", "desc": true}   {"op": "table_add_row", "target": T, "values": [...]}
  {"op": "table_delete_row", "target": T, "row": 2 | "last", "match": "..."}   {"op": "table_add_column", "target": T, "name": "...",
   "formula": "[Qty] * [Price]" | "percent_of": "<column>" | "values": [...]}   {"op": "table_delete_column", "target": T, "name": "..."}
  {"op": "table_total", "target": T}    {"op": "chart", "target": T, "chart": "column|bar|pie|line", "categories": "<column>", "values": ["<column>"]}
Targets T: {"kind": "all"} {"kind": "body"} {"kind": "title"} {"kind": "headings", "level": 1|2} {"kind": "section", "name": "<exact
title from the outline>"} {"kind": "section_body", "name": ...} {"kind": "section", "n": 2} {"kind": "paragraph", "n": 2, "in":
{"kind": "section", "name": ...}} {"kind": "paragraph", "which": "last"} {"kind": "paragraphs", "contains": "..."} {"kind": "table", "n": 1}
{"kind": "table", "which": "all"} {"kind": "list", "n": 1}
Use section titles exactly as the OUTLINE shows them (without their numbers). Do only what was asked; the fewest operations."""


class DocChat:
    def __init__(self, state, planner=None, log=print, render=True, server=None):
        self.state, self.log, self.render = state, log, render
        self.planner = planner
        self.server = server if server is not None else (RN.Server() if render else None)
        self.folder = Path(state["folder"])
        self._maps = {}

    # ------------------------------------------------------------------ files and versions
    @classmethod
    def start(cls, path, chats_dir=None, **kw):
        src = Path(path)
        ext = src.suffix.lower()
        if ext != ".docx":
            raise ValueError("this chat edits Word documents (.docx); decks and workbooks have their own")
        cid = f"chat_{re.sub(r'[^a-z0-9]+', '', src.stem.lower())[:16] or 'doc'}_{time.strftime('%H%M%S')}"
        folder = Path(chats_dir or CHATS) / cid
        folder.mkdir(parents=True, exist_ok=True)
        v0 = folder / f"v0{ext}"
        shutil.copy(src, v0)
        state = {
            "id": cid,
            "kind": "docx",
            "base": str(src),
            "folder": str(folder),
            "cur": 0,
            "turns": [],
            "focus": None,
            "pending": None,
            "versions": [
                {
                    "v": 0,
                    "parent": None,
                    "file": str(v0),
                    "said": None,
                    "done": ["the original file"],
                    "failed": [],
                    "checks": [],
                    "pages": None,
                    "pdf": None,
                }
            ],
        }
        c = cls(state, **kw)
        if c.render:
            c._render(0)
        c.save()
        return c

    @classmethod
    def load(cls, chat_id, chats_dir=None, **kw):
        p = Path(chats_dir or CHATS) / chat_id / "chat.json"
        return cls(json.loads(p.read_text(encoding="utf-8")), **kw)

    def save(self):
        (self.folder / "chat.json").write_text(json.dumps(self.state, ensure_ascii=False, indent=1), encoding="utf-8")

    def close(self):
        if self.server is not None:
            self.server.close()

    @property
    def version(self):
        return self.state["versions"][self.state["cur"]]

    def path(self, v=None):
        return Path(self.state["versions"][self.state["cur"] if v is None else v]["file"])

    def map(self, v=None):
        v = self.state["cur"] if v is None else v
        if v not in self._maps:
            self._maps[v] = DM.docx_map(self.path(v))
        return self._maps[v]

    def _render(self, v):
        ver = self.state["versions"][v]
        r = RN.to_pdf(ver["file"], pdf=str(Path(ver["file"]).with_suffix(".pdf")), update_fields=True, save=True, server=self.server)
        ver["pages"], ver["pdf"], ver["render_ok"] = r.get("pages"), r.get("pdf") if r.get("ok") else None, bool(r.get("ok"))
        ver["words"] = r.get("words")
        self._maps.pop(v, None)  # Word wrote the contents page and fields into the file
        return r

    # ------------------------------------------------------------------ a message
    def say(self, message, files=()):
        t0 = time.perf_counter()
        turn = {"user": message, "intents": [], "ops": [], "v_before": self.state["cur"], "llm": False}
        self._turn = turn
        out, ops, unknown = [], [], []
        pend, self.state["pending"] = self.state.get("pending"), None
        if pend and pend.get("template") and not EDIT_START.search(message) and not QUESTION.search(message):
            message = pend["template"].format(message.strip(" ."))
        focus = self.state.get("focus")
        images = [f for f in files or [] if Path(f).suffix.lower() in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")]
        if images:
            m = self.map()
            t = P.target_in(P.normalize(message), m, focus)
            q = P.quoted(message)
            for img in images:
                ops.append({"op": "image", "path": str(img), **({"anchor": t, "where": "after"} if t else {}), **({"caption": q[0]} if q else {})})
            turn["intents"].append("image")
        prev_change = False
        for i, cl in enumerate(P.clauses(message) or [message]):
            c = POLITE.sub("", FILLER.sub("", cl)).strip()
            if not c:
                continue
            kind, payload = self._clause(c, cl, focus, carry=prev_change)
            turn["intents"].append(kind)
            prev_change = kind == "change"
            if kind == "change":
                ops += payload["ops"]
                focus = payload.get("focus") or focus
            elif kind in ("undo", "redo", "goto", "reset", "versions", "history", "compare", "export"):
                if ops:  # what came before in the message happens first
                    out.append(self._change(ops, message, turn))
                    ops = []
                out.append(payload())
            elif kind in ("question", "ack"):
                out.append(payload)
            elif kind == "ask":
                out.append(payload["question"])
                if payload.get("template"):
                    self.state["pending"] = payload
            else:
                unknown.append(cl)
        if unknown:
            r = self._llm(unknown, message, ops, focus)
            if r.get("ops"):
                ops += r["ops"]
                turn["intents"].append("change")
            if r.get("answer"):
                out.append(r["answer"])
            if r.get("ask"):
                out.append(r["ask"])
        if ops:
            out.append(self._change(ops, message, turn))
            self.state["focus"] = focus
        reply = (
            "\n".join(x for x in out if x).strip()
            or "I'm not sure what to change. You can ask e.g. 'make the headings blue', "
            "'add page numbers', 'shorten the introduction' or 'how many pages?'."
        )
        turn.update(reply=reply, v_after=self.state["cur"], seconds=round(time.perf_counter() - t0, 2))
        self.state["turns"].append(turn)
        self.last_turn = turn
        self.save()
        return reply

    def _meta(self, c, raw):
        """Version commands, export and thanks, the same for every kind of document: (kind, payload) or (None, None)."""
        low = c.lower()
        g = GOTO.search(low)  # "go back to v2" is a jump, not an undo
        if g:
            n = int(next(x for x in g.groups() if x))
            return "goto", lambda: self.goto(n)
        if REDO.search(low):
            return "redo", self.redo
        if UNDO.search(low):
            nm = re.search(r"\b(two|2|three|3|four|4|five|5)\s+(?:changes?|steps?|edits?|times|versions?)\b|\b(twice)\b", low)
            n = {"two": 2, "2": 2, "three": 3, "3": 3, "four": 4, "4": 4, "five": 5, "5": 5, "twice": 2}.get(
                next((g for g in nm.groups() if g), "") if nm else "", 1
            )

            def undo_n(n=n):
                outs = [self.undo() for _ in range(n)]
                return outs[-1] if n == 1 else f"Undone {n} changes: back to v{self.state['cur']}."

            return "undo", undo_n
        if RESET.search(low):
            return "reset", lambda: self.goto(0)
        if VERSIONS.search(low):
            return "versions", self.versions_text
        if HISTORY.search(low):
            return "history", self.history
        if COMPARE.search(low):
            return "compare", lambda: self.compare(low)
        if EXPORT.search(low) and not re.search(r"\b(?:make|add|change)\b.*\b(?:heading|font|title|table)\b", low):
            return "export", self.export
        if ACK.match(low):
            return "ack", "Glad it works. Say 'export' for the PDF, or tell me what to change next."
        return None, None

    def _clause(self, c, raw, focus, carry=False):
        kind, payload = self._meta(c, raw)
        if kind:
            return kind, payload
        low = c.lower()
        if QUESTION.search(low) and not POLITE.match(raw):
            return "question", self.answer(low)
        r = P.parse(c, self.map(), focus, raw=raw, carry=carry)
        if r.get("ops"):
            return "change", r
        if r.get("ask"):
            return "ask", {"question": r["ask"], "template": None}
        return "unknown", c

    # ------------------------------------------------------------------ the model, for what the rules cannot read
    def _llm(self, clauses, message, done_ops, focus):
        if self.planner is None:
            return {"ask": "I could not read: " + "; ".join(f"'{c}'" for c in clauses) + ". Try e.g. 'make the headings blue' or 'add page numbers'."}
        self._turn["llm"] = True
        m = self.map()
        user = (
            f"OUTLINE:\n{DM.outline_text(m)}\nLAST THING TALKED ABOUT: {json.dumps(focus) if focus else 'nothing'}\n"
            f"ALREADY DONE FROM THIS MESSAGE: {[OP.describe(o) for o in done_ops] or 'nothing'}\nWHOLE MESSAGE: {message}\n"
            f"REQUESTS TO TURN INTO OPERATIONS: {clauses}"
        )
        try:
            r = self.planner._call("docs", [{"role": "system", "content": OPS_SYSTEM}, {"role": "user", "content": user}])
            d = parse_json(r.text) or {}
        except Exception as e:  # noqa: BLE001
            return {"ask": f"I could not work that out ({type(e).__name__}). Could you say it differently?"}
        ops = []
        problems = []
        for op in d.get("ops") or []:
            why = self._valid(op, m)
            if why:
                problems.append(why)
            else:
                ops.append(op)
        out = {"ops": ops, "answer": d.get("answer"), "ask": d.get("ask")}
        if problems and not ops and not out["ask"]:
            out["ask"] = "I could not do that: " + "; ".join(problems[:2])
        return out

    def _valid(self, op, m):
        if not isinstance(op, dict) or op.get("op") not in OP.OPS:
            return f"no such edit ({op.get('op') if isinstance(op, dict) else op})"
        for key in ("target", "anchor"):
            t = op.get(key)
            if t is not None:
                if not isinstance(t, dict) or t.get("kind") not in DM.TARGETS:
                    return f"unclear {key}"
                if not DM.find(m, t):
                    return f"{DM.describe_target(m, t)} is not in the document"
        to = (op.get("to") or {}).get("target") if isinstance(op.get("to"), dict) else None
        if op["op"] == "move" and (not to or not DM.find(m, to)):
            return "where to move it is unclear"
        return None

    # ------------------------------------------------------------------ changes
    def _change(self, ops, said, turn):
        """All the operations as ONE new version; each checked; Word renders it; the reply."""
        v0 = self.state["cur"]
        doc = Document(str(self.path(v0)))
        ctx = {"planner": self.planner, "th": DB.theme_of(doc), "info": {}}
        done, failed, checks = [], [], []
        for op in ops:
            before = DM.docx_map(doc)
            ctx["info"] = {}
            try:
                d = OP.apply(doc, op, ctx)
            except OP.OpError as e:
                failed.append(f"{OP.describe(op)}: {e}")
                continue
            except Exception as e:  # noqa: BLE001  (a broken edit must not end the chat)
                failed.append(f"{OP.describe(op)}: {type(e).__name__}: {str(e)[:120]}")
                continue
            after = DM.docx_map(doc)
            try:
                ok, what = OP.check(op, before, after, ctx.get("info"), doc_after=doc)
            except Exception as e:  # noqa: BLE001
                ok, what = False, f"check failed: {type(e).__name__}"
            done.append(d)
            checks.append({"op": OP.describe(op), "ok": bool(ok), "what": what})
            turn["ops"].append(op)
        if not done:
            return "Nothing changed: " + "; ".join(failed or ["that is already how it is"]) + "."
        n = len(self.state["versions"])
        dst = self.folder / f"v{n}.docx"
        doc.save(str(dst))
        ver = {"v": n, "parent": v0, "file": str(dst), "said": said, "done": done, "failed": failed, "checks": checks, "pages": None, "pdf": None}
        self.state["versions"].append(ver)
        self.state["cur"] = n
        page_note = ""
        if self.render:
            r = self._render(n)
            if r.get("ok"):
                p0 = self.state["versions"][v0].get("pages")
                page_note = f" Now {r['pages']} page(s)" + (f" (was {p0})" if p0 and p0 != r["pages"] else "") + "."
                checks += self._render_checks(ops, n)
            else:
                page_note = f" (Word could not render it: {str(r.get('error'))[:80]})"
        bad = [c for c in checks if not c["ok"]]
        msg = f"v{n}: " + "; ".join(done) + "."
        if failed:
            msg += " Couldn't: " + "; ".join(failed) + "."
        msg += (
            f" Checked: {len(checks) - len(bad)}/{len(checks)} OK"
            + (" (" + "; ".join(f"{c['op']}: {c['what']}" for c in bad[:3]) + ")" if bad else "")
            + "."
        )
        return msg + page_note

    def _render_checks(self, ops, v):
        """What only the rendered pages show: fonts really used, the contents page filled in, page numbers on the pages."""
        from ai_pc.office import verify as VF

        ver = self.state["versions"][v]
        out = []
        if not ver.get("pdf"):
            return out
        kinds = {o.get("op") for o in ops}
        fonts = {f for o in ops if o.get("op") in ("style", "theme") for f in [(o.get("set") or {}).get("font")] if f}
        if fonts or "theme" in kinds:
            gf = VF.glyph_fonts(ver["pdf"])
            for f in fonts:
                n = sum(c for name, c in gf.items() if VF._is_family(name, {VF._family(f)}))
                out.append({"op": f"font {f} on the pages", "ok": n > 0, "what": f"{n} characters drawn in {f}"})
        if "toc" in kinds and not any(o.get("remove") for o in ops if o.get("op") == "toc"):
            txt = " ".join(RN.text(ver["pdf"])[:4]).lower()
            filled = "contents" in txt and "right-click to update" not in txt
            out.append({"op": "contents page", "ok": filled, "what": "filled in by Word" if filled else "not filled in"})
        if "page_numbers" in kinds and not any(o.get("remove") for o in ops if o.get("op") == "page_numbers"):
            pages = RN.text(ver["pdf"])
            ok = sum(1 for i, t in enumerate(pages) if re.search(rf"(?:^|\D){i + 1}(?:\D|$)", t[-120:] + " " + t[:120])) >= max(1, len(pages) - 1)
            out.append(
                {"op": "page numbers on the pages", "ok": ok, "what": f"{len(pages)} page(s) numbered" if ok else "numbers not seen on the pages"}
            )
        return out

    # ------------------------------------------------------------------ questions
    def answer(self, q):
        m = self.map()
        ver = self.version
        st = m["stats"]
        many = re.findall(r"\bhow many (pages|words|tables|headings|sections|paragraphs|charts|images|pictures|versions)\b", q)
        if len(many) > 1:  # "how many pages and how many words?"
            return " ".join(self.answer(f"how many {k}") for k in dict.fromkeys(many))
        mm = re.search(r"\bhow many (pages|words|tables|headings|sections|paragraphs|charts|images|pictures|versions)\b", q)
        if mm:
            k = mm.group(1)
            if k == "pages":
                return f"{ver.get('pages') or '?'} page(s)."
            if k == "words":
                return f"About {ver.get('words') or st['words']} words."
            if k == "versions":
                return self.versions_text()
            if k == "sections":
                return f"{sum(1 for s in m['sections'] if s['level'] == min((x['level'] for x in m['sections']), default=1))} section(s)."
            k2 = {"pictures": "images"}.get(k, k)
            return f"{st.get(k2, 0)} {k}."
        if re.search(r"\b(?:title|called|name of (?:the|this) document)\b", q) and not re.search(r"\bsection", q):
            return f"The title is '{m['title']}'."
        if re.search(r"\b(?:headings|sections|outline|structure|contents|chapters)\b", q) and re.search(r"\b(?:what|list|show|which)\b", q):
            return (
                "Sections: " + "; ".join(("  " * (s["level"] - 1)) + s["title"] for s in m["sections"][:30])
                if m["sections"]
                else "The document has no headings."
            )
        if re.search(r"\bfont\b", q):
            try:
                h1 = DM._style_font(Document(str(self.path())).styles["Heading 1"])
            except KeyError:
                h1 = None
            return (
                f"Body text: {m['fonts']['body'] or '?'}"
                + (f" {m['fonts']['body_size']} pt" if m["fonts"]["body_size"] else "")
                + (f"; headings: {h1}" if h1 else "")
                + "."
            )
        if re.search(r"\bmargins?\b|\borientation\b|\bpage size\b|\blandscape\b|\bportrait\b", q):
            pg = m["page"]
            return f"{pg['orientation'].title()}, {pg['width_cm']} x {pg['height_cm']} cm, margins {pg['margins_cm']} cm, {pg['columns']} column(s)."
        if re.search(r"\b(?:table of contents|contents page|toc)\b", q):
            return "Yes, there is a table of contents." if m["toc"] else "No table of contents yet ('add a table of contents' adds one)."
        if re.search(r"\bpage numbers?\b", q):
            return "Yes, the pages are numbered." if m["page_numbers"] else "No page numbers yet."
        if re.search(r"\bheader\b|\bfooter\b", q):
            return f"Header: '{m['header'] or '(none)'}'; footer: '{m['footer'] or '(none)'}'."
        mt = re.search(r"\btable\s*(\d+)?\b", q)
        if mt and m["tables"]:
            t = m["tables"][int(mt.group(1) or 1) - 1] if mt.group(1) and int(mt.group(1)) <= len(m["tables"]) else m["tables"][0]
            return f"Table {t['t']}: {t['rows']} rows x {t['cols']} columns; columns: {', '.join(t['header'])}."
        if re.search(r"\b(?:author|wrote|written by)\b", q):
            return f"Author: {Document(str(self.path())).core_properties.author or 'not set'}."
        return self._llm_answer(q)

    def _llm_answer(self, q):
        if self.planner is None:
            return "I can tell you pages, words, the title, the sections, fonts, margins and tables; what would you like to know?"
        self._turn["llm"] = True
        m = self.map()
        t = P.target_in(P.normalize(q), m, self.state.get("focus"))
        ks = DM.find(m, t) if t else []
        body = (
            "\n".join(m["items"][k]["text"] for k in ks if m["items"][k]["kind"] == "p")[:9000]
            if ks
            else "\n".join(it["text"] for it in m["items"] if it["kind"] == "p")[:9000]
        )
        try:
            r = self.planner._call(
                "fast",
                [
                    {
                        "role": "system",
                        "content": "You answer a client's question about their document in 1-4 short sentences, only "
                        "from the text given. If the text does not say, say so.",
                    },
                    {"role": "user", "content": f"OUTLINE:\n{DM.outline_text(m, 30)}\n\nTEXT:\n{body}\n\nQUESTION: {q}"},
                ],
            )
            return re.sub(r"<think>.*?</think>", "", r.text or "", flags=re.S).strip()[:700] or "I could not find that in the document."
        except Exception as e:  # noqa: BLE001
            return f"I could not answer that just now ({type(e).__name__})."

    # ------------------------------------------------------------------ versions
    def undo(self):
        parent = self.version.get("parent")
        if parent is None:
            return "Nothing to undo: this is the original file."
        was = self.state["cur"]
        self.state["cur"] = parent
        return f"Undone: back to v{parent} (v{was} was: {'; '.join(self.state['versions'][was]['done'])[:160]})."

    def redo(self):
        kids = [v for v in self.state["versions"] if v.get("parent") == self.state["cur"]]
        if not kids:
            return "Nothing to redo."
        self.state["cur"] = kids[-1]["v"]
        return f"Redone: v{kids[-1]['v']} ({'; '.join(kids[-1]['done'])[:160]})."

    def goto(self, n):
        if not 0 <= n < len(self.state["versions"]):
            return f"There is no v{n} (versions 0-{len(self.state['versions']) - 1})."
        self.state["cur"] = n
        v = self.state["versions"][n]
        return f"Now at v{n}" + (f": {'; '.join(v['done'])[:160]}" if n else " (the original file)") + "."

    def versions_text(self):
        lines = [
            f"{'*' if v['v'] == self.state['cur'] else ' '} v{v['v']}: "
            + ("; ".join(v["done"])[:110])
            + (f" ({v['pages']} p.)" if v.get("pages") else "")
            for v in self.state["versions"]
        ]
        return f"{len(self.state['versions'])} version(s):\n" + "\n".join(lines)

    def history(self):
        v = self.version
        if v["v"] == 0:
            return "Nothing changed yet."
        return (
            f"In v{v['v']} (from '{v['said']}'): " + "; ".join(v["done"]) + (f". Not done: {'; '.join(v['failed'])}" if v.get("failed") else "") + "."
        )

    def compare(self, text=""):
        nums = [int(x) for x in re.findall(r"\bv(?:ersion)?\s*(\d+)", text)]
        a = nums[0] if nums else 0
        b = nums[1] if len(nums) > 1 else self.state["cur"]
        if not (0 <= a < len(self.state["versions"]) and 0 <= b < len(self.state["versions"])):
            return "No such version."
        ma, mb = self.map(a), self.map(b)
        ha, hb = [s["title"] for s in ma["sections"]], [s["title"] for s in mb["sections"]]
        added = [h for h in hb if h not in ha]
        gone = [h for h in ha if h not in hb]
        bits = [
            f"words {ma['stats']['words']} -> {mb['stats']['words']}",
            f"pages {self.state['versions'][a].get('pages')} -> {self.state['versions'][b].get('pages')}",
            f"tables {ma['stats']['tables']} -> {mb['stats']['tables']}",
            f"charts {ma['stats']['charts']} -> {mb['stats']['charts']}",
        ]
        if added:
            bits.append("new sections: " + ", ".join(added[:4]))
        if gone:
            bits.append("removed: " + ", ".join(gone[:4]))
        if ma["fonts"]["body"] != mb["fonts"]["body"]:
            bits.append(f"body font {ma['fonts']['body']} -> {mb['fonts']['body']}")
        if ma["toc"] != mb["toc"]:
            bits.append("contents page " + ("added" if mb["toc"] else "removed"))
        if ma["page_numbers"] != mb["page_numbers"]:
            bits.append("page numbers " + ("added" if mb["page_numbers"] else "removed"))
        return f"v{a} -> v{b}: " + "; ".join(bits) + "."

    def export(self):
        v = self.version
        if not v.get("pdf") and self.render:
            self._render(self.state["cur"])
        name = Path(self.state["base"]).stem
        out_docx = self.folder / f"{name}_v{v['v']}.docx"
        shutil.copy(v["file"], out_docx)
        msg = f"Saved: {out_docx}"
        if v.get("pdf"):
            out_pdf = self.folder / f"{name}_v{v['v']}.pdf"
            shutil.copy(v["pdf"], out_pdf)
            msg += f" and {out_pdf} ({v.get('pages')} pages)"
        return msg + "."
