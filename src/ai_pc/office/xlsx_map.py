"""The map of an Excel workbook: its sheets, the table on each (header row, data rows, a totals row, columns with their
kinds and formulas), summary sheets as label / value pairs, charts, conditional formats, filters and frozen panes.

  m = book_map("book.xlsx")      read with openpyxl twice (formulas, and the values Excel saved); never saved by it
  t = find_table(F, V, fmts)     the same detection on any grid (the Excel worker uses it on what Excel reads)
  col(m, sheet, "maths")         a column by its name (case, plurals, a prefix), or by its letter
"""

import datetime as _dt
import re

LABEL_TOTAL = re.compile(r"^\s*(?:grand\s+)?(?:total|totals|sum|average|averages|avg|mean|summary|overall)\b", re.I)
AGG_FN = re.compile(r"=\s*(SUM|AVERAGE|COUNT|COUNTA|MAX|MIN|MEDIAN|SUBTOTAL)\s*\(", re.I)


def letter(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def col_index(s):
    n = 0
    for ch in s.upper():
        n = n * 26 + ord(ch) - 64
    return n


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def is_formula(f):
    return isinstance(f, str) and f.startswith("=")


def _empty(v):
    return v is None or (isinstance(v, str) and not v.strip())


def plain(v):
    """A cell value as JSON can carry it."""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return (
            v.isoformat()[:10]
            if isinstance(v, _dt.date)
            and not isinstance(v, _dt.datetime)
            or (isinstance(v, _dt.datetime) and not v.time().hour and not v.time().minute)
            else v.isoformat()
        )
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return int(v)
    return v


AGG_OWN = re.compile(
    r"^=\s*(?:SUM|AVERAGE|COUNT|COUNTA|MAX|MIN|MEDIAN|SUBTOTAL\(\s*\d+\s*,)\s*\(?\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)\s*$",
    re.I,
)


def _sums_above(frow, vrow, j0, j1, c0, first, prev):
    """A row whose formulas all aggregate their own column over the rows above it (=SUM(G4:G19) under G4:G19), with
    nothing else in it but a label: a totals row even without a 'Total' label."""
    agg = 0
    for j in range(j0, j1 + 1):
        f = frow[j] if j < len(frow) else None
        v = vrow[j] if j < len(vrow) else None
        if _empty(f):
            continue
        if is_formula(f):
            m = AGG_OWN.match(str(f))
            if not (m and m.group(1).upper() == m.group(3).upper() == letter(c0 + j) and int(m.group(2)) <= first + 1 and int(m.group(4)) >= prev):
                return False
            agg += 1
        elif is_num(v):
            return False
    return agg >= 1


def find_table(F, V, fmts=None, r0=1, c0=1):
    """The table on a sheet from its grid of formulas F and values V (lists of rows, top-left at r0, c0): the header row
    (mostly text, above rows of data), the data rows down to an empty row, and a totals row under them (a label like
    Total / Average, or formulas that aggregate under rows that do not). fmts: {column number: number format}.
    None when the sheet holds no table."""
    nrows = len(F)
    for i in range(min(nrows - 1, 15)):
        row = F[i]
        cells = [(j, x) for j, x in enumerate(row) if not _empty(x)]
        if len(cells) < 2:
            continue
        texts = [x for _, x in cells if isinstance(x, str) and not is_formula(x) and not re.fullmatch(r"[-+]?[\d,.]+%?", x.strip())]
        if len(texts) < 0.7 * len(cells):
            continue
        nxt = [x for x in F[i + 1] if not _empty(x)] if i + 1 < nrows else []
        if len(nxt) < 2:
            continue
        j0, j1 = cells[0][0], cells[-1][0]
        header = i

        def filled(r):
            return any(not _empty(F[r][j]) for j in range(j0, j1 + 1) if j < len(F[r]))

        def more_data(r):  # after a blank row: the same table goes on (numbers where it has numbers), not a new block
            if any(not _empty(x) for j, x in enumerate(F[r]) if j < j0 or j > j1):
                return False
            vals = [(F[r][j], V[r][j] if j < len(V[r]) else None) for j in range(j0, j1 + 1) if j < len(F[r]) and not _empty(F[r][j])]
            return any(is_num(v) or is_formula(f) or isinstance(v, (_dt.date, _dt.datetime)) for f, v in vals) or len(vals) < 2

        k, last = header + 1, header
        while k < nrows:
            if filled(k):
                last, k = k, k + 1
                continue
            nxt = next((r for r in range(k + 1, min(nrows, k + 3)) if filled(r)), None)
            if nxt is None or not more_data(nxt):
                break
            k = nxt
        total = None
        if last > header + 1:
            lab = next((F[last][j] for j in range(j0, j1 + 1) if j < len(F[last]) and not _empty(F[last][j])), None)
            aggs = sum(1 for j in range(j0, j1 + 1) if j < len(F[last]) and is_formula(F[last][j]) and AGG_FN.search(str(F[last][j])))
            prev = sum(1 for j in range(j0, j1 + 1) if j < len(F[last - 1]) and is_formula(F[last - 1][j]) and AGG_FN.search(str(F[last - 1][j])))
            if (
                (isinstance(lab, str) and LABEL_TOTAL.match(lab))
                or (aggs >= max(1, (j1 - j0 + 1) // 3) and not prev)
                or _sums_above(F[last], V[last], j0, j1, c0, r0 + header + 1, r0 + last - 1)
            ):
                total = last
                last -= 1
        cols = []
        for j in range(j0, j1 + 1):
            name = F[header][j] if j < len(F[header]) else None
            name = str(name).strip() if not _empty(name) else f"Column {letter(c0 + j)}"
            vals = [V[r][j] if j < len(V[r]) else None for r in range(header + 1, last + 1)]
            forms = [F[r][j] if j < len(F[r]) else None for r in range(header + 1, last + 1)]
            nums = [v for v in vals if is_num(v)]
            filled = [v for v in vals if not _empty(v)]
            fmt = (fmts or {}).get(c0 + j) or ""
            if any(isinstance(v, (_dt.datetime, _dt.date)) for v in filled) and sum(
                isinstance(v, (_dt.datetime, _dt.date)) for v in filled
            ) >= 0.6 * len(filled):
                kind = "date"
            elif filled and len(nums) >= 0.8 * len(filled):
                kind = "percent" if "%" in fmt else "money" if re.search(r"[$€£₨]|Rs|PKR|USD", fmt) else "number"
            else:
                kind = "text"
            f0 = next((f for f in forms if is_formula(f)), None)
            tf = F[total][j] if total is not None and j < len(F[total]) else None
            cols.append(
                {
                    "name": name,
                    "idx": c0 + j,
                    "letter": letter(c0 + j),
                    "kind": kind,
                    "fmt": fmt,
                    "formula": f0,
                    "formulas": sum(1 for f in forms if is_formula(f)),
                    "blank": len(vals) - len(filled),
                    "total": tf if is_formula(tf) else None,
                }
            )
        rows = [[(V[r][j] if j < len(V[r]) else None) for j in range(j0, j1 + 1)] for r in range(header + 1, last + 1)]
        return {
            "header": r0 + header,
            "first": r0 + header + 1,
            "last": r0 + last,
            "total": r0 + total if total is not None else None,
            "c0": c0 + j0,
            "c1": c0 + j1,
            "cols": cols,
            "rows": rows,
            "total_label": (next((F[total][j] for j in range(j0, j1 + 1) if not _empty(F[total][j])), None) if total is not None else None),
        }
    return None


def _pairs(F, V, r0=1):
    out = []
    for i, row in enumerate(F):
        cells = [(j, x) for j, x in enumerate(row) if not _empty(x)]
        if len(cells) >= 2 and isinstance(cells[0][1], str) and not is_formula(cells[0][1]):
            j = cells[1][0]
            out.append(
                {
                    "row": r0 + i,
                    "label": cells[0][1].strip(),
                    "value": plain(V[i][j] if j < len(V[i]) else None),
                    "formula": cells[1][1] if is_formula(cells[1][1]) else None,
                    "cell": f"{letter(j + 1)}{r0 + i}",
                }
            )
    return out


def book_map(path, max_rows=5000):
    import warnings

    import openpyxl

    with warnings.catch_warnings():  # parts openpyxl cannot read (it never writes this file) are not news
        warnings.simplefilter("ignore")
        wf = openpyxl.load_workbook(str(path), data_only=False)
        wv = openpyxl.load_workbook(str(path), data_only=True)
    sheets = []
    for ws in wf.worksheets:
        vs = wv[ws.title]
        n = min(ws.max_row, max_rows)
        F = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=n, max_col=ws.max_column)]
        V = [[c.value for c in row] for row in vs.iter_rows(min_row=1, max_row=n, max_col=ws.max_column)]
        t = find_table(F, V) if F else None
        if t:
            t = find_table(F, V, {c["idx"]: ws.cell(t["first"], c["idx"]).number_format for c in t["cols"]})
            t["rows"] = [[plain(v) for v in r] for r in t["rows"]]
        cfs = 0
        try:
            cfs = sum(len(r.rules) for r in ws.conditional_formatting)
        except Exception:  # noqa: BLE001
            pass
        pivots = len(getattr(ws, "_pivots", []) or [])
        sheets.append(
            {
                "name": ws.title,
                "dims": ws.dimensions,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "pivots": pivots,
                "kind": "pivot" if pivots else "table" if t else ("summary" if any(not _empty(x) for r in F for x in r) else "empty"),
                "table": t,
                "pairs": _pairs(F[: t["header"] - 1], V[: t["header"] - 1]) if t else _pairs(F, V),
                "charts": len(getattr(ws, "_charts", [])),
                "cf": cfs,
                "freeze": ws.freeze_panes,
                "filter": ws.auto_filter.ref,
                "validations": len(ws.data_validations.dataValidation) if ws.data_validations else 0,
                "state": ws.sheet_state,
            }
        )
    tables = [s for s in sheets if s["table"] and not s["pivots"]]
    main = max(tables, key=lambda s: len(s["table"]["rows"]) * len(s["table"]["cols"]))["name"] if tables else (sheets[0]["name"] if sheets else None)
    return {"sheets": sheets, "main": main, "names": [s["name"] for s in sheets], "active": wf.active.title if wf.active else None}


def sheet(m, name=None):
    if name:
        for s in m["sheets"]:
            if s["name"].lower() == str(name).lower():
                return s
        for s in m["sheets"]:
            if s["name"].lower().startswith(str(name).lower()) or str(name).lower() in s["name"].lower():
                return s
        return None
    return next((s for s in m["sheets"] if s["name"] == m["main"]), None)


def _stem(w):
    w = w.lower().strip()
    w = re.sub(r"[^a-z0-9%#]+", " ", w).strip()
    return re.sub(r"(?:ies)$", "y", re.sub(r"(?<=[a-z]{3})(?:es|s)$", "", w))


def col(m, sh, name):
    """A column of a sheet's table by its header (exact, then singular / plural, then a prefix or a word of it), or by
    its letter ("column C")."""
    s = sheet(m, sh) if not isinstance(sh, dict) else sh
    if not s or not s.get("table"):
        return None
    cols = s["table"]["cols"]
    name = str(name or "").strip()
    mm = re.fullmatch(r"(?:col(?:umn)?\s+)?([A-Za-z]{1,2})", name)
    if mm and len(name) <= 9 and (name.lower().startswith("col") or name.isupper()):
        for c in cols:
            if c["letter"] == mm.group(1).upper():
                return c
    low = name.lower()
    for c in cols:
        if c["name"].lower() == low:
            return c
    st = _stem(low)
    for c in cols:
        if _stem(c["name"]) == st:
            return c
    if len(st) >= 3:
        for c in cols:
            if _stem(c["name"]).startswith(st) or st.startswith(_stem(c["name"])) and len(_stem(c["name"])) >= 3:
                return c
        for c in cols:
            if st in _stem(c["name"]).split():
                return c
    return None


def values(s, c):
    """A column's data values."""
    t = s["table"]
    j = c["idx"] - t["c0"]
    return [r[j] if j < len(r) else None for r in t["rows"]]


def records(s):
    t = s["table"]
    return [{c["name"]: (r[c["idx"] - t["c0"]] if c["idx"] - t["c0"] < len(r) else None) for c in t["cols"]} for r in t["rows"]]


def outline(m, rows=3):
    out = []
    for s in m["sheets"]:
        t = s["table"]
        if t:
            out.append(
                f"SHEET '{s['name']}': table at row {t['header']} with {len(t['rows'])} data rows"
                + (f" and a '{t['total_label']}' row" if t["total"] else "")
                + "; columns: "
                + ", ".join(f"{c['letter']} {c['name']} ({c['kind']}{', formula ' + c['formula'] if c['formula'] else ''})" for c in t["cols"])
            )
            for r in t["rows"][:rows]:
                out.append("   " + " | ".join(str(x) for x in r))
        if s["pairs"]:
            out.append(
                f"  {'above the table' if t else 'SHEET ' + repr(s['name']) + ':'} "
                + "; ".join(f"{p['label']} = {p['value']}" for p in s["pairs"][:12])
            )
        if not t and not s["pairs"]:
            out.append(f"SHEET '{s['name']}': empty")
        if s["charts"]:
            out.append(f"  ({s['charts']} chart(s) on this sheet)")
    return "\n".join(out)


# ------------------------------------------------------------------------------------------------ formulas in Python
CALC_FUNCS = {
    "SUM",
    "AVERAGE",
    "AVG",
    "MIN",
    "MAX",
    "COUNT",
    "ROUND",
    "ROUNDUP",
    "ROUNDDOWN",
    "ABS",
    "INT",
    "MOD",
    "SQRT",
    "POWER",
    "IF",
    "IFERROR",
    "AND",
    "OR",
    "NOT",
    "UPPER",
    "LOWER",
    "PROPER",
    "LEN",
    "LEFT",
    "RIGHT",
    "MID",
    "TRIM",
    "CONCAT",
    "CONCATENATE",
    "TODAY",
    "NOW",
    "YEAR",
    "MONTH",
    "DAY",
    "DATE",
    "DAYS",
}
EPOCH = _dt.datetime(1899, 12, 30)


def serial(d):
    """Excel's number for a date (days since 1899-12-30, the time as a fraction)."""
    if isinstance(d, _dt.datetime):
        return (d - EPOCH).total_seconds() / 86400.0
    return float((d - EPOCH.date()).days)


def from_serial(x):
    return EPOCH + _dt.timedelta(days=float(x))


def _xl_str(v):
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def _xl_round(x, d=0, mode="half"):
    import math

    f = 10 ** int(d)
    a = abs(float(x)) * f
    r = math.floor(a + 0.5 + 1e-9) if mode == "half" else math.ceil(a - 1e-9) if mode == "up" else math.floor(a + 1e-9)
    return math.copysign(r / f, float(x))


class _CalcError(Exception):
    pass


def calc(expr, rec):
    """Python's value of an Excel formula written with [Column] names, for one row ({name: value}), to check what Excel
    computes. "#DIV/0!" for a division by zero; None when the formula uses something this does not know."""
    import ast
    import math

    s = str(expr).strip()
    s = s[1:] if s.startswith("=") else s
    names = {}
    lower = {str(k).lower(): v for k, v in rec.items()}
    out = []
    for i, p in enumerate(re.split(r'("(?:[^"]|"")*")', s)):
        if i % 2:
            out.append(repr(p[1:-1].replace('""', '"')))
            continue

        def nm(m):
            k = f"c{len(names)}"
            names[k] = m.group(1).strip().lower()
            return k

        p = re.sub(r"\[([^\]]+)\]", nm, p)
        p = re.sub(r"(\d+(?:\.\d+)?)\s*%", r"(\1/100)", p)
        p = p.replace("<>", "!=")
        p = re.sub(r"(?<![<>!=])=(?!=)", "==", p)
        p = p.replace("^", "**").replace("&", "|")
        p = re.sub(r"\bTRUE\b", "True", p, flags=re.I)
        p = re.sub(r"\bFALSE\b", "False", p, flags=re.I)
        p = re.sub(r"\b([A-Za-z]+)\s*\(", lambda m: m.group(1).upper() + "(", p)
        out.append(p)
    try:
        tree = ast.parse("".join(out), mode="eval")
    except SyntaxError:
        return None

    def num(v):
        if v is None or v == "":
            return 0.0
        if isinstance(v, bool):
            return 1.0 if v else 0.0
        if isinstance(v, (_dt.datetime, _dt.date)):
            return serial(v)
        if isinstance(v, (int, float)):
            return float(v)
        try:
            return float(str(v).replace(",", ""))
        except ValueError:
            raise _CalcError("#VALUE!") from None

    def cmp(a, b, op):
        if isinstance(a, str) or isinstance(b, str):
            a, b = _xl_str(a).lower(), _xl_str(b).lower()
        else:
            a, b = num(a), num(b)
        return {ast.Eq: a == b, ast.NotEq: a != b, ast.Lt: a < b, ast.LtE: a <= b, ast.Gt: a > b, ast.GtE: a >= b}[type(op)]

    def ev(n):
        if isinstance(n, ast.Expression):
            return ev(n.body)
        if isinstance(n, ast.Constant):
            return float(n.value) if isinstance(n.value, (int, float)) and not isinstance(n.value, bool) else n.value
        if isinstance(n, ast.Name):
            if n.id in names:
                if names[n.id] not in lower:
                    raise KeyError(names[n.id])
                v = lower.get(names[n.id])
                return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else v
            raise KeyError(n.id)
        if isinstance(n, ast.UnaryOp):
            if isinstance(n.op, ast.USub):
                return -num(ev(n.operand))
            if isinstance(n.op, ast.UAdd):
                return num(ev(n.operand))
            raise KeyError("unary")
        if isinstance(n, ast.BinOp):
            a, b = ev(n.left), ev(n.right)
            if isinstance(n.op, ast.BitOr):
                return _xl_str(a) + _xl_str(b)
            a, b = num(a), num(b)
            if isinstance(n.op, ast.Add):
                return a + b
            if isinstance(n.op, ast.Sub):
                return a - b
            if isinstance(n.op, ast.Mult):
                return a * b
            if isinstance(n.op, ast.Div):
                if b == 0:
                    raise _CalcError("#DIV/0!")
                return a / b
            if isinstance(n.op, ast.Pow):
                return a**b
            raise KeyError("op")
        if isinstance(n, ast.Compare) and len(n.ops) == 1:
            return cmp(ev(n.left), ev(n.comparators[0]), n.ops[0])
        if isinstance(n, ast.Call) and isinstance(n.func, ast.Name) and n.func.id in CALC_FUNCS:
            f, args = n.func.id, n.args
            if f in ("TODAY", "NOW"):
                return float(serial(_dt.date.today())) if f == "TODAY" else serial(_dt.datetime.now())
            if f in ("YEAR", "MONTH", "DAY"):
                d = from_serial(num(ev(args[0])))
                return float(d.year if f == "YEAR" else d.month if f == "MONTH" else d.day)
            if f == "DATE":
                y, mo, d = (int(num(ev(a))) for a in args[:3])
                return float(serial(_dt.date(y, 1, 1)) + (_dt.date(y + (mo - 1) // 12, (mo - 1) % 12 + 1, 1) - _dt.date(y, 1, 1)).days + d - 1)
            if f == "DAYS":
                return float(int(num(ev(args[0]))) - int(num(ev(args[1]))))
            if f == "IF":
                c = ev(args[0])
                truth = bool(c) if isinstance(c, str) else num(c) != 0
                return ev(args[1]) if truth else (ev(args[2]) if len(args) > 2 else False)
            if f == "IFERROR":
                try:
                    return ev(args[0])
                except _CalcError:
                    return ev(args[1])
            vals = [ev(a) for a in args]
            if f in ("SUM", "AVERAGE", "AVG", "MIN", "MAX", "COUNT"):
                xs = [num(v) for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
                if f == "COUNT":
                    return float(len(xs))
                if not xs:
                    if f in ("AVERAGE", "AVG"):
                        raise _CalcError("#DIV/0!")
                    return 0.0
                return sum(xs) if f == "SUM" else sum(xs) / len(xs) if f in ("AVERAGE", "AVG") else min(xs) if f == "MIN" else max(xs)
            if f in ("ROUND", "ROUNDUP", "ROUNDDOWN"):
                return _xl_round(num(vals[0]), num(vals[1]) if len(vals) > 1 else 0, {"ROUND": "half", "ROUNDUP": "up", "ROUNDDOWN": "down"}[f])
            if f == "ABS":
                return abs(num(vals[0]))
            if f == "INT":
                return float(math.floor(num(vals[0])))
            if f == "MOD":
                return num(vals[0]) % num(vals[1])
            if f == "SQRT":
                return math.sqrt(num(vals[0]))
            if f == "POWER":
                return num(vals[0]) ** num(vals[1])
            if f == "AND":
                return all(num(v) != 0 for v in vals)
            if f == "OR":
                return any(num(v) != 0 for v in vals)
            if f == "NOT":
                return not (num(vals[0]) != 0)
            t = _xl_str(vals[0]) if vals else ""
            if f == "UPPER":
                return t.upper()
            if f == "LOWER":
                return t.lower()
            if f == "PROPER":
                return re.sub(r"[A-Za-z]+", lambda m: m.group(0)[:1].upper() + m.group(0)[1:].lower(), t)
            if f == "LEN":
                return float(len(t))
            if f == "LEFT":
                return t[: int(num(vals[1])) if len(vals) > 1 else 1]
            if f == "RIGHT":
                k = int(num(vals[1])) if len(vals) > 1 else 1
                return t[-k:] if k else ""
            if f == "MID":
                a = int(num(vals[1])) - 1
                return t[a : a + int(num(vals[2]))]
            if f == "TRIM":
                return " ".join(t.split())
            if f in ("CONCAT", "CONCATENATE"):
                return "".join(_xl_str(v) for v in vals)
        raise KeyError("unknown")

    try:
        return ev(tree)
    except _CalcError as e:
        return str(e)
    except (KeyError, IndexError, TypeError, ValueError, OverflowError, ZeroDivisionError):
        return None


def same(a, b, tol=1e-6):
    """Excel's value and Python's agree (numbers within a tolerance, text ignoring case, an error by its code)."""
    if isinstance(a, bool) or isinstance(b, bool):
        return isinstance(a, bool) and isinstance(b, bool) and a == b
    a = serial(a) if isinstance(a, (_dt.datetime, _dt.date)) else a
    b = serial(b) if isinstance(b, (_dt.datetime, _dt.date)) else b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) <= tol * max(1.0, abs(float(a)), abs(float(b)))
    return _xl_str(a).strip().lower() == _xl_str(b).strip().lower()
