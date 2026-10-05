"""Native charts for Word documents: real, editable Office charts (right-click > Edit Data works), made from
python-pptx's chart XML and its embedded workbook, attached to the .docx as a chart part. Word draws them itself when
it renders, so a chart looks exactly like one made in Word.

  add_chart(doc, paragraph, spec, th, width_cm=15, height_cm=8.5)
  spec = {"chart": "column|bar|line|pie|doughnut|area|stacked|stacked_bar", "title": ..., "categories": [...],
          "series": [{"name": ..., "values": [...]}], "number_format": "#,##0"}
"""
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml as docx_parse
from lxml import etree
from pptx.chart.chart import Chart
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.oxml import parse_xml as pptx_parse
from pptx.util import Pt

TYPES = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
         "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA,
         "stacked": XL_CHART_TYPE.COLUMN_STACKED, "stacked_bar": XL_CHART_TYPE.BAR_STACKED}
EMU_PER_CM = 360000


def number_format(spec):
    vals = [v for s in spec["series"] for v in s["values"] if v is not None]
    dec = 0 if all(float(v).is_integer() for v in vals) else (1 if all(float(v * 10).is_integer() for v in vals) else 2)
    return spec.get("number_format") or ("#,##0" if dec == 0 else "#,##0." + "0" * dec)  # "#,##0.##" leaves "22,000."


def chart_data(spec):
    cd = CategoryChartData(number_format=number_format(spec))
    cd.categories = [str(c) for c in spec["categories"]]
    for s in spec["series"]:
        cd.add_series(str(s["name"]), [None if v is None else float(v) for v in s["values"]])
    return cd


def chart_space(spec, th, font=None, size=9):
    """(chartSpace element styled with the theme, xlsx bytes of its data), for a Word document."""
    cd = chart_data(spec)
    cs = pptx_parse(cd.xml_bytes(TYPES[spec.get("chart", "column")]))
    style(Chart(cs, None), spec, th, font, size)
    return cs, cd.xlsx_blob


def style(ch, spec, th, font=None, size=9, text_colour="404040", grid_colour="E3E3E3"):
    """One look for every chart (Word, PowerPoint): the theme's colours and font, a title, a legend only when there
    are several series, values on the bars when there are few, percentages outside the slices of a pie."""
    kind = spec.get("chart", "column")
    num_fmt = number_format(spec)
    colours = [RGBColor.from_string(c) for c in th["chart"]]
    ch.font.size = Pt(size)
    ch.font.name = font or th["body"]
    ch.font.color.rgb = RGBColor.from_string(text_colour)
    if spec.get("title"):
        ch.has_title = True
        tf = ch.chart_title.text_frame
        tf.text = str(spec["title"])
        r = tf.paragraphs[0].runs[0]
        r.font.size, r.font.bold = Pt(size + 2), True
        r.font.name = font or th["body"]  # the title has its own run font (Word draws it in Arial otherwise)
        r.font.color.rgb = RGBColor.from_string(th.get("chart_title") or th["head_color"])
    else:
        ch.has_title = False
    plot = ch.plots[0]
    multi = len(spec["series"]) > 1
    if kind in ("pie", "doughnut"):
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.RIGHT
        ch.legend.include_in_layout = False
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_percentage, dl.show_value, dl.show_category_name = True, False, False
        dl.number_format, dl.number_format_is_linked = "0%", False
        dl.font.size = Pt(size)
        dl.font.color.rgb = RGBColor.from_string(text_colour)
        dl.position = XL_LABEL_POSITION.OUTSIDE_END  # readable on dark and light slices alike
        for i, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = colours[i % len(colours)]
    else:
        ch.has_legend = multi
        if multi:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        for i, s in enumerate(plot.series):
            c = colours[i % len(colours)]
            if kind == "line":
                s.format.line.color.rgb = c
                s.format.line.width = Pt(2.25)
                s.smooth = False
                try:
                    s.marker.format.fill.solid()
                    s.marker.format.fill.fore_color.rgb = c
                except Exception:  # noqa: BLE001
                    pass
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = c
        if kind in ("column", "bar", "stacked", "stacked_bar"):
            try:
                plot.gap_width = 80
                if kind in ("stacked", "stacked_bar"):
                    plot.overlap = 100
            except Exception:  # noqa: BLE001
                pass
        n_points = len(spec["categories"]) * len(spec["series"])
        if kind in ("column", "bar") and n_points <= 12:  # few bars: the values on them read better than a grid
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(size - 0.5)
            plot.data_labels.number_format = num_fmt
            plot.data_labels.number_format_is_linked = False
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        try:
            va = ch.value_axis
            vals = [v for s_ in spec["series"] for v in s_["values"] if v is not None]
            if vals and min(vals) >= 0:  # a line from 1,690,000 makes a 1% change look like a cliff
                va.minimum_scale = 0
            va.tick_labels.number_format = num_fmt
            va.tick_labels.number_format_is_linked = False
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = RGBColor.from_string("E3E3E3")
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(size - 0.5)
            ca = ch.category_axis
            ca.tick_labels.font.size = Pt(size - 0.5)
            ca.format.line.color.rgb = RGBColor.from_string("BFBFBF")
            if grid_colour != "E3E3E3":
                va.major_gridlines.format.line.color.rgb = RGBColor.from_string(grid_colour)
        except Exception:  # noqa: BLE001
            pass
    return ch


def add_chart(doc, paragraph, spec, th, width_cm=15.0, height_cm=8.5):
    """A native chart in `paragraph` (a run with an inline chart). Returns the chart part's name."""
    package = doc.part.package
    n = 1 + sum(1 for pt in package.iter_parts() if str(pt.partname).startswith("/word/charts/chart"))
    cs, xlsx = chart_space(spec, th)
    xlsx_part = Part(PackURI(f"/word/embeddings/Microsoft_Excel_Worksheet{n}.xlsx"), CT.SML_SHEET, xlsx, package)
    chart_part = Part(PackURI(f"/word/charts/chart{n}.xml"), CT.DML_CHART, b"", package)
    x_rid = chart_part.relate_to(xlsx_part, RT.PACKAGE)
    cs.get_or_add_externalData().rId = x_rid
    chart_part._blob = etree.tostring(cs, xml_declaration=True, encoding="UTF-8", standalone=True)
    rid = doc.part.relate_to(chart_part, RT.CHART)
    cx, cy = int(width_cm * EMU_PER_CM), int(height_cm * EMU_PER_CM)
    inline = docx_parse(
        '<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{cx}" cy="{cy}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:docPr id="{4000 + n}" name="Chart {n}"/><wp:cNvGraphicFramePr/>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        f'<c:chart r:id="{rid}"/></a:graphicData></a:graphic></wp:inline></w:drawing>')
    paragraph.add_run()._r.append(inline)
    return str(chart_part.partname)
