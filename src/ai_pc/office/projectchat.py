"""A conversation about a project: several files edited together, with data flowing between them.

  pc = ProjectChat.start(["sales.xlsx"], name="Q1 sales")
  pc.say("clean up the data and add a column Amount = Qty x Unit Price")
  pc.say("write a 2 page report on Q1 sales for the management team from the workbook")
  pc.say("put the pivot of amount by city into the report after the introduction, with a chart of it")
  pc.say("make a 5 slide deck from the report for the board")
  pc.say("in the workbook, delete the cancelled orders")     -> "the report's table and the deck's chart are now out of date"
  pc.say("refresh the report and the deck")
  pc.say("export everything as one pdf with a cover called Q1 Sales Pack, page numbers and a DRAFT watermark")

Each file keeps its own chat (Word, PowerPoint or Excel: versions, checks, its own undo). The project adds:
- routing: each part of a message goes to the file it is about (named, or recognised by its sheets, columns, sections,
  slide titles), else to the file last talked about;
- data links: a table or chart that comes from a workbook is grabbed from Excel exactly as Excel shows it, written as a
  native table or chart, and carries an invisible link; after the workbook changes the reply says which copies are out
  of date, and "refresh" writes them again in place (each cell checked against Excel);
- documents made from other files: a report from a workbook's figures (computed, quoted exactly), a deck from a report
  (its charts from the workbook, not retyped), a handout from a deck (slide pictures and speaker notes);
- PDFs: a pack of several files with a cover, bookmarks, page numbers and a watermark; compress, split, pages;
- one history for the whole project: undo / redo / go to vN move every file back together.
"""

import hashlib
import json
import re
import shutil
import time
from pathlib import Path

from ai_pc.core.config import ROOT
from ai_pc.core.util import parse_json, slug
from ai_pc.office import book_parse as BP
from ai_pc.office import doc_parse as P
from ai_pc.office import pdf_ops as PF
from ai_pc.office import render as RN
from ai_pc.office import xlsx_map as XM
from ai_pc.office.bookchat import BookChat, _fmt
from ai_pc.office.deckchat import DeckChat
from ai_pc.office.docchat import FILLER, GOTO, HISTORY, POLITE, QUESTION, REDO, UNDO, VERSIONS, DocChat

PROJECTS = ROOT / "out" / "docs" / "projects"
KINDS = {".xlsx": "xlsx", ".xlsm": "xlsx", ".docx": "docx", ".pptx": "pptx"}
CHAT = {"xlsx": BookChat, "docx": DocChat, "pptx": DeckChat}
ROLE = {
    "xlsx": r"workbook|spreadsheet|excel|xlsx|sheet",
    "docx": r"report|document|doc|docx|word file|memo|essay|proposal|handout",
    "pptx": r"deck|slides?|presentation|powerpoint|pptx|ppt",
}
NAMES = {"xlsx": "workbook", "docx": "report", "pptx": "deck"}

INSERT = re.compile(r"\b(?:put|insert|add|copy|paste|bring|include|place|drop|show)\b")
DATA = re.compile(r"\b(?:pivot(?: table)?|tables?|summary|summaries|charts?|graphs?|figures|numbers|data|totals?|breakdown)\b")
INTO = re.compile(
    r"\b(?:into|in|to|on|onto|inside)\s+(?:the\s+|a\s+new\s+|a\s+|our\s+|my\s+)?(report|document|doc|word file|memo|handout|deck|slides?|presentation|slide\s+\d+)\b"
)
REFRESH = re.compile(r"\b(?:refresh|update|sync|re-?pull|bring\b.*\bup to date)\b")
DECK_FROM = re.compile(
    r"\b(?:make|create|build|turn|generate|prepare|do)\b.*\b(?:deck|presentation|slides)\b.*\b(?:from|out of|based on|of|summari[sz]ing)\s+(?:the\s+|this\s+|my\s+)?(report|document|doc)\b|"
    r"\bturn\s+(?:the\s+)?(report|document)\s+into\s+(?:a\s+)?(?:\d+[- ]slide\s+)?(?:deck|presentation|slides)\b"
)
REPORT_FROM = re.compile(
    r"\b(?:write|make|create|draft|prepare|produce)\b.*\b(?:report|memo|brief|summary document|document)\b.*\b(?:from|using|based on|out of)\s+(?:the\s+|this\s+|my\s+)?"
    r"(?:workbook|data|spreadsheet|figures|numbers|sheet|sales|excel)\b|\b(?:write|make|create|draft|prepare)\b\s+(?:a\s+)?(?:\d+[- ]page\s+)?report\s+on\b"
)
HANDOUT = re.compile(r"\bhandouts?\b")
PACK = re.compile(
    r"\b(?:export|save|combine|merge|bundle|make|put|turn|print|join)\b.*\b(?:everything|all|both|whole|project|pack|the (?:report|deck|files|documents)(?: and the (?:report|deck|handout))?)\b.*\bpdf\b|"
    r"\bpdf pack\b|\b(?:one|single|a combined)\s+pdf\b"
)
COMPRESS = re.compile(r"\b(?:compress|shrink)\b|\bmake (?:it|the pdf|the pack) smaller\b|\breduce (?:the )?(?:pdf |file )?size\b")
SPLIT = re.compile(r"\bsplit\b.*\bpdf\b|\bsplit (?:it|the pack)\b")
PAGES = re.compile(r"\b(extract|keep|delete|remove|drop|rotate)\s+(?:only\s+)?pages?\s+((?:\d+\s*(?:-|to)?\s*\d*\s*(?:,|and)?\s*)+|the last)\b")
WATERMARK = re.compile(
    r"\b(?:a|an)\s+['\"]?([A-Za-z][\w ]{1,20}?)['\"]?\s+watermark\b|\bwatermark(?:ed)?(?:\s+(?:it|them|the pdf|the pack))?\s+(?:with|saying|reading|that says)\s+['\"]?([\w ]{2,20}?)['\"]?(?=$|,|\s+and\b)|"
    r"\bstamp(?:ed)?\s+(?:it\s+|them\s+)?(?:as\s+|with\s+)?['\"]?(draft|confidential|final|internal|copy)\b",
    re.I,
)
PROJECT_Q = re.compile(
    r"\bwhat(?:'s| is)? in (?:the |this |my )?project\b|\bwhat files\b|\blist (?:the |all )?files\b|\bshow (?:me )?(?:the )?files\b|\bwhich files\b",
    re.I,
)
STALE_Q = re.compile(r"\b(?:up to date|out of date|stale|in sync|current)\b", re.I)
ROUTE_PREFIX = re.compile(r"^\s*(?:in|on|for|to|with)\s+(?:the\s+|my\s+|our\s+)?(?:" + "|".join(ROLE.values()) + r")\s*(?:file)?\s*[,:]?\s*", re.I)
ROUTE_ONLY = re.compile(
    r"^\s*(?:(?:and|then|also)\s+)?(?:in|on|for|to|with|inside)\s+(?:the\s+|my\s+|our\s+)?("
    + "|".join(ROLE.values())
    + r")\s*(?:file)?\s*[,:.]?\s*$",
    re.I,
)
ROUTE_SUFFIX = re.compile(r"\s+(?:in|on|of|for|to)\s+(?:the\s+|my\s+)?(?:" + "|".join(ROLE.values()) + r")\s*\.?\s*$", re.I)

PROJECT_SYSTEM = """You route a client's request in a project of several office files to the one file it is about, rewritten as a request
for that file alone. Reply with ONE JSON object: {"file": "<file key>", "request": "<the request for that file>"} or {"ask": "<short question>"}.
Use the file keys exactly as listed."""


def _core(v):
    """The number in a formatted figure ('Rs 5,634,700' -> '5,634,700'), when it is distinctive enough to find in text."""
    m = re.search(r"\d[\d,]*(?:\.\d+)?", str(v or ""))
    if not m:
        return None
    digits = re.sub(r"\D", "", m.group(0).split(".")[0])
    return m.group(0) if len(digits) >= 4 else None


def prettify(h):
    h = re.sub(r"\s+", " ", str(h or "")).strip()
    m = re.match(r"^(?:Sum|Total) of (.+)$", h)
    if m:
        return m.group(1)
    m = re.match(r"^(Average|Count|Max|Min) of (.+)$", h)
    if m:
        return f"{m.group(1)} {m.group(2)}" if m.group(1) != "Count" else "Count"
    return "" if h.lower() == "row labels" else h


def _hash(obj):
    return hashlib.sha1(json.dumps(obj, sort_keys=True, default=str).encode("utf-8")).hexdigest()[:16]


class ProjectChat:
    def __init__(self, state, planner=None, log=print, server=None):
        self.state, self.planner, self.log = state, planner, log
        self.server = server if server is not None else RN.Server()
        self.folder = Path(state["folder"])
        self.chats = {}
        self._turn = None
        self._pdf_added = False

    # ------------------------------------------------------------------ files
    @classmethod
    def start(cls, files, name=None, projects_dir=None, **kw):
        nm = slug(name or Path(files[0]).stem, 24) or "project"
        folder = Path(projects_dir or PROJECTS) / f"proj_{nm}_{time.strftime('%H%M%S')}"
        folder.mkdir(parents=True, exist_ok=True)
        state = {
            "id": folder.name,
            "name": name or nm,
            "folder": str(folder),
            "files": {},
            "links": {},
            "pdfs": [],
            "versions": [],
            "cur": 0,
            "focus": None,
            "turns": [],
        }
        pc = cls(state, **kw)
        for f in files:
            pc._add_file(f)
        state["versions"].append(
            {
                "v": 0,
                "parent": None,
                "files": {k: pc.chat(k).state["cur"] for k in state["files"]},
                "links": {},
                "pdfs": 0,
                "said": None,
                "done": ["the files as given"],
            }
        )
        pc.save()
        return pc

    @classmethod
    def load(cls, project_id, projects_dir=None, **kw):
        p = Path(projects_dir or PROJECTS) / project_id / "project.json"
        return cls(json.loads(p.read_text(encoding="utf-8")), **kw)

    def save(self):
        (self.folder / "project.json").write_text(json.dumps(self.state, ensure_ascii=False, indent=1), encoding="utf-8")

    def close(self):
        if self.server is not None:
            self.server.close()

    def _add_file(self, path, role=None, label=None):
        src = Path(path)
        kind = KINDS.get(src.suffix.lower())
        if kind is None:
            raise ValueError(f"a project holds .xlsx, .docx and .pptx files, not {src.suffix}")
        base = re.sub(r"[^a-z0-9]+", "_", (role or src.stem).lower()).strip("_")[:20] or kind
        key, k = base, 2
        while key in self.state["files"]:
            key, k = f"{base}{k}", k + 1
        c = CHAT[kind].start(src, chats_dir=self.folder / "chats", planner=self.planner, server=self.server)
        self.state["files"][key] = {"kind": kind, "chat": c.state["id"], "src": str(src), "label": label or src.name, "role": role or NAMES[kind]}
        self.chats[key] = c
        return key

    def chat(self, key):
        if key not in self.chats:
            f = self.state["files"][key]
            self.chats[key] = CHAT[f["kind"]].load(f["chat"], chats_dir=self.folder / "chats", planner=self.planner, server=self.server)
        return self.chats[key]

    @property
    def version(self):
        return self.state["versions"][self.state["cur"]]

    def active(self, kind=None):
        return [k for k in self.version["files"] if kind is None or self.state["files"][k]["kind"] == kind]

    def label(self, key):
        f = self.state["files"][key]
        return f"{f['role']} ({f['label']})" if f["role"] not in f["label"].lower() else f["label"]

    # ------------------------------------------------------------------ which file a request is about
    def route(self, c):
        low = c.lower()
        files = self.active()
        for k in files:  # a file named by its name or its role ("the report", "sales.xlsx")
            f = self.state["files"][k]
            stem = re.sub(r"[_\-]+", " ", Path(f["label"]).stem.lower())
            if re.search(r"(?<!\w)" + re.escape(f["label"].lower()) + r"(?!\w)", low) or (
                len(stem) > 3 and re.search(r"(?<!\w)" + re.escape(stem) + r"(?!\w)", low)
            ):
                return k
        hits = []
        for kind, rx in ROLE.items():
            mm = re.search(r"\b(?:" + rx + r")\b", low)
            if mm:
                hits.append((mm.start(), kind))
        if hits:
            kind = min(hits)[1]
            ks = self.active(kind)
            if len(ks) == 1:
                return ks[0]
            role = next((k for k in ks if self.state["files"][k]["role"] in low), None)
            if role:
                return role
            if self.state.get("focus") in ks:
                return self.state["focus"]
            if ks:
                return ks[-1]
        scores = {}
        for k in files:  # its own words: sheets, columns and values; section titles; slide titles
            kind = self.state["files"][k]["kind"]
            try:
                m = self.chat(k).map()
            except Exception:  # noqa: BLE001
                continue
            words = set()
            if kind == "xlsx":
                s = XM.sheet(m)
                words |= {n.lower() for n in m["names"]}
                if s and s.get("table"):
                    words |= {c_["name"].lower() for c_ in s["table"]["cols"]}
                    for c_ in s["table"]["cols"]:
                        if c_["kind"] == "text":
                            vals = {str(v).lower() for v in XM.values(s, c_) if v}
                            if len(vals) <= 30:
                                words |= vals
            elif kind == "docx":
                words |= {re.sub(r"^\d+(\.\d+)*\.?\s+", "", x["title"]).lower() for x in m["sections"]}
            else:
                words |= {s_["title"].lower() for s_ in m["slides"] if s_["title"]}
            scores[k] = sum(1 for w in words if len(w) > 3 and re.search(r"(?<!\w)" + re.escape(w) + r"(?!\w)", low))
        best = max(scores.items(), key=lambda kv: kv[1]) if scores else (None, 0)
        if best[1] > 0 and list(scores.values()).count(best[1]) == 1:
            return best[0]
        if re.search(r"\bslides?\s+\d+\b", low) and self.active("pptx"):
            return self.active("pptx")[-1]
        f = self.state.get("focus")
        return f if f in files else (files[0] if len(files) == 1 else None)

    @staticmethod
    def strip_route(c):
        c2 = ROUTE_SUFFIX.sub("", ROUTE_PREFIX.sub("", c)).strip()
        return c2 or c

    # ------------------------------------------------------------------ a message
    def say(self, message):
        t0 = time.perf_counter()
        turn = {"user": message, "intents": [], "llm": False, "v_before": self.state["cur"]}
        self._turn = turn
        out, plan = [], []
        self._pdf_added = False
        hint = None
        for cl in P.clauses(message) or [message]:
            c = POLITE.sub("", FILLER.sub("", cl)).strip()
            if not c:
                continue
            low = c.lower()
            ro = ROUTE_ONLY.match(low)
            if ro:  # "in the workbook," on its own: the requests after it are about that file
                kind = next(k for k, rx in ROLE.items() if re.fullmatch(rx, ro.group(1)) or re.search(r"\b(?:" + rx + r")\b", ro.group(1)))
                hint = self.route_kind(low, kind)
                continue
            named = bool(re.search(r"\b(?:" + "|".join(ROLE.values()) + r")\b", low))
            if not named:
                g = GOTO.search(low)
                if g:
                    plan.append(("meta", ("goto", int(next(x for x in g.groups() if x)))))
                    continue
                if REDO.search(low):
                    plan.append(("meta", ("redo", None)))
                    continue
                if UNDO.search(low):
                    plan.append(("meta", ("undo", None)))
                    continue
                if VERSIONS.search(low) or HISTORY.search(low):
                    plan.append(("meta", ("history", None)))
                    continue
            if PROJECT_Q.search(low):
                plan.append(("say", self.files_text()))
                continue
            if STALE_Q.search(low) and (QUESTION.search(low) or re.match(r"\s*(?:is|are|do|does|has|have)\b", low) or cl.strip().endswith("?")):
                plan.append(("say", self.stale_text() or "Everything that came from the workbook is up to date."))
                continue
            x = self.cross(c, cl)
            if x:
                plan.append(("cross", x))
                continue
            key = self.route(c)
            if hint and not re.search(r"\b(?:" + "|".join(ROLE.values()) + r")\b", low):
                key = hint
            if key is None:
                plan.append(("unknown", c))
            else:
                plan.append((key, self.strip_route(c)))
        merged = []
        for item in plan:  # requests in a row for one file are one message to it: one version
            if merged and item[0] in self.state["files"] and merged[-1][0] == item[0]:
                merged[-1] = (item[0], merged[-1][1].rstrip(". ") + ". " + item[1])
            else:
                merged.append(item)
        for kind, payload in merged:
            if kind == "meta":
                what, n = payload
                out.append(
                    self.undo() if what == "undo" else self.redo() if what == "redo" else self.goto(n) if what == "goto" else self.history_text()
                )
                turn["intents"].append(what)
            elif kind == "say":
                out.append(payload)
                turn["intents"].append("question")
            elif kind == "cross":
                turn["intents"].append(payload["op"])
                try:
                    out.append(self.run_cross(payload, message))
                except Exception as e:  # noqa: BLE001
                    import traceback

                    traceback.print_exc()
                    out.append(f"Couldn't {payload['op'].replace('_', ' ')}: {type(e).__name__}: {str(e)[:200]}")
            elif kind == "unknown":
                r = self._llm_route(payload, message)
                if r.get("file") in self.state["files"]:
                    turn["intents"].append("routed")
                    out.append(self._to_file(r["file"], r.get("request") or payload))
                else:
                    turn["intents"].append("ask")
                    out.append(r.get("ask") or f"Which file is that about? ({', '.join(self.label(k) for k in self.active())})")
            else:
                turn["intents"].append("file")
                out.append(self._to_file(kind, payload))
        changed = self._snapshot(message, out)
        stale = self.stale_text()
        if changed and stale:
            out.append(stale)
        reply = "\n".join(x for x in out if x).strip() or "Tell me what to do with " + ", ".join(self.label(k) for k in self.active()) + "."
        turn.update(reply=reply, v_after=self.state["cur"], seconds=round(time.perf_counter() - t0, 2))
        self.state["turns"].append(turn)
        self.last_turn = turn
        self.save()
        return reply

    def _to_file(self, key, text):
        c = self.chat(key)
        r = c.say(text)
        self.state["focus"] = key
        if getattr(c, "last_turn", {}).get("llm"):
            self._turn["llm"] = True
        return f"[{self.label(key)}] {r}"

    def _snapshot(self, said, out):
        files = {k: self.chat(k).state["cur"] for k in self.active()}
        made = [k for k in self.state["files"] if k not in files and k in self.chats and self.state["files"][k].get("made_in") == "pending"]
        for k in made:  # files made during this message
            files[k] = self.chat(k).state["cur"]
            self.state["files"][k]["made_in"] = len(self.state["versions"])
        cur = self.version
        n_pdfs = len(self.state["pdfs"]) if self._pdf_added else cur["pdfs"]
        if (
            files == cur["files"]
            and n_pdfs == cur["pdfs"]
            and json.dumps(self.state["links"], sort_keys=True) == json.dumps(cur.get("links") or {}, sort_keys=True)
        ):
            return False
        n = len(self.state["versions"])
        done = [re.sub(r"\s+Checked:.*$", "", x.split("\n")[0])[:160] for x in out if x and not x.startswith(("Undone", "Redone", "Now at"))]
        self.state["versions"].append(
            {
                "v": n,
                "parent": self.state["cur"],
                "files": files,
                "links": json.loads(json.dumps(self.state["links"])),
                "pdfs": n_pdfs,
                "said": said,
                "done": done,
            }
        )
        self.state["cur"] = n
        return True

    # ------------------------------------------------------------------ history for the whole project
    def _restore(self, v):
        ver = self.state["versions"][v]
        for k, fv in ver["files"].items():
            c = self.chat(k)
            if c.state["cur"] != fv:
                c.state["cur"] = fv
                c.save()
        self.state["links"] = json.loads(json.dumps(ver.get("links") or {}))
        self.state["cur"] = v

    def undo(self):
        parent = self.version.get("parent")
        if parent is None:
            return "Nothing to undo: these are the files as given."
        was = self.state["cur"]
        self._restore(parent)
        gone = [k for k in self.state["versions"][was]["files"] if k not in self.version["files"]]
        return (
            f"Undone: the project is back at v{parent} (v{was} was: {'; '.join(self.state['versions'][was]['done'])[:200]})"
            + (f"; {', '.join(self.label(k) for k in gone)} set aside" if gone else "")
            + "."
        )

    def redo(self):
        kids = [v for v in self.state["versions"] if v.get("parent") == self.state["cur"]]
        if not kids:
            return "Nothing to redo."
        self._restore(kids[-1]["v"])
        return f"Redone: v{kids[-1]['v']} ({'; '.join(kids[-1]['done'])[:200]})."

    def goto(self, n):
        if not 0 <= n < len(self.state["versions"]):
            return f"There is no v{n} (v0 to v{len(self.state['versions']) - 1})."
        self._restore(n)
        return f"Now at project v{n}: {'; '.join(self.version['done'])[:200]}."

    def history_text(self):
        lines = [f"{'*' if v['v'] == self.state['cur'] else ' '} v{v['v']}: {'; '.join(v['done'])[:150]}" for v in self.state["versions"]]
        return f"{len(self.state['versions'])} project version(s):\n" + "\n".join(lines)

    def files_text(self):
        out = []
        for k in self.active():
            f = self.state["files"][k]
            c = self.chat(k)
            n_links = sum(1 for L in self.state["links"].values() if L["dst"] == k)
            out.append(
                f"{self.label(k)}: {f['kind']}, at v{c.state['cur']}" + (f", {n_links} table(s)/chart(s) linked to a workbook" if n_links else "")
            )
        if self.state["pdfs"][: self.version["pdfs"]]:
            out.append("PDFs: " + ", ".join(Path(p["path"]).name for p in self.state["pdfs"][: self.version["pdfs"]]))
        return "Project files: " + "; ".join(out) + "."

    # ------------------------------------------------------------------ data links
    def main_hash(self, wb_key, sheet=None):
        m = self.chat(wb_key).map()
        s = XM.sheet(m)
        main = [r for r in XM.records(s)] if s and s.get("table") else []
        extra = []
        if sheet and sheet != (s or {}).get("name"):
            o = XM.sheet(m, sheet)
            extra = o["table"]["rows"] if o and o.get("table") and o["kind"] != "pivot" else []
        return _hash([main, extra])

    def stale(self):
        out = []
        for lid, L in self.state["links"].items():
            if L["dst"] not in self.active() or L["src"] not in self.active():
                continue
            if self.chat(L["src"]).state["cur"] == L["src_v"]:
                continue
            if self.main_hash(L["src"], L["spec"].get("sheet")) != L["main_hash"]:
                out.append((lid, L))
        return out

    def stale_text(self):
        st = self.stale()
        figs = self.stale_figures()
        if not st and not figs:
            return ""
        by = {}
        for lid, L in st:
            by.setdefault(L["dst"], []).append(f"{L['kind']} '{L['title']}'")
        for dst, what, old, new in figs:
            by.setdefault(dst, []).append(f"the text says {old} for {what} (now {new})")
        return (
            "Out of date since the workbook changed: "
            + "; ".join(f"{self.label(k)}: {', '.join(v)}" for k, v in by.items())
            + ". Say 'refresh "
            + (" and ".join("the " + self.state["files"][k]["role"] for k in by) if len(by) <= 2 else "everything")
            + "' to bring them up to date."
        )

    def grab(self, wb_key, specs):
        """Tables or pivots from the workbook's current version exactly as Excel shows them (pivots refreshed first)."""
        c = self.chat(wb_key)
        tmp = self.folder / f"_grab_{wb_key}.xlsx"
        shutil.copy(c.path(), tmp)  # Excel reads a copy: the version file itself never changes
        r = c._job({"app": "excel", "src": str(tmp.resolve()), "grab": specs})
        tmp.unlink(missing_ok=True)
        if not r.get("ok"):
            raise RuntimeError(f"Excel could not read the workbook: {r.get('error')}")
        outs = []
        for g, spec in zip(r.get("grab") or [], specs):
            if g.get("error"):
                raise RuntimeError(g["error"])
            g["header"] = [prettify(h) or (spec.get("label") or "") for h in g["header"]]
            g["source"] = f"{self.state['files'][wb_key]['label']}, sheet '{g['sheet']}'"
            g["title"] = spec.get("title") or g["sheet"]
            outs.append(g)
        return outs

    def source(self, wb_key, desc):
        """The sheet a request means in a workbook: a sheet named in it, a pivot by its fields, else a new pivot made
        for it ("amount by city and month" -> rows City, columns months, values Amount). (spec, note)."""
        c = self.chat(wb_key)
        m = c.map()
        d = P.normalize(desc)
        named = BP.sheet_in(d + " sheet" if not re.search(r"\bsheet\b", d) else d, m)
        if named and XM.sheet(m, named) and XM.sheet(m, named).get("table") and named != m["main"]:
            s = XM.sheet(m, named)
            return {"sheet": s["name"], "what": "pivot" if s["kind"] == "pivot" else "table", "title": s["name"]}, None
        main = XM.sheet(m)
        t = main["table"]
        bym = re.search(r"\b(?:by|per|for each|across)\s+(.+)$", d)
        bys = BP.cols_in(bym.group(1), main) if bym else []
        datec = next((x for x in t["cols"] if x["kind"] == "date"), None)
        if bym and re.search(r"\bmonth", bym.group(1)) and datec and datec not in bys:
            bys.append(datec)
        vals = [x for x in BP.cols_in(d[: bym.start()] if bym else d, main) if x["kind"] in ("number", "money", "percent") and x not in bys] or (
            [BP.metric(t, additive=True)] if BP.metric(t, additive=True) else []
        )
        pivots = [s for s in m["sheets"] if s["kind"] == "pivot" and s.get("table")]
        if bys:  # a pivot or summary that is already there: its first column is the 'by', its numbers the value
            for s in pivots + [s for s in m["sheets"] if s["kind"] == "table" and s["name"] != m["main"] and s.get("table")]:
                head = [prettify(c_["name"]).lower() for c_ in s["table"]["cols"]]
                if head and head[0] == bys[0]["name"].lower() and (not vals or any(vals[0]["name"].lower() in h for h in head)) and len(bys) == 1:
                    return {
                        "sheet": s["name"],
                        "what": "pivot" if s["kind"] == "pivot" else "table",
                        "title": f"{vals[0]['name'] if vals else 'Rows'} by {bys[0]['name']}",
                    }, None
            rows = [x for x in bys if x["kind"] != "date"][:1] or bys[:1]
            cols = [x for x in bys if x is not rows[0]][:1]
            rname = "month" if rows[0]["kind"] == "date" else rows[0]["name"]
            op = {
                "op": "pivot",
                "sheet": main["name"],
                "rows": [rows[0]["name"]],
                **({"columns": [cols[0]["name"]]} if cols else {}),
                "values": [vals[0]["name"]] if vals else [],
                "fn": "sum",
                **({"period": "month"} if any(x["kind"] == "date" for x in bys) else {}),
                "name": f"{(vals[0]['name'] + ' ') if vals else ''}by {rname}"
                + (f" and {('month' if cols[0]['kind'] == 'date' else cols[0]['name'])}" if cols else ""),
            }
            before = set(m["names"])
            msg = c._change([op], f"(for the project) {desc}", {"ops": []})
            c.save()
            new = [n for n in c.map()["names"] if n not in before]
            if not new:
                raise RuntimeError(f"the workbook could not make that pivot: {msg}")
            title = f"{vals[0]['name'] if vals else 'Rows'} by {rname}" + (
                f" and {'month' if cols[0]['kind'] == 'date' else cols[0]['name']}" if cols else ""
            )
            return {"sheet": new[0], "what": "pivot", "title": title}, f"[{self.label(wb_key)}] {msg}"
        if pivots and re.search(r"\bpivot\b", d):
            return {"sheet": pivots[-1]["name"], "what": "pivot", "title": pivots[-1]["name"]}, None
        if pivots:
            return {"sheet": pivots[-1]["name"], "what": "pivot", "title": pivots[-1]["name"]}, None
        raise RuntimeError("which figures? e.g. 'the pivot', 'amount by city' or 'the Amount by month sheet'")

    # ------------------------------------------------------------------ requests that span files
    def cross(self, c, raw):
        low = P.normalize(c)
        books, docs, decks = self.active("xlsx"), self.active("docx"), self.active("pptx")
        if PACK.search(low) or (re.search(r"\bpdf\b", low) and re.search(r"\b(?:cover|bookmarks?|page numbers)\b", low)):
            return {"op": "pdf_pack", "text": low, "raw": raw}
        if WATERMARK.search(raw) and (self.state["pdfs"][: self.version["pdfs"]] or re.search(r"\bpdf\b", low)):
            return {"op": "pdf_watermark", "text": low, "raw": raw}
        if COMPRESS.search(low) and self.state["pdfs"][: self.version["pdfs"]]:
            return {"op": "pdf_compress"}
        if SPLIT.search(low) and self.state["pdfs"][: self.version["pdfs"]]:
            return {"op": "pdf_split"}
        mp = PAGES.search(low)
        if mp and (re.search(r"\bpdf\b|\bpack\b", low) or (self.state["pdfs"][: self.version["pdfs"]] and not docs and not decks)):
            return {"op": "pdf_pages", "verb": mp.group(1), "pages": mp.group(2), "text": low}
        if HANDOUT.search(low) and decks:
            return {"op": "handout", "deck": self.route_kind(low, "pptx"), "notes": not re.search(r"\bwithout (?:the )?notes\b", low)}
        if DECK_FROM.search(low) and docs:
            return {"op": "deck_from_report", "doc": self.route_kind(low, "docx"), "text": c}
        if REPORT_FROM.search(low) and books and not re.search(r"\b(?:into|in|to)\s+(?:the\s+)?(?:report|document)\b", low):
            return {"op": "report_from_book", "book": books[0], "text": c}
        if REFRESH.search(low) and (self.state["links"] or self.state.get("facts")):
            tg = [
                k
                for k in self.active()
                if self.state["files"][k]["kind"] in ("docx", "pptx")
                and (
                    re.search(r"\b(?:" + ROLE[self.state["files"][k]["kind"]] + r")\b", low)
                    or re.search(r"\b(?:everything|all|links|figures|numbers)\b", low)
                )
            ]
            return {
                "op": "refresh",
                "targets": tg
                or sorted(
                    {L["dst"] for L in self.state["links"].values() if L["dst"] in self.active()}
                    | {k for k in (self.state.get("facts") or {}) if k in self.active()}
                ),
            }
        if books and (docs or decks) and INSERT.search(low) and DATA.search(low):
            mi = INTO.search(low)
            if mi:
                where_word = mi.group(1)
                dst = self.route_kind(where_word, "pptx" if re.search(r"deck|slide|presentation", where_word) else "docx")
                if dst:
                    return {"op": "insert_data", "dst": dst, "book": books[0], "text": low, "into": mi}
        return None

    def route_kind(self, low, kind):
        ks = self.active(kind)
        if not ks:
            return None
        for k in ks:
            if self.state["files"][k]["role"] in low or Path(self.state["files"][k]["label"]).stem.lower() in low:
                return k
        return self.state["focus"] if self.state.get("focus") in ks else ks[-1]

    def run_cross(self, x, message):
        op = x["op"]
        if op == "insert_data":
            return self.insert_data(x, message)
        if op == "refresh":
            return self.refresh(x["targets"], message)
        if op == "report_from_book":
            return self.report_from_book(x["book"], x["text"], message)
        if op == "deck_from_report":
            return self.deck_from_report(x["doc"], x["text"], message)
        if op == "handout":
            return self.handout(x["deck"], x["notes"], message)
        if op.startswith("pdf_"):
            return self.pdf(x, message)
        return f"Not done: {op}"

    def insert_data(self, x, message):
        """'put the pivot of amount by city into the report after the introduction, with a chart of it'."""
        low, dst, wb = x["text"], x["dst"], x["book"]
        kind = self.state["files"][dst]["kind"]
        head = low[: x["into"].start()]
        chart = bool(re.search(r"\b(?:charts?|graphs?|plot)\b", low))
        chart_only = bool(
            re.search(
                r"\b(?:put|insert|add|copy|paste|bring|include|place|drop|show)\s+(?:me\s+)?(?:a|an|the)?\s*(?:\w+\s+)?(?:chart|graph|plot)\b", head
            )
        )
        table = not chart_only or bool(re.search(r"\b(?:tables?|figures|numbers|totals?|breakdown)\b", head))
        what = re.sub(r"\b(?:put|insert|add|copy|paste|bring|include|place|drop|show)\b", " ", low[: x["into"].start()])
        what = (
            re.sub(
                r"\b(?:the|a|an|of|it|and|with|from|workbook|excel|sheet|table|tables|chart|charts|graph|figures|numbers|data|pivot table)\b",
                " ",
                what,
            )
            if not re.search(r"\bby\b", low[: x["into"].start()])
            else low[: x["into"].start()]
        )
        desc = what if re.search(r"\bby\b", what) else ("pivot" if re.search(r"\bpivot\b", low) else what)
        spec, note = self.source(wb, desc if desc.strip() else "pivot")
        data = self.grab(wb, [spec])[0]
        rest = low[x["into"].end() :]
        c = self.chat(dst)
        ops = []
        mk = re.search(r"\b(pie|bar|column|line|doughnut)\b", low)
        if kind == "docx":
            m = c.map()
            t = P.target_in(rest, m, None) if rest.strip() else None
            if t is None and re.search(r"\b(?:introduction|intro|overview|summary|opening|beginning)\b", rest) and m["sections"]:
                first = next((sec for sec in m["sections"] if sec.get("level", 1) == 1), m["sections"][0])
                t = {"kind": "section", "name": re.sub(r"^\d+(\.\d+)*\.?\s+", "", first["title"])}  # the report's own opening section
            wh = (
                "before"
                if re.search(r"\bbefore\b|\babove\b", rest)
                else "after"
                if t
                else ("start" if re.search(r"\bat the (?:start|beginning|top)\b", rest) else "end")
            )
            place = {"anchor": t, "where": wh} if t else {"where": wh}
            if table:
                ops.append({"op": "data_table", "data": data, "link": self._new_link(), **place, "caption": data["title"]})
            if chart:
                ops.append(
                    {
                        "op": "data_chart",
                        "data": data,
                        "link": self._new_link(),
                        **place,
                        "caption": data["title"],
                        **({"chart": mk.group(1)} if mk else {}),
                    }
                )
            if table and chart and t:  # the chart goes under the table, not above it
                ops[1]["anchor"], ops[1]["where"] = t, "after"
        else:
            dm = c.map()
            ma = re.search(r"\bafter\s+slide\s+(\d+)|\bas\s+slide\s+(\d+)|\bbefore\s+slide\s+(\d+)", rest)
            closing = bool(dm["slides"]) and bool(re.search(r"thank|question|contact", dm["slides"][-1]["title"], re.I))
            after = (int(ma.group(1)) if ma.group(1) else int(ma.group(2) or ma.group(3)) - 1) if ma else (dm["count"] - 1 if closing else "end")
            if chart or not table:
                ops.append(
                    {
                        "op": "data_slide",
                        "kind": "chart",
                        "data": data,
                        "link": self._new_link(),
                        "after": after,
                        "title": data["title"],
                        **({"chart": mk.group(1)} if mk else {}),
                    }
                )
            if table:
                ops.append(
                    {
                        "op": "data_slide",
                        "kind": "table",
                        "data": data,
                        "link": self._new_link(),
                        "after": (after + len(ops)) if isinstance(after, int) else after,
                        "title": data["title"],
                    }
                )
        msg = c._change(ops, message, {"ops": []})
        c.save()
        self.state["focus"] = dst
        head = self.chat(wb)
        for o in ops:
            self.state["links"][o["link"]] = {
                "src": wb,
                "spec": spec,
                "dst": dst,
                "kind": "chart" if o["op"] == "data_chart" or o.get("kind") == "chart" else "table",
                "op": {k: v for k, v in o.items() if k not in ("data", "op", "link")},
                "src_v": head.state["cur"],
                "main_hash": self.main_hash(wb, spec.get("sheet")),
                "title": data["title"],
            }
        self._prune_links(dst)  # an edit the file refused leaves no link behind
        return (note + "\n" if note else "") + f"[{self.label(dst)}] {msg}"

    def _prune_links(self, dst):
        """Links whose table or chart is not in the file (a refused edit, an undo inside the file) are forgotten."""
        c = self.chat(dst)
        try:
            if self.state["files"][dst]["kind"] == "docx":
                from docx import Document

                from ai_pc.office.docx_data import links as dlinks

                have = set(dlinks(Document(str(c.path()))))
            else:
                from pptx import Presentation

                from ai_pc.office.pptx_ops import data_links

                have = set(data_links(Presentation(str(c.path()))))
        except Exception:  # noqa: BLE001
            return
        for lid in [k for k, L in self.state["links"].items() if L["dst"] == dst and k not in have]:
            del self.state["links"][lid]

    def _new_link(self):
        n = len(self.state.get("link_seq") or []) + 1
        self.state.setdefault("link_seq", []).append(n)
        return f"aisrc_{n}"

    def refresh(self, targets, message):
        """Every linked table and chart in the target files written again from the workbook as it is now."""
        out = []
        todo = [(lid, L) for lid, L in self.state["links"].items() if L["dst"] in targets and L["dst"] in self.active()]
        figs = [f for f in self.stale_figures() if f[0] in targets]
        if not todo and not figs:
            stale_dsts = [L["dst"] for _, L in self.stale()] + [f[0] for f in self.stale_figures()]
            return "Nothing in those files comes from a workbook." if not stale_dsts else "Nothing there is out of date."
        by_wb = {}
        for lid, L in todo:
            by_wb.setdefault(L["src"], []).append((lid, L))
        fresh = {}
        for wb, items in by_wb.items():
            specs = [L["spec"] for _, L in items]
            for (lid, L), data in zip(items, self.grab(wb, specs)):
                fresh[lid] = data
        for dst in targets:
            items = [(lid, L) for lid, L in todo if L["dst"] == dst]
            ops = []
            for lid, L in items:
                data = fresh[lid]
                ops.append({"op": "data_refresh", "link": lid, "data": data, **({"chart": L["op"]["chart"]} if L["op"].get("chart") else {})})
            seen = set()
            for d_, what, old, new in figs:  # a figure quoted in the text: the old number replaced by the new, everywhere in the file
                if d_ == dst and old not in seen:
                    seen.add(old)
                    ops.append({"op": "replace", "find": old, "with": new})
            if not ops:
                continue
            c = self.chat(dst)
            msg = c._change(ops, message, {"ops": []})
            c.save()
            out.append(f"[{self.label(dst)}] {msg}")
            for lid, L in items:
                L["src_v"] = self.chat(L["src"]).state["cur"]
                L["main_hash"] = self.main_hash(L["src"], L["spec"].get("sheet"))
            F = (self.state.get("facts") or {}).get(dst)
            if F:  # the figures the text now quotes
                F["pairs"], F["src_v"] = self.fact_pairs(F["src"]), self.chat(F["src"]).state["cur"]
        return "\n".join(out)

    # ------------------------------------------------------------------ documents made from other files
    def facts(self, wb_key):
        """The workbook's headline figures, computed (the report quotes these and nothing else)."""
        m = self.chat(wb_key).map()
        s = XM.sheet(m)
        t = s["table"]
        recs = [r for r in XM.records(s) if any(v not in (None, "") for v in r.values())]
        met = BP.metric(t, additive=True)
        vs = BP.value_of_sales(t)
        if vs and (met is None or met["kind"] not in ("money",)):
            name = "Value (Qty x Unit Price)"
            for r in recs:
                a, b = r.get(vs[0]["name"]), r.get(vs[1]["name"])
                r[name] = float(a) * float(b) if isinstance(a, (int, float)) and isinstance(b, (int, float)) else None
            col = {"name": name, "kind": "money", "fmt": vs[1]["fmt"]}
        else:
            col = met
        lines = [f"Rows (records): {len(recs)}"]
        pairs = []
        if col:
            tot = sum(float(r[col["name"]]) for r in recs if isinstance(r.get(col["name"]), (int, float)))
            lines.append(f"Total {col['name']}: {_fmt(tot, col)}")
            pairs.append({"label": f"total {col['name']}", "value": _fmt(tot, col)})
            for g in [c_ for c_ in t["cols"] if c_["kind"] == "text" and 2 <= len({str(r[c_["name"]]) for r in recs if r[c_["name"]]}) <= 12][:4]:
                sums = {}
                for r in recs:
                    if isinstance(r.get(col["name"]), (int, float)) and r.get(g["name"]):
                        sums[r[g["name"]]] = sums.get(r[g["name"]], 0) + float(r[col["name"]])
                order = sorted(sums.items(), key=lambda kv: -kv[1])
                lines.append(f"{col['name']} by {g['name']}: " + "; ".join(f"{k} {_fmt(v, col)} ({v / tot:.0%})" for k, v in order) if tot else "")
                pairs += [{"label": f"{col['name']} of {g['name']} {k}", "value": _fmt(v, col)} for k, v in order]
            dc = next((c_ for c_ in t["cols"] if c_["kind"] == "date"), None)
            if dc:
                sums = {}
                for r in recs:
                    d = str(r.get(dc["name"]) or "")[:7]
                    if re.match(r"\d{4}-\d{2}", d) and isinstance(r.get(col["name"]), (int, float)):
                        sums[d] = sums.get(d, 0) + float(r[col["name"]])
                import datetime as _dt

                lines.append(
                    f"{col['name']} by month: "
                    + "; ".join(f"{_dt.date(int(k[:4]), int(k[5:7]), 1).strftime('%B %Y')} {_fmt(v, col)}" for k, v in sorted(sums.items()))
                )
                pairs += [
                    {"label": f"{col['name']} in {_dt.date(int(k[:4]), int(k[5:7]), 1).strftime('%B %Y')}", "value": _fmt(v, col)}
                    for k, v in sorted(sums.items())
                ]
        self._pairs = pairs
        return "\n".join(x for x in lines if x)

    def fact_pairs(self, wb_key):
        self.facts(wb_key)
        return self._pairs

    def file_text(self, key):
        c = self.chat(key)
        if self.state["files"][key]["kind"] == "docx":
            from docx import Document

            d = Document(str(c.path()))
            return "\n".join([p.text for p in d.paragraphs] + [cell.text for t in d.tables for row in t.rows for cell in row.cells])
        if self.state["files"][key]["kind"] == "pptx":
            from pptx import Presentation

            prs = Presentation(str(c.path()))
            out = []
            for sl in prs.slides:
                for sh in sl.shapes:
                    if sh.has_text_frame:
                        out.append(sh.text_frame.text)
                    if getattr(sh, "has_table", False) and sh.has_table:
                        out += [cell.text for row in sh.table.rows for cell in row.cells]
                if sl.has_notes_slide:
                    out.append(sl.notes_slide.notes_text_frame.text)
            return "\n".join(out)
        return ""

    def stale_figures(self):
        """Figures the text of a report or deck quotes from the workbook that the workbook no longer holds:
        [(file, what, old, new)]."""
        out = []
        for dst, F in (self.state.get("facts") or {}).items():
            if dst not in self.active() or F["src"] not in self.active() or self.chat(F["src"]).state["cur"] == F["src_v"]:
                continue
            now = {p_["label"]: p_["value"] for p_ in self.fact_pairs(F["src"])}
            text = self.file_text(dst)
            for p_ in F["pairs"]:
                old, new = _core(p_["value"]), _core(now.get(p_["label"], ""))
                if old and new and old != new and re.search(r"(?<![\d,.])" + re.escape(old) + r"(?![\d,])", text):
                    out.append((dst, p_["label"], old, new))
        return out

    def report_from_book(self, wb_key, text, message):
        from ai_pc.office.studio import DocStudio

        facts = self.facts(wb_key)
        fpath = self.folder / "facts.md"
        fpath.write_text(f"FACTS (exact, computed from {self.state['files'][wb_key]['label']}):\n{facts}\n", encoding="utf-8")
        pages = None
        mp = re.search(r"\b(\d+|one|two|three|four|five)[- ]pages?\b", text.lower())
        if mp:
            pages = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5}.get(mp.group(1)) or int(mp.group(1))
        req = (
            f"{text}. Write it from the FACTS file only and quote only those figures, exactly as written there. Do not add tables or charts of "
            f"figures: the program adds exact ones from the workbook."
        )
        st = DocStudio(planner=self.planner, log=self.log, look=False, fix_rounds=1)
        sess = st.new(req, files=[str(fpath)], pages=pages, kind="docx")
        dst = self.folder / "files" / Path(sess["docx"]).name
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(sess["docx"], dst)
        key = self._add_file(dst, role="report")
        self.state["files"][key]["made_in"] = "pending"
        self.state["files"][key]["from"] = wb_key
        self.state.setdefault("facts", {})[key] = {"src": wb_key, "pairs": self.fact_pairs(wb_key), "src_v": self.chat(wb_key).state["cur"]}
        self.state["focus"] = key
        c = sess["checks"]["counts"]
        found = sess.get("client_data") or {}
        return (
            f"[{self.label(key)}] new: {sess.get('body_pages') or sess.get('pages')} page(s) written from {self.label(wb_key)}'s figures "
            f"({c.get('pass', 0)} checks pass, {c.get('warn', 0)} warn, {c.get('fail', 0)} fail"
            + (f"; {found.get('found')}/{found.get('checked')} of the figures quoted exactly" if found else "")
            + ")."
        )

    def deck_from_report(self, doc_key, text, message):
        from ai_pc.office.studio import DocStudio

        c = self.chat(doc_key)
        n = None
        mp = re.search(r"\b(\d+)[- ]slides?\b", text.lower())
        if mp:
            n = int(mp.group(1))
        req = f"{text}. Use only what the report says. Do not draw charts or tables of figures: the program adds exact ones from the workbook."
        st = DocStudio(planner=self.planner, log=self.log, look=False, fix_rounds=1)
        sess = st.new_deck(req, files=[str(c.path())], slides=n)
        dst = self.folder / "files" / Path(sess["pptx"]).name
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(sess["pptx"], dst)
        key = self._add_file(dst, role="deck")
        self.state["files"][key]["made_in"] = "pending"
        self.state["files"][key]["from"] = doc_key
        if (self.state.get("facts") or {}).get(doc_key):  # its text comes from the report's: the same figures
            self.state["facts"][key] = json.loads(json.dumps(self.state["facts"][doc_key]))
        self.state["focus"] = key
        out = [
            f"[{self.label(key)}] new: {sess.get('body_pages') or sess.get('pages')} slides from {self.label(doc_key)} "
            f"({sess['checks']['counts'].get('pass', 0)} checks pass, {sess['checks']['counts'].get('fail', 0)} fail)."
        ]
        dk0 = self.chat(key)
        if dk0.version.get("overflow"):  # text PowerPoint measured as too big for its box: fitted by that measure before it is handed over
            out.append(f"[{self.label(key)}] {dk0.say('fix any text that does not fit')}")
        charts = [(lid, L) for lid, L in self.state["links"].items() if L["dst"] == doc_key and L["kind"] == "chart"]
        if charts:  # the report's charts, from the workbook again (exact), before the closing slide
            dk = self.chat(key)
            ops = []
            dm = dk.map()
            closing = bool(re.search(r"thank|question|contact", dm["slides"][-1]["title"], re.I)) if dm["slides"] else False
            pos = dm["count"] - 1 if closing else dm["count"]
            for i, (lid, L) in enumerate(charts):
                data = self.grab(L["src"], [L["spec"]])[0]
                nl = self._new_link()
                ops.append(
                    {
                        "op": "data_slide",
                        "kind": "chart",
                        "data": data,
                        "link": nl,
                        "after": pos + i,
                        "title": data["title"],
                        **({"chart": L["op"]["chart"]} if L["op"].get("chart") else {}),
                    }
                )
                self.state["links"][nl] = {
                    **json.loads(json.dumps(L)),
                    "dst": key,
                    "src_v": self.chat(L["src"]).state["cur"],
                    "main_hash": self.main_hash(L["src"], L["spec"].get("sheet")),
                }
            msg = dk._change(ops, message, {"ops": []})
            dk.save()
            out.append(f"[{self.label(key)}] {msg}")
        return "\n".join(out)

    def handout(self, deck_key, notes, message):
        """A Word handout of a deck: each slide as a picture (PowerPoint draws them) with its speaker notes under it."""
        from pptx import Presentation

        from ai_pc.office import docplan as DPL
        from ai_pc.office import docx_build as DB

        c = self.chat(deck_key)
        png = self.folder / f"_handout_png_{deck_key}"
        shutil.rmtree(png, ignore_errors=True)
        png.mkdir(parents=True)
        tmp = self.folder / f"_handout_{deck_key}.pptx"
        shutil.copy(c.path(), tmp)
        r = c.server.run({"app": "powerpoint", "src": str(tmp.resolve()), "png_dir": str(png.resolve()), "png_width": 1400})
        tmp.unlink(missing_ok=True)
        if not r.get("ok"):
            raise RuntimeError(f"PowerPoint could not draw the slides: {r.get('error')}")
        Presentation(str(c.path()))
        m = c.map()
        blocks = []
        for i, s in enumerate(m["slides"], start=1):
            img = png / f"slide{i:03d}.png"
            blocks.append({"type": "heading", "text": f"Slide {i}: {s['title'] or ''}".strip(": "), "level": 2})
            blocks.append({"type": "image", "path": str(img), "width_cm": 14.5})
            if notes and s.get("notes"):
                blocks.append({"type": "paragraph", "text": s["notes"]})
            if i % 2 == 0 and i < len(m["slides"]):
                blocks.append({"type": "page_break"})
        title = f"{m['title']} - handout"
        plan = DPL.resolve(
            {
                "doctype": "report",
                "title": title,
                "subtitle": f"{len(m['slides'])} slides" + (" with speaker notes" if notes else ""),
                "theme": "corporate",
                "blocks": blocks,
                "cover": False,
                "toc": False,
                "footer": {"page_numbers": True},
            }
        )
        dst = self.folder / "files" / f"{slug(m['title'] or 'deck', 30)}_handout.docx"
        dst.parent.mkdir(parents=True, exist_ok=True)
        DB.build(plan, str(dst))
        key = self._add_file(dst, role="handout")
        self.state["files"][key]["made_in"] = "pending"
        self.state["files"][key]["from"] = deck_key
        h = self.chat(key)
        from docx import Document

        hd = Document(str(h.path()))
        pics = len(hd.inline_shapes)
        body = " ".join(" ".join(p.text.split()) for p in hd.paragraphs)
        noted = sum(1 for s in m["slides"] if s.get("notes") and " ".join(s["notes"].split())[:40] in body)
        ok_pics = pics == len(m["slides"])
        return (
            f"[{self.label(key)}] new: {len(m['slides'])} slide pictures"
            + (f" and {noted} slides' speaker notes" if notes else "")
            + f", {h.version.get('pages') or '?'} page(s). Checked: {'every slide is in it' if ok_pics else f'{pics} pictures for {len(m["slides"])} slides'}."
        )

    # ------------------------------------------------------------------ PDFs
    def _pdf_of(self, key):
        c = self.chat(key)
        v = c.version
        if v.get("pdf") and Path(v["pdf"]).exists():
            return v["pdf"]
        out = self.folder / "pdf" / f"_{key}_v{v['v']}.pdf"
        out.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.folder / f"_pdf_{key}{Path(v['file']).suffix}"
        shutil.copy(v["file"], tmp)
        r = (
            c._job(
                {
                    "app": {"xlsx": "excel", "docx": "word", "pptx": "powerpoint"}[self.state["files"][key]["kind"]],
                    "src": str(tmp.resolve()),
                    "pdf": str(out.resolve()),
                }
            )
            if hasattr(c, "_job")
            else RN.office(
                {
                    "app": {"docx": "word", "pptx": "powerpoint"}[self.state["files"][key]["kind"]],
                    "src": str(tmp.resolve()),
                    "pdf": str(out.resolve()),
                }
            )
        )
        tmp.unlink(missing_ok=True)
        if not r.get("ok") or not out.exists():
            raise RuntimeError(f"no PDF of {self.label(key)}: {r.get('error')}")
        return str(out)

    def _last_pdf(self):
        pdfs = self.state["pdfs"][: self.version["pdfs"]]
        return pdfs[-1] if pdfs else None

    def _add_pdf(self, path, kind, made_from, extra=None):
        self.state["pdfs"] = self.state["pdfs"][: self.version["pdfs"]] + [{"path": str(path), "kind": kind, "from": made_from, **(extra or {})}]
        self._pdf_added = True

    def pdf(self, x, message):
        op = x["op"]
        folder = self.folder / "pdf"
        folder.mkdir(parents=True, exist_ok=True)
        stamp = time.strftime("%H%M%S")
        if op in ("pdf_pack", "pdf_watermark"):
            low, raw = x["text"], x["raw"]
            last = self._last_pdf()
            if op == "pdf_watermark" and last and last["kind"] == "pack":
                opts = dict(last.get("options") or {})
                keys = last["from"]
            else:
                named = [
                    k
                    for k in self.active()
                    if self.state["files"][k]["kind"] in ("docx", "pptx")
                    and (
                        re.search(r"\b" + re.escape(self.state["files"][k]["role"]) + r"\b", low)
                        or Path(self.state["files"][k]["label"]).stem.lower() in low
                    )
                ]
                everything = bool(re.search(r"\b(?:everything|all|both|whole|project)\b", low))
                keys = [k for k in self.active() if self.state["files"][k]["kind"] in ("docx", "pptx")] if everything or not named else named
                order = {"report": 0, "deck": 1, "handout": 2}
                keys.sort(key=lambda k: order.get(self.state["files"][k]["role"], 3))
                opts = {}
                cm = re.search(
                    r"\bcover(?:\s+page)?\s+(?:called|titled|named|saying|reading|with the title)\s+['\"]?(.+?)['\"]?(?=\s*,|\s+(?:and|with)\s+(?:page|a|an|the)\b|\s*$)",
                    raw,
                    re.I,
                )
                if cm or re.search(r"\bcover\b", low):
                    opts["cover"] = {
                        "title": cm.group(1).strip(" .'\"") if cm else self.state["name"],
                        "subtitle": self.state["name"] if cm else None,
                        "lines": [f"{len(keys)} document(s): " + ", ".join(self.state["files"][k]["label"] for k in keys)],
                    }
                if re.search(r"\bpage numbers?\b|\bnumbered\b|\bnumber the pages\b", low):
                    opts["numbers"] = True
            wm = WATERMARK.search(raw)
            if wm:
                opts["watermark"] = next(g for g in wm.groups() if g).strip().upper()
            if not keys:
                raise RuntimeError("there is no report or deck to put in a PDF yet")
            parts = [(self.state["files"][k]["label"].rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title(), self._pdf_of(k)) for k in keys]
            name = slug((opts.get("cover") or {}).get("title") or self.state["name"], 40) or "pack"
            out = folder / f"{name}_{stamp}.pdf"
            r = PF.pack(parts, out, cover=opts.get("cover"), numbers=opts.get("numbers"), watermark=opts.get("watermark"), server=self.server)
            self._add_pdf(out, "pack", keys, {"options": opts, "pages": r["pages"]})
            bad = [ch for ch in r["checks"] if not ch["ok"]]
            return (
                f"[pdf] {out.name}: {r['pages']} pages ({' + '.join(t for t, _ in parts)}"
                + (", a cover" if opts.get("cover") else "")
                + (", page numbers" if opts.get("numbers") else "")
                + (f", a {opts['watermark']} watermark" if opts.get("watermark") else "")
                + f"), {r['kb']} KB. Checked: {len(r['checks']) - len(bad)}/{len(r['checks'])} OK"
                + (" (" + "; ".join(f"{ch['op']}: {ch['what']}" for ch in bad) + ")" if bad else f" ({'; '.join(ch['what'] for ch in r['checks'])})")
                + "."
            )
        last = self._last_pdf()
        if not last:
            raise RuntimeError("there is no PDF yet; say e.g. 'export everything as one pdf'")
        src = last["path"]
        if op == "pdf_compress":
            out = folder / (Path(src).stem.split("_small")[0] + f"_small_{stamp}.pdf")
            r = PF.compress(src, out)
            self._add_pdf(out, "compressed", [Path(src).name])
            bad = [ch for ch in r["checks"] if not ch["ok"]]
            return f"[pdf] {out.name}: {r['kb_before']} KB -> {r['kb_after']} KB. Checked: {len(r['checks']) - len(bad)}/{len(r['checks'])} OK ({'; '.join(ch['what'] for ch in r['checks'])})."
        if op == "pdf_split":
            r = PF.split(src, folder / f"{Path(src).stem}_parts")
            return f"[pdf] split into {len(r['files'])} file(s): {', '.join(Path(f).name for f in r['files'])}. Checked: {r['checks'][0]['what']}."
        if op == "pdf_pages":
            verb = x["verb"]
            out = folder / f"{Path(src).stem}_{verb}_{stamp}.pdf"
            if verb in ("extract", "keep"):
                r = PF.select(src, out, x["pages"])
            elif verb == "rotate":
                r = PF.rotate(src, out, x["pages"], 90)
            else:
                r = PF.delete(src, out, x["pages"])
            self._add_pdf(out, verb, [Path(src).name])
            return f"[pdf] {out.name}: {r['pages']} page(s). Checked: {'; '.join(ch['what'] for ch in r['checks'])}."
        return "Not done."

    # ------------------------------------------------------------------ the model, for what the rules cannot place
    def _llm_route(self, clause, message):
        if self.planner is None:
            return {"ask": f"Which file is '{clause}' about? ({', '.join(self.label(k) for k in self.active())})"}
        self._turn["llm"] = True
        lines = []
        for k in self.active():
            f = self.state["files"][k]
            try:
                m = self.chat(k).map()
                if f["kind"] == "xlsx":
                    what = "sheets: " + ", ".join(m["names"])
                elif f["kind"] == "docx":
                    what = "sections: " + ", ".join(s["title"] for s in m["sections"][:12])
                else:
                    what = "slides: " + ", ".join(s["title"] for s in m["slides"][:12])
            except Exception:  # noqa: BLE001
                what = ""
            lines.append(f"- {k}: {f['kind']} {f['label']} ({f['role']}); {what}")
        try:
            r = self.planner._call(
                "fast",
                [
                    {"role": "system", "content": PROJECT_SYSTEM},
                    {
                        "role": "user",
                        "content": "FILES:\n" + "\n".join(lines) + f"\nLAST FILE: {self.state.get('focus')}\nMESSAGE: {message}\nREQUEST: {clause}",
                    },
                ],
            )
            return parse_json(r.text) or {}
        except Exception as e:  # noqa: BLE001
            return {"ask": f"I could not work out which file that is about ({type(e).__name__})."}
