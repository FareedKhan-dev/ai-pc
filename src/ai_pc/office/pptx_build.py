"""PowerPoint decks from a deck plan, drawn programmatically on blank 16:9 slides (no template look-alikes): a theme's
background, title and accent rule, text sized to fit its box, native charts and tables, speaker notes, slide numbers.

  deck = resolve_deck({"title": ..., "theme": "modern", "accent": None, "meta": {...}, "slides": [...]})
  info = build_deck(deck, "out/docs/x/x.pptx")

Slide layouts (the designer picks one per slide):
  title {title, subtitle}               agenda {items}            section {title, subtitle}
  bullets {title, bullets: [str | {text, items}]}                 two_column / comparison {title, left, right}
      (left / right: {heading, bullets})
  chart {title, chart: {chart, categories, series, title}, takeaway}     table {title, table: {columns, rows}, takeaway}
  stats {title, stats: [{value, label}]}       process {title, steps: [{title, text}]}       image {title, image, bullets, caption}
  quote {text, by}                      closing {title, subtitle, contact}
Every slide may have "notes" (what the presenter says).
"""

import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData  # noqa: F401  (used through charts.chart_data)
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from ai_pc.office import charts as CH
from ai_pc.office import themes
from ai_pc.office.docplan import number, plain, runs

W, H, M = 13.333, 7.5, 0.62
TITLE_Y, TITLE_H, BAR_Y, BODY_Y, BODY_B = 0.42, 0.95, 1.36, 1.72, 6.85
LAYOUTS = {"title", "agenda", "section", "bullets", "two_column", "comparison", "chart", "table", "stats", "process", "image", "quote", "closing"}


def deck_theme(name=None, accent=None):
    dark = name in ("dark", "tech", "night")
    th = themes.get("modern" if dark else name, accent, "other")
    th = dict(th)
    th.update(bg="FFFFFF", text="1F2937", muted="6B7280", card="F3F5F7", line="E5E7EB", dark=False, chart_title=th["head_color"])
    if th["name"] == "elegant":
        th.update(bg="FBF8F4", card="F3ECE6")
    elif th["name"] == "warm":
        th.update(bg="FFFBF5", card="FBEFDF")
    if dark:
        acc = th["accent"] if accent else "38BDF8"
        th.update(
            name="dark",
            bg="0F172A",
            text="F1F5F9",
            muted="94A3B8",
            card="1E293B",
            line="334155",
            accent=acc,
            head_color="F8FAFC",
            chart=[acc, "A78BFA", "F472B6", "FBBF24", "34D399", "F87171"],
            chart_title="F1F5F9",
            dark=True,
        )
        th["table"] = {"header_fill": acc, "header_text": "0F172A", "stripe": "1E293B", "border": "334155", "size": 12}
    return th


def _rgb(h):
    return RGBColor.from_string(str(h).lstrip("#").upper())


def _lum(hx):
    r, g, b = (int(str(hx)[i : i + 2], 16) / 255 for i in (0, 2, 4))
    f = lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def contrast(a, b):
    la, lb = _lum(a), _lum(b)
    return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)


def _mix(hx, to, t):
    a = [int(hx[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(to[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(a, b))


def readable(fg, bg, need=3.0):
    """A colour made just dark (or light) enough to read on another (3:1, the minimum for large text), keeping its hue:
    a teal number on a pale card, white digits on a teal circle (the circle darkens instead)."""
    fg, bg = str(fg).lstrip("#").upper(), str(bg).lstrip("#").upper()
    if contrast(fg, bg) >= need:
        return fg
    toward = "000000" if _lum(bg) > 0.18 else "FFFFFF"
    for k in range(1, 13):
        out = _mix(fg, toward, k * 0.07)
        if contrast(out, bg) >= need:
            return out
    return toward


def resolve_deck(deck):
    d = copy.deepcopy(deck or {})
    d["title"] = " ".join(str(d.get("title") or "").split())
    th = deck_theme(d.get("theme"), d.get("accent"))
    d["theme"], d["th"] = th["name"], th
    slides = []
    for i, s in enumerate(d.get("slides") or []):
        if not isinstance(s, dict):
            continue
        lay = str(s.get("layout") or "bullets").lower().replace(" ", "_").replace("-", "_")
        lay = {
            "title_slide": "title",
            "cover": "title",
            "end": "closing",
            "thank_you": "closing",
            "thanks": "closing",
            "kpi": "stats",
            "numbers": "stats",
            "big_number": "stats",
            "steps": "process",
            "timeline": "process",
            "two_columns": "two_column",
            "columns": "two_column",
            "vs": "comparison",
            "pros_cons": "comparison",
            "graph": "chart",
            "picture": "image",
            "photo": "image",
            "list": "bullets",
            "content": "bullets",
            "text": "bullets",
            "divider": "section",
            "contents": "agenda",
        }.get(lay, lay)
        s = {**s, "layout": lay if lay in LAYOUTS else "bullets", "src": i}
        if s["layout"] == "chart":
            from ai_pc.office.docplan import _clean_block

            c = _clean_block({**(s.get("chart") or {}), "type": "chart"})
            if not c:
                s["layout"] = "bullets"
            else:
                s["chart"] = c
        if s["layout"] == "table":
            from ai_pc.office.docplan import _clean_block

            t = _clean_block({**(s.get("table") or {}), "type": "table"})
            if not t:
                s["layout"] = "bullets"
            else:
                s["table"] = t
        if s["layout"] == "image" and not Path(str(s.get("image") or "")).exists():
            s["layout"] = "bullets"
        if s["layout"] == "bullets":
            s["bullets"] = _items(s.get("bullets") or s.get("items") or ([s["text"]] if s.get("text") else []))
        slides.append(s)
    if slides and slides[0]["layout"] != "title":
        slides.insert(0, {"layout": "title", "title": d["title"], "subtitle": d.get("subtitle"), "src": -1})
    d["slides"] = slides
    return d


def _items(xs):
    out = []
    for x in xs or []:
        if isinstance(x, dict):
            t = " ".join(str(x.get("text") or "").split())
            if t:
                out.append({"text": t, "items": [" ".join(str(y).split()) for y in x.get("items") or [] if str(y).strip()]})
        elif isinstance(x, list) and out:
            out[-1]["items"] += [" ".join(str(y).split()) for y in x if str(y).strip()]
        elif str(x).strip():
            out.append({"text": " ".join(str(x).split()), "items": []})
    return out


# a font's line height (ascent + descent + gap, in em) and average character width (em), measured against what
# PowerPoint reports for its own layout; text boxes add 1.1 line spacing on top
LINE_H = {"segoe ui": 1.33, "segoe ui semibold": 1.33, "calibri": 1.22, "georgia": 1.14, "arial": 1.15, "times new roman": 1.15, "cambria": 1.17}
CHAR_EM = {"segoe ui": 0.53, "segoe ui semibold": 0.56, "calibri": 0.49, "georgia": 0.56, "arial": 0.53, "times new roman": 0.47, "cambria": 0.5}


def text_height(texts, w_in, pt, spacing=1.1, gap_pt=8, char_em=None, font="calibri", bold=False):
    """Estimated height (pt) of the texts (one paragraph each) at `pt` in a box `w_in` wide, words wrapping whole."""
    f = str(font or "calibri").lower()
    lh = LINE_H.get(f, 1.25) * spacing
    em = (char_em or CHAR_EM.get(f, 0.53)) * (1.06 if bold else 1.0)
    cpl = max(6, (w_in * 72 - 10) / (pt * em))
    lines = 0
    for t in texts:
        n, cur = 1, 0
        for w_ in plain(t).split():
            if cur and cur + 1 + len(w_) > cpl:
                n, cur = n + 1, len(w_)
            else:
                cur += (1 if cur else 0) + len(w_)
        lines += n
    return lines * pt * lh + max(0, len(texts) - 1) * gap_pt


def fit_size(texts, w_in, h_in, max_pt, min_pt, spacing=1.1, gap_pt=8, char_em=None, font="calibri", bold=False):
    """The largest size (pt) at which the texts (one paragraph each) fit the box, from the font's measured line height
    and character width (PowerPoint's measurement after rendering is the final word)."""
    f = str(font or "calibri").lower()
    lh = LINE_H.get(f, 1.25) * spacing
    em = (char_em or CHAR_EM.get(f, 0.53)) * (1.06 if bold else 1.0)
    for pt in range(int(max_pt), int(min_pt) - 1, -1):
        cpl = max(6, (w_in * 72 - 10) / (pt * em))
        lines = 0
        for t in texts:  # words wrap whole: a line holds whole words only
            n, cur = 1, 0
            for w_ in plain(t).split():
                if cur and cur + 1 + len(w_) > cpl:
                    n, cur = n + 1, len(w_)
                else:
                    cur += (1 if cur else 0) + len(w_)
            lines += n
        if lines * pt * lh + max(0, len(texts) - 1) * gap_pt <= h_in * 72 - 8:
            return pt
    return int(min_pt)


class _Deck:
    def __init__(self, deck):
        self.d = deck
        self.th = deck["th"]
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0
        self.out = []

    # drawing helpers -----------------------------------------------------------------------------------------------
    def rect(self, slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(fill)
        if line:
            s.line.color.rgb = _rgb(line)
            s.line.width = Pt(1)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = 0.06
        return s

    def text(
        self,
        slide,
        x,
        y,
        w,
        h,
        paras,
        size,
        color=None,
        font=None,
        bold=False,
        align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP,
        bullets=False,
        gap=8,
        italic=False,
        name=None,
    ):
        """A text box. paras: [str] or [(text, level)]; inline **bold** works."""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        first = True
        for item in paras:
            t, lvl = item if isinstance(item, tuple) else (item, 0)
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            p.space_after = Pt(gap)
            p.line_spacing = 1.1
            sz = size if lvl == 0 else max(12, size - 3)
            for seg, f in runs(t):
                r = p.add_run()
                r.text = seg
                r.font.size = Pt(sz)
                r.font.name = font or self.th["body"]
                r.font.bold = bold or bool(f.get("bold"))
                r.font.italic = italic or bool(f.get("italic"))
                r.font.color.rgb = _rgb(color or self.th["text"])
            if bullets:
                self._bullet(p, lvl)
        return tb

    def _bullet(self, p, level):
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(int(Inches(0.30 + 0.38 * level))))
        ppr.set("indent", str(int(-Inches(0.26))))
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont"):
            for el in ppr.findall(qn(tag)):
                ppr.remove(el)
        clr = etree.SubElement(ppr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", self.th["accent"] if level == 0 else self.th["muted"])
        font = etree.SubElement(ppr, qn("a:buFont"))
        font.set("typeface", "Arial")
        ch = etree.SubElement(ppr, qn("a:buChar"))
        ch.set("char", "•" if level == 0 else "\u2013")

    def new_slide(self, s, title=None, number=True):
        slide = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(self.th["bg"])
        if title:
            size = fit_size([title], W - 2 * M, TITLE_H, 34, 22, spacing=1.0, gap_pt=0, font=self.th["head"], bold=True)
            self.text(
                slide,
                M,
                TITLE_Y,
                W - 2 * M,
                TITLE_H,
                [title],
                size,
                self.th["head_color"],
                self.th["head"],
                bold="semibold" not in self.th["head"].lower(),
                anchor=MSO_ANCHOR.BOTTOM,
                gap=0,
                name="Title",
            )
            self.rect(slide, M, BAR_Y, 1.1, 0.055, self.th["accent"])
        if number:
            self.text(slide, W - M - 1.0, H - 0.48, 1.0, 0.3, [str(self.n)], 11, self.th["muted"], align=PP_ALIGN.RIGHT, gap=0, name="SlideNumber")
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])
        self.out.append(
            {"slide": self.n, "layout": s["layout"], "title": plain(title or s.get("title") or s.get("text") or "")[:80], "src": s.get("src")}
        )
        return slide

    def bullet_box(self, slide, items, x, y, w, h, max_pt=28, min_pt=14, scale=1.0, name="Body"):
        [(it["text"], 0) for it in items for _ in [0]]
        flat = []
        for it in items:
            flat.append((it["text"], 0))
            flat += [(x_, 1) for x_ in it.get("items") or []]
        size = fit_size([t for t, _ in flat], w - 0.35, h, max_pt * scale, min_pt, gap_pt=12, font=self.th["body"])
        return self.text(slide, x, y, w, h, flat, size, bullets=True, gap=12 if size >= 20 else 8, name=name)

    # layouts -------------------------------------------------------------------------------------------------------
    def l_title(self, s):
        th = self.th
        slide = self.new_slide(s, number=False)
        self.rect(slide, M, 2.25, 0.12, 2.35, th["accent"])
        title = s.get("title") or self.d["title"]
        size = fit_size([title], W - 2 * M - 0.6, 1.75, 48, 30, spacing=1.0, gap_pt=0, font=self.th["head"], bold=True)
        self.text(
            slide,
            M + 0.4,
            2.1,
            W - 2 * M - 0.6,
            1.75,
            [title],
            size,
            th["head_color"],
            th["head"],
            bold="semibold" not in th["head"].lower(),
            anchor=MSO_ANCHOR.BOTTOM,
            gap=0,
            name="Title",
        )
        if s.get("subtitle") or self.d.get("subtitle"):
            self.text(
                slide, M + 0.4, 3.95, W - 2 * M - 0.6, 0.9, [s.get("subtitle") or self.d.get("subtitle")], 22, th["muted"], gap=0, name="Subtitle"
            )
        meta = self.d.get("meta") or {}
        vals = list(dict.fromkeys(str(meta[k]).strip() for k in ("author", "org", "date") if meta.get(k)))
        if vals:
            self.text(
                slide,
                M + 0.4,
                6.15,
                W - 2 * M - 0.6,
                0.5,
                ["   ·   ".join(vals)],
                16,
                th["text"] if not th["dark"] else th["muted"],
                gap=0,
                name="Meta",
            )

    def l_agenda(self, s):
        th = self.th
        slide = self.new_slide(s, s.get("title") or "Agenda")
        items = [plain(x if isinstance(x, str) else x.get("text", "")) for x in s.get("items") or s.get("bullets") or []][:8]
        rows = max(1, len(items))
        step = min(0.95, (BODY_B - BODY_Y - 0.1) / rows)
        for k, it in enumerate(items):
            y = BODY_Y + 0.1 + k * step
            c = self.rect(slide, M, y, 0.48, 0.48, th["accent"], MSO_SHAPE.OVAL)
            c.text_frame.text = str(k + 1)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.runs[0].font.size, p.runs[0].font.bold = Pt(16), True
            p.runs[0].font.color.rgb = _rgb("FFFFFF" if not th["dark"] else "0F172A")
            p.runs[0].font.name = th["body"]
            self.text(
                slide,
                M + 0.75,
                y - 0.06,
                W - 2 * M - 0.9,
                0.6,
                [it],
                fit_size([it], W - 2 * M - 0.9, 0.6, 26, 16, gap_pt=0, font=self.th["body"]),
                anchor=MSO_ANCHOR.MIDDLE,
                gap=0,
                name=f"Item{k + 1}",
            )

    def l_section(self, s):
        th = self.th
        slide = self.new_slide(s)
        ink = "FFFFFF" if not th["dark"] else "0F172A"
        self.rect(slide, 0, 0, W * 0.38, H, readable(th["accent"], ink))  # the panel darkens until its number reads
        num = str(s.get("number") or "")
        if num:
            self.text(slide, 0.6, 2.4, W * 0.38 - 1.0, 1.6, [num.zfill(2) if num.isdigit() else num], 80, ink, th["head"], bold=True, gap=0)
        self.text(
            slide,
            W * 0.38 + 0.6,
            2.6,
            W * 0.62 - 1.2,
            1.4,
            [s.get("title") or ""],
            fit_size([s.get("title") or ""], W * 0.62 - 1.2, 1.4, 42, 26, gap_pt=0, font=self.th["head"], bold=True),
            th["head_color"],
            th["head"],
            bold=True,
            anchor=MSO_ANCHOR.BOTTOM,
            gap=0,
            name="Title",
        )
        if s.get("subtitle"):
            self.text(slide, W * 0.38 + 0.6, 4.1, W * 0.62 - 1.2, 1.0, [s["subtitle"]], 20, th["muted"], gap=0, name="Subtitle")

    def l_bullets(self, s):
        slide = self.new_slide(s, s.get("title"))
        self.bullet_box(
            slide, s["bullets"] or [{"text": "", "items": []}], M, BODY_Y + 0.1, W - 2 * M, BODY_B - BODY_Y - 0.1, scale=s.get("scale", 1.0)
        )

    def _column(self, slide, part, x, y, w, h, card=False, colour=None, scale=1.0):
        th = self.th
        part = part if isinstance(part, dict) else {"bullets": part if isinstance(part, list) else [part]}
        top = y
        if card:  # the card as tall as its text at the size that fits the slide (a tall empty card looks unfinished)
            items_ = _items(part.get("bullets") or part.get("items") or ([part["text"]] if part.get("text") else []))
            texts = [it["text"] for it in items_] + [x_ for it in items_ for x_ in it.get("items") or []]
            inner = h - 0.25 - (0.65 if part.get("heading") else 0) - 0.3
            sz = fit_size(texts, w - 0.95, inner, 24 * scale, 13, gap_pt=12, font=th["body"])
            need = 0.25 + (0.65 if part.get("heading") else 0) + (text_height(texts, w - 0.95, sz, gap_pt=12, font=th["body"]) + 14) / 72 + 0.3
            h = max(2.2, min(h, need))
            self.rect(slide, x, y, w, h, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE)
            x, w = x + 0.3, w - 0.6
            top = y + 0.25
        if part.get("heading"):
            if card:
                self.rect(slide, x - 0.3, y, w + 0.6, 0.08, colour or th["accent"])
            self.text(
                slide,
                x,
                top,
                w,
                0.55,
                [part["heading"]],
                22,
                readable(colour or th["accent"], th["card"] if card else th["bg"]),
                th["head"],
                bold=True,
                gap=0,
                name="ColumnHeading",
            )
            top += 0.65
        items = _items(part.get("bullets") or part.get("items") or ([part["text"]] if part.get("text") else []))
        if items:
            self.bullet_box(slide, items, x, top, w, (y + h) - top - (0.2 if card else 0), max_pt=24, min_pt=13, scale=scale)

    def l_two_column(self, s, card=False):
        slide = self.new_slide(s, s.get("title"))
        gap = 0.45
        cw = (W - 2 * M - gap) / 2
        self._column(slide, s.get("left") or {}, M, BODY_Y + 0.1, cw, BODY_B - BODY_Y - 0.1, card, None, s.get("scale", 1.0))
        self._column(
            slide,
            s.get("right") or {},
            M + cw + gap,
            BODY_Y + 0.1,
            cw,
            BODY_B - BODY_Y - 0.1,
            card,
            self.th["accent2"] if card else None,
            s.get("scale", 1.0),
        )

    def l_comparison(self, s):
        self.l_two_column(s, card=True)

    def l_chart(self, s):
        th = self.th
        slide = self.new_slide(s, s.get("title"))
        spec = s["chart"]
        take = s.get("takeaway")
        cw = (W - 2 * M) * (0.66 if take else 1.0)
        from pptx.enum.chart import XL_CHART_TYPE  # noqa: F401

        gf = slide.shapes.add_chart(
            CH.TYPES[spec["chart"]], Inches(M), Inches(BODY_Y + 0.05), Inches(cw), Inches(BODY_B - BODY_Y - 0.05), CH.chart_data(spec)
        )
        gf.name = "Chart"
        ch = CH.style(gf.chart, spec, th, th["body"], 16, text_colour=th["text"] if th["dark"] else "404040", grid_colour=th["line"])
        try:
            if ch.has_legend:
                ch.legend.font.size = Pt(16)
                ch.legend.font.color.rgb = _rgb(th["text"] if th["dark"] else "404040")
            if ch.plots[0].has_data_labels:
                ch.plots[0].data_labels.font.size = Pt(15)
        except Exception:  # noqa: BLE001
            pass
        if take:
            x = M + cw + 0.35
            self.rect(slide, x, BODY_Y + 0.4, W - M - x, 3.4, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE)
            self.rect(slide, x, BODY_Y + 0.4, 0.09, 3.4, th["accent"])
            self.text(
                slide,
                x + 0.3,
                BODY_Y + 0.6,
                W - M - x - 0.5,
                3.0,
                [take],
                fit_size([take], W - M - x - 0.5, 3.0, 24, 14, gap_pt=0, font=th["body"]),
                th["text"],
                anchor=MSO_ANCHOR.MIDDLE,
                name="Takeaway",
            )

    def l_table(self, s):
        th = self.th
        slide = self.new_slide(s, s.get("title"))
        t = s["table"]
        cols, rows = t["columns"], t["rows"][:12]
        take = s.get("takeaway")
        tw = (W - 2 * M) * (0.68 if take else 1.0)
        size = 16 if len(rows) <= 5 else 14 if len(rows) <= 8 else 12
        rh = min(0.55, (BODY_B - BODY_Y - 0.2) / (len(rows) + 1))
        gt = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(M), Inches(BODY_Y + 0.1), Inches(tw), Inches(rh * (len(rows) + 1)))
        gt.name = "Table"
        tbl = gt.table
        numeric = []
        for j in range(len(cols)):
            cells = [r[j] for r in rows if str(r[j]).strip()]
            numeric.append(j > 0 and bool(cells) and sum(1 for c in cells if number(c)) >= 0.7 * len(cells))
        weight = [max(4, max(len(plain(str(x))) for x in [cols[j]] + [r[j] for r in rows])) for j in range(len(cols))]
        for j, wgt in enumerate(weight):
            tbl.columns[j].width = Emu(int(Inches(tw) * wgt / sum(weight)))
        tt = th["table"]
        for i in range(len(rows) + 1):
            for j in range(len(cols)):
                cell = tbl.cell(i, j)
                val = cols[j] if i == 0 else rows[i - 1][j]
                cell.text = ""
                p = cell.text_frame.paragraphs[0]
                r = p.add_run()
                r.text = plain(str(val))
                r.font.size = Pt(size)
                r.font.name = th["body"]
                r.font.bold = i == 0
                r.font.color.rgb = _rgb(tt["header_text"] if i == 0 else th["text"])
                p.alignment = PP_ALIGN.RIGHT if numeric[j] else PP_ALIGN.LEFT
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(tt["header_fill"] if i == 0 else (tt.get("stripe") or th["card"]) if i % 2 == 0 else th["bg"])
                cell.margin_left = cell.margin_right = Inches(0.1)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if take:
            x = M + tw + 0.35
            self.text(
                slide,
                x,
                BODY_Y + 0.3,
                W - M - x,
                4.0,
                [take],
                fit_size([take], W - M - x, 4.0, 22, 14, gap_pt=0, font=th["body"]),
                th["text"],
                anchor=MSO_ANCHOR.TOP,
                name="Takeaway",
            )

    def l_stats(self, s):
        th = self.th
        slide = self.new_slide(s, s.get("title"))
        stats = [x for x in s.get("stats") or [] if isinstance(x, dict)][:4] or [{"value": "", "label": ""}]
        n = len(stats)
        gap = 0.35
        cw = (W - 2 * M - gap * (n - 1)) / n
        for k, st in enumerate(stats):
            x = M + k * (cw + gap)
            self.rect(slide, x, BODY_Y + 0.5, cw, 3.6, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE)
            self.rect(slide, x, BODY_Y + 0.5, cw, 0.09, th["accent"] if k % 2 == 0 else th["accent2"])
            val = plain(str(st.get("value") or ""))
            token = " " not in val.strip()
            one = fit_size([val], cw - 0.4, 1.0, 60, 14 if token else 26, spacing=1.0, gap_pt=0, font=th["head"], bold=True)  # one line
            tb = self.text(
                slide,
                x + 0.2,
                BODY_Y + 0.85,
                cw - 0.4,
                1.5,
                [val],
                one if (one > 26 or token) else fit_size([val], cw - 0.4, 1.45, 44, 22, spacing=1.0, gap_pt=0, font=th["head"], bold=True),
                readable(th["accent"] if k % 2 == 0 else th["accent2"], th["card"]),
                th["head"],
                bold=True,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
                gap=0,
                name=f"Stat{k + 1}",
            )
            if token:  # "5,634,700" never breaks in the middle: one line, sized to fit
                tb.text_frame.word_wrap = False
            lab = str(st.get("label") or "")
            self.text(
                slide,
                x + 0.25,
                BODY_Y + 2.45,
                cw - 0.5,
                1.5,
                [lab],
                fit_size([lab], cw - 0.5, 1.5, 20, 12, gap_pt=0, font=th["body"]),
                th["muted"],
                align=PP_ALIGN.CENTER,
                gap=0,
                name=f"StatLabel{k + 1}",
            )
        if s.get("note"):
            self.text(slide, M, 6.2, W - 2 * M, 0.55, [s["note"]], 16, th["muted"], italic=True, gap=0, name="Note")

    def l_process(self, s):
        th = self.th
        slide = self.new_slide(s, s.get("title"))
        steps = [x if isinstance(x, dict) else {"title": str(x)} for x in s.get("steps") or []][:6] or [{"title": ""}]
        n = len(steps)
        cw = (W - 2 * M) / n
        y = BODY_Y + 0.7
        self.rect(slide, M + cw / 2, y + 0.33, cw * (n - 1), 0.05, th["line"])
        for k, st in enumerate(steps):
            cx = M + k * cw + cw / 2
            ink = "FFFFFF" if not th["dark"] else "0F172A"
            c = self.rect(slide, cx - 0.36, y, 0.72, 0.72, readable(th["accent"] if k % 2 == 0 else th["accent2"], ink), MSO_SHAPE.OVAL)
            c.text_frame.text = str(k + 1)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.runs[0].font.size, p.runs[0].font.bold = Pt(20), True
            p.runs[0].font.color.rgb = _rgb(ink)
            p.runs[0].font.name = th["body"]
            t = plain(str(st.get("title") or ""))
            self.text(
                slide,
                M + k * cw + 0.12,
                y + 0.95,
                cw - 0.24,
                0.9,
                [t],
                fit_size([t], cw - 0.24, 0.9, 22, 13, gap_pt=0, font=th["head"], bold=True),
                th["head_color"],
                th["head"],
                bold=True,
                align=PP_ALIGN.CENTER,
                gap=0,
                name=f"Step{k + 1}",
            )
            if st.get("text"):
                tx = str(st["text"])
                self.text(
                    slide,
                    M + k * cw + 0.12,
                    y + 1.9,
                    cw - 0.24,
                    2.4,
                    [tx],
                    fit_size([tx], cw - 0.24, 2.4, 18, 11, gap_pt=0, font=th["body"]),
                    th["muted"],
                    align=PP_ALIGN.CENTER,
                    gap=0,
                    name=f"StepText{k + 1}",
                )

    def l_image(self, s):
        th = self.th
        slide = self.new_slide(s, s.get("title"))
        items = _items(s.get("bullets") or [])
        bx, bw = (M + (W - 2 * M) * 0.45, (W - 2 * M) * 0.55) if items else (M, W - 2 * M)
        by, bh = BODY_Y + 0.1, BODY_B - BODY_Y - (0.55 if s.get("caption") else 0.1)
        from PIL import Image

        with Image.open(s["image"]) as im:
            iw, ih = im.size
        sc = min(bw / iw, bh / ih)
        w_, h_ = iw * sc, ih * sc
        pic = slide.shapes.add_picture(str(s["image"]), Inches(bx + (bw - w_) / 2), Inches(by + (bh - h_) / 2), Inches(w_), Inches(h_))
        pic.name = "Image"
        if s.get("caption"):
            self.text(slide, bx, BODY_B - 0.45, bw, 0.4, [s["caption"]], 13, th["muted"], align=PP_ALIGN.CENTER, italic=True, gap=0, name="Caption")
        if items:
            self.bullet_box(slide, items, M, BODY_Y + 0.1, (W - 2 * M) * 0.42, BODY_B - BODY_Y - 0.1, max_pt=22)

    def l_quote(self, s):
        th = self.th
        slide = self.new_slide(s)
        self.text(slide, M + 0.3, 0.6, 2.2, 2.4, ["“"], 140, th["accent"], "Georgia", gap=0, name="QuoteMark")
        q = str(s.get("text") or s.get("quote") or "")
        self.text(
            slide,
            M + 1.2,
            2.2,
            W - 2 * M - 2.0,
            2.8,
            [q],
            fit_size([q], W - 2 * M - 2.0, 2.8, 36, 20, gap_pt=0, font=th["head"]),
            th["head_color"],
            th["head"],
            italic=True,
            anchor=MSO_ANCHOR.MIDDLE,
            gap=0,
            name="Quote",
        )
        if s.get("by"):
            self.text(slide, M + 1.2, 5.2, W - 2 * M - 2.0, 0.5, [f"{s['by']}"], 18, th["muted"], gap=0, name="By")

    def l_closing(self, s):
        th = self.th
        slide = self.new_slide(s, number=False)
        self.rect(slide, 0, H - 0.25, W, 0.25, th["accent"])
        t = s.get("title") or "Thank you"
        self.text(
            slide,
            M,
            2.3,
            W - 2 * M,
            1.4,
            [t],
            fit_size([t], W - 2 * M, 1.4, 54, 32, spacing=1.0, gap_pt=0, font=th["head"], bold=True),
            th["head_color"],
            th["head"],
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.BOTTOM,
            gap=0,
            name="Title",
        )
        sub = s.get("subtitle") or ""
        if sub:
            self.text(slide, M, 3.85, W - 2 * M, 0.8, [sub], 22, th["muted"], align=PP_ALIGN.CENTER, gap=0, name="Subtitle")
        if s.get("contact"):
            self.text(slide, M, 4.8, W - 2 * M, 0.8, [s["contact"]], 16, th["muted"], align=PP_ALIGN.CENTER, gap=0, name="Contact")

    def run(self, path):
        for s in self.d["slides"]:
            getattr(self, "l_" + s["layout"])(s)
        cp = self.prs.core_properties
        cp.title = self.d["title"][:250]
        cp.author = str((self.d.get("meta") or {}).get("author") or "")[:250]
        cp.comments = "Made by the AI PC document agent"
        cp.keywords = f"theme:{self.th['name']};accent:{self.th['accent']}"  # later edits draw new slides in the same look
        self.prs.save(str(path))
        return {"path": str(path), "slides": self.out}


def build_deck(deck, path):
    return _Deck(deck if "th" in deck else resolve_deck(deck)).run(path)
