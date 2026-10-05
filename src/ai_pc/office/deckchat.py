"""A conversation about a PowerPoint deck (the Word chat's machinery: one version per message, undo / redo / go to vN /
versions / compare / export, questions, the cheap model for what the rules cannot read).

  c = DeckChat.start("out/docs/deck/deck.pptx")
  c.say("delete slide 4, move slide 7 to the end and make the titles dark blue")
  c.say("add a slide about financing options after slide 5")
  c.say("make slide 3 shorter, add speaker notes to every slide and translate slide 2 into urdu")

Every version is rendered by PowerPoint in the background; the slides that changed are measured (text taller than its box
is an overflow) and the reply says so.
"""

import json
import re
import shutil
import time
from pathlib import Path

from pptx import Presentation

from ai_pc.core.util import parse_json
from ai_pc.office import doc_parse as P
from ai_pc.office import pptx_ops as PO
from ai_pc.office import render as RN
from ai_pc.office.docchat import CHATS, QUESTION, DocChat

LAYOUT_WORDS = [
    ("chart", r"\bchart|graph\b"),
    ("agenda", r"\bagenda\b"),
    ("quote", r"\bquote\b"),
    ("closing", r"\b(?:thank you|thanks|closing|q ?& ?a|questions)\b"),
    ("comparison", r"\b(?:comparison|compare|vs\.?|versus|pros and cons|before and after)\b"),
    ("process", r"\b(?:steps|process|timeline|how to|roadmap)\b"),
    ("stats", r"\b(?:stats|statistics|numbers|kpis?|key figures|at a glance)\b"),
]
DECK_SYSTEM = """You turn a client's request about their PowerPoint deck into edit operations for a program. Reply with ONE JSON
object: {"ops": [...], "ask": "<short question back, only if it cannot be done as asked>", "answer": "<only if it was a question>"}.
Operations (slides are numbered from 1; T = {"kind": "slide", "n": 3} | {"kind": "slides", "ns": [2, 5]} | {"kind": "all"} |
{"kind": "slide", "which": "last"}):
  {"op": "slide_delete", "target": T}    {"op": "slide_move", "slide": 5, "to": 2 | "end" | "start"}    {"op": "slide_duplicate", "slide": 3}
  {"op": "slide_add", "after": 4 | "end" | "start", "about": "<what the slide says>", "layout": "bullets|stats|chart|process|comparison|quote|agenda|closing"}
  {"op": "title_set", "slide": 3, "text": "..."}    {"op": "replace", "find": "...", "with": "..."}
  {"op": "rewrite", "target": T, "instruction": "..."}    {"op": "translate", "target": T, "language": "..."}
  {"op": "style", "target": T, "part": "titles|body|all", "set": {"font", "size" (pt or "+2"), "color", "bold", "italic"}}
  {"op": "background", "target": T, "color": "..."}    {"op": "notes", "target": T, "write": true | "text": "..."}
  {"op": "shrink", "target": T}    {"op": "bullet_add", "slide": 3, "text": "..."}    {"op": "bullet_delete", "slide": 3, "n": 2 | "last"}
Do only what was asked, with the fewest operations."""


def slides_in(c, m):
    """Slide numbers named in a clause: "slide 3", "slides 2 and 5", "slides 2-4", "the last slide", "the title slide",
    "the chart slide", "every slide"."""
    if re.search(r"\b(?:every|each|all(?: the)?) slides?\b|\bthe (?:whole |entire )?deck\b|\beverywhere\b|\ball of them\b", c):
        return list(range(1, m["count"] + 1)), "all"
    mm = re.search(r"\bslides?\s+(\d+)\s*(?:-|to|through)\s*(\d+)\b", c)
    if mm:
        a, b = int(mm.group(1)), int(mm.group(2))
        return [k for k in range(min(a, b), max(a, b) + 1) if 1 <= k <= m["count"]], "slides"
    mm = re.search(r"\bslides?\s+((?:\d+\s*(?:,|and|&)\s*)*\d+)\b", c)
    if mm:
        return [int(x) for x in re.findall(r"\d+", mm.group(1)) if 1 <= int(x) <= m["count"]], "slides"
    mm = re.search(r"\b(first|second|third|fourth|fifth|sixth|seventh|eighth|ninth|tenth|last|\d+(?:st|nd|rd|th)) slide\b", c)
    if mm:
        w = mm.group(1)
        k = m["count"] if w == "last" else P.ORD.get(w) or int(re.sub(r"\D", "", w) or 1)
        return ([k] if 1 <= k <= m["count"] else []), "slides"
    if re.search(r"\b(?:title|cover|opening) slide\b", c):
        ks = [s["n"] for s in m["slides"] if s.get("cover")] or [
            s["n"] for s in m["slides"] if m["title"] and s["title"].strip().lower() == m["title"].strip().lower()
        ]
        return ks[:1] or [1], "slides"
    if re.search(r"\b(?:thank you|thanks|closing|final|end|last) slide\b", c):
        ks = [s["n"] for s in m["slides"] if re.search(r"\bthank|\bquestions?\b|q ?& ?a|شکریہ", s["title"], re.I)]
        return ks[-1:] or [m["count"]], "slides"
    for word, test in (("chart", "chart"), ("table", "table"), ("picture", "picture"), ("image", "picture"), ("photo", "picture")):
        if re.search(rf"\bthe {word} slide\b", c):
            ks = [s["n"] for s in m["slides"] if s.get(test)]
            if ks:
                return ks[:1], "slides"
    mm = re.search(r"\bthe (\w+(?: \w+)?) slide\b", c)
    if mm:
        ks = [s["n"] for s in m["slides"] if mm.group(1) in s["title"].lower()]
        if ks:
            return ks[:1], "slides"
    return [], None


def parse_deck(clause, m, focus=None, raw=None, carry=False):
    c = P.normalize(clause)
    raw = raw or clause
    q = P.quoted(raw)
    cm = P.QUOTE.sub(" QQ ", c)
    ns, how = slides_in(cm, m)
    named = [int(x) for x in re.findall(r"\bslides?\s+(\d+)", cm)] + [
        int(x) for x in re.findall(r"\bslides?\s+\d+\s*(?:-|to|through|,|and|&)\s*(\d+)", cm)
    ]
    gone = [k for k in named if not 1 <= k <= m["count"]]
    if gone:
        return {"ops": [], "focus": focus, "ask": f"The deck has {m['count']} slides; there is no slide {gone[0]}."}
    if (
        not ns
        and re.search(
            r"\b(?:it|its|this|them|their|these|those)\b|\bthe (?:copy|duplicate|new (?:one|slide))\b|\bthat\b(?!\s+(?:is|are|was|were|does|do|doesn'?t|don'?t|did|didn'?t|has|have|had|can|can'?t|could|will|won'?t|would|should|not|fits?|says?|shows?|i|you|we|they|the|a|an)\b)",
            cm,
        )
        and focus
    ):
        ns, how = list(focus.get("ns") or []), "slides"
    # "and make the titles blue" after "delete slide 4": every title, not slide 4's
    if not ns and carry and focus and not re.search(r"\b(?:titles|headings|subtitles|all (?:the )?(?:text|bullets)|every\w*|background)\b", cm):
        ns, how = list(focus.get("ns") or []), "slides"
    fresh = bool(ns) and bool(focus) and ns == list(focus.get("ns") or []) and bool(focus.get("fresh")) and not slides_in(cm, m)[0]
    target = {"kind": "all"} if how == "all" else ({"kind": "slides", "ns": ns} if len(ns) > 1 else {"kind": "slide", "n": ns[0]}) if ns else None
    out = {"ops": [], "focus": {"ns": ns} if ns else focus, "ask": None}
    remove = bool(re.search(r"\b(?:remove|delete|drop|get rid of|take out|cut)\b", c))

    def done(*ops):
        out["ops"] = [dict(o, _fresh=True) for o in ops] if fresh else list(ops)
        return out

    # slides: delete, move, swap, duplicate, add
    if remove and re.search(r"\bslides?\b", c) and ns and not re.search(r"\bbullet|notes?\b|\bbackground\b|\btitle\b(?! slide)", c):
        return done({"op": "slide_delete", "target": target})
    mm = re.search(r"\bswap\s+slides?\s+(\d+)\s+(?:and|with)\s+(\d+)", c)
    if mm:
        return done({"op": "slide_swap", "a": int(mm.group(1)), "b": int(mm.group(2))})
    mm = re.search(
        r"\bmove\s+(?:the\s+)?(?:slide\s+(\d+)|(\w+(?: you)?) slide|(it|this|that))\s+(?:back\s+)?(?:to\s+)?(?:(the end|the start|the beginning|the top|the front|the back)|(?:position|place)\s+(\d+)|"
        r"(after|before)\s+slide\s+(\d+))",
        c,
    )
    if mm:
        n = int(mm.group(1)) if mm.group(1) else (ns[0] if ns else None)
        if n:
            out["focus"] = {"ns": [n]}
            if mm.group(4):
                return done({"op": "slide_move", "slide": n, "to": "end" if mm.group(4) in ("the end", "the back") else "start"})
            if mm.group(5):
                return done({"op": "slide_move", "slide": n, "to": int(mm.group(5))})
            return done({"op": "slide_move", "slide": n, mm.group(6): int(mm.group(7))})
    if re.search(r"\b(?:duplicate|copy|clone)\b", c) and ns and not re.search(r"\bthe (?:copy|duplicate)\b", c):
        out["focus"] = {"ns": [ns[0] + 1], "fresh": True}  # "...and make it red": the copy
        return done(*[{"op": "slide_duplicate", "slide": k} for k in ns[:1]])
    mm = re.search(
        r"\b(?:add|insert|create|make|put)\b.*?\b(?:a |an |one |another )?(?:new )?(\w+ )?slide\b(?:\s+(?:about|on|with|showing|for|explaining|covering|listing|comparing)\s+(.+?))?"
        r"(?=\s+(?:after|before|at the end|at the start|at the beginning|as slide)\b|$)",
        c,
    )
    if mm and not re.search(r"\bnotes?\b|\bbullet\b", c):
        about = (mm.group(2) or "").strip(" .")
        kind_word = (mm.group(1) or "").strip()
        lay = next((ln for ln, rx in LAYOUT_WORDS if re.search(rx, c)), None)
        pos = "end"
        mp = re.search(r"\b(after|before)\s+slide\s+(\d+)|\bas slide\s+(\d+)|\bat the (start|beginning)\b", c)
        if mp:
            pos = (
                int(mp.group(2))
                if mp.group(1) == "after"
                else int(mp.group(2)) - 1
                if mp.group(1) == "before"
                else int(mp.group(3)) - 1
                if mp.group(3)
                else "start"
            )
        out["focus"] = {
            "ns": [m["count"] + 1 if pos == "end" else 1 if pos == "start" else int(pos) + 1],
            "fresh": True,
        }  # "...and make it shorter": the new slide
        if lay == "closing" and not about:
            return done({"op": "slide_add", "after": pos, "layout": "closing", "title": "Thank you", "subtitle": "Questions?"})
        if lay == "agenda" and not about:
            items = [s["title"] for s in m["slides"][1:] if s["title"] and not re.search(r"thank|agenda|question", s["title"], re.I)][:7]
            return done({"op": "slide_add", "after": pos if pos != "end" else 1, "layout": "agenda", "items": items})
        if q and lay == "quote":
            return done({"op": "slide_add", "after": pos, "layout": "quote", "text": q[0]})
        if about or kind_word:
            return done({"op": "slide_add", "after": pos, "about": about or f"a {kind_word} slide for this deck", **({"layout": lay} if lay else {})})
    # titles and text
    mm = re.search(r"\b(?:change|rename|set|make)\b.*?\btitle of (?:slide\s+)?(\d+)\b|\bslide\s+(\d+)(?:'s)? title\b", c)
    if mm and q:
        return done({"op": "title_set", "slide": int(mm.group(1) or mm.group(2)), "text": q[-1]})
    mm = re.search(
        r"\b(?:replace|change|rename)\s+(.+?)\s+(?:with|to|into|by)\s+(.+?)(?:\s+(?:everywhere|throughout|in the deck|on all slides))?$", raw, re.I
    )
    if mm and not re.search(r"\b(?:font|size|colou?r|title|background|theme|text)\b", mm.group(1).lower()):
        a, b = (q[0], q[1]) if len(q) >= 2 else (mm.group(1).strip(" '\"."), mm.group(2).strip(" '\"."))
        deck_text = " ".join(t["text"] for s in m["slides"] for t in s["texts"]).lower()
        if a and b and (len(q) >= 2 or a.lower() in deck_text):
            return done({"op": "replace", "find": a, "with": b})
    mm = re.search(r"\btranslate\b.*?\b(?:in|into|to)\s+(" + P.LANGS + r")\b|\b(?:in|into)\s+(" + P.LANGS + r")\b", c)
    if mm and re.search(r"\btranslat|convert|write\b", c):
        return done({"op": "translate", "target": target or {"kind": "all"}, "language": mm.group(1) or mm.group(2)})
    mm = re.search(
        r"\b(shorter|simpler|punchier|clearer|more formal|more persuasive|fewer words|less text|more concise|catchier|friendlier|stronger)\b|"
        r"\b(rewrite|rephrase|reword|shorten|simplify|tighten|improve)\b",
        c,
    )
    if mm and (target or re.search(r"\b(?:deck|presentation|slides|text)\b", c)) and not re.search(r"\bnotes?\b|\bfont|size\b|\bbigger|smaller\b", c):
        how = mm.group(1) or mm.group(2)
        ins = {
            "fewer words": "fewer words: short bullet phrases",
            "less text": "less text: short bullet phrases",
            "shorten": "shorter",
            "rewrite": "rewrite it better",
            "rephrase": "rephrase in fresh words",
            "reword": "rephrase in fresh words",
            "simplify": "simpler words",
            "tighten": "tighter, fewer words",
            "improve": "clearer and stronger",
        }.get(how, f"make it {how}")
        part = "titles" if re.search(r"\btitles?\b|\bheadings?\b", c) else "all" if re.search(r"\beverything\b|\ball (?:the )?text\b", c) else "body"
        return done({"op": "rewrite", "target": target or {"kind": "all"}, "instruction": ins, "part": part})
    if re.search(r"\b(?:speaker )?notes?\b", c) and re.search(r"\b(?:add|write|create|give|make|generate|put)\b", c):
        if q:
            return done({"op": "notes", "target": target or {"kind": "all"}, "text": q[0]})
        if re.search(r"\b(?:have|has) (?:none|no notes?)\b|\bwithout (?:any )?notes\b|\bmissing\b|\bdon'?t have\b|\bno notes\b", c):
            empty = [s["n"] for s in m["slides"] if not s["notes"]]
            if not empty:
                out["ask"] = "Every slide already has speaker notes. Say 'rewrite the notes on slide N' or 'write new notes for every slide'."
                return out
            return done(
                {"op": "notes", "target": {"kind": "slides", "ns": empty} if len(empty) > 1 else {"kind": "slide", "n": empty[0]}, "write": True}
            )
        return done({"op": "notes", "target": target or {"kind": "all"}, "write": True})
    if re.search(r"\b(?:overflow\w*|doesn'?t fit|does not fit|fit (?:on|in|the box)|too much text|spilling|cut off|shrink)\b", c):
        return done({"op": "shrink", "target": target or {"kind": "all"}})
    mm = re.search(r"\badd\b.*?\bbullet\b.*?(?:slide\s+(\d+))?\s*[:,-]\s*(.+)$", raw, re.I)
    if mm or (q and re.search(r"\badd\b.*\bbullet", c)):
        n = int(mm.group(1)) if mm and mm.group(1) else (ns[0] if ns else None)
        text = q[0] if q else (mm.group(2).strip(" .") if mm else "")
        if n and text:
            return done({"op": "bullet_add", "slide": n, "text": text})
    mm = re.search(r"\b(?:remove|delete|drop)\b.*?\b(?:the )?(first|second|third|fourth|fifth|last|\d+(?:st|nd|rd|th)) bullet\b", c)
    if mm and ns:
        w = mm.group(1)
        return done({"op": "bullet_delete", "slide": ns[0], "n": "last" if w == "last" else P.ORD.get(w) or int(re.sub(r"\D", "", w))})
    mm = re.search(r"\bbackground\b", c)
    if mm:
        cols = list(P.COLOUR_RX.finditer(c))
        if cols:
            ops = [{"op": "background", "target": target or {"kind": "all"}, "color": cols[0].group(1)}]
            more = re.search(r"\band\b.*?\b(titles?|headings?|text|body|bullets?)\b", c[cols[0].end() :])
            if more and len(cols) > 1:
                part = "titles" if more.group(1).startswith(("title", "heading")) else "body"
                ops.append({"op": "style", "target": target or {"kind": "all"}, "part": part, "set": {"color": cols[1].group(1)}})
            return done(*ops)
    s = P.style_set(c)
    s = {k: v for k, v in s.items() if k in ("font", "size", "color", "bold", "italic")}
    if s:
        part = (
            "titles"
            if re.search(r"\btitles?\b|\bheadings?\b", c)
            else "table"
            if re.search(r"\btables?\b", c)
            else "body"
            if re.search(r"\b(?:body|bullets?|text|content)\b", c) and not re.search(r"\ball text\b|\bevery ?thing\b", c)
            else "all"
        )
        if part == "table" and not ns:
            ks = [s_["n"] for s_ in m["slides"] if s_["table"]]
            target = {"kind": "slides", "ns": ks} if len(ks) > 1 else {"kind": "slide", "n": ks[0]} if ks else None
        out["focus"] = {"ns": ns, "part": part} if ns else focus
        return done({"op": "style", "target": target or {"kind": "all"}, "part": part, "set": s})
    return out


GROWS = {"rewrite", "translate", "bullet_add", "title_set", "slide_add", "replace"}  # edits that can make text outgrow its box


class DeckChat(DocChat):
    @classmethod
    def start(cls, path, chats_dir=None, **kw):
        src = Path(path)
        if src.suffix.lower() != ".pptx":
            raise ValueError("DeckChat edits PowerPoint decks (.pptx)")
        cid = f"deck_{re.sub(r'[^a-z0-9]+', '', src.stem.lower())[:16] or 'deck'}_{time.strftime('%H%M%S')}"
        folder = Path(chats_dir or CHATS) / cid
        folder.mkdir(parents=True, exist_ok=True)
        v0 = folder / "v0.pptx"
        shutil.copy(src, v0)
        state = {
            "id": cid,
            "kind": "pptx",
            "base": str(src),
            "folder": str(folder),
            "cur": 0,
            "turns": [],
            "focus": None,
            "pending": None,
            "versions": [
                {
                    "v": 0,
                    "parent": None,
                    "file": str(v0),
                    "said": None,
                    "done": ["the original file"],
                    "failed": [],
                    "checks": [],
                    "pages": None,
                    "pdf": None,
                }
            ],
        }
        c = cls(state, **kw)
        if c.render:
            c._render(0)
        c.save()
        return c

    def map(self, v=None):
        v = self.state["cur"] if v is None else v
        if v not in self._maps:
            self._maps[v] = PO.deck_map(self.path(v))
        return self._maps[v]

    def _render(self, v, slides=None):
        ver = self.state["versions"][v]
        extra = {"measure": True, **({"measure_slides": slides} if slides else {})}
        r = RN.to_pdf(ver["file"], pdf=str(Path(ver["file"]).with_suffix(".pdf")), server=self.server, **extra)
        ver["pages"], ver["pdf"], ver["render_ok"] = r.get("slides"), r.get("pdf") if r.get("ok") else None, bool(r.get("ok"))
        over = [
            x for x in r.get("shapes") or [] if x["name"] not in PO.DECO and (x["text_height"] > x["height"] + 3 or x["text_width"] > x["width"] + 3)
        ]
        ver["overflow"] = sorted({x["slide"] for x in over})
        ver["overflow_fit"] = {
            f"{x['slide']}:{x['name']}": [round(x["text_height"] / max(1.0, x["height"]), 3), round(x["text_width"] / max(1.0, x["width"]), 3)]
            for x in over
        }
        return r

    def say(self, message, files=()):
        f = self.state.get("focus")
        if isinstance(f, dict) and f.get("fresh"):  # a new message counts slides as they are now
            self.state["focus"] = {k: v for k, v in f.items() if k != "fresh"}
        return super().say(message, files)

    def _llm_answer(self, q):
        if self.planner is None:
            return "Ask me how many slides, what is on slide N, the slide titles, which slides have notes or charts."
        self._turn["llm"] = True
        m = self.map()
        ns, _ = slides_in(P.normalize(q), m)
        ss = [m["slides"][k - 1] for k in ns] if ns else m["slides"]
        body = "\n".join(
            f"SLIDE {s['n']}: " + " | ".join(t["text"].replace("\n", " / ") for t in s["texts"]) + (f"\n  NOTES: {s['notes']}" if s["notes"] else "")
            for s in ss
        )[:9000]
        try:
            r = self.planner._call(
                "fast",
                [
                    {
                        "role": "system",
                        "content": "You answer a client's question about their slide deck in 1-4 short sentences, only from "
                        "the slides given. If the slides do not say, say so.",
                    },
                    {"role": "user", "content": f"DECK: {m['title']} ({m['count']} slides)\n{body}\n\nQUESTION: {q}"},
                ],
            )
            return re.sub(r"<think>.*?</think>", "", r.text or "", flags=re.S).strip()[:700] or "I could not find that in the deck."
        except Exception as e:  # noqa: BLE001
            return f"I could not answer that just now ({type(e).__name__})."

    def _clause(self, c, raw, focus, carry=False):
        kind, payload = super()._meta(c, raw)
        if kind:
            return kind, payload
        low = c.lower()
        if QUESTION.search(low):
            return "question", self.answer(low)
        r = parse_deck(c, self.map(), focus, raw=raw, carry=carry)
        sh = [o for o in r.get("ops") or [] if o["op"] == "shrink" and (o.get("target") or {}).get("kind") == "all"]
        if sh and self.render and not re.search(r"\b(?:every|all) (?:the )?slides?\b|\beverywhere\b", low):
            self._render(self.state["cur"])  # measure every slide now
            over = self.version.get("overflow") or []
            if not over:
                return "question", "All text already fits its boxes (PowerPoint measured every slide); nothing to shrink."
            sh[0]["target"] = {"kind": "slides", "ns": over} if len(over) > 1 else {"kind": "slide", "n": over[0]}
            sh[0]["fit"] = self.version.get("overflow_fit") or {}
        if r.get("ops"):
            return "change", r
        if r.get("ask"):
            return "ask", {"question": r["ask"], "template": None}
        return "unknown", c

    def _llm(self, clauses, message, done_ops, focus):
        if self.planner is None:
            return {"ask": "I could not read: " + "; ".join(f"'{c}'" for c in clauses) + ". Try e.g. 'delete slide 3' or 'make the titles blue'."}
        self._turn["llm"] = True
        m = self.map()
        outline = "\n".join(
            f"{s['n']}. [{s['layout']}] {s['title'][:70]}" + (" (chart)" if s["chart"] else "") + (" (table)" if s["table"] else "")
            for s in m["slides"]
        )
        user = f"DECK: {m['title']}\nSLIDES:\n{outline}\nLAST SLIDES TALKED ABOUT: {json.dumps(focus) if focus else 'none'}\nWHOLE MESSAGE: {message}\nREQUESTS: {clauses}"
        try:
            r = self.planner._call("docs", [{"role": "system", "content": DECK_SYSTEM}, {"role": "user", "content": user}])
            d = parse_json(r.text) or {}
        except Exception as e:  # noqa: BLE001
            return {"ask": f"I could not work that out ({type(e).__name__})."}
        ops = [o for o in d.get("ops") or [] if isinstance(o, dict) and o.get("op") in PO.OPS]
        return {"ops": ops, "answer": d.get("answer"), "ask": d.get("ask") if not ops else None}

    def _change(self, ops, said, turn):
        v0 = self.state["cur"]
        prs = Presentation(str(self.path(v0)))
        ctx = {"planner": self.planner, "th": PO.deck_theme(prs), "info": {}}
        done, failed, checks, touched = [], [], [], set()
        ids0 = [el.get("id") for el in prs.slides._sldIdLst]  # the slides as numbered when the client wrote the message
        for op in ops:
            before = PO.deck_map(prs)
            ctx["info"] = {}
            try:
                op = self._renumber(op, ids0, [el.get("id") for el in prs.slides._sldIdLst])
                d = PO.apply(prs, op, ctx)
            except PO.OpError as e:
                failed.append(f"{PO.describe(op)}: {e}")
                continue
            except Exception as e:  # noqa: BLE001
                failed.append(f"{PO.describe(op)}: {type(e).__name__}: {str(e)[:120]}")
                continue
            after = PO.deck_map(prs)
            ok, what = PO.check(op, before, after, ctx.get("info"), prs)
            done.append(d)
            checks.append({"op": PO.describe(op), "ok": bool(ok), "what": what})
            turn["ops"].append(op)
            for k in PO.find(after, op.get("target")) if op.get("target") else []:
                touched.add(k)
            for key in ("slide", "a", "b"):
                if op.get(key):
                    touched.add(int(op[key]))
            if (ctx.get("info") or {}).get("added"):
                touched.add(ctx["info"]["added"]["at"])
            if op.get("op") in ("data_slide", "data_refresh") and op.get("link"):  # the slide that holds the linked chart or table
                k_, _ = PO.linked(prs, op["link"])
                if k_:
                    touched.add(k_)
        if not done:
            return "Nothing changed: " + "; ".join(failed or ["that is already how it is"]) + "."
        n = len(self.state["versions"])
        dst = self.folder / f"v{n}.pptx"
        prs.save(str(dst))
        self.state["versions"].append(
            {"v": n, "parent": v0, "file": str(dst), "said": said, "done": done, "failed": failed, "checks": checks, "pages": None, "pdf": None}
        )
        self.state["cur"] = n
        note = ""
        if self.render:
            r = self._render(n, sorted(touched) if touched and len(touched) < 8 else None)
            over = self.state["versions"][n].get("overflow") or []
            grew = [
                o
                for o in turn["ops"]
                if o.get("op") in GROWS
                or (o.get("op") == "style" and ((o.get("set") or {}).get("size") is not None or (o.get("set") or {}).get("font")))
            ]
            fitting = any(o.get("op") == "shrink" and o.get("fit") for o in turn["ops"])
            rounds = 0
            # text still too big for its box after this edit: shrink by PowerPoint's own measure and measure again, a few
            # rounds (a two-line number at 85% is still two lines: line breaks come in steps)
            while r.get("ok") and over and (grew or fitting) and rounds < 3:
                rounds += 1
                prs2 = Presentation(str(dst))
                try:
                    d2 = PO.apply(
                        prs2,
                        {"op": "shrink", "target": {"kind": "slides", "ns": over}, "fit": self.state["versions"][n].get("overflow_fit") or {}},
                        ctx,
                    )
                except PO.OpError:
                    break
                prs2.save(str(dst))
                self._maps.pop(n, None)
                if rounds == 1 or not fitting:
                    done.append(d2.replace("text made smaller", "so it fits, made smaller"))
                r = self._render(n, over)
                over = self.state["versions"][n].get("overflow") or []
            if rounds > 1:
                done.append(f"({rounds + (1 if fitting else 0)} measured rounds)")
            if r.get("ok"):
                over = self.state["versions"][n].get("overflow") or []
                checks.append({"op": "text fits its boxes", "ok": not over, "what": f"overflow on slide(s) {over}" if over else "no overflow"})
                note = f" Now {r.get('slides')} slide(s)."
        bad = [c for c in checks if not c["ok"]]
        msg = f"v{n}: " + "; ".join(done) + "."
        if failed:
            msg += " Couldn't: " + "; ".join(failed) + "."
        msg += (
            f" Checked: {len(checks) - len(bad)}/{len(checks)} OK"
            + (" (" + "; ".join(f"{c['op']}: {c['what']}" for c in bad[:3]) + ")" if bad else "")
            + "."
        )
        return msg + note

    @staticmethod
    def _renumber(op, ids0, ids):
        """An operation's slide numbers (as the client saw them) turned into where those slides are now, after the earlier
        operations of the same message added, moved or deleted slides."""
        if op.get("_fresh"):  # numbered after the edits before it in the message already
            return {k: v for k, v in op.items() if k != "_fresh"}

        def now(n):
            try:
                n = int(n)
            except (TypeError, ValueError):
                return n
            if not 1 <= n <= len(ids0):
                return n
            sid = ids0[n - 1]
            if sid not in ids:
                raise PO.OpError(f"slide {n} was removed earlier in this message")
            return ids.index(sid) + 1

        op = json.loads(json.dumps(op))
        t = op.get("target")
        if isinstance(t, dict):
            if t.get("n") is not None:
                t["n"] = now(t["n"])
            if t.get("ns"):
                t["ns"] = [now(x) for x in t["ns"]]
        for k in ("slide", "a", "b"):
            if op.get(k) is not None:
                op[k] = now(op[k])
        for k in ("after", "before"):
            if isinstance(op.get(k), int):
                op[k] = now(op[k])
        return op

    def answer(self, q):
        parts = [
            p.strip()
            for p in re.split(r"\?\s*|\s+and\s+(?=(?:what|which|how|is|are|does|do|who|where|when)\b)|,\s*(?=(?:what|which|how)\b)", q)
            if p and p.strip()
        ]
        outs = []
        for p in parts or [q]:
            a = self._answer_one(p)
            if a and a not in outs:
                outs.append(a)
        return " ".join(outs)

    def _answer_one(self, q):
        m = self.map()
        if re.search(r"\bhow many slides\b", q):
            return f"{m['count']} slides."
        mm = re.search(r"\b(?:what(?:'s| is)? on|show me|what does) slide (\d+)", q)
        if mm and 1 <= int(mm.group(1)) <= m["count"]:
            s = m["slides"][int(mm.group(1)) - 1]
            return (
                f"Slide {s['n']} ('{s['title']}'): "
                + " | ".join(t["text"].replace("\n", " / ") for t in s["texts"][1:4])[:400]
                + (" [chart]" if s["chart"] else "")
                + (" [table]" if s["table"] else "")
            )
        if re.search(r"\b(?:titles|list (?:the )?slides|slide titles|outline|structure|what are the slides)\b", q):
            return "; ".join(f"{s['n']}. {s['title']}" for s in m["slides"])
        if re.search(r"\bnotes?\b", q):
            have = [s["n"] for s in m["slides"] if s["notes"]]
            return f"Slides with speaker notes: {have or 'none'}."
        if re.search(r"\bcharts?\b", q):
            return f"Slides with charts: {[s['n'] for s in m['slides'] if s['chart']] or 'none'}."
        if re.search(r"\boverflow|fit\b", q):
            over = self.version.get("overflow") or []
            return f"Text runs out of its box on slide(s) {over}." if over else "All text fits its boxes."
        return self._llm_answer(q)

    def compare(self, text=""):
        nums = [int(x) for x in re.findall(r"\bv(?:ersion)?\s*(\d+)", text)]
        a = nums[0] if nums else 0
        b = nums[1] if len(nums) > 1 else self.state["cur"]
        ma, mb = self.map(a), self.map(b)
        ta, tb = [s["title"] for s in ma["slides"]], [s["title"] for s in mb["slides"]]
        return (
            f"v{a} -> v{b}: slides {ma['count']} -> {mb['count']}; words {ma['words']} -> {mb['words']}"
            + (f"; new: {', '.join(t[:30] for t in tb if t not in ta)[:200]}" if any(t not in ta for t in tb) else "")
            + (f"; gone: {', '.join(t[:30] for t in ta if t not in tb)[:200]}" if any(t not in tb for t in ta) else "")
            + "."
        )

    def export(self):
        v = self.version
        name = Path(self.state["base"]).stem
        out = self.folder / f"{name}_v{v['v']}.pptx"
        shutil.copy(v["file"], out)
        msg = f"Saved: {out}"
        if v.get("pdf"):
            pdf = self.folder / f"{name}_v{v['v']}.pdf"
            shutil.copy(v["pdf"], pdf)
            msg += f" and {pdf}"
        return msg + "."
