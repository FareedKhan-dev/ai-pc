"""The writer: the cheap model plans and writes the document's content as plan blocks (docplan.py); code does the rest.

  inputs = read_inputs(files)                   what the client gave: text, tables, images (truncated for the prompt)
  plan, info = design(request, inputs, planner) one outline call, then every section written in parallel (long
                                                documents), or one call for a short one (letter, CV, invoice...)

Writing rules in every prompt: specific and concrete, short paragraphs that lead with the point, the client's own
terms and data used exactly; no invented statistics, studies, authors or links (approximate figures are labelled as
such); tables and charts only where numbers help.
"""
import datetime as dt
import json
import re
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from ai_pc.core.util import parse_json

SHORT = {"letter", "application", "cv", "invoice", "quotation", "notice", "memo", "certificate", "agenda"}
SHORT_RX = re.compile(r"\b(?:letter|application|leave|cover letter|resume|résumé|cv|curriculum vitae|invoice|bill\b|quotation|quote for|notice|memo|"
                      r"memorandum|certificate|agenda)\b", re.I)
WORDS_PER_PAGE = 420

BLOCK_GUIDE = """Block types (inline **bold** and *italic* allowed in text):
  {"type": "paragraph", "text": "...", "style": "lead|note"(optional)}
  {"type": "bullets", "items": ["...", {"text": "...", "items": ["sub", "sub"]}], "numbered": false}
  {"type": "heading", "text": "...", "level": 2}          (sub-headings inside a section only, level 2 or 3)
  {"type": "table", "columns": ["...", ...], "rows": [[...], ...], "caption": "...", "total": true|false}
  {"type": "chart", "chart": "column|bar|line|pie|doughnut|area|stacked", "title": "...", "categories": [...],
   "series": [{"name": "...", "values": [numbers]}], "caption": "..."}
  {"type": "callout", "title": "...", "text": "..."}       (one key message, at most one per section)
  {"type": "quote", "text": "...", "by": "..."}            {"type": "kv", "items": [["Label", "Value"], ...]}
  {"type": "references", "items": ["...", ...]}            {"type": "entry", "title", "org", "place", "dates", "bullets": [...]}"""

RULES = """Writing rules:
- Specific and concrete; lead each paragraph with its point; 50-120 words per paragraph; vary sentence length.
- Use the client's own terms, names and data exactly. Never invent precise statistics, surveys, studies, authors, quotes
  or links. Where no data was given, use clearly approximate or illustrative figures and say so ("about", "roughly",
  "illustrative"). References only to real, well-known organisations or publications; if unsure, leave them out.
- No filler or cliches ("in today's fast-paced world", "delve", "tapestry", "it is important to note", "in conclusion").
- Tables and charts only where numbers help the reader; every table and chart gets a short caption. Totals are added by
  code: do not add a total row yourself (set "total": true instead).
- Write for the audience and in the language asked (Urdu in Urdu script only if the client asked for Urdu)."""

OUTLINE_SYSTEM = """You are a senior editor planning a document for a client. Reply with ONE JSON object:
{"doctype": "report|proposal|assignment|essay|thesis|research|manual|business_plan|minutes|contract|newsletter|other",
 "title": "...", "subtitle": "...", "language": "en|ur", "audience": "...",
 "theme": "corporate|academic|modern|elegant|minimal|warm", "accent": "<colour word or #hex ONLY if the client named a colour, else null>",
 "target_pages": <number>, "page": {"size": "A4|Letter", "orientation": "portrait|landscape"},
 "meta": {"author": "...", "org": "...", "date": "...", "course": "...", "instructor": "...", "roll_no": "...", "university": "...",
          "department": "...", "class": "...", "client": "..."},
 "sections": [{"heading": "...", "goal": "<what this section must say, one sentence>", "words": <number>,
               "elements": ["paragraphs", "bullets", "numbered steps", "table: <what it shows>", "chart: <what it shows>", "callout", "quote"]}],
 "data_from_client": ["<every number, name or fact the client gave, to be used exactly>"]}
Rules:
- meta: ONLY what the client gave (leave out unknown keys); the date defaults to today.
- theme: academic for university work (assignments, essays, theses), corporate for business, modern for tech/startups,
  elegant for formal/creative, minimal for plain internal documents, warm for food/travel/NGO topics.
- target_pages: what the client asked; else a sensible length (2-4 pages for a short report, 6-10 for a full one).
- sections: 4-9 for a few pages; words per section add up to about target_pages x %d (less where tables/charts are);
  headings short (2-6 words) and specific; an executive summary / introduction first and a conclusion or
  recommendations last when the document type expects them; a References section only for academic work.
- elements: tables or charts only where data helps; at most one chart per two pages.""" % WORDS_PER_PAGE

SECTION_SYSTEM = """You write ONE section of a document for a client (the whole outline is given so sections do not repeat
each other). Reply with ONE JSON object: {"blocks": [...]}. Do NOT repeat the section's own heading.
""" + BLOCK_GUIDE + "\n" + RULES

SHORT_SYSTEM = """You write a complete short document for a client. Reply with ONE JSON object:
{"doctype": "letter|application|cv|invoice|quotation|notice|memo|certificate|agenda",
 "title": "...", "subtitle": "...", "language": "en|ur", "theme": "corporate|academic|modern|elegant|minimal|warm",
 "accent": "<colour word or #hex ONLY if the client named a colour, else null>",
 "meta": {...}, "blocks": [...]}
By type (use these exact block types):
- letter / application: blocks = address {lines: [recipient lines]}, date {text}, subject {text}, salutation {text},
  paragraph..., closing {text: "Yours sincerely,|Yours faithfully,|Yours obediently,", name, title}. Formal, polite, concise;
  Pakistani conventions when the request suggests Pakistan ("Respected Sir/Madam", "Yours obediently" for students).
- cv: title = the person's name, subtitle = their headline; meta = {location, phone, email, linkedin, website}; blocks =
  heading {text, level: 1} sections (Profile, Experience, Education, Skills, Projects, Certifications, Languages) with
  paragraph, entry {title, org, place, dates, bullets: [achievements with numbers]}, kv {items: [[label, value]]}.
  One page unless the experience needs two. Use only the client's real details; leave out what they did not give.
- invoice / quotation: meta = {org, address, phone, email, ntn, number, date, due, bill_to}; blocks = invoice_items
  {items: [{desc, qty, price}], currency: "Rs.", tax_rate: "<e.g. 18%, only if asked or standard for the client>",
  discount}, then a paragraph with payment terms. Code does all the arithmetic: give only quantities and unit prices.
- notice / memo / agenda: title, meta {org, date}, then paragraph / bullets / table blocks; short and clear.
- certificate: title (e.g. "Certificate of Achievement"), subtitle ("This is to certify that"), meta {org, recipient,
  date}, blocks = paragraph {text: what it is for, align: "center"}, signature {name, title, align: "center"}.
""" + BLOCK_GUIDE + "\n" + RULES


def _ask(planner, system, user, tier="docs", log=print):
    msgs = [{"role": "system", "content": system}, {"role": "user", "content": user}]
    r = planner._call(tier, msgs)
    d = parse_json(r.text)
    if d is None:
        r = planner._call(tier, msgs + [{"role": "assistant", "content": (r.text or "")[:3000]},
                                        {"role": "user", "content": "Reply with the single JSON object only."}])
        d = parse_json(r.text)
    return d if isinstance(d, dict) else {}


def read_inputs(files, limit=12000):
    """What the client gave, for the prompt: {"text", "tables": [{"name", "columns", "rows"}], "images": [paths]}."""
    text, tables, images = [], [], []
    for f in files or []:
        p = Path(f)
        ext = p.suffix.lower()
        try:
            if ext in (".txt", ".md", ".csv", ".json"):
                text.append(f"--- {p.name}\n" + p.read_text(encoding="utf-8", errors="replace")[:limit])
            elif ext == ".docx":
                import docx
                d = docx.Document(str(p))
                text.append(f"--- {p.name}\n" + "\n".join(x.text for x in d.paragraphs if x.text.strip())[:limit])
                for i, t in enumerate(d.tables[:5]):
                    rows = [[c.text.strip() for c in r.cells] for r in t.rows[:40]]
                    if rows:
                        tables.append({"name": f"{p.name} table {i + 1}", "columns": rows[0], "rows": rows[1:]})
            elif ext in (".xlsx", ".xlsm"):
                import openpyxl
                wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
                for ws in wb.worksheets[:4]:
                    rows = [[("" if v is None else v) for v in r] for r in ws.iter_rows(max_row=60, values_only=True)]
                    rows = [r for r in rows if any(str(v).strip() for v in r)]
                    if rows:
                        tables.append({"name": f"{p.name} / {ws.title}", "columns": [str(x) for x in rows[0]], "rows": rows[1:]})
            elif ext == ".pdf":
                import pypdfium2 as pdfium
                doc = pdfium.PdfDocument(str(p))
                text.append(f"--- {p.name}\n" + "\n".join(doc[i].get_textpage().get_text_range() for i in range(min(len(doc), 15)))[:limit])
                doc.close()
            elif ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
                images.append(str(p))
        except Exception as e:  # noqa: BLE001
            text.append(f"--- {p.name}: could not be read ({type(e).__name__})")
    joined = "\n\n".join(text)
    return {"text": joined[:limit], "tables": tables, "images": images}


def _inputs_text(inputs, limit=9000):
    parts = []
    if inputs.get("text"):
        parts.append("CLIENT'S FILES (text):\n" + inputs["text"][:limit])
    for t in inputs.get("tables") or []:
        rows = "\n".join(" | ".join(str(c) for c in r) for r in t["rows"][:40])
        parts.append(f"CLIENT'S TABLE {t['name']}:\n" + " | ".join(str(c) for c in t["columns"]) + "\n" + rows)
    if inputs.get("images"):
        parts.append("CLIENT'S IMAGES (can be placed with an image block {\"type\": \"image\", \"path\", \"caption\"}): "
                     + ", ".join(inputs["images"]))
    return "\n\n".join(parts)[:limit + 3000]


def words_per_page(th):
    """Words of running text a page of this theme holds, less a share for headings, tables and charts (calibrated on
    a 5-page corporate report: 1,953 words on 5 pages with a table and a chart)."""
    return round(560 * (11.0 / th["size"]) ** 2 * (1.15 / th["line"]) * 0.75)


def pages_asked(request):
    r = str(request).lower()
    m = re.search(r"\b(\d{1,3})\s*(?:-\s*)?(?:pages?|pgs?|sides?)\b", r)
    if m:
        return int(m.group(1))
    words = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10, "twelve": 12, "fifteen": 15, "twenty": 20}
    m = re.search(r"\b(" + "|".join(words) + r")\s*(?:-\s*)?pages?\b", r)
    if m:
        return words[m.group(1)]
    m = re.search(r"\b(\d{3,5})\s*words?\b", r)
    if m:
        return max(1, round(int(m.group(1)) / WORDS_PER_PAGE))
    return None


def design(request, inputs, planner, log=print, theme=None, pages=None):
    """(plan, info) for a request: the content written by the model, ready for docplan.resolve."""
    t0 = time.perf_counter()
    today = dt.date.today().strftime("%d %B %Y")
    given = _inputs_text(inputs)
    want_pages = pages or pages_asked(request)
    info = {"calls": 0}
    if SHORT_RX.search(str(request)) and not re.search(r"\breport\b|\bproposal\b|\bassignment\b|\bessay\b", str(request), re.I):
        d = _ask(planner, SHORT_SYSTEM, f"TODAY: {today}\nREQUEST: {request}\n\n{given}", log=log)
        info["calls"] += 1
        if d.get("blocks"):
            plan = {**d, "blocks": d.get("blocks") or [], "target_pages": want_pages}
            if theme:
                plan["theme"] = theme
            info["seconds"] = round(time.perf_counter() - t0, 1)
            info["mode"] = "short"
            return plan, info
    o = _ask(planner, OUTLINE_SYSTEM, f"TODAY: {today}\nREQUEST: {request}\n\n{given}", log=log)
    info["calls"] += 1
    sections = [s for s in o.get("sections") or [] if isinstance(s, dict) and s.get("heading")]
    if not sections:
        raise ValueError("the outline came back empty")
    target = want_pages or int(o.get("target_pages") or 3)
    from ai_pc.office import themes as TH
    th = TH.get(theme or o.get("theme"), None, o.get("doctype"))
    wpp = words_per_page(th)
    budget = max(150, round(target * wpp - 120))
    total = sum(int(s.get("words") or 0) for s in sections) or budget
    for s in sections:  # the words asked for, shared out as the outline weighted them
        s["words"] = max(80, round(budget * int(s.get("words") or budget / len(sections)) / total))
    outline_txt = "\n".join(f"{i + 1}. {s['heading']}: {s.get('goal', '')} (~{s['words']} words; {', '.join(s.get('elements') or [])})"
                            for i, s in enumerate(sections))
    head = (f"TODAY: {today}\nREQUEST: {request}\nDOCUMENT: {o.get('doctype')} '{o.get('title')}' for {o.get('audience') or 'the client'}, "
            f"language {o.get('language') or 'en'}\nOUTLINE:\n{outline_txt}\nDATA FROM THE CLIENT: {json.dumps(o.get('data_from_client') or [], ensure_ascii=False)[:1500]}\n\n{given}")

    def write(i):
        s = sections[i]
        d = _ask(planner, SECTION_SYSTEM, head + f"\n\nWRITE SECTION {i + 1}: '{s['heading']}' — {s.get('goal', '')}\n"
                 f"About {s['words']} words. Elements wanted: {', '.join(s.get('elements') or ['paragraphs'])}.", log=log)
        return [b for b in d.get("blocks") or [] if isinstance(b, dict)]
    with ThreadPoolExecutor(max_workers=6) as ex:
        bodies = list(ex.map(write, range(len(sections))))
    info["calls"] += len(sections)
    blocks = []
    for s, body in zip(sections, bodies):
        blocks.append({"type": "heading", "text": s["heading"], "level": 1})
        for b in body:
            if b.get("type") == "heading" and str(b.get("text", "")).strip().lower() == s["heading"].strip().lower():
                continue  # the model repeated the section heading
            if b.get("type") == "heading":
                b["level"] = max(2, min(3, int(b.get("level") or 2)))
            blocks.append(b)
    plan = {"doctype": o.get("doctype") or "report", "title": o.get("title"), "subtitle": o.get("subtitle"), "language": o.get("language"),
            "theme": theme or o.get("theme"), "accent": o.get("accent"), "target_pages": target, "page": o.get("page") or {},
            "meta": {k: v for k, v in (o.get("meta") or {}).items() if v}, "blocks": blocks}
    info.update(mode="long", sections=len(sections), seconds=round(time.perf_counter() - t0, 1), outline=sections, words_per_page=wpp,
                data=[str(x) for x in o.get("data_from_client") or []][:30])
    return plan, info


EXPAND_SYSTEM = """You lengthen one section of a document that came out shorter than the client asked. Keep everything that
is there (same facts, same structure) and add substance: examples, explanation, implications, a table or list where it
helps. Reply with ONE JSON object: {"blocks": [...]} — the WHOLE section again, without its heading.
""" + BLOCK_GUIDE + "\n" + RULES

SHORTEN_SYSTEM = """You shorten one section of a document that came out longer than the client asked. Keep the key points,
data, tables and charts; cut repetition and padding. Reply with ONE JSON object: {"blocks": [...]} — the WHOLE section
again, without its heading.
""" + BLOCK_GUIDE + "\n" + RULES


def resize(plan, request, planner, factor, log=print):
    """The plan with its level-1 sections rewritten longer (factor > 1) or shorter, in parallel."""
    blocks = plan["blocks"]
    starts = [i for i, b in enumerate(blocks) if b.get("type") == "heading" and int(b.get("level") or 1) == 1]
    if not starts:
        return plan, 0
    spans = [(a, starts[k + 1] if k + 1 < len(starts) else len(blocks)) for k, a in enumerate(starts)]
    system = EXPAND_SYSTEM if factor > 1 else SHORTEN_SYSTEM

    def words(bs):
        return sum(len(json.dumps(b, ensure_ascii=False).split()) for b in bs)

    def redo(span):
        a, z = span
        body = blocks[a + 1:z]
        want = max(80, round(words(body) * factor))
        d = _ask(planner, system, f"REQUEST: {request}\nSECTION '{blocks[a].get('text')}' NOW:\n{json.dumps(body, ensure_ascii=False)[:9000]}\n\n"
                 f"Make it about {want} words.", log=log)
        new = [b for b in d.get("blocks") or [] if isinstance(b, dict) and not (b.get("type") == "heading" and int(b.get("level") or 1) == 1)]
        return new or body
    with ThreadPoolExecutor(max_workers=6) as ex:
        bodies = list(ex.map(redo, spans))
    out = []
    for (a, z), body in zip(spans, bodies):
        out.append(blocks[a])
        out.extend(body)
    return {**plan, "blocks": blocks[:starts[0]] + out}, len(spans)


DECK_RX = re.compile(r"\b(?:presentation|slides?|slide ?deck|deck|ppt|pptx|powerpoint|pitch)\b", re.I)

DECK_SYSTEM = """You are a presentation designer writing a slide deck for a client. Reply with ONE JSON object:
{"title": "...", "subtitle": "...", "theme": "corporate|modern|elegant|minimal|warm|dark|academic",
 "accent": "<colour word or #hex ONLY if the client named a colour, else null>",
 "meta": {"author": "...", "org": "...", "date": "..."}, "slides": [...]}
Slide layouts (pick the best one for each message):
  {"layout": "title", "title": "...", "subtitle": "..."}
  {"layout": "agenda", "items": ["...", ...]}
  {"layout": "section", "number": 1, "title": "...", "subtitle": "..."}
  {"layout": "bullets", "title": "...", "bullets": ["...", {"text": "...", "items": ["...", "..."]}]}
  {"layout": "two_column", "title": "...", "left": {"heading": "...", "bullets": [...]}, "right": {"heading": "...", "bullets": [...]}}
  {"layout": "comparison", ...as two_column, for options, pros and cons, before and after}
  {"layout": "chart", "title": "...", "chart": {"chart": "column|bar|line|pie|doughnut|area|stacked", "categories": [...],
   "series": [{"name": "...", "values": [numbers]}]}, "takeaway": "<the one thing the chart shows>"}
  {"layout": "table", "title": "...", "table": {"columns": [...], "rows": [[...]]}, "takeaway": "..."}
  {"layout": "stats", "title": "...", "stats": [{"value": "45%", "label": "..."}], "note": "..."}   (2-4 big numbers)
  {"layout": "process", "title": "...", "steps": [{"title": "...", "text": "..."}]}                (3-6 steps)
  {"layout": "quote", "text": "...", "by": "..."}
  {"layout": "image", "title": "...", "image": "<a path from the client's images only>", "bullets": [...], "caption": "..."}
  {"layout": "closing", "title": "Thank you", "subtitle": "...", "contact": "..."}
Every slide also has "notes": what the presenter says (2-4 sentences).
Design rules:
- A slide title is the message, as a short sentence ("Solar pays for itself in about 4 years"), not a topic label.
- One idea per slide; at most 6 bullets of at most 12 words; no paragraphs on slides (details go in the notes).
- Vary the layouts: numbers -> stats or chart; steps -> process; options -> comparison; one strong line -> quote.
- Title slide first; an agenda for decks of 8+ slides; section slides only for 12+ slides; a closing slide last.
- Slides: the number the client asked for; else 8-12.
- The client's own names, numbers and data exactly; never invent statistics, studies or quotes (approximate figures are
  labelled as approximate in the notes or a stats note); write in the client's language."""


def slides_asked(request):
    m = re.search(r"\b(\d{1,2})\s*(?:-\s*)?(?:slides?|pages?)\b", str(request).lower())
    return int(m.group(1)) if m else None


def design_deck(request, inputs, planner, log=print, theme=None, slides=None):
    """(deck plan, info) for a presentation request: one call writes the whole deck."""
    t0 = time.perf_counter()
    today = dt.date.today().strftime("%d %B %Y")
    n = slides or slides_asked(request)
    d = _ask(planner, DECK_SYSTEM, f"TODAY: {today}\nREQUEST: {request}\n" + (f"SLIDES: exactly {n}\n" if n else "") + "\n" + _inputs_text(inputs), log=log)
    if not d.get("slides"):
        raise ValueError("the deck came back empty")
    if theme:
        d["theme"] = theme
    d["target_slides"] = n
    return d, {"calls": 1, "mode": "deck", "seconds": round(time.perf_counter() - t0, 1)}


BOOK_RX = re.compile(r"\b(?:excel|spreadsheet|xlsx|workbook|worksheet|google sheets?|tracker|ledger|(?:attendance|expense|budget|grade|marks?|salary|"
                     r"inventory|stock|time|fee|result) ?sheet|calculator in excel)\b", re.I)

BOOK_SYSTEM = """You design an Excel workbook for a client. Reply with ONE JSON object:
{"title": "...", "theme": "corporate|modern|minimal|warm|elegant", "accent": "<colour ONLY if the client named one, else null>", "sheets": [...]}
Sheet kinds:
  {"name": "Data", "kind": "table",
   "columns": [{"name": "Date", "type": "date|text|number|integer|money|percent", "formula": "<optional, e.g. [Qty] * [Price]>",
                "options": ["<optional dropdown values>"]}],
   "rows": [[one value per column; null for computed columns], ...],
   "total": {"<column>": "sum|average|count|max|min"},
   "highlight": [{"column": "...", "rule": "<|>|<=|>=|=|top|bottom|scale|bar", "value": <number or text>}]}
  {"name": "Summary", "kind": "summary", "title": "...",
   "metrics": [{"label": "...", "fn": "sum|average|count|max|min", "of": "<column>", "where": {"<column>": "<value>"}, "source": "<table sheet>"}],
   "groups": [{"title": "...", "source": "<table sheet>", "by": "<column>", "values": [{"of": "<column>", "fn": "sum|average|count"}],
               "chart": "column|bar|pie|line"},
              {"title": "...", "source": "<table sheet>", "columns": ["<column>", "<column>", ...], "fn": "average|sum|max|min",
               "label": "<what the columns are, e.g. Subject>", "chart": "column|bar"}]}
  (a group with "by" sums rows per value of a column, e.g. sales by region; a group with "columns" compares several columns,
   e.g. the class average of each subject)
Rules:
- NEVER write cell addresses (A1, C2:C9) or formulas with them. Computed columns name other columns in brackets:
  "[Qty] * [Price]", "IF([Marks] >= 50, \\"Pass\\", \\"Fail\\")", "ROUND([Obtained] / [Total] * 100, 1)". Totals, metrics and groups
  are declared; code writes every formula, so they stay live when the client edits the data.
- Use ALL of the client's data exactly. For a template or tracker to fill in, give 3-6 realistic example rows plus dropdown
  options where a column has fixed choices (status, category, present/absent).
- Types: money for amounts (numbers, no currency text), percent for rates (0.18), dates as YYYY-MM-DD, integer for counts.
- Add a summary sheet when there is something to sum up (totals, by category, month or person) with one chart for the main
  breakdown; skip it for a plain list."""


def design_book(request, inputs, planner, log=print, theme=None):
    """(workbook plan, info) for a spreadsheet request: one call."""
    t0 = time.perf_counter()
    today = dt.date.today().strftime("%d %B %Y")
    d = _ask(planner, BOOK_SYSTEM, f"TODAY: {today}\nREQUEST: {request}\n\n" + _inputs_text(inputs, 14000), log=log)
    if not d.get("sheets"):
        raise ValueError("the workbook came back empty")
    if theme:
        d["theme"] = theme
    return d, {"calls": 1, "mode": "workbook", "seconds": round(time.perf_counter() - t0, 1)}
