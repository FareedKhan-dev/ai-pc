"""Excel workbooks from a workbook plan, with live formulas that code writes (the model never writes a cell address).

  wb = resolve_book({"title": ..., "theme": "corporate", "sheets": [...]})
  info = build_book(wb, "out/docs/x/x.xlsx")   -> {"path", "checks": [{"sheet", "cell", "expect", "what"}], "sheets": [...]}

Sheets:
  table   {"name", "kind": "table", "columns": [{"name", "type": text|number|integer|money|percent|date, "formula": "[Qty] * [Price]",
           "options": [list for a dropdown]}], "rows": [[...]], "total": {"Amount": "sum", "Qty": "sum"},
           "highlight": [{"column", "rule": "<|>|<=|>=|=|top|bottom|scale|bar", "value"}]}
           Formulas name other columns in [brackets] and may use Excel functions (IF, ROUND, SUM, AVERAGE, MAX, MIN,
           COUNTIF, IFERROR...): code turns [Qty] into C2, C3... row by row.
  summary {"name", "kind": "summary", "metrics": [{"label", "fn": sum|average|count|max|min, "of": column, "where": {column: value},
           "source": sheet, "format"}], "groups": [{"title", "source", "by": column, "values": [{"of", "fn"}], "chart": column|bar|pie|line}]}
           Every number is a SUMIFS / AVERAGEIFS / COUNTIFS formula on the source table: change the data, the summary follows.
Each sheet: header styled from the theme, frozen header row, filters, number formats, column widths from the content;
charts are native Excel charts. The checks list holds totals computed here in Python, for verify to compare with what
Excel computes.
"""

import copy
import datetime as dt
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ai_pc.office import themes
from ai_pc.office.docplan import number

FORMATS = {
    "money": "#,##0;[Red]-#,##0",
    "money2": "#,##0.00;[Red]-#,##0.00",
    "number": "#,##0.##",
    "integer": "#,##0",
    "percent": "0.0%",
    "date": "dd-mmm-yyyy",
    "text": "@",
}
FN = {"sum": "SUM", "average": "AVERAGE", "avg": "AVERAGE", "mean": "AVERAGE", "count": "COUNTA", "max": "MAX", "min": "MIN"}
IFS = {"sum": "SUMIFS", "average": "AVERAGEIFS", "avg": "AVERAGEIFS", "mean": "AVERAGEIFS", "count": "COUNTIFS", "max": "MAXIFS", "min": "MINIFS"}


def py_value(expr, row):
    """The value of an arithmetic formula ("[Qty] * [Price] - [Discount]") for one row, computed in Python to check
    Excel's; None when the formula uses anything beyond numbers, columns, + - * / and brackets."""
    import ast

    names = {}

    def sub(m):
        k = f"c{len(names)}"
        names[k] = m.group(1).strip().lower()
        return k

    try:
        tree = ast.parse(re.sub(r"\[([^\]]+)\]", sub, str(expr).strip().lstrip("=")), mode="eval")
    except SyntaxError:
        return None
    vals = {k.lower(): v for k, v in row.items()}

    def ev(n):
        if isinstance(n, ast.Expression):
            return ev(n.body)
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)):
            return float(n.value)
        if isinstance(n, ast.Name) and n.id in names:
            v = vals.get(names[n.id])
            if isinstance(v, bool) or not isinstance(v, (int, float)):
                raise ValueError
            return float(v)
        if isinstance(n, ast.UnaryOp) and isinstance(n.op, (ast.USub, ast.UAdd)):
            return -ev(n.operand) if isinstance(n.op, ast.USub) else ev(n.operand)
        if isinstance(n, ast.BinOp) and isinstance(n.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
            a, b = ev(n.left), ev(n.right)
            return a + b if isinstance(n.op, ast.Add) else a - b if isinstance(n.op, ast.Sub) else a * b if isinstance(n.op, ast.Mult) else a / b
        raise ValueError

    try:
        return ev(tree)
    except (ValueError, ZeroDivisionError):
        return None


def _agg(fn, vals):
    """What Excel's SUM / AVERAGE / MAX / MIN / COUNTA give for the values (only the one asked is computed)."""
    if fn == "COUNTA":
        return len(vals)
    nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if not nums:
        return 0
    return {"SUM": lambda: sum(nums), "AVERAGE": lambda: sum(nums) / len(nums), "MAX": lambda: max(nums), "MIN": lambda: min(nums)}[fn]()


def _sheet_name(n, used):
    n = re.sub(r"[\[\]:*?/\\]", " ", str(n or "Sheet")).strip()[:31] or "Sheet"
    base, k = n, 2
    while n.lower() in used:
        n = f"{base[:28]} {k}"
        k += 1
    used.add(n.lower())
    return n


def _value(v, typ):
    """A cell value of the column's type: numbers from "Rs 1,200" or "12%", dates from common spellings."""
    if v is None or v == "":
        return None
    if typ in ("number", "integer", "money", "percent"):
        p = number(v)
        if p is None:
            return v
        x = p[0] / 100 if (typ == "percent" and (p[2] == "%" or p[0] > 1)) else p[0]
        return int(x) if typ == "integer" or (float(x).is_integer() and typ != "percent") else x
    if typ == "date" and isinstance(v, str):
        for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y", "%d-%b-%Y"):
            try:
                return dt.datetime.strptime(v.strip(), f).date()
            except ValueError:
                continue
    return v


def resolve_book(book):
    b = copy.deepcopy(book or {})
    th = themes.get(b.get("theme"), b.get("accent"), "report")
    b["th"] = th
    used, sheets = set(), []
    for s in b.get("sheets") or []:
        if not isinstance(s, dict):
            continue
        kind = "summary" if s.get("kind") == "summary" or s.get("metrics") or s.get("groups") else "table"
        s = {**s, "kind": kind, "name": _sheet_name(s.get("name") or ("Summary" if kind == "summary" else "Data"), used)}
        if kind == "table":
            cols = []
            for c in s.get("columns") or []:
                c = c if isinstance(c, dict) else {"name": str(c)}
                typ = str(c.get("type") or "text").lower()
                typ = {
                    "currency": "money",
                    "amount": "money",
                    "float": "number",
                    "decimal": "number",
                    "int": "integer",
                    "pct": "percent",
                    "percentage": "percent",
                    "string": "text",
                }.get(typ, typ)
                cols.append({**c, "name": str(c.get("name") or f"Column {len(cols) + 1}").strip(), "type": typ if typ in FORMATS else "text"})
            if not cols:
                continue
            n = len(cols)
            rows = []
            for r in s.get("rows") or []:
                r = list(r) if isinstance(r, (list, tuple)) else [r]
                r = (r + [None] * n)[:n]
                rows.append([None if cols[j].get("formula") else _value(r[j], cols[j]["type"]) for j in range(n)])
            s["columns"], s["rows"] = cols, rows
        sheets.append(s)
    # summaries first in the workbook when there is one (what a reader opens), data after
    b["sheets"] = sorted(sheets, key=lambda s: s["kind"] != "summary")
    return b


class _Book:
    def __init__(self, book):
        self.b = book
        self.th = book["th"]
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.tables = {}  # sheet name -> {"cols": {name: letter}, "first": 2, "last": n, "types": {name: type}}
        self.checks = []
        tt = self.th["table"]
        self.head_font = Font(name=self.th["body"], bold=True, color=tt["header_text"], size=11)
        self.head_fill = PatternFill("solid", fgColor=tt["header_fill"])
        self.body_font = Font(name=self.th["body"], size=11)
        self.thin = Side(style="thin", color=tt["border"])

    def _fmt(self, typ, values=()):
        if typ == "money" and any(isinstance(v, float) and not float(v).is_integer() for v in values):
            return FORMATS["money2"]
        return FORMATS.get(typ, "General")

    def formula(self, expr, letters, row):
        """'[Qty] * [Price]' -> '=C5*D5' (column names in brackets, case-insensitive)."""
        lower = {k.lower(): v for k, v in letters.items()}

        def sub(m):
            name = m.group(1).strip().lower()
            if name not in lower:
                raise KeyError(m.group(1))
            return f"{lower[name]}{row}"

        f = re.sub(r"\[([^\]]+)\]", sub, str(expr).strip().lstrip("="))
        parts = re.split(r'("[^"]*")', f)  # Excel's own style, "=B2*C2": no spaces around operators (text in quotes untouched)
        f = "".join(p if i % 2 else re.sub(r"\s*([-+*/^&=<>,()])\s*", r"\1", p) for i, p in enumerate(parts))
        return "=" + f

    def table(self, s):
        ws = self.wb.create_sheet(s["name"])
        cols, rows = s["columns"], s["rows"]
        s["_calc"] = []
        for r in rows:  # formula columns: Python's value where the formula is plain arithmetic (to check Excel against)
            named = {c["name"]: r[j] for j, c in enumerate(cols)}
            for c in cols:
                if c.get("formula"):
                    named[c["name"]] = py_value(c["formula"], named)
            s["_calc"].append([named[c["name"]] for c in cols])
        letters = {c["name"]: get_column_letter(j + 1) for j, c in enumerate(cols)}
        first, last = 2, 1 + len(rows)
        for j, c in enumerate(cols):
            cell = ws.cell(1, j + 1, c["name"])
            cell.font, cell.fill = self.head_font, self.head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        bad = []
        for i, r in enumerate(rows):
            for j, c in enumerate(cols):
                if c.get("formula"):
                    try:
                        v = self.formula(c["formula"], letters, first + i)
                    except KeyError as e:
                        bad.append(f"{c['name']}: no column {e}")
                        v = None
                else:
                    v = r[j]
                cell = ws.cell(first + i, j + 1, v)
                cell.font = self.body_font
                cell.number_format = self._fmt(c["type"], [x[j] for x in rows])
                cell.border = Border(bottom=self.thin)
                if c["type"] in ("number", "integer", "money", "percent"):
                    cell.alignment = Alignment(horizontal="right")
                if i % 2 == 1 and self.th["table"].get("stripe"):
                    cell.fill = PatternFill("solid", fgColor=self.th["table"]["stripe"])
        # totals row (formulas over the column; what they must equal is computed here for the check)
        tot = s.get("total") or {}
        if isinstance(tot, list):
            tot = {c: "sum" for c in tot}
        if tot and rows:
            tr = last + 1
            fns = {str(f).lower() for f in tot.values()}  # say what the row holds: totals, averages, or both
            label = "Total" if fns <= {"sum"} else "Average" if fns <= {"average", "avg", "mean"} else "Summary"
            ws.cell(tr, 1, label).font = Font(name=self.th["body"], bold=True)
            for name, fn in tot.items():
                key = next((c for c in letters if c.lower() == str(name).lower()), None)
                if key is None:
                    continue
                L = letters[key]
                f = FN.get(str(fn).lower(), "SUM")
                cell = ws.cell(tr, list(letters).index(key) + 1, f"={f}({L}{first}:{L}{last})")
                cell.font = Font(name=self.th["body"], bold=True)
                cell.number_format = self._fmt(next(c["type"] for c in cols if c["name"] == key), [x[list(letters).index(key)] for x in rows])
                cell.border = Border(
                    top=Side(style="thin", color=self.th["table"]["header_fill"]), bottom=Side(style="double", color=self.th["table"]["header_fill"])
                )
                j = list(letters).index(key)
                vals = [x[j] for x in s["_calc"] if isinstance(x[j], (int, float)) and not isinstance(x[j], bool)]
                if vals and len(vals) == len(rows):
                    expect = _agg(f, vals)
                    self.checks.append({"sheet": s["name"], "cell": f"{L}{tr}", "expect": round(expect, 6), "what": f"{fn} of {key}"})
        # dropdowns, highlights, layout
        for c in cols:
            if c.get("options") and rows:
                opts = ",".join(str(o).replace(",", " ") for o in c["options"])[:250]
                dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{letters[c['name']]}{first}:{letters[c['name']]}{max(last, first + 200)}")
        for h in s.get("highlight") or []:
            key = next((c for c in letters if c.lower() == str(h.get("column", "")).lower()), None)
            if not key or not rows:
                continue
            rng = f"{letters[key]}{first}:{letters[key]}{last}"
            rule = str(h.get("rule") or "").lower()
            red, green = PatternFill("solid", fgColor="FDE2E1"), PatternFill("solid", fgColor="E3F4E6")
            ops = {"<": "lessThan", ">": "greaterThan", "<=": "lessThanOrEqual", ">=": "greaterThanOrEqual", "=": "equal"}
            if rule in ops and h.get("value") is not None:
                v = h["value"]
                val = f'"{v}"' if not isinstance(v, (int, float)) and number(v) is None else str(number(v)[0] if number(v) else v)
                ws.conditional_formatting.add(rng, CellIsRule(operator=ops[rule], formula=[val], fill=red if rule in ("<", "<=") else green))
            elif rule == "scale":
                ws.conditional_formatting.add(
                    rng,
                    ColorScaleRule(
                        start_type="min",
                        start_color="F8696B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color="63BE7B",
                    ),
                )
            elif rule == "bar":
                ws.conditional_formatting.add(rng, DataBarRule(start_type="min", end_type="max", color=self.th["accent"]))
            elif rule in ("top", "bottom"):
                k = int(h.get("value") or 3)
                fn = "LARGE" if rule == "top" else "SMALL"
                ws.conditional_formatting.add(
                    rng,
                    FormulaRule(
                        formula=[
                            f"{letters[key]}{first}>={fn}(${letters[key]}${first}:${letters[key]}${last},{k})"
                            if rule == "top"
                            else f"{letters[key]}{first}<={fn}(${letters[key]}${first}:${letters[key]}${last},{k})"
                        ],
                        fill=green if rule == "top" else red,
                    ),
                )
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
        ws.page_setup.orientation = "landscape" if len(cols) > 6 else "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"
        for j, c in enumerate(cols):
            vals = [str(c["name"])] + [("" if r[j] is None else f"{r[j]:,.2f}" if isinstance(r[j], float) else str(r[j])) for r in rows[:300]]
            width = max(len(v) for v in vals) + 3
            if c.get("formula"):
                width = max(width, 12)
            ws.column_dimensions[get_column_letter(j + 1)].width = max(8, min(48, width))
        ws.sheet_view.showGridLines = False
        calc_ok = {
            c["name"]
            for j, c in enumerate(cols)
            if all(isinstance(x[j], (int, float)) or x[j] is None for x in s["_calc"])
            and (not c.get("formula") or all(x[j] is not None for x in s["_calc"]))
        }
        self.tables[s["name"]] = {
            "cols": letters,
            "first": first,
            "last": last,
            "types": {c["name"]: c["type"] for c in cols},
            "rows": s["_calc"],
            "colnames": [c["name"] for c in cols],
            "formulas": {c["name"] for c in cols if c.get("formula") and c["name"] not in calc_ok},
        }
        return bad

    def summary(self, s):
        ws = self.wb.create_sheet(s["name"])
        th = self.th
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 20
        title = s.get("title") or self.b.get("title") or s["name"]
        ws.cell(1, 1, title).font = Font(name=th["head"], size=18, bold=True, color=th["head_color"])
        ws.row_dimensions[1].height = 28
        row = 3
        bad = []
        default_src = next((n for n, t in self.tables.items()), None)

        def ref(src, col):
            t = self.tables.get(src)
            key = next((c for c in t["cols"] if c.lower() == str(col).lower()), None) if t else None
            if key is None:
                raise KeyError(f"{src}!{col}")
            L = t["cols"][key]
            q = f"'{src}'" if re.search(r"\W", src) else src
            return f"{q}!${L}${t['first']}:${L}${t['last']}", key, t

        for m in s.get("metrics") or []:
            src = m.get("source") if m.get("source") in self.tables else default_src
            fn = str(m.get("fn") or "sum").lower()
            try:
                rng, key, t = ref(src, m.get("of") or "")
                where = m.get("where") or {}
                if where:
                    crit = []
                    for wc, wv in where.items():
                        wr, _, _ = ref(src, wc)
                        crit += [wr, f'"{wv}"' if not isinstance(wv, (int, float)) else str(wv)]
                    f = f"={IFS.get(fn, 'SUMIFS')}({rng}, {', '.join(crit)})" if fn not in ("count",) else f"=COUNTIFS({', '.join(crit)})"
                    if IFS.get(fn) in ("SUMIFS", "AVERAGEIFS", "MAXIFS", "MINIFS"):
                        f = f"={IFS[fn]}({rng}, {', '.join(crit)})"
                else:
                    f = f"={FN.get(fn, 'SUM')}({rng})"
            except KeyError as e:
                bad.append(f"metric '{m.get('label')}': no column {e}")
                continue
            ws.cell(row, 1, str(m.get("label") or f"{fn} of {key}")).font = Font(name=th["body"], size=12, color=th["muted"])
            c = ws.cell(row, 2, f)
            c.font = Font(name=th["head"], size=14, bold=True, color=th["accent"])
            c.number_format = (
                FORMATS["integer"]
                if fn == "count"
                else self._fmt(m.get("format") or t["types"].get(key, "number"), [r[t["colnames"].index(key)] for r in t["rows"]])
            )
            c.alignment = Alignment(horizontal="right")
            if key not in t["formulas"]:
                jk = t["colnames"].index(key)

                def keep(r):
                    return all(
                        str(r[t["colnames"].index(next(c for c in t["colnames"] if c.lower() == str(wc).lower()))]) == str(wv)
                        for wc, wv in where.items()
                    )

                vals = [r[jk] for r in t["rows"] if keep(r) and (fn == "count" and r[jk] not in (None, "") or isinstance(r[jk], (int, float)))]
                if vals:
                    e = _agg(FN.get(fn, "SUM"), vals)
                    self.checks.append({"sheet": s["name"], "cell": f"B{row}", "expect": round(e, 6), "what": str(m.get("label"))})
            row += 1
        row += 1
        chart_col = 5
        for g in s.get("groups") or []:
            src = g.get("source") if g.get("source") in self.tables else default_src
            t = self.tables.get(src)
            if not t:
                continue
            if not g.get("by") and not g.get("columns") and len([v for v in g.get("values") or [] if isinstance(v, dict) and v.get("of")]) > 1:
                vs = [v for v in g["values"] if isinstance(v, dict) and v.get("of")]  # "values" without "by": the model means the columns
                g = {**g, "columns": [v["of"] for v in vs], "fn": vs[0].get("fn") or "average"}
            if g.get("columns") and not g.get("by"):  # several columns side by side ("average per subject")
                row, chart_col = self._compare(ws, g, src, t, row, chart_col, ref, bad)
                continue
            try:
                by_rng, by_key, _ = ref(src, g.get("by") or "")
            except KeyError as e:
                bad.append(f"group '{g.get('title')}': no column {e}")
                continue
            vals_spec = [
                v
                for v in g.get("values")
                or [{"of": next((c for c in t["colnames"] if t["types"][c] in ("money", "number", "integer")), None), "fn": "sum"}]
                if isinstance(v, dict) and v.get("of")
            ]
            j_by = t["colnames"].index(by_key)
            groups = list(dict.fromkeys(str(r[j_by]) for r in t["rows"] if r[j_by] not in (None, "")))
            if not groups:
                continue
            ws.cell(row, 1, str(g.get("title") or f"By {by_key}")).font = Font(name=th["head"], size=13, bold=True, color=th["head_color"])
            row += 1
            head = row
            for j, h in enumerate([by_key] + [f"{v.get('fn', 'sum').title()} of {v['of']}" for v in vals_spec]):
                c = ws.cell(row, 1 + j, h)
                c.font, c.fill = self.head_font, self.head_fill
                c.alignment = Alignment(horizontal="center")
                ws.column_dimensions[get_column_letter(1 + j)].width = max(ws.column_dimensions[get_column_letter(1 + j)].width or 0, len(h) + 4)
            for gi, gv in enumerate(groups):
                r_ = row + 1 + gi
                ws.cell(r_, 1, gv).font = self.body_font
                for j, v in enumerate(vals_spec):
                    try:
                        rng, key, _ = ref(src, v["of"])
                    except KeyError as e:
                        bad.append(f"group value: no column {e}")
                        continue
                    fn = str(v.get("fn") or "sum").lower()
                    f = f"=COUNTIFS({by_rng}, A{r_})" if fn == "count" else f"={IFS.get(fn, 'SUMIFS')}({rng}, {by_rng}, A{r_})"
                    c = ws.cell(r_, 2 + j, f)
                    c.font = self.body_font
                    c.number_format = self._fmt(t["types"].get(key, "number"), [x[t["colnames"].index(key)] for x in t["rows"]])
                    if fn == "sum" and key not in t["formulas"]:
                        jj = t["colnames"].index(key)
                        e = sum(x[jj] for x in t["rows"] if str(x[j_by]) == gv and isinstance(x[jj], (int, float)))
                        self.checks.append(
                            {"sheet": s["name"], "cell": f"{get_column_letter(2 + j)}{r_}", "expect": round(e, 6), "what": f"{key} for {gv}"}
                        )
            last = row + len(groups)
            if g.get("chart") and vals_spec:
                kind = str(g["chart"]).lower()
                ch = PieChart() if kind in ("pie", "doughnut") else LineChart() if kind == "line" else BarChart()
                if isinstance(ch, BarChart):
                    ch.type = "bar" if kind == "bar" else "col"
                    ch.gapWidth = 80
                ch.title = str(g.get("title") or f"By {by_key}")
                ch.style = 10
                data = Reference(ws, min_col=2, min_row=head, max_row=last)
                cats = Reference(ws, min_col=1, min_row=head + 1, max_row=last)
                ch.add_data(data, titles_from_data=True)
                ch.set_categories(cats)
                if not isinstance(ch, PieChart):  # openpyxl leaves the axes "deleted" for current Excel: no labels without this
                    ch.x_axis.delete = False
                    ch.y_axis.delete = False
                    ch.y_axis.scaling.min = 0  # from zero: a cut axis exaggerates small differences
                    ch.y_axis.majorGridlines = ch.y_axis.majorGridlines
                    ch.y_axis.numFmt = "#,##0"
                ch.height, ch.width = 7.5, 15
                if isinstance(ch, PieChart):
                    ch.dataLabels = DataLabelList()
                    ch.dataLabels.showPercent = True
                    ch.dataLabels.showVal = ch.dataLabels.showCatName = ch.dataLabels.showSerName = ch.dataLabels.showLegendKey = False
                else:
                    ch.legend = None if len(vals_spec) == 1 else ch.legend
                    try:
                        from openpyxl.chart.shapes import GraphicalProperties

                        ch.series[0].graphicalProperties = GraphicalProperties(solidFill=th["accent"])
                    except Exception:  # noqa: BLE001
                        pass
                ws.add_chart(ch, f"{get_column_letter(chart_col)}{head - 1}")
            row = last + 3 if not g.get("chart") else max(last + 3, head + 17)
        return bad

    def _compare(self, ws, g, src, t, row, chart_col, ref, bad):
        """One row per chosen column with its sum / average / max / min (e.g. the class average of each subject), and a
        chart of them."""
        th = self.th
        fn = str(g.get("fn") or "average").lower()
        cols = []
        for c in g.get("columns") or []:
            try:
                rng, key, _ = ref(src, c)
                cols.append((key, rng))
            except KeyError as e:
                bad.append(f"group '{g.get('title')}': no column {e}")
        if not cols:
            return row, chart_col
        ws.cell(row, 1, str(g.get("title") or f"{fn.title()} by column")).font = Font(name=th["head"], size=13, bold=True, color=th["head_color"])
        row += 1
        head = row
        m = re.search(r"\b(?:by|per|of each)\s+([A-Za-z][\w ]{1,20})$", str(g.get("title") or ""))
        label = str(g.get("label") or (m.group(1).strip().title() if m else "Item"))
        for j, h in enumerate([label, fn.title()]):
            c = ws.cell(row, 1 + j, h)
            c.font, c.fill = self.head_font, self.head_fill
            c.alignment = Alignment(horizontal="center")
        for i, (key, rng) in enumerate(cols):
            r_ = row + 1 + i
            ws.cell(r_, 1, key).font = self.body_font
            c = ws.cell(r_, 2, f"={FN.get(fn, 'AVERAGE')}({rng})")
            c.font = self.body_font
            c.number_format = "#,##0.0" if fn in ("average", "avg", "mean") else self._fmt(t["types"].get(key, "number"))
            jk = t["colnames"].index(key)
            if key not in t["formulas"]:
                vals = [x[jk] for x in t["rows"] if isinstance(x[jk], (int, float)) and not isinstance(x[jk], bool)]
                if vals:
                    self.checks.append(
                        {"sheet": ws.title, "cell": f"B{r_}", "expect": round(_agg(FN.get(fn, "AVERAGE"), vals), 6), "what": f"{fn} of {key}"}
                    )
        last = row + len(cols)
        if g.get("chart", "column"):
            kind = str(g.get("chart") or "column").lower()
            ch = PieChart() if kind in ("pie", "doughnut") else LineChart() if kind == "line" else BarChart()
            if isinstance(ch, BarChart):
                ch.type = "bar" if kind == "bar" else "col"
                ch.gapWidth = 80
            ch.title = str(g.get("title") or "")
            ch.style = 10
            ch.add_data(Reference(ws, min_col=2, min_row=head, max_row=last), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=last))
            ch.height, ch.width = 7.5, 15
            if not isinstance(ch, PieChart):
                ch.x_axis.delete = False
                ch.y_axis.delete = False
                ch.y_axis.scaling.min = 0
                ch.legend = None
                try:
                    from openpyxl.chart.shapes import GraphicalProperties

                    ch.series[0].graphicalProperties = GraphicalProperties(solidFill=th["accent"])
                except Exception:  # noqa: BLE001
                    pass
            ws.add_chart(ch, f"{get_column_letter(chart_col)}{head - 1}")
            return max(last + 3, head + 17), chart_col
        return last + 3, chart_col

    def run(self, path):
        problems = []
        for s in [x for x in self.b["sheets"] if x["kind"] == "table"]:
            problems += self.table(s)
        for s in [x for x in self.b["sheets"] if x["kind"] == "summary"]:
            problems += self.summary(s)
        order = [s["name"] for s in self.b["sheets"]]
        self.wb._sheets.sort(key=lambda ws: order.index(ws.title) if ws.title in order else 99)
        self.wb.active = 0
        self.wb.properties.title = str(self.b.get("title") or "")[:250]
        self.wb.properties.creator = "AI PC document agent"
        self.wb.save(str(path))
        return {
            "path": str(path),
            "checks": self.checks,
            "problems": problems,
            "sheets": [{"name": s["name"], "kind": s["kind"], "rows": len(s.get("rows") or [])} for s in self.b["sheets"]],
        }


def build_book(book, path):
    return _Book(book if "th" in book else resolve_book(book)).run(path)
