"""References for assignments and theses: DOIs looked up at Crossref and ISBNs at Open Library (both free, no key), then
written in APA 7, MLA 9, Harvard, IEEE or Chicago author-date, as a Word list (hanging indent, titles in italics,
sorted or numbered as the style wants) plus BibTeX (.bib) and RIS (.ris) files that Zotero, Mendeley and EndNote import.

  'cite 10.1038/nature14539 and isbn 9780262035613 in apa'   'references in ieee for doi 10.1145/3065386'
"""
import re
from pathlib import Path

from ai_pc.hub.http import Api, HubError

NAME, LABEL = "cite", "References and citations (APA, MLA, Harvard, IEEE, Chicago; BibTeX, RIS)"
EXAMPLES = ["cite 10.1038/nature14539 and isbn 9780262035613 in apa", "references in ieee for doi 10.1145/3065386"]
STYLES = {"apa": "APA 7", "mla": "MLA 9", "harvard": "Harvard", "ieee": "IEEE", "chicago": "Chicago (author-date)"}
UA = {"User-Agent": "AI-PC/1.0 (personal desktop reference tool)"}
TRANSPORT = None  # tests put a fake here


def doi_lookup(doi):
    x = Api("https://api.crossref.org", headers=UA, service="crossref", transport=TRANSPORT).get(f"works/{doi}")["message"]
    parts = ((x.get("issued") or x.get("published-print") or x.get("published-online") or {}).get("date-parts") or [[None]])[0]
    return {"kind": "book" if x.get("type") in ("book", "monograph", "edited-book") else "chapter" if x.get("type") == "book-chapter" else "article",
            "title": (x.get("title") or [""])[0].strip(), "authors": [(a.get("given", ""), a.get("family") or a.get("name", "")) for a in x.get("author") or []],
            "year": parts[0], "journal": (x.get("container-title") or [""])[0], "volume": x.get("volume"), "issue": x.get("issue"),
            "pages": (x.get("page") or "").replace("-", "–"), "publisher": x.get("publisher"), "doi": x.get("DOI", doi)}


def isbn_lookup(isbn):
    api = Api("https://openlibrary.org", headers=UA, service="openlibrary", transport=TRANSPORT)
    x = api.get(f"isbn/{isbn}.json")
    authors = []
    for a in x.get("authors") or []:
        name = api.get(f"{a['key']}.json").get("name", "")
        given, _, family = name.rpartition(" ")
        authors.append((given, family or name))
    year = re.search(r"\d{4}", str(x.get("publish_date") or ""))
    return {"kind": "book", "title": (x.get("title") or "") + (f": {x['subtitle']}" if x.get("subtitle") else ""), "authors": authors,
            "year": int(year.group(0)) if year else None, "publisher": (x.get("publishers") or [""])[0], "isbn": isbn, "edition": x.get("edition_name")}


def _init(given):
    return " ".join(f"{p[0]}." for p in re.split(r"[\s-]+", given) if p)


def _names(authors, style):
    if not authors:
        return ""
    if style == "apa":
        names = [f"{f}, {_init(g)}".strip(", ") for g, f in authors]
        return names[0] if len(names) == 1 else ", ".join(names[:-1]) + ", & " + names[-1] if len(names) <= 20 else ", ".join(names[:19]) + ", ... " + names[-1]
    if style == "mla":
        g, f = authors[0]
        first = f"{f}, {g}".strip(", ")
        return first if len(authors) == 1 else f"{first}, and {authors[1][0]} {authors[1][1]}" if len(authors) == 2 else f"{first}, et al"
    if style == "harvard":
        names = [f"{f}, {_init(g)}".strip(", ") for g, f in authors]
        return names[0] if len(names) == 1 else ", ".join(names[:-1]) + " and " + names[-1] if len(names) <= 3 else f"{names[0]} et al."
    if style == "ieee":
        names = [f"{_init(g)} {f}".strip() for g, f in authors]
        return names[0] if len(names) == 1 else ", ".join(names[:-1]) + ", and " + names[-1] if len(names) <= 6 else f"{names[0]} et al."
    g, f = authors[0]  # chicago
    first = f"{f}, {g}".strip(", ")
    rest = [f"{gg} {ff}".strip() for gg, ff in authors[1:]]
    return first if not rest else first + ", " + ", ".join(rest[:-1]) + (", and " if len(rest) > 1 else " and ") + rest[-1] if len(authors) <= 10 else first + ", et al."


def reference(r, style, n=None):
    """One reference as runs [(text, italic)]."""
    a, y, t = _names(r["authors"], style), r.get("year") or "n.d.", r["title"].rstrip(".")
    doi = f"https://doi.org/{r['doi']}" if r.get("doi") else ""
    vol, iss, pg = r.get("volume"), r.get("issue"), r.get("pages")
    book = r["kind"] == "book"
    if style == "apa":
        if book:
            return [(f"{a} ({y}). ", False), (t, True), (f". {r.get('publisher') or ''}".rstrip() + "." + (f" {doi}" if doi else ""), False)]
        return [(f"{a} ({y}). {t}. ", False), (r["journal"], True), (", ", False)] + ([(str(vol), True)] if vol else []) + \
            [((f"({iss})" if iss else "") + (f", {pg}" if pg else "") + "." + (f" {doi}" if doi else ""), False)]
    if style == "mla":
        if book:
            return [(f"{a}. ", False), (t, True), (f". {r.get('publisher') or ''}, {y}.", False)]
        return [(f'{a}. "{t}." ', False), (r["journal"], True), ((f", vol. {vol}" if vol else "") + (f", no. {iss}" if iss else "") + f", {y}" +
                                                                (f", pp. {pg}" if pg else "") + "." + (f" {doi}." if doi else ""), False)]
    if style == "harvard":
        if book:
            return [(f"{a} ({y}) ", False), (t, True), (f". {r.get('publisher') or ''}.", False)]
        return [(f"{a} ({y}) '{t}', ", False), (r["journal"], True), ((f", {vol}" if vol else "") + (f"({iss})" if iss else "") +
                                                                     (f", pp. {pg}" if pg else "") + "." + (f" doi:{r['doi']}." if r.get("doi") else ""), False)]
    if style == "ieee":
        lead = f"[{n}] " if n else ""
        if book:
            return [(f"{lead}{a}, ", False), (t, True), (f". {r.get('publisher') or ''}, {y}.", False)]
        return [(f'{lead}{a}, "{t}," ', False), (r["journal"], True), ((f", vol. {vol}" if vol else "") + (f", no. {iss}" if iss else "") +
                                                                       (f", pp. {pg}" if pg else "") + f", {y}" + (f", doi: {r['doi']}." if r.get("doi") else "."), False)]
    if book:  # chicago author-date
        return [(f"{a}. {y}. ", False), (t, True), (f". {r.get('publisher') or ''}.", False)]
    return [(f'{a}. {y}. "{t}." ', False), (r["journal"], True), ((f" {vol}" if vol else "") + (f" ({iss})" if iss else "") + (f": {pg}" if pg else "") +
                                                                   "." + (f" {doi}." if doi else ""), False)]


def text_of(runs):
    return "".join(t for t, _ in runs)


def bibtex(r, key):
    f = {"title": r["title"], "author": " and ".join(f"{fam}, {g}".strip(", ") for g, fam in r["authors"]), "year": r.get("year"),
         "journal": r.get("journal") if r["kind"] != "book" else None, "volume": r.get("volume"), "number": r.get("issue"),
         "pages": (r.get("pages") or "").replace("–", "--") or None, "publisher": r.get("publisher") if r["kind"] == "book" else None,
         "doi": r.get("doi"), "isbn": r.get("isbn")}
    body = ",\n".join(f"  {k} = {{{v}}}" for k, v in f.items() if v)
    return f"@{'book' if r['kind'] == 'book' else 'article'}{{{key},\n{body}\n}}\n"


def ris(r):
    lines = [f"TY  - {'BOOK' if r['kind'] == 'book' else 'JOUR'}"] + [f"AU  - {fam}, {g}".rstrip(", ") for g, fam in r["authors"]] + \
        [f"TI  - {r['title']}", f"PY  - {r.get('year') or ''}"]
    for tag, k in (("JO", "journal"), ("VL", "volume"), ("IS", "issue"), ("PB", "publisher"), ("DO", "doi"), ("SN", "isbn")):
        if r.get(k):
            lines.append(f"{tag}  - {r[k]}")
    if r.get("pages"):
        sp, _, ep = r["pages"].partition("–")
        lines += [f"SP  - {sp}"] + ([f"EP  - {ep}"] if ep else [])
    return "\n".join(lines + ["ER  - ", ""])


def make(ids, style, out):
    recs = []
    for kind, value in ids:
        try:
            recs.append(doi_lookup(value) if kind == "doi" else isbn_lookup(value))
        except HubError as e:
            raise ValueError(f"{kind.upper()} {value} was not found ({e})") from e
    if style != "ieee":
        recs.sort(key=lambda r: ((r["authors"][0][1] if r["authors"] else r["title"]).lower(), r.get("year") or 0))
    refs = [reference(r, style, i) for i, r in enumerate(recs, 1)]
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    from docx import Document
    from docx.shared import Inches, Pt
    d = Document()
    d.add_heading("References", level=1)
    for runs in refs:
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.space_after = Pt(6)
        for t, it in runs:
            p.add_run(t).italic = it
    d.save(out / "references.docx")
    keys = []
    for r in recs:
        k = re.sub(r"\W", "", (r["authors"][0][1] if r["authors"] else "ref").lower()) + str(r.get("year") or "")
        while k in keys:
            k += "a"
        keys.append(k)
    (out / "references.bib").write_text("\n".join(bibtex(r, k) for r, k in zip(recs, keys)), encoding="utf-8")
    (out / "references.ris").write_text("\n".join(ris(r) for r in recs), encoding="utf-8")
    return recs, refs, [out / "references.docx", out / "references.bib", out / "references.ris"]


def parse(text, ctx):
    c = text.lower()
    dois = re.findall(r"\b(10\.\d{4,9}/[^\s,;]+[^\s,;.)])", text)
    isbns = [re.sub(r"[-\s]", "", x) for x in re.findall(r"\bisbn(?:-1[03])?:?\s*([\d][\d\s-]{8,16}[\dxX])", text, re.I)]
    if not (dois or isbns) or not re.search(r"\b(?:cite|citation|reference|bibliography|works cited|apa|mla|harvard|ieee|chicago)\b|\bdoi\b|\bisbn\b", c):
        return None
    style = next((s for s in STYLES if re.search(rf"\b{s}\b", c)), "apa")
    return {"op": "cite", "ids": [("doi", d) for d in dois] + [("isbn", i) for i in isbns], "style": style}


def run(op, ctx):
    recs, refs, files = make(op["ids"], op["style"], Path(ctx["out"]) / "references")
    from docx import Document
    back = [p.text for p in Document(files[0]).paragraphs][1:]
    ok = back == [text_of(r) for r in refs]
    return (f"{len(refs)} references in {STYLES[op['style']]}:\n" + "\n".join(text_of(r) for r in refs) +
            f"\nSaved: {', '.join(str(f) for f in files)} ({'checked' if ok else 'NOT the same when read back'}; titles in italics in the Word file).")
