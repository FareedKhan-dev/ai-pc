"""Statistics on an Excel or CSV sheet, the way SPSS reports them for a thesis: descriptives and frequencies, Pearson
correlations, independent-samples t-tests (Student and Welch, Cohen's d), one-way ANOVA (eta squared), crosstabs with
chi-square (Cramer's V), linear regression (B, SE, t, p, R squared, F) and Cronbach's alpha. Each answer is written as
an APA-style sentence and table in a Word report, with p-values from the t, F and chi-square distributions computed
here (numpy plus the incomplete beta and gamma functions; checked against textbook values in the tests).

  'describe survey.xlsx'   't-test of Score by Gender in survey.xlsx'   'anova of Score by Class in survey.xlsx'
  'correlation between Hours and Score in survey.xlsx'   'crosstab Gender by Choice in survey.xlsx'
  'regression of Score on Hours, Sleep in survey.xlsx'   'reliability of Q1, Q2, Q3, Q4 in survey.xlsx'
"""

import math
import re
from pathlib import Path

import numpy as np

NAME, LABEL = "stats", "Statistics like SPSS: descriptives, t-test, ANOVA, chi-square, correlation, regression, alpha"
EXAMPLES = ["describe survey.xlsx", "t-test of Score by Gender in survey.xlsx", "regression of Score on Hours, Sleep in survey.xlsx"]


# ---------------------------------------------------------------- distributions
def _betacf(a, b, x):
    m, qab, qap, qam = 1, a + b, a + 1, a - 1
    c, d = 1.0, 1 - qab * x / qap
    d = 1 / (d if abs(d) > 1e-300 else 1e-300)
    h = d
    for m in range(1, 300):
        m2 = 2 * m
        aa = m * (b - m) * x / ((qam + m2) * (a + m2))
        d = 1 + aa * d
        d = 1 / (d if abs(d) > 1e-300 else 1e-300)
        c = 1 + aa / (c if abs(c) > 1e-300 else 1e-300)
        h *= d * c
        aa = -(a + m) * (qab + m) * x / ((a + m2) * (qap + m2))
        d = 1 + aa * d
        d = 1 / (d if abs(d) > 1e-300 else 1e-300)
        c = 1 + aa / (c if abs(c) > 1e-300 else 1e-300)
        de = d * c
        h *= de
        if abs(de - 1) < 3e-14:
            break
    return h


def betai(a, b, x):
    """The regularized incomplete beta function I_x(a, b)."""
    if x <= 0:
        return 0.0
    if x >= 1:
        return 1.0
    bt = math.exp(math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b) + a * math.log(x) + b * math.log(1 - x))
    return bt * _betacf(a, b, x) / a if x < (a + 1) / (a + b + 2) else 1 - bt * _betacf(b, a, 1 - x) / b


def gammap(a, x):
    """The regularized lower incomplete gamma function P(a, x)."""
    if x <= 0:
        return 0.0
    if x < a + 1:
        s, term, n = 1 / a, 1 / a, a
        for _ in range(1000):
            n += 1
            term *= x / n
            s += term
            if abs(term) < abs(s) * 1e-15:
                break
        return s * math.exp(-x + a * math.log(x) - math.lgamma(a))
    b, c, d = x + 1 - a, 1e300, 1 / (x + 1 - a)
    h = d
    for i in range(1, 1000):
        an = -i * (i - a)
        b += 2
        d = an * d + b
        d = 1 / (d if abs(d) > 1e-300 else 1e-300)
        c = b + an / (c if abs(c) > 1e-300 else 1e-300)
        h *= d * c
        if abs(d * c - 1) < 1e-15:
            break
    return 1 - math.exp(-x + a * math.log(x) - math.lgamma(a)) * h


def p_t(t, df):
    """Two-tailed p of Student's t."""
    return betai(df / 2, 0.5, df / (df + t * t))


def p_f(f, d1, d2):
    return betai(d2 / 2, d1 / 2, d2 / (d2 + d1 * f)) if f > 0 else 1.0


def p_chi2(x, df):
    return 1 - gammap(df / 2, x / 2)


# ---------------------------------------------------------------- data
def sheet(path):
    p = Path(path)
    if p.suffix.lower() == ".csv":
        import csv

        with p.open(encoding="utf-8-sig", newline="") as f:
            rows = [r for r in csv.reader(f) if any(c.strip() for c in r)]
    else:
        from openpyxl import load_workbook

        rows = [
            list(r)
            for r in load_workbook(p, data_only=True, read_only=True).active.iter_rows(values_only=True)
            if any(c not in (None, "") for c in r)
        ]
    heads = [str(h).strip() for h in rows[0]]
    cols = {h: [r[i] if i < len(r) else None for r in rows[1:]] for i, h in enumerate(heads)}
    return cols


def numeric(vals):
    out = []
    for v in vals:
        try:
            out.append(float(str(v).replace(",", "")) if v not in (None, "") else math.nan)
        except ValueError:
            return None
    return np.array(out)


def col(data, name):
    hit = next((k for k in data if k.lower() == name.lower()), None)
    if hit is None:
        raise ValueError(f"no column called '{name}' (columns: {', '.join(data)})")
    return hit


def pfmt(p):
    return "p < .001" if p < 0.001 else f"p = {p:.3f}".replace("0.", ".")


# ---------------------------------------------------------------- tests
def describe(data):
    num, cat = [], []
    for k, v in data.items():
        x = numeric(v)
        if x is not None and np.isfinite(x).sum() > 1:
            x = x[np.isfinite(x)]
            n, m, s = len(x), x.mean(), x.std(ddof=1)
            z = (x - m) / s if s else x * 0
            num.append(
                {
                    "var": k,
                    "N": n,
                    "Mean": m,
                    "SD": s,
                    "Min": x.min(),
                    "Median": float(np.median(x)),
                    "Max": x.max(),
                    "Skewness": float((z**3).mean() * n * n / ((n - 1) * (n - 2))) if n > 2 and s else 0.0,
                }
            )
        else:
            vals = [str(t).strip() for t in v if t not in (None, "")]
            counts = {}
            for t in vals:
                counts[t] = counts.get(t, 0) + 1
            cat.append({"var": k, "N": len(vals), "freq": sorted(counts.items(), key=lambda kv: -kv[1])})
    return num, cat


def ttest(data, dv, iv):
    y, g = numeric(data[dv]), [str(t).strip() for t in data[iv]]
    groups = [k for k in dict.fromkeys(t for t in g if t and t != "None")]
    if len(groups) != 2:
        raise ValueError(f"a t-test needs exactly two groups in {iv}; it has {len(groups)}: {', '.join(groups[:6])}")
    a = np.array([v for v, k in zip(y, g) if k == groups[0] and np.isfinite(v)])
    b = np.array([v for v, k in zip(y, g) if k == groups[1] and np.isfinite(v)])
    n1, n2, m1, m2, s1, s2 = len(a), len(b), a.mean(), b.mean(), a.var(ddof=1), b.var(ddof=1)
    sp = ((n1 - 1) * s1 + (n2 - 1) * s2) / (n1 + n2 - 2)
    t = (m1 - m2) / math.sqrt(sp * (1 / n1 + 1 / n2))
    df = n1 + n2 - 2
    tw = (m1 - m2) / math.sqrt(s1 / n1 + s2 / n2)
    dfw = (s1 / n1 + s2 / n2) ** 2 / ((s1 / n1) ** 2 / (n1 - 1) + (s2 / n2) ** 2 / (n2 - 1))
    d = (m1 - m2) / math.sqrt(sp)
    return {
        "groups": groups,
        "n": (n1, n2),
        "mean": (m1, m2),
        "sd": (math.sqrt(s1), math.sqrt(s2)),
        "t": t,
        "df": df,
        "p": p_t(t, df),
        "t_welch": tw,
        "df_welch": dfw,
        "p_welch": p_t(tw, dfw),
        "d": d,
    }


def anova(data, dv, iv):
    y, g = numeric(data[dv]), [str(t).strip() for t in data[iv]]
    groups = {}
    for v, k in zip(y, g):
        if np.isfinite(v) and k and k != "None":
            groups.setdefault(k, []).append(v)
    allv = np.concatenate([np.array(v) for v in groups.values()])
    gm = allv.mean()
    ssb = sum(len(v) * (np.mean(v) - gm) ** 2 for v in groups.values())
    ssw = sum(((np.array(v) - np.mean(v)) ** 2).sum() for v in groups.values())
    d1, d2 = len(groups) - 1, len(allv) - len(groups)
    f = (ssb / d1) / (ssw / d2)
    return {
        "groups": {k: (len(v), float(np.mean(v)), float(np.std(v, ddof=1))) for k, v in groups.items()},
        "F": f,
        "df": (d1, d2),
        "p": p_f(f, d1, d2),
        "eta2": ssb / (ssb + ssw),
    }


def correlation(data, a, b):
    x, y = numeric(data[a]), numeric(data[b])
    ok = np.isfinite(x) & np.isfinite(y)
    x, y = x[ok], y[ok]
    r = float(np.corrcoef(x, y)[0, 1])
    n = len(x)
    t = r * math.sqrt((n - 2) / max(1e-15, 1 - r * r))
    return {"r": r, "n": n, "p": p_t(t, n - 2)}


def crosstab(data, a, b):
    xa, xb = [str(t).strip() for t in data[a]], [str(t).strip() for t in data[b]]
    rows, cols = list(dict.fromkeys(xa)), list(dict.fromkeys(xb))
    obs = np.zeros((len(rows), len(cols)))
    for p, q in zip(xa, xb):
        obs[rows.index(p), cols.index(q)] += 1
    exp = obs.sum(1, keepdims=True) * obs.sum(0, keepdims=True) / obs.sum()
    chi = float(((obs - exp) ** 2 / exp).sum())
    df = (len(rows) - 1) * (len(cols) - 1)
    v = math.sqrt(chi / (obs.sum() * (min(obs.shape) - 1))) if min(obs.shape) > 1 else 0.0
    return {"rows": rows, "cols": cols, "obs": obs, "chi2": chi, "df": df, "p": p_chi2(chi, df), "V": v, "low_expected": int((exp < 5).sum())}


def regression(data, dv, ivs):
    y = numeric(data[dv])
    xs = [numeric(data[c]) for c in ivs]
    ok = np.isfinite(y)
    for x in xs:
        ok &= np.isfinite(x)
    yv = y[ok]
    X = np.column_stack([np.ones(ok.sum())] + [x[ok] for x in xs])
    beta, *_ = np.linalg.lstsq(X, yv, rcond=None)
    n, k = X.shape
    res = yv - X @ beta
    sse, sst = float(res @ res), float(((yv - yv.mean()) ** 2).sum())
    mse = sse / (n - k)
    se = np.sqrt(np.diag(mse * np.linalg.inv(X.T @ X)))
    r2 = 1 - sse / sst
    f = ((sst - sse) / (k - 1)) / mse
    return {
        "names": ["(Constant)"] + list(ivs),
        "B": beta,
        "SE": se,
        "t": beta / se,
        "p": [p_t(t, n - k) for t in beta / se],
        "R2": r2,
        "adjR2": 1 - (1 - r2) * (n - 1) / (n - k),
        "F": f,
        "df": (k - 1, n - k),
        "pF": p_f(f, k - 1, n - k),
        "n": n,
    }


def alpha(data, items):
    m = np.column_stack([numeric(data[c]) for c in items])
    m = m[np.isfinite(m).all(1)]
    k = m.shape[1]
    return {"alpha": k / (k - 1) * (1 - m.var(0, ddof=1).sum() / m.sum(1).var(ddof=1)), "k": k, "n": m.shape[0]}


# ---------------------------------------------------------------- the report
def report(title, sentences, tables, dest):
    from docx import Document

    d = Document()
    d.add_heading(title, level=1)
    for s in sentences:
        d.add_paragraph(s)
    for name, head, rows in tables:
        d.add_paragraph(name).runs[0].italic = True
        t = d.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, h in enumerate(head):
            t.rows[0].cells[i].text = str(h)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = f"{v:.2f}" if isinstance(v, float) else str(v)
    d.save(dest)
    return dest


def parse(text, ctx):
    from ai_pc.apps.appschat import find_file

    f = find_file(text, ctx, {".xlsx", ".csv"})
    if not f:
        return None
    t = re.sub(r"\s+in\s+\S+\.(?:xlsx|csv)\b.*$", "", text, flags=re.I).strip()
    pats = (
        ("ttest", r"^\s*(?:an?\s+)?(?:independent\s+)?t-?test\s+(?:of|on)\s+(\w+)\s+by\s+(\w+)"),
        ("anova", r"^\s*(?:one-?way\s+)?anova\s+(?:of|on)\s+(\w+)\s+by\s+(\w+)"),
        ("corr", r"^\s*correlation\s+(?:between|of)\s+(\w+)\s+and\s+(\w+)"),
        ("cross", r"^\s*(?:crosstab|chi-?square(?:\s+test)?)\s+(?:of\s+)?(\w+)\s+(?:by|and|with)\s+(\w+)"),
        ("reg", r"^\s*(?:linear\s+)?regression\s+of\s+(\w+)\s+on\s+(.+)$"),
        ("alpha", r"^\s*(?:reliability|cronbach'?s?\s+alpha)\s+of\s+(.+)$"),
        ("describe", r"^\s*(?:describe|descriptives?|summary statistics|frequencies)\b"),
    )
    for kind, rx in pats:
        m = re.match(rx, t, re.I)
        if m:
            return {"op": kind, "file": f, "args": [g.strip() for g in m.groups()]}
    return None


def run(op, ctx):
    data = sheet(op["file"])
    out = Path(ctx["out"]) / "stats"
    out.mkdir(parents=True, exist_ok=True)
    k, a = op["op"], op["args"]
    if k == "describe":
        num, cat = describe(data)
        sents = [f"{r['var']}: M = {r['Mean']:.2f}, SD = {r['SD']:.2f} (N = {r['N']})." for r in num] + [
            f"{c['var']}: " + ", ".join(f"{v} {n} ({100 * n / c['N']:.1f}%)" for v, n in c["freq"][:6]) + "." for c in cat
        ]
        tabs = [
            (
                "Descriptive statistics",
                ["Variable", "N", "Mean", "SD", "Min", "Median", "Max", "Skewness"],
                [[r["var"], r["N"], r["Mean"], r["SD"], r["Min"], r["Median"], r["Max"], r["Skewness"]] for r in num],
            )
        ]
        dest = report("Descriptive statistics", sents, tabs, out / "descriptives.docx")
    elif k == "ttest":
        dv, iv = col(data, a[0]), col(data, a[1])
        r = ttest(data, dv, iv)
        sents = [
            f"An independent-samples t-test compared {dv} between {r['groups'][0]} (M = {r['mean'][0]:.2f}, SD = {r['sd'][0]:.2f}) and {r['groups'][1]} "
            f"(M = {r['mean'][1]:.2f}, SD = {r['sd'][1]:.2f}): t({r['df']}) = {r['t']:.2f}, {pfmt(r['p'])}, d = {r['d']:.2f}. "
            f"Without assuming equal variances (Welch): t({r['df_welch']:.1f}) = {r['t_welch']:.2f}, {pfmt(r['p_welch'])}."
        ]
        tabs = [("Group statistics", ["Group", "N", "Mean", "SD"], [[g, n, m, s] for g, n, m, s in zip(r["groups"], r["n"], r["mean"], r["sd"])])]
        dest = report(f"t-test: {dv} by {iv}", sents, tabs, out / f"ttest_{dv}_{iv}.docx")
    elif k == "anova":
        dv, iv = col(data, a[0]), col(data, a[1])
        r = anova(data, dv, iv)
        sents = [
            f"A one-way ANOVA found {'a' if r['p'] < 0.05 else 'no'} significant difference in {dv} between the {iv} groups: "
            f"F({r['df'][0]}, {r['df'][1]}) = {r['F']:.2f}, {pfmt(r['p'])}, eta squared = {r['eta2']:.2f}."
        ]
        tabs = [("Group statistics", ["Group", "N", "Mean", "SD"], [[g, n, m, s] for g, (n, m, s) in r["groups"].items()])]
        dest = report(f"ANOVA: {dv} by {iv}", sents, tabs, out / f"anova_{dv}_{iv}.docx")
    elif k == "corr":
        x, y = col(data, a[0]), col(data, a[1])
        r = correlation(data, x, y)
        sents = [f"{x} and {y} were {'positively' if r['r'] > 0 else 'negatively'} correlated, r({r['n'] - 2}) = {r['r']:.2f}, {pfmt(r['p'])}."]
        dest = report(f"Correlation: {x} and {y}", sents, [], out / f"correlation_{x}_{y}.docx")
    elif k == "cross":
        x, y = col(data, a[0]), col(data, a[1])
        r = crosstab(data, x, y)
        sents = [
            f"A chi-square test of independence: chi2({r['df']}, N = {int(r['obs'].sum())}) = {r['chi2']:.2f}, {pfmt(r['p'])}, Cramer's V = {r['V']:.2f}."
            + (f" {r['low_expected']} cell(s) have expected counts below 5: read the result with care." if r["low_expected"] else "")
        ]
        tabs = [(f"{x} by {y}", [x] + r["cols"], [[rw] + [int(v) for v in r["obs"][i]] for i, rw in enumerate(r["rows"])])]
        dest = report(f"Crosstab: {x} by {y}", sents, tabs, out / f"crosstab_{x}_{y}.docx")
    elif k == "reg":
        dv = col(data, a[0])
        ivs = [col(data, v.strip()) for v in re.split(r",|\band\b", a[1]) if v.strip()]
        r = regression(data, dv, ivs)
        sents = [
            f"A linear regression of {dv} on {', '.join(ivs)} explained {100 * r['R2']:.1f}% of its variance (R2 = {r['R2']:.2f}, adjusted R2 = "
            f"{r['adjR2']:.2f}), F({r['df'][0]}, {r['df'][1]}) = {r['F']:.2f}, {pfmt(r['pF'])}."
        ]
        tabs = [
            (
                "Coefficients",
                ["", "B", "SE", "t", "p"],
                [[nm, float(b), float(s), float(t), round(p, 3)] for nm, b, s, t, p in zip(r["names"], r["B"], r["SE"], r["t"], r["p"])],
            )
        ]
        dest = report(f"Regression: {dv}", sents, tabs, out / f"regression_{dv}.docx")
    else:
        items = [col(data, v.strip()) for v in re.split(r",|\band\b", a[0]) if v.strip()]
        r = alpha(data, items)
        judge = (
            "excellent"
            if r["alpha"] >= 0.9
            else "good"
            if r["alpha"] >= 0.8
            else "acceptable"
            if r["alpha"] >= 0.7
            else "questionable"
            if r["alpha"] >= 0.6
            else "poor"
        )
        sents = [f"Cronbach's alpha for the {r['k']} items was {r['alpha']:.2f} ({judge}), N = {r['n']}."]
        dest = report("Reliability", sents, [], out / "reliability.docx")
    return "\n".join(sents) + f"\nReport: {dest}"
