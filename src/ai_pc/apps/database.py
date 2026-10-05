"""Databases from spreadsheets: an Excel workbook or CSV becomes a Microsoft Access database (.accdb, written by DAO, the
engine Access itself uses; Access does not open) or a SQLite database (.db), one table per sheet with each column's
type found from its values; tables listed and read-only questions in SQL answered, with the answer saved to Excel.
Checks: every table holds as many rows as its sheet.

  'make an access database from sales.xlsx'   'sqlite database from customers.csv'   'tables in sales.accdb'
  'query sales.accdb: SELECT Customer, SUM(Amount) AS Total FROM Sales GROUP BY Customer'
"""
import csv
import datetime as dt
import re
import sqlite3
from pathlib import Path

NAME, LABEL = "database", "Databases: Excel/CSV to Access (.accdb) or SQLite; questions in SQL"
EXAMPLES = ["make an access database from sales.xlsx", "query sales.accdb: SELECT Customer, SUM(Amount) AS Total FROM Sales GROUP BY Customer"]
DB_LANG, DB_V120 = ";LANGID=0x0409;CP=1252;COUNTRY=0", 128


def sheets_of(path):
    """{table name: (headers, rows)} from an .xlsx (every sheet) or a .csv."""
    p = Path(path)
    out = {}
    if p.suffix.lower() == ".csv":
        with p.open(encoding="utf-8-sig", newline="") as f:
            rows = [r for r in csv.reader(f) if any(c.strip() for c in r)]
        out[_name(p.stem)] = (rows[0], rows[1:]) if rows else ([], [])
        return out
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
        if rows:
            out[_name(ws.title)] = ([str(h or f"Column{i + 1}") for i, h in enumerate(rows[0])], rows[1:])
    return out


def _name(s):
    return re.sub(r"\W+", "_", str(s)).strip("_")[:60] or "Table1"


def kinds(headers, rows):
    """Each column's type from its values: number, date or text."""
    out = []
    for i in range(len(headers)):
        vals = [r[i] for r in rows if i < len(r) and r[i] not in (None, "")]
        if vals and all(isinstance(v, (int, float)) or re.fullmatch(r"-?[\d,]+(?:\.\d+)?", str(v)) for v in vals):
            out.append("number")
        elif vals and all(isinstance(v, (dt.date, dt.datetime)) for v in vals):
            out.append("date")
        else:
            out.append("text")
    return out


def _value(v, kind):
    if v in (None, ""):
        return None
    if kind == "number":
        return float(str(v).replace(",", "")) if not isinstance(v, (int, float)) else v
    if kind == "date":
        return v
    return str(v)


def to_sqlite(src, dest):
    data = sheets_of(src)
    cx = sqlite3.connect(dest)
    for t, (heads, rows) in data.items():
        ks = kinds(heads, rows)
        cols = [_name(h) for h in heads]
        cx.execute(f'DROP TABLE IF EXISTS "{t}"')
        cx.execute(f'CREATE TABLE "{t}" (' + ", ".join(f'"{c}" {dict(number="REAL", date="TEXT", text="TEXT")[k]}' for c, k in zip(cols, ks)) + ")")
        cx.executemany(f'INSERT INTO "{t}" VALUES ({",".join("?" * len(cols))})',
                       [[(_value(r[i], k).isoformat() if k == "date" and _value(r[i], k) else _value(r[i], k)) if i < len(r) else None
                         for i, k in enumerate(ks)] for r in rows])
    cx.commit()
    cx.close()
    return data


def to_access(src, dest):
    import win32com.client
    data = sheets_of(src)
    dest = Path(dest)
    if dest.exists():
        dest.unlink()
    eng = win32com.client.Dispatch("DAO.DBEngine.120")
    db = eng.CreateDatabase(str(dest), DB_LANG, DB_V120)
    try:
        for t, (heads, rows) in data.items():
            ks = kinds(heads, rows)
            cols = [_name(h) for h in heads]
            db.Execute(f"CREATE TABLE [{t}] (" + ", ".join(f"[{c}] {dict(number='DOUBLE', date='DATETIME', text='LONGTEXT')[k]}" for c, k in zip(cols, ks)) + ")")
            rs = db.OpenRecordset(t, 2)  # dbOpenDynaset
            for r in rows:
                rs.AddNew()
                for i, (c, k) in enumerate(zip(cols, ks)):
                    v = _value(r[i], k) if i < len(r) else None
                    if v is not None:
                        rs.Fields.Item(c).Value = v
                rs.Update()
            rs.Close()
    finally:
        db.Close()
    return data


def query(path, sql, limit=1000):
    """(headers, rows) of a read-only question (SELECT only)."""
    if not re.match(r"^\s*select\b", sql, re.I) or re.search(r";\s*\S", sql) or re.search(r"\b(?:insert|update|delete|drop|alter|create)\b", sql, re.I):
        raise ValueError("only a single SELECT question is run here (nothing that changes the data)")
    p = Path(path)
    if p.suffix.lower() in (".accdb", ".mdb"):
        import win32com.client
        db = win32com.client.Dispatch("DAO.DBEngine.120").OpenDatabase(str(p), False, True)  # read only
        try:
            rs = db.OpenRecordset(sql, 4)  # dbOpenSnapshot
            heads = [rs.Fields.Item(i).Name for i in range(rs.Fields.Count)]
            rows = []
            while not rs.EOF and len(rows) < limit:
                rows.append([rs.Fields.Item(i).Value for i in range(rs.Fields.Count)])
                rs.MoveNext()
            rs.Close()
        finally:
            db.Close()
        return heads, rows
    cx = sqlite3.connect(f"file:{p}?mode=ro", uri=True)
    try:
        cur = cx.execute(sql)
        return [d[0] for d in cur.description], [list(r) for r in cur.fetchmany(limit)]
    finally:
        cx.close()


def tables(path):
    p = Path(path)
    if p.suffix.lower() in (".accdb", ".mdb"):
        import win32com.client
        db = win32com.client.Dispatch("DAO.DBEngine.120").OpenDatabase(str(p), False, True)
        try:
            names = [db.TableDefs(i).Name for i in range(db.TableDefs.Count) if not db.TableDefs(i).Name.startswith("MSys")]
            counts = {}
            for n in names:  # the recordset is held while its field is read (a released one invalidates its fields)
                rs = db.OpenRecordset(f"SELECT COUNT(*) FROM [{n}]", 4)
                counts[n] = rs.Fields.Item(0).Value
                rs.Close()
            return counts
        finally:
            db.Close()
    cx = sqlite3.connect(f"file:{p}?mode=ro", uri=True)
    try:
        return {n: cx.execute(f'SELECT COUNT(*) FROM "{n}"').fetchone()[0] for (n,) in cx.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    finally:
        cx.close()


def parse(text, ctx):
    from ai_pc.apps.appschat import find_file
    c = text.lower()
    m = re.match(r"^\s*(?:query|ask|sql)\s+(.+?):\s*(select\b.+)$", text, re.I | re.S)
    if m:
        f = find_file(m.group(1), ctx, {".accdb", ".mdb", ".db", ".sqlite"}) or (m.group(1).strip() if Path(m.group(1).strip()).exists() else None)
        return {"op": "query", "file": f, "sql": m.group(2).strip()} if f else None
    if re.search(r"\btables?\s+in\b", c):
        f = find_file(text, ctx, {".accdb", ".mdb", ".db", ".sqlite"})
        return {"op": "tables", "file": f} if f else None
    if re.search(r"\b(?:access|sqlite|accdb)\b.*\bdatabase\b|\bdatabase\b.*\bfrom\b", c):
        f = find_file(text, ctx, {".xlsx", ".csv"})
        return {"op": "make", "file": f, "kind": "sqlite" if "sqlite" in c else "access"} if f else None
    return None


def run(op, ctx):
    out = Path(ctx["out"]) / "databases"
    out.mkdir(parents=True, exist_ok=True)
    if op["op"] == "make":
        stem = Path(op["file"]).stem
        dest = out / f"{stem}.{'accdb' if op['kind'] == 'access' else 'db'}"
        data = (to_access if op["kind"] == "access" else to_sqlite)(op["file"], dest)
        got = tables(dest)
        ok = all(got.get(t) == len(rows) for t, (_, rows) in data.items())
        return (f"{'Access' if op['kind'] == 'access' else 'SQLite'} database: {dest} with " + ", ".join(f"{t} ({len(r)} rows)" for t, (_, r) in data.items()) +
                (" (checked: every table holds every row)." if ok else f" NOT right: {got}."))
    if op["op"] == "tables":
        return "Tables: " + ", ".join(f"{t} ({n} rows)" for t, n in tables(op["file"]).items()) + "."
    heads, rows = query(op["file"], op["sql"])
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(heads)
    for r in rows:
        ws.append([v if not hasattr(v, "isoformat") else str(v) for v in r])
    dest = out / "answer.xlsx"
    wb.save(dest)
    show = "\n".join(" | ".join(str(v) for v in r) for r in rows[:15])
    return f"{len(rows)} rows ({', '.join(heads)}):\n{show}\nSaved to {dest}."
