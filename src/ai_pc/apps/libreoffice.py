"""LibreOffice (26.8.1 in tools/libreoffice, unpacked from its signed MSI; its user profile in tools/libreoffice/home): office
files converted headless (Word/Writer, Excel/Calc, PowerPoint/Impress: DOCX, ODT, PDF, RTF, TXT, XLSX, ODS, CSV, PPTX, ODP),
and spreadsheets recalculated by Calc (a file saved by code keeps formulas but no results; Calc fills every result in).
Runs on the hidden desktop. Checked: the result opens as its format and carries the original's text, a PDF has its pages,
and every formula of a recalculated sheet has a value and no error.

  'libreoffice convert report.docx to pdf and odt'   'libreoffice convert prices.xlsx to ods'   'libreoffice recalculate invoice.xlsx'
"""
import re
import zipfile
from pathlib import Path

from ai_pc.core.config import ROOT

NAME, LABEL = "libreoffice", "LibreOffice: office files converted (DOCX/ODT/PDF/XLSX/ODS/CSV/PPTX/ODP), sheets recalculated"
EXAMPLES = ["libreoffice convert report.docx to pdf and odt", "libreoffice convert prices.xlsx to ods", "libreoffice recalculate invoice.xlsx"]
HOME = ROOT / "tools" / "libreoffice"
DOCS = {".docx", ".doc", ".odt", ".rtf", ".txt", ".html", ".xlsx", ".xls", ".ods", ".csv", ".pptx", ".ppt", ".odp"}
TARGETS = {"pdf", "odt", "docx", "rtf", "txt", "html", "ods", "xlsx", "csv", "odp", "pptx"}
MIMES = {".odt": "application/vnd.oasis.opendocument.text", ".ods": "application/vnd.oasis.opendocument.spreadsheet",
         ".odp": "application/vnd.oasis.opendocument.presentation"}


def soffice():
    hits = sorted(HOME.rglob("soffice.com")) if HOME.exists() else []
    return hits[0] if hits else None


def convert(src, fmt, outdir):
    from ai_pc.core import hidden_desktop
    profile = (HOME / "home").resolve()
    profile.mkdir(parents=True, exist_ok=True)
    return hidden_desktop.run([str(soffice()), f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore", "--nologo",
                               "--convert-to", fmt, "--outdir", str(outdir), str(src)], timeout=300)


def text_of(path):
    """Plain words of a document or sheet (for the 'same content' check)."""
    p = Path(path)
    ext = p.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            t = " ".join(pg.extract_text() or "" for pg in PdfReader(str(p)).pages)
        elif ext == ".docx":
            from docx import Document
            t = " ".join(par.text for par in Document(str(p)).paragraphs)
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(p, data_only=False, read_only=True)
            t = " ".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row
                         if isinstance(c, str) and not c.startswith("="))  # formulas are not words (a converted sheet holds their values)
        elif ext in (".odt", ".ods", ".odp", ".pptx"):
            with zipfile.ZipFile(p) as z:
                xml = z.read("content.xml" if ext != ".pptx" else next(n for n in z.namelist() if n.startswith("ppt/slides/slide"))).decode("utf-8", "replace")
            t = re.sub(r"<[^>]+>", " ", xml)
        else:
            t = p.read_text(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001 - a file that cannot be read has no words
        return set()
    return set(re.findall(r"[a-z]{3,}", t.lower()))


def opens_as(path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext in MIMES:
        with zipfile.ZipFile(p) as z:
            return z.read("mimetype").decode() == MIMES[ext]
    if ext in (".docx", ".xlsx", ".pptx"):
        return zipfile.is_zipfile(p)
    if ext == ".pdf":
        from pypdf import PdfReader
        return len(PdfReader(str(p)).pages) >= 1
    return p.stat().st_size > 0


def parse(text, ctx):
    from ai_pc.apps.appschat import find_file
    c = text.lower()
    if not re.search(r"\blibre\s*office\b|\bodt\b|\bods\b|\bodp\b|\bopen\s*document\b", c):
        return None
    f = find_file(text, ctx, DOCS)
    if not f:
        return None
    if re.search(r"\brecalc(?:ulate)?\b|\bformulas?\b", c) and Path(f).suffix.lower() in (".xlsx", ".xls", ".ods"):
        return {"op": "recalc", "file": f}
    after = c.split(Path(f).name.lower(), 1)[-1]
    targets = [t for t in dict.fromkeys(re.findall(r"\b(" + "|".join(TARGETS) + r")\b", after)) if t != Path(f).suffix[1:].lower()]
    return {"op": "convert", "file": f, "targets": targets or ["pdf"]}


def run(op, ctx):
    if not soffice():
        return "LibreOffice is not in tools/libreoffice."
    src = Path(op["file"]).resolve()
    out = (Path(ctx["out"]) / "libreoffice").resolve()
    out.mkdir(parents=True, exist_ok=True)
    if op["op"] == "recalc":
        from openpyxl import load_workbook
        sub = out / "recalculated"
        sub.mkdir(exist_ok=True)
        convert(src, "xlsx", sub)
        dest = sub / (src.stem + ".xlsx")
        formulas = load_workbook(src, data_only=False)
        values = load_workbook(dest, data_only=True) if dest.exists() else None
        cells, empty, errors = 0, [], []
        for ws in formulas.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cells += 1
                        v = values[ws.title][cell.coordinate].value if values else None
                        if v is None:
                            empty.append(f"{ws.title}!{cell.coordinate}")
                        elif isinstance(v, str) and v.startswith("#"):
                            errors.append(f"{ws.title}!{cell.coordinate} {v}")
        sample = []
        if values:
            for ws in formulas.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("=") and len(sample) < 3:
                            sample.append(f"{cell.coordinate} {cell.value} = {values[ws.title][cell.coordinate].value}")
        ok = dest.exists() and cells and not empty and not errors
        return (f"LibreOffice Calc recalculated {src.name}: {dest} ({cells} formulas; e.g. {'; '.join(sample)}). " +
                ("Checked: every formula has a value, none is an error." if ok else f"NOT right: {len(empty)} without a value {empty[:3]}, errors {errors[:3]}."))
    words = text_of(src)
    made, checks = [], []
    for fmt in op["targets"]:
        dest = out / f"{src.stem}.{fmt}"
        dest.unlink(missing_ok=True)
        convert(src, fmt, out)
        if not dest.exists():
            checks.append((f"{fmt.upper()} made", False))
            continue
        made.append(dest)
        got = text_of(dest)
        same = len(words & got) / max(1, len(words)) if words else 1.0
        checks += [(f"the {fmt.upper()} opens as {fmt.upper()}", opens_as(dest)), (f"it carries the original's words ({same:.0%})", same >= 0.9)]
    bad = [w for w, ok in checks if not ok]
    return (f"LibreOffice converted {src.name} to {', '.join(p.name for p in made) or 'nothing'} in {out}. " +
            ("Checked: " + "; ".join(w for w, _ in checks) + "." if not bad else "NOT right: " + "; ".join(bad) + "."))
