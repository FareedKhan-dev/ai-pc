"""Text from pictures and scanned PDFs with Windows' own OCR engine (Windows.Media.Ocr; nothing to install): each page
read line by line, saved as text (or a Word file), and the things a receipt or bill usually holds picked out (total,
date, NTN, phone). Small pictures are enlarged first (the engine reads text about 20+ pixels high best); PDF pages are
drawn at 200 dpi.

  'read the text in receipt.jpg'   'ocr scan.pdf to word'   'what is the total on bill.png?'
"""

import json
import re
import subprocess
from pathlib import Path

from PIL import Image

NAME, LABEL = "ocr", "Text from pictures and scans (Windows OCR)"
EXAMPLES = ["read the text in receipt.jpg", "ocr scan.pdf to word", "what is the total on bill.png?"]
PS = Path(__file__).parent / "ps" / "ocr.ps1"
PICTURES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".gif", ".webp"}
AMOUNT = r"(?:rs\.?|pkr)?\s*([\d,]+(?:\.\d{1,2})?)"


def read_picture(path, lang="en-US", work=None):
    """One picture -> {"text", "lines": [{"text", "words"}], "angle"}; enlarged first when small."""
    p = Path(path)
    im = Image.open(p)
    src = p
    if p.suffix.lower() not in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"} or max(im.size) < 1400:
        scale = max(1.0, 1800 / max(im.size))
        work = Path(work or p.parent)
        work.mkdir(parents=True, exist_ok=True)
        src = work / f"{p.stem}.ocr.png"
        im.convert("RGB").resize((int(im.width * scale), int(im.height * scale)), Image.LANCZOS).save(src)
    r = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(PS), "-Path", str(src.resolve()), "-Lang", lang],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=120,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    if r.returncode != 0 or not r.stdout.strip():
        raise RuntimeError(f"Windows OCR failed: {(r.stderr or r.stdout).strip()[:300]}")
    out = json.loads(r.stdout.strip().splitlines()[-1])
    out["lines"] = out.get("lines") or []
    return out


def rows(lines):
    """Lines on the same row joined left to right: the engine gives a receipt's 'Grand Total' and its amount as two lines."""
    items = []
    for ln in lines:
        ws = ln.get("words") or []
        if not ws:
            continue
        items.append((sum(w["y"] + w["h"] / 2 for w in ws) / len(ws), max(w["h"] for w in ws), min(w["x"] for w in ws), ln["text"]))
    items.sort()
    out = []
    for y, h, x, t in items:
        if out and abs(out[-1]["y"] - y) < 0.6 * max(h, out[-1]["h"]):
            out[-1]["parts"].append((x, t))
        else:
            out.append({"y": y, "h": h, "parts": [(x, t)]})
    return ["  ".join(t for _, t in sorted(r["parts"])) for r in out]


def read(path, lang="en-US", work=None):
    """A picture or a PDF -> {"pages": [text], "text"}."""
    p = Path(path)
    pages = []
    if p.suffix.lower() == ".pdf":
        import pypdfium2 as pdfium

        work = Path(work or p.parent)
        work.mkdir(parents=True, exist_ok=True)
        doc = pdfium.PdfDocument(str(p))
        for i in range(len(doc)):
            img = work / f"{p.stem}.p{i + 1}.png"
            doc[i].render(scale=200 / 72).to_pil().save(img)
            pages.append("\n".join(rows(read_picture(img, lang, work)["lines"])))
        doc.close()
    else:
        pages.append("\n".join(rows(read_picture(p, lang, work)["lines"])))
    return {"pages": pages, "text": "\n\n".join(pages)}


def fields(text):
    """What a receipt or bill usually holds: total, date, NTN, phone."""
    out = {}
    totals = []
    for line in text.splitlines():
        if re.search(r"\b(?:grand total|total|net payable|amount due|balance due|net amount|payable)\b", line, re.I) and not re.search(
            r"\bsub ?-?total\b|\btotal (?:items|qty|quantity)\b", line, re.I
        ):
            for m in re.finditer(AMOUNT, line, re.I):
                v = m.group(1).replace(",", "")
                if re.fullmatch(r"\d+(?:\.\d+)?", v) and float(v) > 0:
                    totals.append(float(v))
    if totals:
        out["total"] = max(totals)
    m = re.search(r"\b(\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}|\d{1,2}\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{4})\b", text, re.I)
    if m:
        out["date"] = m.group(1)
    m = re.search(r"\bntn\b\D{0,5}(\d{7}-?\d?)", text, re.I)
    if m:
        out["ntn"] = m.group(1)
    m = re.search(r"((?:\+92|0)3\d{2}[- ]?\d{7}|0\d{2,3}[- ]?\d{6,8})", text)
    if m:
        out["phone"] = m.group(1)
    return out


def parse(text, ctx):
    from ai_pc.apps.appschat import find_file

    c = text.lower()
    if not re.search(
        r"\bocr\b|\b(?:read|extract|get|copy|pull)\b.*\b(?:text|words|writing)\b|\btext (?:from|in|out of|on)\b|\b(?:what is|what's) the total\b"
        r"|\btotal on\b",
        c,
    ):
        return None
    f = find_file(text, ctx, PICTURES | {".pdf"})
    if not f:
        return None
    return {
        "op": "ocr",
        "file": f,
        "to": "docx" if re.search(r"\b(?:word|docx)\b", c) else "txt",
        "fields": bool(re.search(r"\btotal\b|\breceipt\b|\bbill\b|\binvoice\b", c)),
    }


def run(op, ctx):
    out = Path(ctx["out"])
    r = read(op["file"], work=out / "ocr_work")
    stem = Path(op["file"]).stem
    words = len(r["text"].split())
    if not words:
        return f"No text found in {Path(op['file']).name} (is it a photo of text, upright and in focus?)."
    if op.get("to") == "docx":
        from docx import Document

        d = Document()
        for i, page in enumerate(r["pages"], 1):
            if len(r["pages"]) > 1:
                d.add_heading(f"Page {i}", level=2)
            for line in page.splitlines():
                d.add_paragraph(line)
        dest = out / f"{stem}.text.docx"
        d.save(dest)
        back = "\n".join(p.text for p in Document(dest).paragraphs)
        ok = all(w in back for w in r["text"].split()[:20])
    else:
        dest = out / f"{stem}.text.txt"
        dest.write_text(r["text"], encoding="utf-8")
        ok = dest.read_text(encoding="utf-8") == r["text"]
    reply = (
        f"Read {words} words from {Path(op['file']).name}"
        + (f" ({len(r['pages'])} pages)" if len(r["pages"]) > 1 else "")
        + f"; saved to {dest} ({'checked' if ok else 'NOT the same when read back'})."
    )
    if op.get("fields"):
        f = fields(r["text"])
        if f:
            reply += " Found: " + ", ".join(f"{k} {v:,.2f}" if isinstance(v, float) else f"{k} {v}" for k, v in f.items()) + "."
    first = " / ".join(r["text"].splitlines()[:3])
    return reply + (f" It starts: {first[:160]}" if first else "")
