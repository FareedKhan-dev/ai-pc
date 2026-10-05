"""A user's folders as they really look, made inside the project for tests (never the real folders): a messy Downloads
with photos (EXIF dates in August and September; a WhatsApp one dated only by its name), a screenshot, a CV, invoices,
PDFs, a spreadsheet, old installers and a new one, unfinished downloads (one abandoned, one still downloading), the
same photo and PDF downloaded twice, an empty folder, a zip, a short video and a song; and an empty Documents and
Pictures."""
import datetime as dt
import os
import shutil
import subprocess
from pathlib import Path

from ai_pc.core.paths import ROOT


def _when(path, days_ago):
    t = (dt.datetime.now() - dt.timedelta(days=days_ago)).timestamp()
    os.utime(path, (t, t))


def _photo(path, w, h, colour, taken=None):
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (w, h), colour)
    d = ImageDraw.Draw(im)
    d.rectangle([w // 4, h // 4, 3 * w // 4, 3 * h // 4], outline="white", width=6)
    ex = Image.Exif()
    if taken:
        ex[306] = taken.strftime("%Y:%m:%d %H:%M:%S")
        ex.get_ifd(0x8769)[36867] = taken.strftime("%Y:%m:%d %H:%M:%S")
    im.save(path, "JPEG", quality=90, exif=ex)


def build(base):
    base = Path(base)
    shutil.rmtree(base, ignore_errors=True)
    dl, docs, pics = base / "Downloads", base / "Documents", base / "Pictures"
    for d in (dl, docs, pics):
        d.mkdir(parents=True)
    y = dt.datetime.now().year
    shots = [(f"IMG_{y}0812_101500.jpg", dt.datetime(y, 8, 12, 10, 15), "#2E7D32"), (f"IMG_{y}0819_183000.jpg", dt.datetime(y, 8, 19, 18, 30), "#1565C0"),
             (f"IMG_{y}0903_090500.jpg", dt.datetime(y, 9, 3, 9, 5), "#C62828"), (f"IMG_{y}0915_164200.jpg", dt.datetime(y, 9, 15, 16, 42), "#6A1B9A"),
             (f"WhatsApp Image {y}-08-25 at 20.11.03.jpeg", None, "#EF6C00")]  # WhatsApp strips EXIF: its date is only in its name
    for name, when, colour in shots:
        _photo(dl / name, 2400, 1600, colour, when)
        _when(dl / name, 20)
    shutil.copy2(dl / f"IMG_{y}0812_101500.jpg", dl / f"IMG_{y}0812_101500 (1).jpg")  # downloaded twice
    from PIL import Image
    Image.new("RGB", (1366, 768), "#ECEFF1").save(dl / "Screenshot 2026-09-12 104455.png")
    import docx
    d = docx.Document()
    d.add_heading("Fareed Hassan Khan", 0)
    d.add_paragraph("Data analyst. Python, Excel, Power BI.")
    d.save(dl / "CV - Fareed Hassan.docx")
    for i in (1, 2):
        d = docx.Document()
        d.add_heading(f"Invoice INV-00{i}", 1)
        d.add_paragraph(f"Ali Traders. Amount due: Rs {i * 25000:,}.")
        d.save(dl / f"Invoice INV-00{i}.docx")
    pdfs = [p for p in (ROOT / "out" / "docs").rglob("*.pdf") if "_tests" not in p.parts and p.stat().st_size < 600_000][:2]
    for name, src_pdf in zip(("Bank statement September.pdf", "Electricity bill Aug.pdf"), pdfs):
        shutil.copy2(src_pdf, dl / name)
    shutil.copy2(dl / "Electricity bill Aug.pdf", dl / "Electricity bill Aug (1).pdf")  # the same bill twice
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["Month", "Spend"])
    wb.active.append(["August", 42000])
    wb.save(dl / "budget.xlsx")
    for name, days in (("vlc-3.0.21-win64.exe", 120), ("python-3.12.4-amd64.exe", 95), ("ZoomInstaller.exe", 3)):
        (dl / name).write_bytes(b"MZ" + name.encode() + b"\0" * 2048)  # a stand-in, not a program
        _when(dl / name, days)
    (dl / "movie.mp4.crdownload").write_bytes(b"\0" * 4096)  # still downloading (it changed just now)
    (dl / "big-file.zip.part").write_bytes(b"\0" * 1024)  # given up on two days ago
    _when(dl / "big-file.zip.part", 2)
    with __import__("zipfile").ZipFile(dl / "photos-from-ahmed.zip", "w") as z:
        z.writestr("ahmed/IMG_1.jpg", (dl / f"IMG_{y}0903_090500.jpg").read_bytes())
        z.writestr("ahmed/notes.txt", "From the wedding.")
    (dl / "New folder").mkdir()
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    subprocess.run(["ffmpeg", "-y", "-v", "error", "-f", "lavfi", "-i", "testsrc=duration=3:size=320x240:rate=25", "-f", "lavfi", "-i", "sine=frequency=440:duration=3",
                    "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", str(dl / "clip from phone.mp4")], check=True, creationflags=flags)
    subprocess.run(["ffmpeg", "-y", "-v", "error", "-f", "lavfi", "-i", "sine=frequency=330:duration=4", "-c:a", "libmp3lame", str(dl / "song.mp3")], check=True,
                   creationflags=flags)
    (dl / "notes.txt").write_text("Call the bank on Monday.\n", encoding="utf-8")
    return {"root": base, "downloads": dl, "documents": docs, "pictures": pics}


if __name__ == "__main__":
    b = build(ROOT / "out" / "_tests" / "win" / "sandbox")
    print(b, sorted(p.name for p in b["downloads"].iterdir()))
