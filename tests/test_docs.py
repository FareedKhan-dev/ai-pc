"""The Word document engine, offline (no model calls): plans built into .docx, rendered by Word in the background, and
checked on the rendered pages.

  .venv\\Scripts\\python.exe tests\\test_docs.py
About a minute: each document is one Word render (Word runs hidden; open documents are never touched).
"""
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")

from harness.office import docplan as DP  # noqa: E402
from harness.office import docx_build as DB  # noqa: E402
from harness.office import render as RN  # noqa: E402
from harness.office import verify as VF  # noqa: E402

OUT = ROOT / "out" / "docs" / "_tests"
FAILS = []


def check(name, ok, detail=""):
    print(("ok   " if ok else "FAIL ") + name + (f"  ({detail})" if detail and not ok else ""))
    if not ok:
        FAILS.append(name)


def make(name, plan):
    rp = DP.resolve(plan)
    path = OUT / f"{name}.docx"
    DB.build(rp, path)
    r = RN.to_pdf(path, update_fields=True, save=True)
    rep = VF.check(rp, {}, r) if r.get("ok") else {"counts": {"fail": 1}, "results": [], "facts": {}}
    return rp, r, rep


def bad(rep):
    return [f"{x['id']}: {x['why'][:90]}" for x in rep["results"] if x["status"] in ("warn", "fail")]


def deck_and_book():
    from harness.office import pptx_build as PB
    from harness.office import xlsx_build as XB
    deck = {"title": "Solar for Lahore Homes", "subtitle": "Costs and savings", "theme": "modern", "meta": {"author": "Test", "org": "Test"},
            "slides": [{"layout": "bullets", "title": "Bills doubled since 2022", "bullets": ["Tariffs rose", {"text": "Peak bills **Rs 40,000**", "items": ["Slab rates"]}]},
                       {"layout": "stats", "title": "At a glance", "stats": [{"value": "Rs 9–11 lakh", "label": "installed cost"}, {"value": "3.5 yrs", "label": "payback"}]},
                       {"layout": "chart", "title": "Savings peak in summer", "chart": {"chart": "column", "categories": ["Jan", "May", "Jul"],
                                                                                        "series": [{"name": "Saving", "values": [18000, 34000, 32000]}]},
                        "takeaway": "Summer pays most of it back."},
                       {"layout": "bullets", "title": "An overloaded slide with far too much text on it",
                        "bullets": [("A long bullet with many words that wraps over several lines " * 2).strip()] * 9},
                       {"layout": "closing", "title": "Thank you"}]}
    rd = PB.resolve_deck(deck)
    check("a deck starts with a title slide", rd["slides"][0]["layout"] == "title")
    PB.build_deck(rd, OUT / "deck.pptx")
    png = OUT / "deck_png"
    png.mkdir(exist_ok=True)
    for f in png.glob("*.png"):
        f.unlink()
    r = RN.to_pdf(OUT / "deck.pptx", measure=True, png_dir=str(png), png_width=800)
    check("PowerPoint renders the deck (PDF, slide images, measurements)", r.get("ok") and r.get("slides") == 6
          and len(list(png.glob("*.png"))) == 6 and r.get("shapes"), r.get("error"))
    rep = VF.check_deck(rd, {}, r)
    flagged = [x["page"] for x in rep["results"] if x.get("fix", {}).get("op") == "shrink"]
    check("only the overloaded slide can be flagged as overflowing", set(flagged) <= {5},
          "; ".join(x["why"] for x in rep["results"] if x["status"] != "pass"))
    ok_slides = [x for x in rep["results"] if x["type"] == "slide" and x["page"] in (2, 3, 4)]
    check("the normal slides fit their boxes", ok_slides and all(x["status"] == "pass" for x in ok_slides), "; ".join(x["why"] for x in ok_slides))
    book = {"title": "Sales", "sheets": [
        {"name": "Sales", "kind": "table", "columns": [{"name": "Region"}, {"name": "Qty", "type": "integer"}, {"name": "Price", "type": "money"},
                                                       {"name": "Amount", "type": "money", "formula": "[Qty] * [Price]"},
                                                       {"name": "Result", "formula": 'IF([Amount] >= 10000, "big", "small")'}],
         "rows": [["North", 3, "4,500"], ["South", 5, 2200], ["North", 2, 3100]], "total": {"Qty": "sum", "Amount": "sum"}},
        {"name": "Summary", "kind": "summary", "metrics": [{"label": "Total", "fn": "sum", "of": "Amount"},
                                                           {"label": "North", "fn": "sum", "of": "Amount", "where": {"Region": "North"}}],
         "groups": [{"title": "By region", "by": "Region", "values": [{"of": "Amount", "fn": "sum"}], "chart": "column"}]}]}
    rb = XB.resolve_book(book)
    built = XB.build_book(rb, OUT / "book.xlsx")
    check("the workbook's formulas are written by code (no unknown columns)", not built["problems"], str(built["problems"]))
    r = RN.to_pdf(OUT / "book.xlsx", save=True, read=[{"sheet": c["sheet"], "cell": c["cell"]} for c in built["checks"]])
    rep = VF.check_book(rb, built, r)
    check("Excel recalculates with no error cells and its totals equal Python's", rep["counts"]["fail"] == 0 and len(built["checks"]) >= 5,
          "; ".join(x["why"] for x in rep["results"]))
    import openpyxl
    ws = openpyxl.load_workbook(OUT / "book.xlsx")["Sales"]
    check("computed columns are live formulas", str(ws["D2"].value) == "=B2*C2" and str(ws["E2"].value).startswith("=IF(D2>=10000"),
          f"{ws['D2'].value} {ws['E2'].value}")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    para = "Solar panels on homes in Lahore now pay for themselves in about four years, because tariffs rose while panel prices fell. "
    long_rows = [[f"Item {i}", f"Description of item number {i}", str(i * 3), f"{i * 1250:,}"] for i in range(1, 71)]
    report = {"doctype": "report", "title": "Engine Test Report", "subtitle": "Every block on real pages", "theme": "corporate",
              "meta": {"author": "Test", "org": "AI PC"}, "target_pages": 5, "blocks": [
                  {"type": "heading", "text": "Summary", "level": 1}, {"type": "paragraph", "text": para * 3, "style": "lead"},
                  {"type": "bullets", "numbered": True, "items": ["One", "Two", "Three"]},
                  {"type": "paragraph", "text": "A second list:"}, {"type": "bullets", "numbered": True, "items": ["Again one", "Again two"]},
                  {"type": "heading", "text": "Costs", "level": 1}, {"type": "paragraph", "text": para * 2},
                  {"type": "table", "caption": "Costs", "columns": ["Size", "Panels", "Total (Rs.)"],
                   "rows": [["3 kW", "Rs. 330,000", "600,000"], ["5 kW", "Rs. 550,000", "930,000"]], "total": True},
                  {"type": "heading", "text": "Savings", "level": 1},
                  {"type": "chart", "chart": "line", "title": "Monthly savings (Rs.)", "categories": ["Jan", "Feb", "Mar", "Apr"],
                   "series": [{"name": "Savings", "values": [18000, 21000, 26000, 31000]}], "caption": "Savings by month"},
                  {"type": "heading", "text": "Inventory", "level": 1},
                  {"type": "table", "caption": "A long table", "columns": ["Code", "Description", "Qty", "Value"], "rows": long_rows},
                  {"type": "heading", "text": "Closing", "level": 1}, {"type": "paragraph", "text": para * 2}]}
    rp, r, rep = make("report", report)
    check("report renders", r.get("ok"), r.get("error"))
    check("report: every check passes", not bad(rep), "; ".join(bad(rep)))
    tbl = next(b for b in rp["blocks"] if b["type"] == "table")
    check("table total summed by code", tbl.get("total_row") and tbl["total_row"][1] == "Rs. 880,000" and tbl["total_row"][2] == "1,530,000",
          str(tbl.get("total_row")))
    long_t = next(x for x in rep["results"] if x["type"] == "table" and "repeated" in x["why"])
    check("a 70-row table repeats its header on every page", long_t["status"] == "pass", long_t["why"])
    toc = next((x for x in rep["results"] if x["id"] == "toc"), None)
    check("contents page numbers match the headings' pages", toc and toc["status"] == "pass", toc and toc["why"])
    with zipfile.ZipFile(OUT / "report.docx") as z:
        check("the chart is a native Word chart with its data", any(n.startswith("word/charts/chart") for n in z.namelist())
              and any(n.startswith("word/embeddings/") and n.endswith(".xlsx") for n in z.namelist()))
    txt = " ".join(RN.text(r["pdf"]))
    check("the second numbered list starts again at 1", "1.\tAgain one" in txt or "1. Again one" in txt or "1.Again one" in txt or
          ("Again one" in txt and txt[txt.index("Again one") - 4:txt.index("Again one")].strip().startswith("1")), txt[txt.find("Again one") - 8:txt.find("Again one")])
    fonts = VF.glyph_fonts(r["pdf"])
    check("every glyph in the theme's fonts", all(VF._is_family(f, {"calibri"}) or VF._is_family(f, VF.OK_FONTS) for f in fonts), str(fonts))

    inv = {"doctype": "invoice", "theme": "corporate", "meta": {"org": "Bright Pixels", "number": "INV-1", "bill_to": "Ali Traders\n45 Mall Road"},
           "blocks": [{"type": "invoice_items", "currency": "Rs.", "tax_rate": "18%", "discount": "5%", "items": [
               {"desc": "Logo", "qty": 1, "price": 25000}, {"desc": "Posts", "qty": 10, "price": 2500}, {"desc": "Photos", "qty": 3.5, "price": 4000}]}]}
    rp, r, rep = make("invoice", inv)
    c = next(b for b in rp["blocks"] if b["type"] == "invoice_items")["computed"]
    # 25,000 + 25,000 + 14,000 = 64,000; -5% = 3,200; tax 18% of 60,800 = 10,944; total 71,744
    check("invoice arithmetic", (c["subtotal"], c["discount"], c["tax"], c["total"]) == (64000, 3200, 10944, 71744), str(c))
    check("invoice: every check passes (the total is on the page)", not bad(rep), "; ".join(bad(rep)))
    check("bill-to keeps its line break", "Ali Traders\n45 Mall Road" in rp["meta"]["bill_to"])

    acad = {"doctype": "assignment", "title": "Inflation in Pakistan", "theme": "academic", "target_pages": 2,
            "meta": {"university": "University of the Punjab", "course": "ECON-201", "author": "Ahmed Raza", "roll_no": "2023-ECO-045"},
            "blocks": [{"type": "heading", "text": "Introduction", "level": 1}, {"type": "paragraph", "text": para * 4},
                       {"type": "heading", "text": "Causes", "level": 1}, {"type": "heading", "text": "Money supply", "level": 2},
                       {"type": "paragraph", "text": para * 3}, {"type": "heading", "text": "References", "level": 1},
                       {"type": "references", "items": ["State Bank of Pakistan (2025). Annual Report."]}]}
    rp, r, rep = make("assignment", acad)
    heads = [b["text"] for b in rp["blocks"] if b["type"] == "heading"]
    check("academic headings numbered (not References)", heads == ["1 Introduction", "2 Causes", "2.1 Money supply", "References"], str(heads))
    check("academic: Times New Roman everywhere", all(VF._is_family(f, {"timesnewroman"}) or VF._is_family(f, VF.OK_FONTS) for f in VF.glyph_fonts(r["pdf"])))
    check("assignment cover lists the course and roll number", "ECON-201" in " ".join(RN.text(r["pdf"])[:1]) and "2023-ECO-045" in " ".join(RN.text(r["pdf"])[:1]))

    cv = {"doctype": "cv", "title": "Test Person", "subtitle": "Analyst", "theme": "modern", "meta": {"email": "t@example.com"},
          "blocks": [{"type": "heading", "text": "Profile", "level": 1}, {"type": "paragraph", "text": para * 2}]
          + [x for i in range(5) for x in ({"type": "heading", "text": f"Section {i}", "level": 1},
                                            {"type": "entry", "title": f"Role {i}", "org": "Org", "dates": "2020 - 2024", "bullets": ["Did a thing well", "Did another thing"]})]
          + [{"type": "heading", "text": "Languages", "level": 1}, {"type": "kv", "items": [["Urdu", "Native"], ["English", "Fluent"]]}]}
    rp, r, rep = make("cv", cv)
    g = next(x for x in rep["results"] if x["id"] == "global")
    if r.get("pages", 1) > 1:
        check("a CV spilling onto an almost empty page is flagged with a fix", g.get("fix", {}).get("op") == "fit_pages", g["why"])
        rp2, r2, rep2 = make("cv_tight", {**cv, "tight": 1})
        check("tightened, the CV fits on one page", r2.get("pages") == 1, f"{r2.get('pages')} pages")
    else:
        rp2, r2, rep2 = make("cv_tight", {**cv, "blocks": cv["blocks"] + [{"type": "heading", "text": "More", "level": 1},
                                                                           {"type": "paragraph", "text": para}]})
        check("(cv fitted on one page already)", True)
    deck_and_book()
    print(f"\n{'ALL OK' if not FAILS else f'{len(FAILS)} FAILED: ' + '; '.join(FAILS)}")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
