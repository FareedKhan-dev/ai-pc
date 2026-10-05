"""Multi-turn conversations about Word documents, each turn checked on the document it produced (like
tests/conversations.py for video). Word renders every version in the background.

  .venv\\Scripts\\python.exe tests\\doc_conversations.py [report] [assignment] [messy] [--offline]
--offline: no model (the turns that need one are skipped). Chats go to out/docs/_tests/chats (wiped at each run).
"""
import json
import re
import shutil
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

from harness.office import docmap as DM  # noqa: E402
from harness.office.docchat import DocChat  # noqa: E402

CHATS = ROOT / "out" / "docs" / "_tests" / "chats"
REPORT = ROOT / "out/docs/rooftop-solar-for-homes-in-lahore_135936/rooftop-solar-for-homes-in-lahore.docx"
ASSIGN = ROOT / "out/docs/causes-of-inflation-in-pakistan_140241/causes-of-inflation-in-pakistan.docx"


class X:
    def __init__(self, c, v0, v1, reply, turn):
        self.c, self.v0, self.v1, self.reply, self.turn = c, v0, v1, reply, turn
        self.doc = Document(str(c.path()))
        self.m = DM.docx_map(self.doc)
        self.m0 = c.map(v0)

    def new(self):
        return self.v1 != self.v0

    def says(self, *words):
        return all(re.search(w, self.reply, re.I) for w in words)

    def style(self, name):
        try:
            return self.doc.styles[name]
        except KeyError:
            return None

    def heads(self):
        return [h["text"] for h in self.m["headings"]]

    def text(self):
        return " ".join(it["text"] for it in self.m["items"] if not it.get("toc"))

    def words(self, m=None):
        return (m or self.m)["stats"]["words"]

    def section_words(self, name, m=None):
        ks = DM.find(m or self.m, {"kind": "section_body", "name": name})
        return sum(len((m or self.m)["items"][k]["text"].split()) for k in ks)


def col(style):
    c = style.font.color
    return str(c.rgb) if c is not None and c.type is not None else None


def font(style):
    return DM._style_font(style)


def messy_doc(path):
    """A document typed by hand: bold lines instead of headings, mixed fonts and sizes, typos, no styles."""
    d = Document()
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(10)
    parts = [("Introduction", None), ("Climate change is a big problem for pakistan. It effect the agriculture and the water supply of the country. "
                                      "Many people is affected every year by floods and heatwaves.", "Comic Sans MS"),
             ("Causes", None), ("The main causes is the burning of fossil fuels and cutting of forests. Pakistan contribute less then one percent of "
                                "global emissions but it suffer more.", None),
             ("Effects on Agriculture", None), ("Crops like wheat and rice needs water at the right time. When the monsoon come late the yield drop. "
                                                "Farmers dont have insurance so they loose money.", "Calibri"),
             ("Conclusion", None), ("Pakistan need to adapt by building dams, planting trees and helping farmers with new seeds.", None)]
    for text, f in parts:
        p = d.add_paragraph()
        r = p.add_run(text)
        if f is None and len(text.split()) < 5:
            r.bold = True
            r.font.size = Pt(13)
        elif f:
            r.font.name = f
            r.font.size = Pt(11)
    d.save(str(path))
    return path


# (message, intents the turn must include (None = any), check(x) -> bool, needs the model)
CONVERSATIONS = {
    "report": (REPORT, [
        ("how many pages and how many words?", "question", lambda x: x.says(r"\d+ page", r"\d+ words") and not x.new(), False),
        ("make the headings dark red and use Georgia for them", "change",
         lambda x: col(x.style("Heading 1")) == "8B0000" and font(x.style("Heading 1")) == "Georgia", False),
        ("and make them a bit bigger", "change",
         lambda x: x.style("Heading 1").font.size.pt > Document(str(x.c.path(x.v0))).styles["Heading 1"].font.size.pt, False),
        ("add page numbers at the bottom right, page x of y", "change", lambda x: x.m["page_numbers"] and "NUMPAGES" in x.doc.sections[0].footer._element.xml, False),
        ("put 'Confidential' in the header", "change", lambda x: "Confidential" in x.m["header"], False),
        ("sort the cost table by estimated cost, highest first, and add a column Share = share of Estimated Cost (PKR)", "change",
         lambda x: "Share" in x.m["tables"][0]["header"] and x.m["tables"][0]["data"][1][2] >= x.m["tables"][0]["data"][2][2], False),
        ("make a bar chart from the table", "change", lambda x: x.m["stats"]["charts"] == x.m0["stats"]["charts"] + 1, False),
        ("delete the simple payback math section", "change", lambda x: not any("Simple Payback" in h for h in x.heads()), False),
        ("move recommendations before net metering explained", "change",
         lambda x: [h for h in x.heads() if "Recommendations" in h or "Net Metering Explained" in h][0].endswith("Recommendations"), False),
        ("replace Lahore with Karachi everywhere", "change", lambda x: "Lahore" not in x.text() and "Karachi" in x.text(), False),
        ("undo", "undo", lambda x: "Lahore" in x.text(), False),
        ("redo", "redo", lambda x: "Lahore" not in x.text(), False),
        ("make the payback period section shorter", "change", lambda x: x.section_words("Payback Period") < x.section_words("Payback Period", x.m0) * 0.9, True),
        ("add a section on financing options after payback period", "change",
         lambda x: any("Financing" in h for h in x.heads()) and x.heads().index(next(h for h in x.heads() if "Financing" in h)) ==
         x.heads().index(next(h for h in x.heads() if "Payback Period" in h)) + 1, True),
        ("number the headings", "change", lambda x: sum(1 for h in x.heads() if re.match(r"^\d", h)) >= 5, False),
        ("what did you change?", "history", lambda x: x.says("number") and not x.new(), False),
        ("go back to v2", "goto", lambda x: x.c.state["cur"] == 2, False),
        ("export", "export", lambda x: x.says(r"\.docx", r"\.pdf"), False),
    ]),
    "assignment": (ASSIGN, [
        ("make the introduction shorter and more formal", "change", lambda x: x.section_words("Introduction") < x.section_words("Introduction", x.m0) * 0.9, True),
        ("translate the conclusion into urdu", "change", lambda x: sum(1 for ch in x.text() if "؀" <= ch <= "ۿ") > 100, True),
        ("add a section on the role of the IMF after monetary causes", "change",
         lambda x: any(re.search(r"\bIMF\b", h) for h in x.heads()) and any(re.match(r"^3 .*IMF", h) for h in x.heads()), True),
        ("bold the key terms in the introduction", "change", lambda x: x.new(), True),
        ("add an abstract at the start", "change", lambda x: re.sub(r"^\d+\s+", "", x.heads()[0]).lower().startswith("abstract"), True),
        ("what does the cost-push section say?", "question", lambda x: x.says("cost") and not x.new(), True),
        ("make it look professional with dark green headings", "change",
         lambda x: col(x.style("Heading 1")) == "1B5E20" and font(x.style("Normal")) == "Calibri", False),
        ("remove the table of contents and set 1 inch margins in landscape", "change",
         lambda x: not x.m["toc"] and x.m["page"]["orientation"] == "landscape" and abs(x.m["page"]["margins_cm"][0] - 2.54) < 0.05, False),
        ("undo the last two changes", "undo", lambda x: x.c.state["cur"] == x.v0 - 2 or x.c.state["cur"] < x.v0, False),
        ("how many versions are there?", None, lambda x: x.says("version") and not x.new(), False),
        ("export", "export", lambda x: x.says(r"\.pdf"), False),
    ]),
    "messy": ("messy", [
        ("what are the sections?", "question", lambda x: x.says("Introduction") and x.says("Conclusion"), False),
        ("fix the headings", "change", lambda x: sum(1 for it in x.m["headings"] if not it.get("visual")) == 4, False),
        ("format it for university: times new roman 12, 1.5 spacing, justified", "change",
         lambda x: font(x.style("Normal")) == "Times New Roman" and x.style("Normal").paragraph_format.line_spacing == 1.5, False),
        ("add a cover page with my name Ahmed Raza, roll number 2023-CS-117, course Programming Fundamentals, teacher Dr. Sana Malik", "change",
         lambda x: "2023-CS-117" in x.text() and "Ahmed Raza" in x.text(), False),
        ("add a table of contents and page numbers", "change", lambda x: x.m["toc"] and x.m["page_numbers"], False),
        ("number the headings", "change", lambda x: sum(1 for h in x.heads() if re.match(r"^\d", h)) >= 3, False),
        ("fix the grammar and spelling", "change", lambda x: "Pakistan contribute less then" not in x.text(), True),
        ("what font is used?", "question", lambda x: x.says("Times New Roman"), False),
        ("export", "export", lambda x: x.says(r"\.pdf"), False),
    ]),
}


def run(names, planner):
    shutil.rmtree(CHATS, ignore_errors=True)
    CHATS.mkdir(parents=True, exist_ok=True)
    rows = []
    t_all = time.perf_counter()
    for name in names:
        src, turns = CONVERSATIONS[name]
        if src == "messy":
            src = messy_doc(CHATS / "messy_essay.docx")
        c = DocChat.start(src, chats_dir=CHATS, planner=planner)
        print(f"\n=== {name} ({Path(src).name})", flush=True)
        try:
            for msg, want, check, needs_model in turns:
                if needs_model and planner is None:
                    print(f"skip (needs the model)  {msg}")
                    continue
                v0 = c.state["cur"]
                t = time.perf_counter()
                err = None
                try:
                    reply = c.say(msg)
                except Exception as e:  # noqa: BLE001
                    import traceback
                    traceback.print_exc()
                    reply, err = "", f"{type(e).__name__}: {e}"
                    c.last_turn = {"intents": ["crash"], "ops": []}
                secs = time.perf_counter() - t
                x = X(c, v0, c.state["cur"], reply, c.last_turn)
                intent_ok = want is None or want in (c.last_turn.get("intents") or [])
                try:
                    check_ok = bool(check(x))
                except Exception as e:  # noqa: BLE001
                    check_ok, err = False, err or f"check {type(e).__name__}: {e}"
                ok = intent_ok and check_ok and not err
                rows.append({"chat": name, "msg": msg, "ok": ok, "intents": c.last_turn.get("intents"), "seconds": round(secs, 2),
                             "llm": bool(c.last_turn.get("llm")), "reply": reply[:500], "error": err})
                print(f"{'OK ' if ok else 'BAD'} [{','.join(c.last_turn.get('intents') or [])}{'+llm' if c.last_turn.get('llm') else ''}] "
                      f"v{v0}->v{c.state['cur']} {secs:5.1f}s  {msg}", flush=True)
                if not ok:
                    print(f"      intent_ok={intent_ok} check_ok={check_ok} {err or ''}\n      {reply[:600]}")
        finally:
            c.close()
    n, ok = len(rows), sum(r["ok"] for r in rows)
    usd = planner.cost()[1] if planner is not None else 0.0
    print(f"\n{ok}/{n} turns OK; {sum(r['llm'] for r in rows)} used the model for reading the request; "
          f"{sum(r['seconds'] for r in rows) / max(1, n):.1f} s per turn; {time.perf_counter() - t_all:.0f} s in all; ${usd:.4f}")
    (CHATS / "report.json").write_text(json.dumps({"ok": ok, "turns": n, "rows": rows}, ensure_ascii=False, indent=1), encoding="utf-8")
    return ok, n


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    planner = None
    if "--offline" not in sys.argv:
        from harness.planner import ChatPlanner
        planner = ChatPlanner()
    ok, n = run(args or list(CONVERSATIONS), planner)
    sys.exit(0 if ok == n else 1)
