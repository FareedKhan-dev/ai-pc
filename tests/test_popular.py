"""The world's most popular apps, one basic feature set each (harness/apps): Unity, Premiere Pro / DaVinci Resolve / Final
Cut Pro, Photoshop and the ones after them, through their functions and through the chat.

  .venv\\Scripts\\python.exe tests\\test_popular.py
No network and no keys (services are fakes); everything is written under out\\_tests\\popular.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from PIL import Image  # noqa: E402

from harness.apps.appschat import AppsChat  # noqa: E402

OUT = ROOT / "out" / "_tests" / "popular"
RESULTS = []
NOWIN = getattr(subprocess, "CREATE_NO_WINDOW", 0)


def check(name, ok, detail=""):
    RESULTS.append((name, bool(ok)))
    print(("ok   " if ok else "FAIL ") + name + ("" if ok or not detail else f"\n     why: {str(detail)[:700]}"))


def clear(folder):
    """The last run's files removed: git's object files are read-only, and an editor watching a new git folder may hold it
    for a moment (then it is moved aside and removed next time)."""
    import os
    import stat
    for old in folder.parent.glob(folder.name + ".old*"):
        shutil.rmtree(old, onexc=lambda f, p, e: (os.chmod(p, stat.S_IWRITE), f(p)) if not isinstance(e, PermissionError) or os.path.exists(p) else None)
    for _ in range(5):
        if not folder.exists():
            return
        try:  # the \\?\ form: Android builds make paths longer than Windows' 260-character limit
            shutil.rmtree("\\\\?\\" + str(folder.resolve()), onexc=lambda f, p, e: (os.chmod(p, stat.S_IWRITE), f(p)))
        except OSError:
            time.sleep(1)
    if folder.exists():
        folder.rename(folder.with_name(f"{folder.name}.old{int(time.time())}"))


def chat(files=(), **extra):
    return AppsChat.start(files=files, chats_dir=OUT / "chats", extra=extra)


def t_unity():
    from harness.apps import unity
    c = chat()
    r = c.say("unity game: coin collector called 'Coin Run', 12 coins, red player, speed 7, 60 seconds")
    folder = Path(c.state["memo"]["unity"]["folder"])
    cfg = json.loads((folder / "Assets" / "AIPC" / "game.json").read_text(encoding="utf-8"))
    check("unity: a project with scripts, the scene builder, packages and settings", r.startswith("Unity project made") and "Checked" in r and
          all((folder / p).exists() for p in ("Assets/Scripts/PlayerController.cs", "Assets/Scripts/GameManager.cs", "Assets/Editor/AIPCBuilder.cs",
                                              "Packages/manifest.json", "ProjectSettings/ProjectVersion.txt")), r)
    check("unity: the words became the game's settings", (cfg["title"], cfg["coins"], cfg["player_color"], cfg["speed"], cfg["seconds"]) ==
          ("Coin Run", 12, "#E53935", 7.0, 60), cfg)
    r = c.say("make the unity player green and 20 coins")
    cfg = json.loads((folder / "Assets" / "AIPC" / "game.json").read_text(encoding="utf-8"))
    check("unity: changed by words, the same project updated", r.startswith("Unity project updated") and cfg["player_color"] == "#43A047" and cfg["coins"] == 20, r)
    r = c.say("build the unity game")
    check("unity: building says how when no Unity editor is installed (or builds the .exe when one is)",
          ("Unity Hub" in r) if unity.editor() is None else ("checked: the .exe exists" in r), r)


def t_premiere():
    from harness.apps import premiere
    clips = []
    for n, (secs, size) in enumerate(((6, "1280x720"), (8, "1280x720")), 1):
        p = OUT / f"clip{n}.mp4"
        subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", f"testsrc=size={size}:rate=30:duration={secs}", "-f", "lavfi", "-i",
                        f"sine=duration={secs}", "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", str(p)], check=True, creationflags=NOWIN)
        clips.append(p)
    song = OUT / "song.mp3"
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", "sine=frequency=330:duration=20", str(song)], check=True, creationflags=NOWIN)
    tl, files = premiere.make([(clips[0], 1, 4), (clips[1], None, None)], OUT / "tl", "Shop promo", song)
    check("premiere: the plan (3 s + 8 s = 330 frames at 30 fps, 1280x720)", tl["total_f"] == 330 and tl["fps"] == 30 and (tl["width"], tl["height"]) == (1280, 720), tl)
    bad = [w for w, ok in premiere.check(tl, files) if not ok]
    check("premiere: Premiere XML, FCPXML, EDL and OTIO written and read back", not bad, bad)
    r = chat([*clips, song]).say("premiere timeline: clip1.mp4 1-4, clip2.mp4; music song.mp3; called 'Shop promo'")
    check("premiere: by chat (Premiere, Resolve, Final Cut told how to import)", "11.0 s at 30 fps" in r and "Checked" in r and "Import Timeline" in r, r)


def t_photoshop():
    from harness.apps import photoshop
    bg = OUT / "shop.jpg"
    Image.new("RGB", (1600, 1000), (40, 90, 160)).save(bg)
    logo = OUT / "logo.png"
    lg = Image.new("RGBA", (300, 300), (0, 0, 0, 0))
    from PIL import ImageDraw
    ImageDraw.Draw(lg).ellipse((10, 10, 290, 290), fill=(255, 193, 7, 255))
    lg.save(logo)
    r = chat([bg, logo]).say("photoshop file 1080x1080: background shop.jpg; logo logo.png top-left; title 'Eid Sale' white; subtitle '20% off everything' yellow")
    psd = next((OUT / "chats").rglob("Eid_Sale.psd"))
    im = Image.open(psd)
    check("photoshop: a PSD with Background, Logo, Title and Subtitle layers, read back by Pillow's PSD reader", r.startswith("PSD with 4 layers") and "Checked" in r
          and [x[0] for x in im.layers] == ["Background", "Logo", "Title", "Subtitle"] and im.size == (1080, 1080), r)
    r = chat([bg, logo]).say("psd with layers shop.jpg, logo.png")
    check("photoshop: pictures stacked as layers", r.startswith("PSD with 2 layers") and "Checked" in r, r)


def t_aftereffects():
    r = chat().say("after effects title: 'Eid Sale' subtitle 'Up to 50% off' blue background, 4 seconds")
    check("after effects: Lottie played by lottie-web (title hidden at first, shown at the end), the .jsx parses", r.startswith("Title animation 'Eid Sale'")
          and "Checked" in r and "NOT right" not in r, r)


def t_obs():
    from popular_fakes import FakeObs
    from harness.apps import obs
    fake = FakeObs(folder=OUT)
    try:
        obs.Client("wrong", port=fake.port).connect()
        check("obs: a wrong password is refused, saying where the right one is", False)
    except RuntimeError as e:
        check("obs: a wrong password is refused, saying where the right one is", "Show Connect Info" in str(e), e)
    logo = OUT / "obslogo.png"
    Image.new("RGBA", (64, 64), (255, 0, 0, 255)).save(logo)
    c = chat([logo], clients={"obs": obs.Client("obs-pass", port=fake.port)})
    r = c.say("obs scenes")
    check("obs: scenes listed over obs-websocket (password challenge passed)", r == "OBS scenes: Main, Gaming; showing now: Main.", r)
    r = c.say("switch obs to scene Gaming")
    check("obs: a scene switched and read back", r.startswith("OBS now shows scene 'Gaming'") and fake.current == "Gaming", r)
    r = c.say("add text 'Live now' to obs scene Main")
    check("obs: a text source added and read back", "read back from OBS" in r and any(i["inputKind"] == "text_gdiplus_v3" for i in fake.inputs.values()), r)
    r = c.say("add image obslogo.png to obs scene Main")
    check("obs: a picture source added and read back", "read back from OBS" in r, r)
    r1, r2 = c.say("start recording in obs"), c.say("stop recording in obs")
    check("obs: recording started and stopped, the file named", r1 == "OBS is recording." and "Recording saved:" in r2 and "(the file is there)" in r2, (r1, r2))
    r = c.say("start streaming in obs")
    check("obs: going live waits for a yes", r.startswith("Ready to START STREAMING") and not fake.streaming, r)
    r = c.say("yes")
    check("obs: and goes live after it", r == "OBS is LIVE." and fake.streaming, r)
    r = c.say("obs screenshot")
    check("obs: a screenshot of what OBS shows", "(saved)" in r, r)


def t_google():
    from popular_fakes import FakeGoogle
    from harness.apps import gworkspace
    fg = FakeGoogle()
    c = chat(clients={"gworkspace": gworkspace.Client("g-token", fg)})
    r = c.say("google doc 'Meeting notes': # Decisions\\nPrices stay the same.\\n# Next steps\\nCall Haier on Monday.")
    d = next(iter(fg.docs.values()))
    check("google docs: a doc with headings, read back", "read back: every line is there" in r and d["styles"] == ["HEADING_1", "HEADING_1"], (r, d))
    r = c.say("google slides 'Sales update': Q3: sales up 12%; 40 new customers | Q4 plan: open a second shop; hire two people")
    deck = next(iter(fg.decks.values()))
    check("google slides: a title slide and a slide per section, read back", "3 slides read back" in r and deck["texts"].get("t0") == "Sales update"
          and deck["texts"].get("aipc_b2") == "open a second shop\nhire two people", (r, deck["texts"]))
    r = c.say("google form quiz 'Science test': 1. What is H2O? a) Water* b) Salt 2. The sun is a star. True* 3. Capital of Pakistan? answer: Islamabad")
    form = next(iter(fg.forms.values()))
    check("google forms: a marked quiz (3 graded questions), with the link to send", "3 questions read back, marked as a quiz" in r and
          form["settings"]["quizSettings"]["isQuiz"] and "viewform" in r, r)


def t_github():
    from popular_fakes import FakeGitHub
    from harness.apps import github
    fg = FakeGitHub(OUT / "remotes")
    (OUT / "remotes").mkdir(parents=True, exist_ok=True)
    proj = OUT / "shop-site"
    proj.mkdir(parents=True, exist_ok=True)
    (proj / "index.html").write_text("<h1>Khan Electronics</h1>", encoding="utf-8")
    c = chat(clients={"github": github.Client("gh-token", fg)})
    r = c.say("create github repo 'shop-site'")
    check("github: a new repository shown first (private)", r.startswith("Ready to create the private GitHub repository 'shop-site'"), r)
    r = c.say("yes")
    check("github: created and read back", "Created ayesha/shop-site (private, read back)" in r, r)
    c.say(f"push {proj} to github repo shop-site")
    r = c.say("yes")
    check("github: a project committed and pushed; GitHub's latest commit is the folder's", "checked: GitHub's latest commit is this folder's" in r, r)
    check("github: the token never lands in the project's git settings", "gh-token" not in (proj / ".git" / "config").read_text(encoding="utf-8"))
    c.say("open github issue in shop-site: 'Checkout button broken'")
    r = c.say("yes")
    check("github: an issue opened", r.startswith("Opened issue #1"), r)
    r = c.say("github issues in shop-site")
    check("github: issues listed", "#1 Checkout button broken" in r, r)


def t_services():
    from popular_fakes import FakeGraph, FakeSalesforce, FakeSpotify
    from harness.apps import mstodo, salesforce, spotify
    import datetime as dt
    fs, fsf, fg = FakeSpotify(), FakeSalesforce(), FakeGraph()
    c = chat(clients={"spotify": spotify.Client({"client_id": "x", "access_token": "sp-token", "refresh_token": "r", "expires_at": 9e12}, fs),
                      "salesforce": salesforce.Client({"client_id": "x", "access_token": "old-token", "refresh_token": "r", "instance_url": fsf.INSTANCE}, fsf),
                      "mstodo": mstodo.Client("ms-token", fg)})
    r = c.say("spotify playlist 'Road trip': Blinding Lights by The Weeknd; Shape of You by Ed Sheeran; No Such Song by Nobody")
    check("spotify: a private playlist from song names; the one not found is named; read back", "2 of 3 songs added; not found: No Such Song by Nobody" in r
          and "read back: 2 songs" in r and not next(iter(fs.playlists.values()))["public"], r)
    r = c.say("what's playing on spotify")
    check("spotify: now playing (nothing)", r == "Nothing is playing on Spotify right now.", r)
    r = c.say("add salesforce lead: Ali Raza, Ali Traders, ali@x.com, 0300-1234567")
    check("salesforce: a new lead shown first", r.startswith("Ready to add a Salesforce lead: FirstName Ali, LastName Raza, Company Ali Traders"), r)
    r = c.say("yes")
    check("salesforce: added (after renewing an expired session, on the newest API version), read back", "read back: the same details" in r and fsf.session == "sf-token-2", r)
    r = c.say("salesforce leads by status")
    check("salesforce: leads counted by status", r == "Leads by status: Open - Not Contacted 1.", r)
    now = dt.datetime.now()
    r = c.say("microsoft todo: call Haier about the TV order by friday")
    t = next(iter(fg.tasks.values()))
    check("microsoft to do: a task with its due date, read back", "(read back)" in r and t["title"] == "call Haier about the TV order"
          and t["dueDateTime"]["timeZone"] == "Pakistan Standard Time", (r, t))
    r = c.say("complete microsoft task call Haier")
    check("microsoft to do: ticked off", r.startswith("Ticked off 'call Haier about the TV order' (completed)"), r)
    r = c.say("onenote page 'Meeting with Ali' in Work: prices agreed at 84,000")
    check("onenote: a page written into the Work section", "written in section Work" in r and "prices agreed at 84,000" in fg.pages[0], r)


def t_devtools():
    from openpyxl import Workbook
    from harness.apps import postman
    from harness.hub.http import FakeTransport
    wb = Workbook()
    ws = wb.active
    ws.append(["Customer", "Item", "Amount"])
    for row in (("Ali Traders", "TV", 85000), ("Bilal & Sons", "Mount", 3500), ("Ali Traders", "Cable", 1200), ("Sara Stores", "TV", 84000)):
        ws.append(row)
    src = OUT / "sales.xlsx"
    wb.save(src)
    r = chat([src]).say("notebook analysing sales.xlsx: total Amount by Customer")
    nb = json.loads(next((OUT / "chats").rglob("sales_analysis.ipynb")).read_text(encoding="utf-8"))
    outs = [o for c in nb["cells"] if c["cell_type"] == "code" for o in c["outputs"]]
    totals = next(o for c in nb["cells"] if c["cell_type"] == "code" and "groupby" in "".join(c["source"]) for o in c["outputs"])
    check("jupyter: a notebook run cell by cell with pandas, its tables and chart saved in it", "Checked" in r and any("image/png" in o.get("data", {}) for o in outs)
          and "86200" in totals["data"]["text/plain"] and "text/html" in totals["data"], r)
    site = OUT / "flask-site"
    site.mkdir(parents=True, exist_ok=True)
    (site / "requirements.txt").write_text("flask==3.1.2\n", encoding="utf-8")
    (site / "app.py").write_text("from flask import Flask\napp = Flask(__name__)\n", encoding="utf-8")
    r = chat().say(f"dockerfile for {site}")
    d = (site / "Dockerfile").read_text(encoding="utf-8")
    check("docker: a Flask project's Dockerfile (pinned image, non-root, port 5000), .dockerignore and compose.yaml, checked", "flask project" in r and
          "Checked" in r and "FROM python:3.12-slim" in d and "USER app" in d and (site / "compose.yaml").exists(), r)
    postman.TRANSPORT = FakeTransport({("GET", "/posts/1"): (200, {"id": 1}), ("GET", "/users"): (200, [{"id": 1}]), ("POST", "/posts"): (201, {"id": 101})})
    c = chat()
    r = c.say('postman collection for https://jsonplaceholder.typicode.com: GET /posts/1, GET /users, POST /posts {"title": "Hello"}')
    col = json.loads(next((OUT / "chats").rglob("*.postman_collection.json")).read_text(encoding="utf-8"))
    check("postman: a v2.1 collection with an environment, a test on every request", "checked: v2.1 format" in r and len(col["item"]) == 3
          and all(i["event"][0]["listen"] == "test" for i in col["item"]), r)
    r = c.say("run the postman collection")
    check("postman: a run that would POST waits for a yes", r.startswith("Ready to run the collection: it sends POST /posts"), r)
    r = c.say("yes")
    check("postman: run like Newman, every request passed, a report", r.startswith("Ran 3 requests: 3 passed, 0 failed") and "Report:" in r, r)


def t_godot():
    c = chat()
    r = c.say("godot game: platformer called 'Jump Hero', 8 coins, green player, jump 600")
    check("godot: a platformer made and played headless (no script errors; landed, moved right, jumped; every coin there)",
          r.startswith("Godot game made") and "Played headless and checked" in r and "NOT right" not in r, r)
    r = c.say("make the godot player red")
    cfg = json.loads((Path(c.state["memo"]["godot"]["folder"]) / "game.json").read_text(encoding="utf-8"))
    check("godot: changed by words and played again", r.startswith("Godot game updated") and cfg["player_color"] == "#e53935" and "NOT right" not in r, r)


def t_illustrator():
    r = chat().say("illustrator poster 1080x1350: title 'Grand Opening' white, subtitle 'Saturday 10 am' yellow, blue background, badge 'FREE GIFTS'")
    check("illustrator: a layered vector poster (Background, Title, Subtitle, Badge) and a one-page PDF of its shape", r.startswith("Vector poster with 4 layers")
          and "Checked" in r and "NOT right" not in r, r)
    r = chat().say("illustrator logo 'Khan Electronics' with initials 'KE' green")
    check("illustrator: a vector logo with its mark and name as layers", r.startswith("Vector logo with 2 layers (Mark, Name)") and "Checked" in r, r)


def t_revit():
    r = chat().say("revit model of a 7 marla double story house with 3 bedrooms")
    check("revit: words become an IFC building model (2 storeys, walls cut round doors and windows, rooms with areas), reopened, built and validated",
          r.startswith("BIM model") and "2 storey(s)" in r and "BED ROOM" in r and "Checked" in r and "NOT right" not in r, r)
    from harness.cad import cadparse, floorplan
    plan = OUT / "house.plan.json"
    plan.write_text(json.dumps(floorplan.plan(cadparse.house_brief("a 5 marla single story house with 2 bedrooms"))), encoding="utf-8")
    r = chat(files=[str(plan)]).say("ifc model from house.plan.json")
    check("revit: a saved CAD plan becomes a one-storey IFC model", r.startswith("BIM model") and "1 storey(s)" in r and "Checked" in r and "NOT right" not in r, r)


def t_lightroom():
    import numpy as np
    from PIL import Image
    yy, xx = np.mgrid[0:300, 0:450]
    photo = OUT / "beach.jpg"
    Image.fromarray(np.stack([xx / 450 * 200 + 30, yy / 300 * 180 + 40, 120 + 60 * np.sin(xx / 40)], -1).clip(0, 255).astype(np.uint8)).save(photo, quality=95)
    r = chat(files=[str(photo)]).say("lightroom preset 'Warm Film': warm vintage, exposure +0.3, vignette -20 for beach.jpg")
    check("lightroom: a develop preset (.xmp) Lightroom reads, and the photo came out brighter, warmer, with darker corners",
          r.startswith("Lightroom preset 'Warm Film'") and all(w in r for w in ("brighter", "warmer", "darker corners")) and "NOT right" not in r, r)
    r = chat(files=[str(photo)]).say("lightroom black and white preset for beach.jpg")
    check("lightroom: a black and white look, the photo really grey", "came out black and white" in r and "NOT right" not in r, r)


def t_anki():
    r = chat().say("anki deck 'Capitals': France = Paris; Japan = Tokyo; Pakistan = Islamabad, both ways")
    check("anki: a deck from words, both ways (3 notes, 6 cards), read back from the package", "Anki deck 'Capitals'" in r and "3 notes, 6 cards" in r
          and "Checked" in r and "NOT right" not in r, r)
    from openpyxl import Workbook
    wb = Workbook()
    for row in (["Term", "Meaning"], ["Mitosis", "Cell division into two identical cells"], ["Osmosis", "Water moving through a membrane"]):
        wb.active.append(row)
    wb.save(OUT / "biology.xlsx")
    r = chat(files=[str(OUT / "biology.xlsx")]).say("anki deck from biology.xlsx")
    check("anki: a deck from an Excel sheet (header skipped)", "Anki deck 'Biology'" in r and "2 notes, 2 cards" in r and "NOT right" not in r, r)
    r = chat().say("anki cloze: The {mitochondria} is the powerhouse of the cell; Water boils at {100} degrees at {sea level}")
    check("anki: cloze cards, one per hidden part", "2 notes, 3 cards" in r and "NOT right" not in r, r)


def t_arduino():
    c = chat()
    r = c.say("arduino traffic light on a nano, green 5 seconds")
    check("arduino: a sketch from words, really compiled for the Nano by arduino-cli, with wiring", r.startswith("Arduino sketch TrafficLight for the Nano")
          and "Checked: it compiles" in r and "Wiring: pin 10" in r, r)
    r = c.say("upload it to the arduino")
    check("arduino: upload is never done without a yes (no board here: it says so)", ("No Arduino found" in r or "Say 'yes'" in r) and "Uploaded" not in r, r)
    bad = OUT / "Broken.ino"
    bad.write_text("void setup() {\n  pinMode(13, OUTPUT)\n}\nvoid loop() {}\n", encoding="utf-8")
    r = chat(files=[str(bad)]).say("arduino compile Broken.ino")
    check("arduino: a sketch with a mistake is caught (the compiler's error shown)", "NOT right: it compiles" in r and "error" in r.lower(), r[-300:])


def t_msproject():
    r = chat().say(
        "project plan 'House build' starting 2 november 2026: foundation 10 days by Mason team; walls 3 weeks after foundation by Mason team; roof 7 days after walls; "
        "plumbing 5 days after walls by Plumber; electrical 6 days after walls; finishing 12 days after roof, plumbing and electrical; handover milestone after finishing")
    check("ms project: a plan from words scheduled on working days (44 days, ends Thu 31 Dec), critical path, Project XML in schema order",
          "44 working days" in r and "Thu 31 Dec 2026" in r and "Critical path: foundation -> walls -> roof -> finishing -> handover" in r
          and "plumbing 2 days" in r and "Checked" in r and "NOT right" not in r, r)
    r = chat().say("project plan: a 2 days after b; b 3 days after a")
    check("ms project: tasks that wait for each other in a loop are refused", "loop" in r and "a -> b -> a" in r, r)


def t_prusaslicer():
    r = chat().say("slice cube 20 mm for ender 3 in pla, 20% infill")
    check("prusaslicer: a cube sliced to Ender-3 G-code by PrusaSlicer (on the bed, 20 mm tall, 20% infill, .3mf too)",
          r.startswith("Sliced cube_20mm.stl for the Ender-3") and "g of filament" in r and "the print is as tall as the model (20 mm)" in r
          and "NOT right" not in r, r)
    r = chat().say("3d print box 60x40x10 mm on a prusa core one, petg, supports, brim, 2 copies")
    check("prusaslicer: a box for the Prusa CORE One in PETG with supports, a brim and 2 copies", "for the Prusa CORE One" in r and "2 copies" in r
          and "supports on, a brim" in r and "NOT right" not in r, r)
    r = chat().say("slice box 300x300x10 mm for ender 3")
    check("prusaslicer: a model too big for the bed is refused with a way out", r.startswith("It is too big for the Ender-3's bed"), r)


def t_freecad():
    c = chat()
    r = c.say("solidworks plate 100x60x5 mm with 4 holes 6 mm, rounded corners 5 mm, aluminium")
    check("freecad: a plate with 4 holes and rounded corners built in FreeCAD, STEP read back the same, weight in aluminium",
          r.startswith("Made the part: plate 100 x 60 x 5 mm, 4 holes 6 mm, rounded corners 5 mm; aluminium") and "Checked in FreeCAD" in r
          and "NOT right" not in r, r)
    r = c.say("make the holes 8 mm")
    check("freecad: changed by words (holes 8 mm), rebuilt and checked", r.startswith("Changed the part") and "4 holes 8 mm" in r and "NOT right" not in r, r)
    r = c.say("slice it for ender 3")
    check("freecad -> prusaslicer: 'slice it' prints the part just made", "Sliced plate_100_x_60_x_5_mm.stl for the Ender-3" in r and "NOT right" not in r, r)
    r = chat().say("fusion flange 80 mm diameter 10 mm thick, 6 holes 8 mm on 60 mm circle, centre hole 30 mm, steel")
    check("freecad: a flange with 6 bolt holes on a circle and a centre hole", "flange 80 mm across" in r and "6 x 8 mm" in r and "1 x 30 mm" in r
          and "NOT right" not in r, r)


def t_matlab():
    r = chat().say("matlab plot sin(x) and cos(x) from 0 to 2*pi")
    check("matlab: a plot script run headless by Octave, the figure saved, values match Python", "Figure: figure.png" in r and "match Python's" in r
          and "NOT right" not in r, r)
    r = chat().say("matlab solve 2x + 3y = 13, x - y = 1")
    check("matlab: linear equations solved (x = 3.2, y = 2.2), checked against every equation", "x = 3.200000" in r and "y = 2.200000" in r
          and "NOT right" not in r, r)
    r = chat().say("matlab fft of 50 Hz and 120 Hz")
    check("matlab: an FFT finds exactly the 50 Hz and 120 Hz in the signal", "Peaks at 50 Hz" in r and "Peaks at 120 Hz" in r and "NOT right" not in r, r)
    r = chat().say("matlab rc circuit R 1k C 100uF 5 V")
    check("matlab: an RC circuit by ode45 reaches 63.2% at one time constant", "(63.2%)" in r and "NOT right" not in r, r)
    bad = OUT / "broken_model.m"
    bad.write_text("a = [1 2 3];\nb = a * [4 5 6];\n", encoding="utf-8")
    r = chat(files=[str(bad)]).say("matlab run broken_model.m")
    check("matlab: a .m with a mistake is caught (Octave's error shown)", "NOT right: the script ran with no errors" in r and "nonconformant" in r, r[-300:])


def t_musescore():
    import os
    user_dirs = [Path.home() / "Documents" / "MuseScore4", Path(os.environ["LOCALAPPDATA"]) / "MuseScore"]
    before = [p.exists() for p in user_dirs]
    c = chat()
    r = c.say("musescore 'Little Song': C4 D4 E4 F4 G4/2 G4/2 A4 A4 A4 A4 G4/1 at 100 bpm on piano")
    check("musescore: a melody becomes .mscz, sheet-music PDF (with its title), PNG and MP3, made unseen on the hidden desktop",
          r.startswith("MuseScore score 'Little Song'") and "holds all 11 notes" in r and "the MP3 lasts as long as the music" in r and "NOT right" not in r, r)
    c.say("melody 'Violin Line': E4 F#4 G4 A4 B4/2 B4/2 at 90 bpm on violin")
    r = c.say("musescore it")
    check("music -> musescore: 'musescore it' turns the melody just made into sheet music", "MuseScore score 'Violin_Line'" in r and "holds all 6 notes" in r
          and "NOT right" not in r, r)
    check("musescore: nothing written to Documents or AppData (its folders stay in tools/musescore-portable)",
          [p.exists() for p in user_dirs] == before, str(list(zip(user_dirs, before))))


def t_audacity():
    import os
    import wave
    import numpy as np
    sr, parts, rng = 44100, [], np.random.default_rng(3)
    for f in (180, 220, 200, 240, 190, 210):  # six quiet 'phrases' with long pauses, like a raw interview
        t = np.arange(int(sr * 1.6)) / sr
        parts += [sum(np.sin(2 * np.pi * f * k * t) / k for k in range(1, 6)) * (0.5 + 0.5 * np.sin(2 * np.pi * 4 * t)) * 0.03, rng.normal(0, 0.0003, int(sr * 2.5))]
    for name, x in (("interview.wav", np.concatenate(parts)), ("tune.wav", np.sin(2 * np.pi * 440 * np.arange(sr * 4) / sr) * 0.3)):
        with wave.open(str(OUT / name), "wb") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(sr)
            w.writeframes((x * 32767).astype("<i2").tobytes())
    user_dirs = [Path(os.environ["APPDATA"]) / "audacity", Path(os.environ["LOCALAPPDATA"]) / "audacity", Path.home() / "Documents" / "Audacity"]
    before = [p.exists() for p in user_dirs]
    c = chat(files=[str(OUT / "interview.wav"), str(OUT / "tune.wav")])
    r = c.say("audacity podcast cleanup interview.wav")
    check("audacity: the real Audacity, hidden, cleans a recording for a podcast (pauses cut, podcast loudness, no clipping), measured by FFmpeg",
          "Audacity cleaned up for a podcast" in r and "no pause over 1.2 s left" in r and "podcast standard" in r and "NOT right" not in r, r)
    r = c.say("audacity clip interview.wav 2 to 9")
    check("audacity: a 2-9 s clip is exactly 7 s", "exactly the 7 s asked" in r and "NOT right" not in r, r)
    r = c.say("make it 10% slower")
    check("audacity: 'make it 10% slower' works on the clip just made (7.0 s -> 7.8 s, same pitch)", "slower: 7.0 s -> 7.8 s" in r and "NOT right" not in r, r)
    r = c.say("audacity pitch up 2 semitones of tune.wav")
    check("audacity: pitch up 2 semitones, measured by FFT (440 Hz -> 494 Hz), same length", "440 Hz -> 494 Hz" in r and "NOT right" not in r, r)
    check("audacity: nothing left in AppData or Documents, no Audacity left running", [p.exists() for p in user_dirs] == before and
          "Audacity.exe" not in subprocess.run(["tasklist"], capture_output=True, text=True, creationflags=NOWIN).stdout, str(list(zip(user_dirs, before))))


def t_calibre():
    import os
    from docx import Document
    d = Document()
    d.add_heading("The Garden Book", 0)
    for i in range(1, 4):
        d.add_heading(f"Chapter {i}", 1)
        d.add_paragraph(("Plants grow best with care and water in the morning. " * 12).strip())
    d.save(OUT / "garden.docx")
    user_dirs = [Path(os.environ["APPDATA"]) / "calibre", Path(os.environ["LOCALAPPDATA"]) / "calibre-cache", Path.home() / "Documents" / "Calibre Library"]
    before = [p.exists() for p in user_dirs]
    r = chat(files=[str(OUT / "garden.docx")]).say('calibre convert garden.docx to epub and kindle, title "The Garden Book" by Ali Khan')
    check("calibre: a Word file becomes EPUB and Kindle AZW3, each opening as its format with all the text, title and author",
          "garden.epub, garden.azw3" in r and r.count("100% of the words") == 2 and "title and author read back from the AZW3" in r and "NOT right" not in r, r)
    epub = next(OUT.glob("chats/*/calibre/garden.epub"), None) or next(OUT.rglob("calibre/garden.epub"))
    r = chat(files=[str(epub)]).say('calibre set title "Gardens of Lahore" author "Sara Malik" on garden.epub')
    check("calibre: book details changed on a copy and read back", "title Gardens of Lahore" in r and "NOT right" not in r, r)
    check("calibre: nothing written to AppData or Documents (settings in tools/calibre/home)", [p.exists() for p in user_dirs] == before,
          str(list(zip(user_dirs, before))))


def t_kicad():
    import os
    import shutil
    user_dirs = [Path(os.environ["APPDATA"]) / "kicad", Path(os.environ["LOCALAPPDATA"]) / "kicad", Path.home() / "Documents" / "KiCad"]
    before = [p.exists() for p in user_dirs]
    r = chat().say("kicad schematic red led on 5v")
    check("kicad: an LED circuit from words (220 ohm, 13.6 mA) as a KiCad project; ERC clean; KiCad's netlist exactly as meant",
          "R1 220" in r and "LED current 13.6 mA" in r and "Electrical Rules Check finds no errors" in r and "joins exactly the pins" in r and "NOT right" not in r, r)
    r = chat().say("kicad 555 blinker 2 hz on 9v")
    check("kicad: a 555 blinker (1.97 Hz) with KiCad's own NE555P symbol, 7 nets exactly as meant", "1.97 Hz" in r and "TRIG: 4 pins" in r
          and "NOT right" not in r, r)
    board = OUT / "board"
    board.mkdir(exist_ok=True)
    for f in (ROOT / "tools" / "kicad" / "share" / "kicad" / "demos" / "ecc83").glob("ecc83-pp.*"):
        shutil.copy(f, board / f.name)
    r = chat(files=[str(board / "ecc83-pp.kicad_pcb")]).say("kicad gerbers for ecc83-pp.kicad_pcb")
    check("kicad: Gerbers + drill zipped for a PCB maker after a clean DRC (KiCad's demo board)", "Design Rules Check found no errors" in r
          and "every layer a maker needs" in r, r)
    check("kicad: nothing written to AppData or Documents (settings in tools/kicad/home)", [p.exists() for p in user_dirs] == before,
          str(list(zip(user_dirs, before))))


def t_krita():
    import os
    leftovers = [Path(os.environ["LOCALAPPDATA"]) / "kritarc", Path(os.environ["LOCALAPPDATA"]) / "kritadisplayrc", Path(os.environ["APPDATA"]) / "krita"]
    before = [p.exists() for p in leftovers]
    r = chat().say("krita poster 1080x1350: title 'Art Fair' white, subtitle 'Sunday 4 pm' yellow")
    check("krita: layers written as OpenRaster, turned into Krita's own .kra by Krita (hidden), its render identical to the layers",
          "Art_Fair.kra" in r and "every layer by name" in r and "mean difference 0.00" in r and "NOT right" not in r, r)
    r = chat().say("krita canvas 1600x1000 with layers sketch, ink, colours")
    check("krita: a painter's canvas with named empty layers over paper", "layers: Colours, Ink, Sketch, Paper" in r and "NOT right" not in r, r)
    check("krita: AppData left as it was (Krita's resources live in tools/krita/home)", [p.exists() for p in leftovers] == before,
          str(list(zip(leftovers, before))))


def t_openscad():
    r = chat().say("openscad box with lid 80x60x40 mm, 2 mm walls")
    check("openscad: a parametric box with a lid rendered headless to STL, manifold, the size asked", "box 80 x 60 x 40 mm inside with a lid" in r
          and "the STL is the size asked" in r and "NOT right" not in r, r)
    r = chat().say("openscad gear 20 teeth module 2")
    check("openscad: an involute spur gear, 44 mm across, watertight (checked by PrusaSlicer too)", "44 mm across" in r and "watertight" in r
          and "NOT right" not in r, r)
    c = chat()
    r = c.say("openscad keychain 'ALI' 50x18 mm")
    check("openscad: a keychain with raised text", "keychain 'ALI' 50 x 18 mm" in r and "NOT right" not in r, r)
    r = c.say("slice it for ender 3")
    check("openscad -> prusaslicer: 'slice it' prints the keychain just made", "Sliced keychain_ALI_50_x_18_mm.stl for the Ender-3" in r and "NOT right" not in r, r)


def t_libreoffice():
    import os
    from docx import Document
    from openpyxl import Workbook, load_workbook
    d = Document()
    d.add_heading("Quarterly Report", 0)
    d.add_paragraph("Sales grew strongly in Lahore and Karachi while costs stayed level across every branch.")
    d.save(OUT / "report.docx")
    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Qty", "Price", "Total"])
    for i, (item, q, p) in enumerate((("Fan", 3, 4500), ("Iron", 2, 3200), ("Kettle", 5, 2100)), start=2):
        ws.append([item, q, p, f"=B{i}*C{i}"])
    ws.append(["Grand total", None, None, "=SUM(D2:D4)"])
    wb.save(OUT / "invoice.xlsx")
    profile = [Path(os.environ["APPDATA"]) / "LibreOffice"]
    before = [p.exists() for p in profile]
    r = chat(files=[str(OUT / "report.docx")]).say("libreoffice convert report.docx to pdf and odt")
    check("libreoffice: a Word file converted headless to PDF and ODT, each opening as itself with the original's words",
          "report.pdf, report.odt" in r and "the ODT opens as ODT" in r and "NOT right" not in r, r)
    r = chat(files=[str(OUT / "invoice.xlsx")]).say("libreoffice recalculate invoice.xlsx")
    dest = next(OUT.rglob("recalculated/invoice.xlsx"), None)
    total = load_workbook(dest, data_only=True).active["D5"].value if dest else None
    check("libreoffice: Calc fills in every formula of a sheet saved by code (grand total 3*4500 + 2*3200 + 5*2100 = 30400)",
          "every formula has a value" in r and total == 3 * 4500 + 2 * 3200 + 5 * 2100, f"{r} | D5={total}")
    check("libreoffice: its user profile stays in tools/libreoffice/home (nothing in AppData)", [p.exists() for p in profile] == before, str(profile))


def t_gimp():
    import os
    import numpy as np
    from PIL import Image
    yy, xx = np.mgrid[0:600, 0:900]
    Image.fromarray(np.stack([xx / 900 * 180 + 40, yy / 600 * 120 + 60, 140 + 40 * np.sin(xx / 30)], -1).clip(0, 255).astype("uint8")).save(OUT / "dull.jpg", quality=92)
    leftovers = [Path(os.environ["LOCALAPPDATA"]) / "GIMP", Path(os.environ["APPDATA"]) / "GIMP"]
    before = [p.exists() for p in leftovers]
    c = chat(files=[str(OUT / "dull.jpg")])
    r = c.say("gimp enhance dull.jpg")
    check("gimp: a dull photo enhanced by GIMP itself (hidden): contrast widened, sharper", "contrast range widened" in r and "sharper" in r
          and "NOT right" not in r, r)
    r = c.say("gimp black and white dull.jpg resize to 450 wide")
    check("gimp: black and white and resized to 450 px", "black and white; 450 px wide" in r and "NOT right" not in r, r)
    r = chat().say("gimp poster 1080x1350: title 'Book Fair' white, subtitle 'Saturday' yellow")
    check("gimp: a layered poster saved as GIMP's .xcf, reopened by GIMP with every layer, matching the layers", "Book_Fair.xcf" in r
          and "every layer by name" in r and "NOT right" not in r, r)
    check("gimp: AppData left as it was (profile, cache and fonts in tools/gimp-portable/home)", [p.exists() for p in leftovers] == before,
          str(list(zip(leftovers, before))))


def t_qgis():
    import os
    (OUT / "shops.csv").write_text("name,lat,lon\nLahore,31.5204,74.3587\nKarachi,24.8607,67.0011\nIslamabad,33.6844,73.0479\nPeshawar,34.0151,71.5249\n"
                                   "Quetta,30.1798,66.9750\nMultan,30.1575,71.5249\n", encoding="utf-8")
    leftovers = [Path(os.environ["APPDATA"]) / "QGIS", Path(os.environ["LOCALAPPDATA"]) / "QGIS"]
    before = [p.exists() for p in leftovers]
    c = chat(files=[str(OUT / "shops.csv")])
    r = c.say("qgis map of shops.csv")
    check("qgis: a map of 6 places over QGIS's world map: project, PNG, PDF; every marker drawn and the labels readable by OCR",
          "6 markers for 6 places" in r and "labels are readable" in r and "NOT right" not in r, r)
    r = c.say("qgis buffer 5 km around shops.csv")
    check("qgis: 5 km buffers in metres (UTM), each pi r^2 within 1%", "pi r^2 within 1%" in r and "NOT right" not in r, r)
    r = c.say("qgis convert shops.csv to geopackage")
    check("qgis: converted to GeoPackage with all 6 features", "all 6 features" in r, r)
    check("qgis: AppData left as it was (profile and cache in tools/qgis/home)", [p.exists() for p in leftovers] == before, str(list(zip(leftovers, before))))


def t_shotcut():
    for name, src, secs in (("intro.mp4", "testsrc2=size=1280x720:rate=30", 6), ("main.mp4", "testsrc=size=1280x720:rate=30", 12)):
        if not (OUT / name).exists():  # moving test patterns (testsrc draws its own frame counter), so each moment looks different
            subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", f"{src}:duration={secs}", "-f", "lavfi", "-i", f"sine=frequency=330:duration={secs}",
                            "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", str(OUT / name)], capture_output=True, creationflags=NOWIN)
    if not (OUT / "song.mp3").exists():
        subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", "sine=frequency=220:duration=30", "-c:a", "libmp3lame", str(OUT / "song.mp3")],
                       capture_output=True, creationflags=NOWIN)
    files = [str(OUT / n) for n in ("intro.mp4", "main.mp4", "song.mp3")]
    r = chat(files=files).say("shotcut timeline: intro.mp4 0-4, main.mp4 2-8; music song.mp3; title 'My Trip'")
    check("shotcut: a timeline as a Shotcut .mlt rendered by its engine: 10 s, sound, every cut right, the title readable",
          "rendered by Shotcut's engine" in r and "10.0 s, as long as the cuts" in r and "every cut shows the right moment" in r and "title 'My Trip' is readable" in r
          and "NOT right" not in r, r)


def t_latex():
    from docx import Document
    (OUT / "notes.md").write_text("# Introduction\nSolar pumps cut **diesel costs** in *Punjab* [@khan2024].\n\n## Method\nWe compared 40 farms:\n- 20 solar\n- 20 diesel\n\n"
                                  "The saving is\n$$ S = (c_d - c_s) \\times h $$\n\n# Results\n| Farm | Cost (Rs) |\n|---|---|\n| Solar | 12,000 |\n| Diesel | 95,000 |\n\n"
                                  "# Conclusion\nSolar pays back in two seasons [@khan2024; @ali2023].\n", encoding="utf-8")
    (OUT / "refs.bib").write_text("@article{khan2024, author = {Ali Khan}, title = {Solar Irrigation in Punjab}, journal = {Pakistan Journal of Energy}, year = {2024}}\n"
                                  "@book{ali2023, author = {Sara Ali}, title = {Water and Power}, publisher = {Lahore Press}, year = {2023}}\n", encoding="utf-8")
    r = chat(files=[str(OUT / "notes.md"), str(OUT / "refs.bib")]).say("latex paper from notes.md titled 'Solar Pumps in Punjab' by Ali Khan with refs.bib")
    check("latex: a paper from Markdown (sections, list, equation, table, citations) compiled by Tectonic, references resolved",
          "LaTeX paper made" in r and "every heading is there (4 of 4)" in r and "no unresolved reference" in r and "NOT right" not in r, r)
    d = Document()
    d.add_heading("Water Use in Sindh", 0)
    for h, t in (("Background", "Sindh relies on the Indus for 90% of its water."), ("Findings", "Canal losses reach 40% & more in summer.")):
        d.add_heading(h, 1)
        d.add_paragraph(t)
    d.save(OUT / "chapters.docx")
    r = chat(files=[str(OUT / "chapters.docx")]).say("latex thesis from chapters.docx titled 'Water Use in Sindh' by Sara")
    check("latex: a thesis (chapters, contents) from a Word file, % and & handled", "LaTeX thesis made" in r and "NOT right" not in r, r)
    (OUT / "broken.tex").write_text("\\documentclass{article}\n\\begin{document}\nHello \\undefinedthing{x}\n\\end{document}\n", encoding="utf-8")
    r = chat(files=[str(OUT / "broken.tex")]).say("latex compile broken.tex")
    check("latex: a .tex with a mistake is caught, with its line", "NOT right: it compiles with no errors" in r and "broken.tex:3" in r, r[-300:])


def t_rstats():
    import os
    import random
    from openpyxl import Workbook
    random.seed(4)
    wb = Workbook()
    ws = wb.active
    ws.append(["Score", "Gender", "Class", "Hours", "Sleep"])
    for i in range(60):
        g, h, s = ("Male" if i % 2 else "Female"), random.uniform(1, 10), random.uniform(5, 9)
        ws.append([round(40 + 4 * h + 2 * s + (5 if g == "Female" else 0) + random.gauss(0, 6), 1), g, ["A", "B", "C"][i % 3], round(h, 1), round(s, 1)])
    wb.save(OUT / "survey.xlsx")
    leftovers = [Path(os.environ["APPDATA"]) / "R", Path(os.environ["LOCALAPPDATA"]) / "R", Path.home() / "Documents" / "R"]
    before = [p.exists() for p in leftovers]
    c = chat(files=[str(OUT / "survey.xlsx")])
    r = c.say("r t-test of Score by Gender in survey.xlsx")
    check("r: a t-test run by R (base-R script for RStudio), t and p equal to the statistics program's own", "R ttest: t = " in r
          and "agree with the statistics program" in r and "NOT right" not in r, r)
    r = c.say("rstudio regression of Score on Hours, Sleep in survey.xlsx")
    check("r: a regression by R, R^2 and coefficients equal to the statistics program's", "R regression: R2 = " in r and "NOT right" not in r, r)
    r = c.say("descriptives of survey.xlsx")
    check("r: without 'R' the statistics program still answers", "R describe" not in r and "Score" in r, r[:200])
    check("r: nothing in AppData or Documents (R's home and library in tools/r/home)", [p.exists() for p in leftovers] == before, str(list(zip(leftovers, before))))


def t_drawio():
    import os
    leftovers = [Path(os.environ["APPDATA"]) / "draw.io", Path(os.environ["LOCALAPPDATA"]) / "draw.io-updater"]
    before = [p.exists() for p in leftovers]
    r = chat().say("drawio flowchart: Start -> Take order -> In stock? -yes-> Pack -> Ship -> End; In stock? -no-> Order from supplier -> Pack")
    check("drawio: a flowchart exported by draw.io itself to a one-page PDF, PNG (every label read by OCR, no arrow across a box) and SVG",
          "the PDF has 1 page(s)" in r and "reads 7 of 7 box labels" in r and "the SVG parses" in r and "NOT right" not in r, r)
    check("drawio: nothing in AppData (its settings in tools/drawio/home), no update check", [p.exists() for p in leftovers] == before,
          str(list(zip(leftovers, before))))


def t_handbrake():
    clip = OUT / "trip.mp4"
    if not clip.exists():
        subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", "testsrc2=size=1920x1080:rate=30:duration=40", "-f", "lavfi", "-i",
                        "sine=frequency=440:duration=40", "-c:v", "libx264", "-preset", "veryfast", "-b:v", "12M", "-pix_fmt", "yuv420p", "-c:a", "aac",
                        "-shortest", str(clip)], capture_output=True, creationflags=NOWIN)
    c = chat(files=[str(clip)])
    r = c.say("handbrake compress trip.mp4 for discord")
    check("handbrake: a 60 MB, 40 s 1080p video made Discord-ready by HandBrake's own preset: under 10 MB, same length, sound kept",
          "Social 10 MB 1 Minute 540p60" in r and "under 10 MB" in r and "as long as the original" in r and "NOT right" not in r, r)
    r = c.say("handbrake trip.mp4 for gmail")
    check("handbrake: Gmail-ready, under 25 MB at 720p", "under 25 MB" in r and "NOT right" not in r, r)
    r = c.say("handbrake trip.mp4 fast on the intel gpu")
    check("handbrake: the Intel GPU preset falls back to software H.265 when the GPU refuses it (and says so)", "NOT right" not in r and
          ("software H.265 preset was used" in r or "H.265 QSV 1080p" in r), r)


def t_visualstudio():
    import os
    leftovers = [Path(os.environ["APPDATA"]) / "NuGet", Path(os.environ["LOCALAPPDATA"]) / "NuGet", Path.home() / ".dotnet", Path.home() / ".nuget"]
    before = [p.exists() for p in leftovers]
    r = chat().say("visual studio web api called 'Shop' with tests")
    check("visualstudio: a C# solution (core library, ASP.NET Core API, xUnit) built by the .NET SDK; 5 tests pass; the API answers correctly over HTTP",
          "builds with 0 errors" in r and "5 of 5 passed" in r and "the API runs" in r and "NOT right" not in r, r)
    r = chat().say("c# console app called Stock")
    check("visualstudio: a console app builds, its tests pass, and it runs and prints the right stock value", "Rs 22,000" in r and "NOT right" not in r, r)
    r = chat().say("dotnet winforms app called Till")
    check("visualstudio: a WinForms app builds and opens its window (on the hidden desktop)", "opens its window 'Till Inventory'" in r and "NOT right" not in r, r)
    broken = OUT / "visualstudio_broken"
    if not broken.exists():
        shutil_copy = __import__("shutil").copytree
        shutil_copy(next(OUT.rglob("visualstudio/Stock")), broken)
    core = broken / "Stock.Core" / "Inventory.cs"
    core.write_text(core.read_text(encoding="utf-8").replace("public decimal TotalValue =>", "public decimal TotalValue =>> "), encoding="utf-8")
    r = chat(files=[str(next(broken.glob("Stock.sln*")))]).say("visual studio build and test Stock.slnx")
    check("visualstudio: a solution with a mistake: the build fails and the error is given with its file and line", "FAILED" in r and "Inventory.cs:" in r, r[:300])
    check("visualstudio: nothing in AppData or the user folder (.NET and NuGet homes in tools/dotnet/home)", [p.exists() for p in leftovers] == before,
          str(list(zip(leftovers, before))))


def t_intellij():
    import shutil
    leftovers = [Path.home() / ".m2"]
    before = [p.exists() for p in leftovers]
    r = chat().say("intellij java project called 'Shop'")
    check("intellij: a Java 25 Maven project compiled, 5 JUnit tests passing, the jar packaged and run (Rs 22,000)",
          "5 of 5 tests passed" in r and "Rs 22,000" in r and "NOT right" not in r, r)
    broken = OUT / "java_broken"
    shutil.rmtree(broken, ignore_errors=True)
    shutil.copytree(next(OUT.rglob("intellij/Shop")), broken, ignore=shutil.ignore_patterns("target"))
    inv = next(broken.rglob("Inventory.java"))
    inv.write_text(inv.read_text(encoding="utf-8").replace("public long totalValue() {", "public long totalValue() { int x = \"text\";"), encoding="utf-8")
    r = chat(files=[str(broken / "pom.xml")]).say("intellij build and test pom.xml")
    check("intellij: a project with a type error: the build fails, the error given with its file and line", "FAILED" in r and "Inventory.java:" in r, r[:300])
    check("intellij: nothing in the user folder (Maven's repository in tools/maven/home)", [p.exists() for p in leftovers] == before, str(leftovers))


def t_clion():
    import winreg

    def kitware():
        try:
            winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Kitware"))
            return True
        except OSError:
            return False
    before = kitware()
    r = chat().say("clion c++ project called 'Shop'")
    check("clion: a CMake C++20 project built by clang (LLVM-MinGW) with Ninja, 5 CTest tests passing, the .exe run (Rs 22,000) and standalone",
          "CTest ran the tests: 5 of 5 passed" in r and "Rs 22,000" in r and "stands alone" in r and "NOT right" not in r, r)
    broken = OUT / "cpp_broken"
    shutil.rmtree(broken, ignore_errors=True)
    shutil.copytree(next(OUT.rglob("clion/Shop")), broken, ignore=shutil.ignore_patterns("build"))
    src = broken / "src" / "inventory.cpp"
    src.write_text(src.read_text(encoding="utf-8").replace("std::int64_t sum = 0;", "std::int64_t sum = \"text\";"), encoding="utf-8")
    r = chat(files=[str(broken / "CMakeLists.txt")]).say("cmake build CMakeLists.txt")
    check("clion: a project with a type error: the build fails, the error given with its file and line", "FAILED" in r and "inventory.cpp:" in r, r[:300])
    check("clion: CMake wrote nothing to the registry (HKCU\\Software\\Kitware)", kitware() == before, "")


def t_goland():
    leftovers = [Path(os.environ["APPDATA"]) / "go", Path(os.environ["LOCALAPPDATA"]) / "go-build", Path(os.environ["LOCALAPPDATA"]) / "go", Path.home() / "go"]
    before = [p.exists() for p in leftovers]
    r = chat().say("goland go project called 'Shop'")
    check("goland: a Go 1.27 module vetted, 5 tests passing, built to one .exe that prints Rs 22,000, gofmt-clean",
          "go test: 5 of 5 tests passed" in r and "Rs 22,000" in r and "gofmt finds every file formatted" in r and "NOT right" not in r, r)
    broken = OUT / "go_broken"
    shutil.rmtree(broken, ignore_errors=True)
    shutil.copytree(next(OUT.rglob("goland/Shop")), broken, ignore=shutil.ignore_patterns("*.exe"))
    src = broken / "inventory" / "inventory.go"
    src.write_text(src.read_text(encoding="utf-8").replace("var sum int64\n", "var sum int64 = \"text\"\n"), encoding="utf-8", newline="\n")
    r = chat(files=[str(broken / "go.mod")]).say("golang build go.mod")
    check("goland: a module with a type error: the build fails, the error given with its file and line", "FAILED" in r and "inventory.go:" in r, r[:300])
    check("goland: nothing in AppData or the user folder (caches, settings and telemetry in tools/go/home)", [p.exists() for p in leftovers] == before,
          str(leftovers))


def t_rustrover():
    leftovers = [Path.home() / ".cargo", Path.home() / ".rustup"]
    before = [p.exists() for p in leftovers]
    r = chat().say("rustrover rust project called 'Shop'")
    check("rustrover: a Rust 1.99 Cargo package: rustfmt, clippy clean, 5 unit tests passing, release .exe prints Rs 22,000 and needs only Windows DLLs",
          "clippy finds nothing" in r and "5 of 5 unit tests passed" in r and "Rs 22,000" in r and "only Windows' own DLLs" in r and "NOT right" not in r, r)
    broken = OUT / "rust_broken"
    shutil.rmtree(broken, ignore_errors=True)
    shutil.copytree(next(OUT.rglob("rustrover/Shop")), broken, ignore=shutil.ignore_patterns("target"))
    src = broken / "src" / "lib.rs"
    src.write_text(src.read_text(encoding="utf-8").replace("let mut out = String::new();", "let mut out: i64 = String::new();"), encoding="utf-8", newline="\n")
    r = chat(files=[str(broken / "Cargo.toml")]).say("cargo build Cargo.toml")
    check("rustrover: a package with a type error: the build fails, the error given with its file and line", "FAILED" in r and "lib.rs:" in r, r[:300])
    check("rustrover: nothing in the user folder (~/.cargo, ~/.rustup; Cargo's home is tools/rust/home)", [p.exists() for p in leftovers] == before, str(leftovers))


def t_pycharm():
    c = chat()
    r = c.say("pycharm python project called 'Shop'")
    check("pycharm: a Python project with its own .venv (tools/python 3.11), 5 unittest tests passing, python -m shop prints Rs 22,000",
          "5 of 5 tests passed" in r and "Rs 22,000" in r and "NOT right" not in r, r)
    r = c.say("vscode set up the last project")
    check("vscode: the last project set up for VS Code: tasks/settings/launch/extensions JSON read back, build and test tasks really pass",
          "read back as JSON" in r and "the test task's command passes" in r and "NOT right" not in r, r)
    vs = next(OUT.rglob("pycharm/Shop/.vscode"))
    tasks = json.loads((vs / "tasks.json").read_text(encoding="utf-8"))["tasks"]
    check("vscode: tasks.json has the default build and test tasks", {t["label"] for t in tasks} >= {"build", "test"}
          and any(t.get("group", {}).get("kind") == "test" for t in tasks), str([t["label"] for t in tasks]))
    r = c.say("which vscode extensions do I have")
    check("vscode: the installed extensions listed from code --list-extensions", re.search(r"VS Code has \d+ extension", r) is not None, r[:200])
    r = c.say("install the gitlens extension in vscode")
    check("vscode: installing an extension waits for a yes (or says it is already there)", "Say 'yes'" in r or "already has" in r, r)
    c.say("no")
    broken = OUT / "py_broken"
    shutil.rmtree(broken, ignore_errors=True)
    shutil.copytree(next(OUT.rglob("pycharm/Shop")), broken, ignore=shutil.ignore_patterns(".venv", "__pycache__"))
    src = broken / "shop" / "inventory.py"
    src.write_text(src.read_text(encoding="utf-8").replace("return sum(p.value for p", "return 1 + sum(p.value for p"), encoding="utf-8")
    r = chat(files=[str(broken / "pyproject.toml")]).say("pycharm test pyproject.toml")
    check("pycharm: a project with a wrong sum: the tests fail, the failing test given with its file and line", "FAILED" in r and "test_inventory.py:" in r, r[:300])


def t_nodejs():
    r = chat().say("webstorm typescript project called 'Shop'")
    check("nodejs: a TypeScript project run by Node itself (no npm downloads), 5 tests passing with node --test, main.ts prints Rs 22,000",
          "5 of 5 tests passed" in r and "Rs 22,000" in r and "NOT right" not in r, r)
    broken = OUT / "ts_broken"
    shutil.rmtree(broken, ignore_errors=True)
    shutil.copytree(next(OUT.rglob("nodejs/Shop")), broken)
    src = broken / "src" / "inventory.ts"
    src.write_text(src.read_text(encoding="utf-8").replace("sum += p.price * p.quantity;", "sum += p.price;"), encoding="utf-8", newline="\n")
    r = chat(files=[str(broken / "package.json")]).say("node test package.json")
    check("nodejs: a project with a wrong sum: the tests fail, the failing test named", "FAILED" in r and "total value is price times quantity" in r, r[:300])


def t_php():
    r = chat().say("php website called 'Shop'")
    check("php: a PHP 8.4 website: every file linted, 5 tests passing, the page served on 127.0.0.1 with Rs 22,000",
          "5 of 5 passed" in r and "HTTP 200" in r and "Rs 22,000" in r and "NOT right" not in r, r)
    broken = OUT / "php_broken"
    shutil.rmtree(broken, ignore_errors=True)
    shutil.copytree(next(OUT.rglob("php/Shop")), broken)
    src = broken / "src" / "Inventory.php"
    src.write_text(src.read_text(encoding="utf-8").replace("return array_sum(", "return 1 + array_sum("), encoding="utf-8", newline="\n")
    (broken / "src" / "Broken.php").write_text("<?php\nfunction ( {\n", encoding="utf-8")
    r = chat(files=[str(broken / "composer.json")]).say("php test composer.json")
    check("php: a project with a syntax error and a wrong sum: FAILED, the syntax error and the failing check with file and line",
          "FAILED" in r and "Broken.php:2" in r and "run.php:" in r, r[:400])


def t_sevenzip():
    docs = OUT / "zipme" / "docs"
    docs.mkdir(parents=True, exist_ok=True)
    (docs / "report.txt").write_text("Quarterly report\n", encoding="utf-8")
    (docs / "data.csv").write_text("x,y\n" + "\n".join(f"{i},{i * i}" for i in range(5000)), encoding="utf-8")
    r = chat().say(f"7zip pack {docs} with password Lahore123")
    check("sevenzip: a folder packed into .7z with AES-256: tests clean, every file with its size, unreadable (even names) without the password",
          "tests the archive clean" in r and "all 2 file(s)" in r and "not even its file names" in r and "Lahore123" not in r and "NOT right" not in r, r)
    archive = next(OUT.rglob("7zip/docs.7z"))
    r = chat(files=[str(archive)]).say("7zip extract docs.7z password Lahore123")
    check("sevenzip: unpacked with the password, every file back with its size", "all 2 file(s) unpacked" in r and "NOT right" not in r, r)


def t_obsidian():
    r = chat().say("obsidian vault for my exam prep with notes on Physics, Chemistry and Maths")
    check("obsidian: a vault from words: Home, 3 linked notes, daily note, canvas; every [[link]] and card resolves, settings JSON",
          "3 notes (Physics, Chemistry, Maths)" in r and "point at notes that exist" in r and "NOT right" not in r, r)
    (OUT / "course.md").write_text("# Course\n\n## Week 1: Basics\nVariables.\n\n## Week 2: Loops\nfor, while.\n", encoding="utf-8")
    r = chat(files=[str(OUT / "course.md")]).say("obsidian vault from course.md")
    check("obsidian: a vault from an outline: one note per ## heading, the # title names the vault, links resolve",
          "vault 'Course'" in r and "2 notes (Week 1 Basics, Week 2 Loops)" in r and "NOT right" not in r, r)


def t_lmms():
    leftovers = [Path.home() / ".lmmsrc.xml", Path.home() / "Documents" / "lmms", Path.home() / "lmms"]
    before = [p.exists() for p in leftovers]
    c = chat()
    r = c.say("lmms song: A4 B4 C5 D5 E5/2 D5/2 at 120 bpm with drums")
    check("lmms: a song (melody + kick drum) written as an LMMS project and rendered by LMMS: length, no clipping, 6 of 6 notes at pitch",
          "6 of 6 measured by FFT" in r and "does not clip" in r and "NOT right" not in r, r)
    c.say("melody 'Tune': C4 E4 G4 C5/2 G4 E4/2 at 100 bpm")
    r = c.say("lmms it as mp3")
    check("lmms: 'lmms it' renders the melody just written to MP3, every note at its pitch", "Tune.mp3" in r and "6 of 6" in r and "NOT right" not in r, r)
    check("lmms: nothing in the user folder (LMMS's settings come from tools/lmms/home)", [p.exists() for p in leftovers] == before, str(leftovers))


def t_visio():
    r = chat().say("visio flowchart: Start -> Take order -> Paid? -> Pack -> Ship -> End")
    check("visio: a flowchart written as .vsdx (6 shapes, 5 glued connectors, well-formed parts) that LibreOffice's Visio reader opens with all 6 labels",
          "6 shapes and 5 glued connectors" in r and "6 of 6 box labels" in r and "NOT right" not in r, r)


def t_autohotkey():
    c = chat()
    r = c.say("autohotkey script: brb -> be right back; addr -> House 12, Street 4, Lahore; ctrl+alt+n opens notepad; win+shift+d types the date")
    check("autohotkey: 2 text expansions + 2 hotkeys written as an AutoHotkey v2 script that AutoHotkey loads with no errors (nothing run)",
          "loads it with no errors" in r and "all 4 item(s)" in r and "NOT right" not in r, r)
    r = c.say("autohotkey start it")
    check("autohotkey: starting a script (live hotkeys) waits for a yes", "Say 'yes'" in r, r)
    c.say("no")
    (OUT / "bad.ahk").write_text("#Requires AutoHotkey v2.0\nx := (1 +\n", encoding="utf-8")
    r = chat(files=[str(OUT / "bad.ahk")]).say("autohotkey check bad.ahk")
    check("autohotkey: a script with an error: FAILED with its file and line", "FAILED" in r and "bad.ahk:2" in r, r)


def t_mongodb():
    r = chat().say("mongodb database for a shop with products, customers and orders")
    check("mongodb: a real local MongoDB 9.0: validated collections + unique name, bad documents refused, aggregation equals Python, JSON backup",
          "refuses a negative price" in r and "refuses a second product" in r and "equals the same sums done in Python" in r and "NOT right" not in r, r)
    (OUT / "msales.csv").write_text("date,city,product,qty,amount\n2026-10-01,Lahore,Fan,2,9000.50\n2026-10-01,Karachi,Iron,1,3200\n", encoding="utf-8")
    r = chat(files=[str(OUT / "msales.csv")]).say("mongodb import msales.csv")
    check("mongodb: a CSV imported as a collection with dates and numbers typed, every row in", "date date" in r and "qty int" in r and "all 2 documents" in r, r)
    check("mongodb: the server is shut down after the job", "mongod.exe" not in subprocess.run(["tasklist"], capture_output=True, text=True, creationflags=NOWIN).stdout, "")


def t_rawtherapee():
    from PIL import ImageDraw, ImageFilter
    leftovers = [Path(os.environ["LOCALAPPDATA"]) / "RawTherapee", Path(os.environ["APPDATA"]) / "RawTherapee"]
    before = [p.exists() for p in leftovers]
    im = Image.new("RGB", (900, 600), (110, 115, 105))
    d = ImageDraw.Draw(im)
    for i in range(12):
        d.rectangle([40 + i * 70, 80 + (i % 3) * 120, 90 + i * 70, 260 + (i % 3) * 120], fill=(120 + i * 6, 100 + (i * 13) % 50, 95 + (i * 7) % 40))
    im.filter(ImageFilter.GaussianBlur(2)).save(OUT / "dull.jpg", quality=92)
    r = chat(files=[str(OUT / "dull.jpg")]).say("rawtherapee dull.jpg brighter and vivid")
    check("rawtherapee: a photo developed by RawTherapee with a .pp3 profile from words: measured brighter and more colourful, same size",
          "Checked: brighter" in r and "more colourful" in r and "NOT right" not in r, r)
    r = chat(files=[str(OUT / "dull.jpg")]).say("rawtherapee dull.jpg black and white as tiff")
    check("rawtherapee: black and white as TIFF, measured colourless", ".tif" in r and "black and white (colour" in r and "NOT right" not in r, r)
    check("rawtherapee: nothing in AppData (its settings and cache in tools/rawtherapee/home)", [p.exists() for p in leftovers] == before, str(leftovers))


def t_keepassxc():
    (OUT / "logins.csv").write_text("name,url,username,password\nGmail,https://mail.google.com,ali@example.com,Gm-Secret-111\n"
                                    "Daraz,https://www.daraz.pk,ali,\nBank,https://bank.example,ali.k,Bk-Secret-222\n", encoding="utf-8")
    c = chat(files=[str(OUT / "logins.csv")])
    r = c.say("keepassxc vault called Home from logins.csv with password Test-Master-Pass-9")
    check("keepassxc: an encrypted vault from a browser export: 3 entries with user names and passwords (1 generated), wrong master password refused",
          "all 3 entries are in it" in r and "a wrong master password cannot open it" in r and "1 given new strong passwords" in r and "NOT right" not in r, r)
    log = (c.folder / "chat.json").read_text(encoding="utf-8")
    check("keepassxc: the master password is hidden in the saved chat log", "Test-Master-Pass-9" not in log and "[hidden]" in log, "")
    r = chat().say("keepassxc generate a password of 24 characters")
    check("keepassxc: a 24-character password generated by KeePassXC with all four character groups", "24 characters with lower case" in r, r[:60])


def t_powerbi():
    from harness.apps import powerbi
    (OUT / "sales.csv").write_text("date,city,product,qty,amount\n2026-10-01,Lahore,Fan,2,9000\n2026-10-01,Karachi,Iron,1,3200\n"
                                   "2026-10-02,Lahore,Kettle,3,6300\n", encoding="utf-8")
    r = chat(files=[str(OUT / "sales.csv")]).say("power bi report from sales.csv")
    check("powerbi: a Power BI Project from a CSV: model with DAX measures, a page with card/chart/table, every file valid against Microsoft's schemas",
          "validate against Microsoft's published schemas" in r and "Total amount = 18,500" in r and "NOT right" not in r, r)
    page = json.loads(next(OUT.rglob("powerbi/Sales/Sales.Report/definition/pages/page1/page.json")).read_text(encoding="utf-8"))
    page.pop("displayOption")
    check("powerbi: the schema check is real: a page without its required displayOption is caught", any("displayOption" in p for p in powerbi.validate(page)), "")


def t_flutter():
    leftovers = [Path(os.environ["APPDATA"]) / ".flutter", Path(os.environ["APPDATA"]) / ".dart-tool", Path(os.environ["LOCALAPPDATA"]) / "Pub",
                 Path(os.environ["LOCALAPPDATA"]) / ".dartServer", Path.home() / ".dart-tool", Path.home() / ".pub-cache", Path.home() / ".flutter-devtools"]
    before = [p.exists() for p in leftovers]
    r = chat().say("flutter app called 'Shop'")
    check("flutter: a Flutter app analysed clean, 6 tests (incl. a widget test reading Rs 22,000 off the screen), built for the web and served on 127.0.0.1",
          "no issues" in r and "6 of 6 tests passed" in r and "main.dart.js" in r and "NOT right" not in r, r)
    broken = OUT / "flutter_broken"
    shutil.rmtree(broken, ignore_errors=True)
    shutil.copytree(next(OUT.rglob("flutter/Shop")), broken, ignore=shutil.ignore_patterns("build", ".dart_tool"))
    src = broken / "lib" / "inventory.dart"
    src.write_text(src.read_text(encoding="utf-8").replace("int get totalValue =>", "int get totalValue => 'text'.length +"), encoding="utf-8", newline="\n")
    src.write_text(src.read_text(encoding="utf-8").replace("final old = _items[", "final int old = _items["), encoding="utf-8", newline="\n")
    r = chat(files=[str(broken / "pubspec.yaml")]).say("flutter test pubspec.yaml")
    check("flutter: a project with a type error: FAILED with the file and line", "FAILED" in r and "inventory.dart:" in r, r[:300])
    check("flutter: nothing in AppData or the user folder (pub cache, settings, analytics in tools/flutter_home)", [p.exists() for p in leftovers] == before,
          str(leftovers))


def t_androidstudio():
    import re
    leftovers = [Path.home() / ".android", Path.home() / ".gradle", Path.home() / ".m2", Path.home() / ".kotlin"]
    before = [p.exists() for p in leftovers]
    c = chat()
    r = c.say("android studio app called 'Shop'")
    check("androidstudio: a real APK built by Gradle (AGP, SDK 36), 5 unit tests passing, read back by aapt2, signature verified by apksigner, dex checked",
          "5 of 5 unit tests passed" in r and "aapt2 reads the APK" in r and "apksigner verifies" in r and "NOT right" not in r, r)
    r = c.say("install it on my phone")
    check("androidstudio: installing waits for a phone: none connected, so it says how to connect one and asks no yes",
          "No Android phone is connected" in r and "Say 'yes'" not in r, r)
    broken = OUT / "android_broken"
    shutil.rmtree(broken, ignore_errors=True)
    shutil.copytree(next(OUT.rglob("androidstudio/Shop")), broken, ignore=shutil.ignore_patterns("build", ".gradle"))
    inv = next(broken.rglob("Inventory.java"))
    inv.write_text(inv.read_text(encoding="utf-8").replace("public long totalValue() {", "public long totalValue() { int x = \"text\";"), encoding="utf-8")
    r = chat(files=[str(broken / "settings.gradle.kts")]).say("android build settings.gradle.kts")
    check("androidstudio: a project with a type error: the build fails, the error given with its file and line", "FAILED" in r and "Inventory.java:" in r, r[:300])
    running = subprocess.run(["tasklist"], capture_output=True, text=True, creationflags=NOWIN).stdout
    check("androidstudio: nothing in the user folder (~/.android, ~/.gradle, ~/.m2, ~/.kotlin); adb not left running",
          [p.exists() for p in leftovers] == before and not re.search(r"(?im)^adb\.exe", running), str(leftovers))


def t_postgres():
    leftovers = [Path(os.environ["APPDATA"]) / "postgresql", Path.home() / ".psql_history"]
    before = [p.exists() for p in leftovers]
    r = chat().say("postgres database for a shop with products, customers and orders")
    check("postgres: a real local PostgreSQL 18: shop tables with keys, wrong rows refused, report sums equal Python's, pg_dump backup",
          "refuses an order for a customer that does not exist" in r and "equal the same sums done in Python" in r and "NOT right" not in r, r)
    (OUT / "sales.csv").write_text("date,city,product,qty,amount\n2026-10-01,Lahore,Fan,2,9000\n2026-10-01,Karachi,Iron,1,3200\n"
                                   "2026-10-02,Lahore,Kettle,3,6300\n", encoding="utf-8")
    r = chat(files=[str(OUT / "sales.csv")]).say("postgres import sales.csv")
    check("postgres: a CSV imported as a table with worked-out types (timestamp, bigint, numeric), every row in", "all 3 rows are in it" in r
          and "date timestamp" in r and "qty bigint" in r, r)
    (OUT / "pg_evil.sql").write_text(f"\\! echo hacked > {(OUT / 'pg_hacked.txt').as_posix()}\nSELECT 1;\n", encoding="utf-8")
    r = chat(files=[str(OUT / "pg_evil.sql")]).say("postgres run pg_evil.sql")
    check("postgres: a .sql file runs in psql's restricted mode: its \\! shell command is refused", "FAILED" in r and "restricted" in r
          and not (OUT / "pg_hacked.txt").exists(), r[:200])
    check("postgres: the server is stopped after the job", "postgres.exe" not in subprocess.run(["tasklist"], capture_output=True, text=True,
                                                                                              creationflags=NOWIN).stdout, "")
    check("postgres: nothing in AppData or the user folder (history and settings in tools/postgres/home)", [p.exists() for p in leftovers] == before,
          str(leftovers))


def t_mysql():
    leftovers = [Path(os.environ["APPDATA"]) / "MySQL", Path(os.environ["LOCALAPPDATA"]) / "MySQL", Path(r"C:\ProgramData\MySQL"), Path.home() / ".mysql_history"]
    before = [p.exists() for p in leftovers]
    r = chat().say("mysql database for a shop with products, customers and orders")
    check("mysql: a real local MySQL 9.7: shop tables with keys, wrong rows (customer, price, status) refused, report sums equal Python's, mysqldump",
          "refuses an order for a customer that does not exist" in r and "refuses an order status" in r and "equal the same sums done in Python" in r
          and "NOT right" not in r, r)
    (OUT / "mysales.csv").write_text("date,city,product,qty,amount\n2026-10-01,Lahore,Fan,2,9000.50\n2026-10-01,Karachi,Iron,1,3200\n"
                                     "2026-10-02,Lahore,Kettle,3,6300\n", encoding="utf-8")
    r = chat(files=[str(OUT / "mysales.csv")]).say("mysql import mysales.csv")
    check("mysql: a CSV imported as a table with worked-out types (datetime, bigint, decimal), every row in", "all 3 rows are in it" in r
          and "date datetime" in r and "qty bigint" in r and "amount decimal" in r, r)
    (OUT / "evil.sql").write_text(f"system echo hacked > {(OUT / 'hacked.txt').as_posix()}\nSELECT 1;\n", encoding="utf-8")
    r = chat(files=[str(OUT / "evil.sql")]).say("mysql run evil.sql")
    check("mysql: a .sql file can only send SQL: its 'system' line is refused, no shell command runs", "FAILED" in r and not (OUT / "hacked.txt").exists(), r[:200])
    check("mysql: the server is shut down after the job; nothing in AppData, ProgramData or the user folder",
          "mysqld.exe" not in subprocess.run(["tasklist"], capture_output=True, text=True, creationflags=NOWIN).stdout
          and [p.exists() for p in leftovers] == before, str(leftovers))


def main():
    t0 = time.time()
    clear(OUT)
    OUT.mkdir(parents=True)
    for t in (t_unity, t_premiere, t_photoshop, t_aftereffects, t_obs, t_google, t_github, t_services, t_devtools, t_godot, t_illustrator, t_revit, t_lightroom, t_anki, t_arduino, t_msproject, t_prusaslicer, t_freecad, t_matlab, t_musescore, t_audacity, t_calibre, t_kicad, t_krita, t_openscad, t_libreoffice, t_gimp, t_rawtherapee, t_qgis, t_shotcut, t_latex, t_rstats, t_drawio, t_visio, t_handbrake, t_sevenzip, t_obsidian, t_lmms, t_autohotkey, t_keepassxc, t_clion, t_goland, t_rustrover, t_pycharm, t_nodejs, t_php, t_visualstudio, t_intellij, t_androidstudio, t_flutter, t_postgres, t_mysql, t_mongodb, t_powerbi):
        try:
            t()
        except Exception:  # noqa: BLE001 - one app failing does not hide the others
            import traceback
            check(f"{t.__name__}: ran without an error", False, traceback.format_exc()[-900:])
    bad = [n for n, ok in RESULTS if not ok]
    print(f"\n{'ALL PASS' if not bad else f'{len(bad)} FAILED'}  ({len(RESULTS) - len(bad)}/{len(RESULTS)}, {time.time() - t0:.0f} s)")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
