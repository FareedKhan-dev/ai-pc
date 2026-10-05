"""Engine tests for data that crosses files and for PDFs (no model; Office renders in the background).

  .venv\\Scripts\\python.exe tests\\integration\\test_links.py
About a minute: Excel reads a workbook as it shows it; Word and PowerPoint get linked tables and charts and refresh
them; a deck survives deleting and adding slides; PDFs are packed, numbered, watermarked, compressed and split.
"""
import json
import shutil
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")
import warnings  # noqa: E402

warnings.simplefilter("ignore")
from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from pptx import Presentation  # noqa: E402

from ai_pc.office import docmap as DM  # noqa: E402
from ai_pc.office import docplan as DP  # noqa: E402
from ai_pc.office import docx_build as DB  # noqa: E402
from ai_pc.office import docx_data as DD  # noqa: E402
from ai_pc.office import docx_ops as DO  # noqa: E402
from ai_pc.office import pdf_ops as PF  # noqa: E402
from ai_pc.office import pptx_build as PB  # noqa: E402
from ai_pc.office import pptx_ops as PO  # noqa: E402
from ai_pc.office import render as RN  # noqa: E402

OUT = ROOT / "out" / "docs" / "_tests" / "links"
FAILS = []


def check(name, ok, detail=""):
    print(f"{'ok  ' if ok else 'FAIL'} {name}" + (f"  ({detail})" if not ok and detail else ""))
    if not ok:
        FAILS.append(name)


def main():
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir(parents=True)
    srv = RN.Server()
    try:
        # ---------------------------------------------------------------- Excel shows it, the bridge reads it so
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["City", "Month", "Amount"])
        rows = [("Lahore", "Jan", 120000), ("Karachi", "Jan", 98000), ("Lahore", "Feb", 135000), ("Islamabad", "Feb", 64000), ("Karachi", "Feb", 77000)]
        for r in rows:
            ws.append(list(r))
        for c in ws["C"][1:]:
            c.number_format = '"Rs "#,##0'
        book = OUT / "sales.xlsx"
        wb.save(str(book))
        r = srv.run({"app": "excel", "src": str(book.resolve()), "save": True, "ops": [{"op": "pivot", "rows": ["City"], "values": ["Amount"], "fn": "sum"}]})
        check("Excel makes a pivot and checks it against Python", r.get("ok") and r["ops"][0].get("ok"), json.dumps(r.get("ops"))[:300])
        r = srv.run({"app": "excel", "src": str(book.resolve()), "grab": [{"sheet": "Pivot by City"}, {"sheet": "Sales", "what": "table"}]})
        piv, tab = r["grab"]
        check("the pivot comes back as Excel shows it", piv["rows"] == [["Islamabad", "Rs 64,000"], ["Karachi", "Rs 175,000"], ["Lahore", "Rs 255,000"]]
              and piv["total"] == ["Grand Total", "Rs 494,000"], json.dumps(piv)[:300])
        check("with the numbers behind it", [v[1] for v in piv["values"]] == [64000.0, 175000.0, 255000.0])
        check("a table comes back whole", tab["header"] == ["City", "Month", "Amount"] and len(tab["rows"]) == 5 and tab["rows"][0][2] == "Rs 120,000", json.dumps(tab)[:200])

        # ---------------------------------------------------------------- Word: linked table and chart, refreshed in place
        plan = DP.resolve({"doctype": "report", "title": "Links", "theme": "corporate", "cover": False, "toc": False, "blocks": [
            {"type": "heading", "text": "Introduction", "level": 1}, {"type": "paragraph", "text": "Sales grew in every city. " * 6},
            {"type": "heading", "text": "Outlook", "level": 1}, {"type": "paragraph", "text": "The second quarter looks strong. " * 6}]})
        rep = OUT / "report.docx"
        DB.build(plan, str(rep))
        doc = Document(str(rep))
        ctx = {"planner": None, "info": {}}
        m0 = DM.docx_map(doc)
        t_op = {"op": "data_table", "data": piv, "link": "aisrc_1", "anchor": {"kind": "section", "name": "Introduction"}, "where": "after", "caption": "Amount by city"}
        c_op = {"op": "data_chart", "data": piv, "link": "aisrc_2", "anchor": {"kind": "section", "name": "Introduction"}, "where": "after", "caption": "Amount by city", "chart": "pie"}
        DO.apply(doc, t_op, ctx)
        DO.apply(doc, c_op, ctx)
        m1 = DM.docx_map(doc)
        ok_t, what_t = DO.check(t_op, m0, m1, {}, doc_after=doc)
        ok_c, what_c = DO.check(c_op, m0, m1, {}, doc_after=doc)
        check("a linked Word table shows exactly Excel's cells", ok_t, what_t)
        check("a linked Word chart holds exactly the values", ok_c, what_c)
        heads = [it["text"] for it in m1["items"]]
        intro, outlook = heads.index("Introduction"), heads.index("Outlook")
        caps = [i for i, x in enumerate(heads) if x.startswith(("Table 1", "Figure 1"))]
        check("both land in the Introduction, table then chart", len(caps) == 2 and all(intro < i < outlook for i in caps) and heads[caps[0]].startswith("Table"), str(heads))
        piv2 = json.loads(json.dumps(piv))
        piv2["rows"][0][1], piv2["values"][0][1] = "Rs 99,000", 99000.0
        for op in ({"op": "data_refresh", "link": "aisrc_1", "data": piv2}, {"op": "data_refresh", "link": "aisrc_2", "data": piv2, "chart": "pie"}):
            DO.apply(doc, op, ctx)
            ok, what = DO.check(op, None, None, {}, doc_after=doc)
            check(f"refreshed in place ({'table' if op['link'] == 'aisrc_1' else 'chart'})", ok, what)
        doc.save(str(OUT / "report_linked.docx"))
        d2 = Document(str(OUT / "report_linked.docx"))
        check("the links survive saving", DD.links(d2) == ["aisrc_1", "aisrc_2"], str(DD.links(d2)))
        caps2 = [it["text"] for it in DM.docx_map(d2)["items"] if it["text"].startswith(("Table ", "Figure "))]
        check("one caption each after the refresh", caps2 == ["Table 1: Amount by city", "Figure 1: Amount by city"], str(caps2))
        r = srv.run({"app": "word", "src": str((OUT / "report_linked.docx").resolve()), "pdf": str((OUT / "report_linked.pdf").resolve()), "update_fields": True})
        check("Word draws the refreshed figure", r.get("ok") and "Rs 99,000" in " ".join(RN.text(str(OUT / "report_linked.pdf"))))

        # ---------------------------------------------------------------- PowerPoint: data slides, deletions, copies
        deck = OUT / "deck.pptx"
        PB.build_deck({"title": "Q1", "theme": "modern", "slides": [{"layout": "title", "title": "Q1"}, {"layout": "bullets", "title": "Context", "bullets": ["Growth", "Risks"]},
                                                                       {"layout": "bullets", "title": "Plan", "bullets": ["Hire", "Expand"]},
                                                                       {"layout": "closing", "title": "Thank you"}]}, str(deck))
        prs = Presentation(str(deck))
        ctx = {"planner": None, "info": {}, "th": PO.deck_theme(prs)}
        PO.apply(prs, {"op": "slide_delete", "target": {"kind": "slide", "n": 2}}, ctx)
        ctx["info"] = {}
        s_op = {"op": "data_slide", "kind": "chart", "data": piv, "link": "aisrc_3", "after": 1, "title": "Amount by city"}
        PO.apply(prs, s_op, ctx)
        ok, what = PO.check(s_op, None, None, ctx["info"], prs)
        check("a data slide's chart holds exactly the values", ok, what)
        takeaway = next((sh.text_frame.text for sh in prs.slides[1].shapes if sh.name == "Takeaway"), "")
        check("its takeaway is computed from the numbers", "Lahore leads with Rs 255,000, 52% of the total" in takeaway, takeaway)
        ctx["info"] = {}
        t2 = {"op": "data_slide", "kind": "table", "data": tab, "link": "aisrc_4", "after": 2, "title": "Sales"}
        PO.apply(prs, t2, ctx)
        ok, what = PO.check(t2, None, None, ctx["info"], prs)
        check("a data slide's table shows exactly Excel's cells", ok, what)
        PO.apply(prs, {"op": "slide_duplicate", "slide": 2}, ctx)
        a = next(sh for sh in prs.slides[1].shapes if getattr(sh, "has_chart", False) and sh.has_chart).chart
        b = next(sh for sh in prs.slides[2].shapes if getattr(sh, "has_chart", False) and sh.has_chart).chart
        check("a copied chart slide gets its own chart, look and data kept", a.part is not b.part and list(a.plots[0].series[0].values) == list(b.plots[0].series[0].values)
              and b.part.chart_workbook.xlsx_part is not None)
        tab2 = json.loads(json.dumps(tab))
        tab2["rows"] = tab2["rows"][:3]
        op = {"op": "data_refresh", "link": "aisrc_4", "data": tab2}
        PO.apply(prs, op, ctx)
        ok, what = PO.check(op, None, None, {}, prs)
        check("a table with fewer rows is drawn again in its place", ok, what)
        prs.save(str(OUT / "deck_linked.pptx"))
        names = zipfile.ZipFile(str(OUT / "deck_linked.pptx")).namelist()
        check("no part is written twice after deleting and adding slides", len(names) == len(set(names)), str(len(names) - len(set(names))))
        r = srv.run({"app": "powerpoint", "src": str((OUT / "deck_linked.pptx").resolve()), "pdf": str((OUT / "deck_linked.pdf").resolve())})
        check("PowerPoint opens and draws the deck", r.get("ok") and r.get("slides") == len(Presentation(str(OUT / "deck_linked.pptx")).slides), str(r.get("error")))

        # ---------------------------------------------------------------- PDFs
        pk = PF.pack([("Report", str(OUT / "report_linked.pdf")), ("Deck", str(OUT / "deck_linked.pdf"))], OUT / "pack" / "pack.pdf",
                     cover={"title": "Q1 Pack", "subtitle": "Test"}, numbers=True, watermark="DRAFT", server=srv)
        check("a pack: pages, bookmarks, page numbers, watermark", all(c["ok"] for c in pk["checks"]), "; ".join(c["what"] for c in pk["checks"]))
        check("its numbers go where each page has room", "right" in (pk.get("number_places") or {}) or "centre" in (pk.get("number_places") or {}), str(pk.get("number_places")))
        cp = PF.compress(pk["path"], OUT / "pack" / "small.pdf")
        check("compressed: smaller, same pages and text", all(c["ok"] for c in cp["checks"]), "; ".join(c["what"] for c in cp["checks"]))
        sp = PF.split(pk["path"], OUT / "pack" / "parts")
        check("split by bookmarks keeps every page", sp["checks"][0]["ok"] and len(sp["files"]) == 3, sp["checks"][0]["what"])
        se = PF.select(pk["path"], OUT / "pack" / "sel.pdf", "2-3")
        de = PF.delete(pk["path"], OUT / "pack" / "del.pdf", "1")
        ro = PF.rotate(pk["path"], OUT / "pack" / "rot.pdf", "1", 90)
        check("pages selected, deleted and rotated", se["pages"] == 2 and de["pages"] == pk["pages"] - 1 and ro["checks"][0]["ok"])
    finally:
        srv.close()
    print(f"\n{'ALL OK' if not FAILS else f'{len(FAILS)} FAILED: ' + '; '.join(FAILS)}")
    return not FAILS


if __name__ == "__main__":
    sys.exit(0 if main() else 1)
