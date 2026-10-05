"""A two-page design saved as PowerPoint the way Canva exports one (1920 x 1080 px pages, every colour and font set on the
shape itself): a cropped photo, titles in a Google font named the PowerPoint way ('Montserrat Bold'), mixed text runs, a
rounded button with centred text, an oval, a group, a line, a freeform triangle, bullets, letter spacing, capitals, a
link, and a chart (which has no web equivalent and must be cut from the page's picture)."""

from pathlib import Path

PX = 9525  # EMU per px


def make(path, photo):
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    def px(v):
        return Emu(int(v * PX))

    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1920), px(1080)
    blank = prs.slide_layouts[6]

    def bg(slide, hexs):
        cs = slide._element.find(qn("p:cSld"))
        b = etree.SubElement(cs, qn("p:bg"))
        cs.insert(0, b)
        pr = etree.SubElement(b, qn("p:bgPr"))
        sf = etree.SubElement(pr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=hexs)
        etree.SubElement(pr, qn("a:effectLst"))

    def box(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
        tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = px(0)
        tf.margin_top = tf.margin_bottom = px(0)
        for i, (runs, align, spacing) in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if spacing:
                p.line_spacing = spacing
            for text, font, size, color, extra in runs:
                r = p.add_run()
                r.text = text
                r.font.name, r.font.size = font, Pt(size)
                r.font.color.rgb = RGBColor.from_string(color)
                rpr = r._r.get_or_add_rPr()
                for k, v in (extra or {}).items():
                    if k == "link":
                        r.hyperlink.address = v
                    else:
                        rpr.set(k, str(v))
        return tb

    s1 = prs.slides.add_slide(blank)
    bg(s1, "FFF7ED")
    pic = s1.shapes.add_picture(str(photo), px(0), px(0), px(1920), px(420))
    pic.crop_top, pic.crop_bottom = 0.1, 0.1
    box(s1, 160, 470, 1600, 110, [([("Eid Mubarak Sale", "Montserrat Bold", 66, "7C2D12", None)], PP_ALIGN.CENTER, 1.0)])
    box(
        s1,
        260,
        600,
        1400,
        50,
        [([("Up to 40% off on ", "Montserrat", 24, "9A3412", None), ("all phones", "Montserrat Bold", 24, "EA580C", None)], PP_ALIGN.CENTER, 1.0)],
    )
    btn = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(810), px(690), px(300), px(80))
    btn.fill.solid()
    btn.fill.fore_color.rgb = RGBColor.from_string("EA580C")
    btn.line.fill.background()
    btn.adjustments[0] = 0.5
    tf = btn.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = px(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Shop now"
    r.font.name, r.font.size, r.font.bold = "Montserrat", Pt(21), True
    r.font.color.rgb = RGBColor.from_string("FFFFFF")
    r._r.get_or_add_rPr().set("cap", "all")
    r._r.get_or_add_rPr().set("spc", "150")
    oval = s1.shapes.add_shape(MSO_SHAPE.OVAL, px(1640), px(820), px(160), px(160))
    oval.fill.solid()
    oval.fill.fore_color.rgb = RGBColor.from_string("FDBA74")
    oval.line.fill.background()
    grp = s1.shapes.add_group_shape()
    for i, col in enumerate(("FB923C", "F97316", "C2410C")):
        rr = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(160 + i * 70), px(840), px(50), px(100 - i * 20))
        rr.fill.solid()
        rr.fill.fore_color.rgb = RGBColor.from_string(col)
        rr.line.fill.background()
    ln = s1.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(160), px(800), px(1760), px(800))
    ln.line.color.rgb = RGBColor.from_string("FDBA74")
    ln.line.width = px(4)
    ff = s1.shapes.build_freeform(px(1500), px(980))
    ff.add_line_segments([(px(1580), px(860)), (px(1560), px(980))])
    tri = ff.convert_to_shape()
    tri.fill.solid()
    tri.fill.fore_color.rgb = RGBColor.from_string("9A3412")
    tri.line.fill.background()
    lst = box(
        s1,
        500,
        850,
        700,
        150,
        [
            ([("Free delivery in Lahore", "Arial", 18, "431407", None)], PP_ALIGN.LEFT, None),
            ([("Cash on delivery", "Arial", 18, "431407", None)], PP_ALIGN.LEFT, None),
            ([("Visit our shop", "Arial", 18, "C2410C", {"u": "sng", "link": "https://example.com/shop"})], PP_ALIGN.LEFT, None),
        ],
    )
    for p in list(lst.text_frame.paragraphs)[:2]:
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", str(px(28)))
        ppr.set("indent", str(-px(28)))
        etree.SubElement(ppr, qn("a:buChar"), char="•")

    s2 = prs.slides.add_slide(blank)
    bg(s2, "FFFFFF")
    box(s2, 160, 120, 1600, 90, [([("Our best sellers", "Montserrat Bold", 48, "111827", None)], PP_ALIGN.LEFT, 1.0)])
    cd = CategoryChartData()
    cd.categories = ["Phones", "Laptops", "TVs"]
    cd.add_series("Sold", (120, 80, 45))
    s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, px(160), px(260), px(900), px(600), cd)
    box(
        s2,
        1140,
        280,
        620,
        300,
        [([("Phones lead our sales this month, with laptops close behind.", "Arial", 24, "374151", None)], PP_ALIGN.LEFT, 1.2)],
    )
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path
