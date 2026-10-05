"""A long conversation about a project of several files, each turn checked on the files it produced: a messy client
workbook becomes a cleaned workbook, a report written from its figures, a deck made from the report, exact tables and
charts linked from the workbook into both, a PDF pack, a handout; the workbook changes, the copies are found out of date
and refreshed; undo and go back move every file together.

  .venv\\Scripts\\python.exe tests\\integration\\project_conversations.py [--offline]
--offline: no model (the turns that need one are skipped, and so are the turns that need what they make).
"""

import json
import re
import shutil
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")
import warnings  # noqa: E402

warnings.simplefilter("ignore")
from docx import Document  # noqa: E402
from messy_books import sales_book  # noqa: E402
from pptx import Presentation  # noqa: E402

from ai_pc.office import pptx_ops as PO  # noqa: E402
from ai_pc.office import xlsx_map as XM  # noqa: E402
from ai_pc.office.bookchat import _fmt  # noqa: E402
from ai_pc.office.docx_data import links as doc_links  # noqa: E402
from ai_pc.office.projectchat import ProjectChat  # noqa: E402

OUT = ROOT / "out" / "docs" / "_tests" / "projects"


class W:
    def __init__(self, pc, v0, reply, turn):
        self.pc, self.v0, self.reply, self.turn = pc, v0, reply, turn

    def says(self, *words):
        return all(re.search(w, self.reply, re.I) for w in words)

    def key(self, role):
        return next((k for k in self.pc.active() if self.pc.state["files"][k]["role"] == role), None)

    def book(self):
        c = self.pc.chat(self.key("workbook"))
        return XM.book_map(c.path())

    def recs(self):
        m = self.book()
        return [r for r in XM.records(XM.sheet(m)) if any(v not in (None, "") for v in r.values())]

    def total(self):
        m = self.book()
        s = XM.sheet(m)
        col = XM.col(None, s, "Amount")
        return sum(float(r["Amount"]) for r in self.recs() if isinstance(r.get("Amount"), (int, float))), col

    def doc(self, role="report"):
        k = self.key(role)
        return Document(str(self.pc.chat(k).path())) if k else None

    def deck(self):
        k = self.key("deck")
        return Presentation(str(self.pc.chat(k).path())) if k else None

    def links_to(self, role):
        k = self.key(role)
        return [lid for lid, L in self.pc.state["links"].items() if L["dst"] == k]

    def checks_ok(self, role):
        v = self.pc.chat(self.key(role)).version
        return bool(v.get("checks")) and all(c["ok"] for c in v["checks"])

    def pdfs(self):
        return self.pc.state["pdfs"][: self.pc.version["pdfs"]]


def _lahore(x):
    rs = [r for r in x.recs() if r.get("City") == "Lahore" and isinstance(r.get("Amount"), (int, float))]
    tot, col = x.total()
    return _fmt(sum(float(r["Amount"]) for r in rs), col)


def _heading_colour(doc):
    st = doc.styles["Heading 1"]
    c = st.font.color
    return str(c.rgb) if c is not None and c.type is not None else None


def _title_colours(prs):
    return {PO._rgb_of(r.font.color) for k in range(1, len(prs.slides) + 1) for r in PO._part_runs(prs.slides[k - 1], "titles")}


TURNS = [
    ("what files are in the project?", "question", lambda x: x.says("workbook"), False),
    (
        "clean up the data and add a column Amount = Qty x Unit Price with a total row",
        "file",
        lambda x: (
            "Amount" in [c["name"] for c in XM.sheet(x.book())["table"]["cols"]]
            and XM.sheet(x.book())["table"]["total"]
            and len(x.recs()) == 16
            and all(abs(float(r["Amount"]) - float(r["Qty"]) * float(r["Unit Price"])) < 1e-6 for r in x.recs())
        ),
        False,
    ),
    (
        "write a two page report on Q1 sales for the management team from the workbook",
        "report_from_book",
        lambda x: x.key("report") is not None and _fmt(x.total()[0], x.total()[1]).replace("Rs ", "") in " ".join(p.text for p in x.doc().paragraphs),
        True,
    ),
    (
        "put the amount by city into the report after the introduction, with a pie chart of it",
        "insert_data",
        lambda x: len(x.links_to("report")) == 2 and x.checks_ok("report") and set(doc_links(x.doc())) >= set(x.links_to("report")),
        True,
    ),
    (
        "make a 5 slide deck from the report for the board",
        "deck_from_report",
        lambda x: x.key("deck") is not None and len(x.links_to("deck")) == 1 and x.checks_ok("deck"),
        True,
    ),
    (
        "add a chart of amount by month to the deck after slide 2",
        "insert_data",
        lambda x: len(x.links_to("deck")) == 2 and PO.linked(x.deck(), x.links_to("deck")[-1])[0] == 3 and x.checks_ok("deck"),
        True,
    ),
    (
        "make the report's headings dark blue and make the deck's titles dark green",
        "file",
        lambda x: _heading_colour(x.doc()) == "1F3864" and _title_colours(x.deck()) == {"1B5E20"},
        True,
    ),
    ("how much did Lahore sell?", "file", lambda x: x.says(re.escape(_lahore(x))), False),
    (
        "in the workbook, delete the cancelled orders",
        "file",
        lambda x: len(x.recs()) == 15 and (x.key("report") is None or x.says("out of date", "the text says")),
        False,
    ),
    ("is everything up to date?", "question", lambda x: x.says("out of date") or x.key("report") is None, False),
    (
        "refresh everything",
        "refresh",
        lambda x: (
            not x.pc.stale()
            and not x.pc.stale_figures()
            and x.checks_ok("report")
            and x.checks_ok("deck")
            and _fmt(x.total()[0], x.total()[1]).replace("Rs ", "") in x.pc.file_text(x.key("report"))
        ),
        True,
    ),
    ("what is the total amount now?", "file", lambda x: x.says(re.escape(_fmt(x.total()[0], x.total()[1]))), False),
    (
        "export the report and the deck as one pdf with a cover called Q1 Sales Pack, page numbers and a DRAFT watermark",
        "pdf_pack",
        lambda x: x.pdfs() and x.pdfs()[-1]["kind"] == "pack" and x.says(r"Checked: 4/4 OK"),
        True,
    ),
    ("make the pdf smaller", "pdf_compress", lambda x: x.pdfs()[-1]["kind"] == "compressed" and x.says(r"Checked: 2/2 OK"), True),
    ("undo", "undo", lambda x: x.pdfs()[-1]["kind"] == "pack", True),
    (
        "make a handout from the deck with the speaker notes",
        "handout",
        lambda x: x.key("handout") is not None and x.says("every slide is in it"),
        True,
    ),
    ("what files are in the project?", "question", lambda x: x.says("report", "deck", "handout"), True),
    ("history", "history", lambda x: x.says(r"project version"), False),
    ("go back to v2", "goto", lambda x: x.pc.state["cur"] == 2, False),
]


def run(planner):
    shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir(parents=True, exist_ok=True)
    book = sales_book(OUT / "sales_q1.xlsx")
    pc = ProjectChat.start([book], name="Q1 sales", projects_dir=OUT, planner=planner)
    rows, t_all, _made = [], time.perf_counter(), True
    try:
        for msg, want, check, needs_model in TURNS:
            if needs_model and planner is None:
                print(f"skip (needs the model or what it makes)  {msg}")
                continue
            v0 = pc.state["cur"]
            t = time.perf_counter()
            err = None
            try:
                reply = pc.say(msg)
            except Exception as e:  # noqa: BLE001
                import traceback

                traceback.print_exc()
                reply, err = "", f"{type(e).__name__}: {e}"
                pc.last_turn = {"intents": ["crash"]}
            secs = time.perf_counter() - t
            intent_ok = want is None or want in (pc.last_turn.get("intents") or [])
            try:
                check_ok = bool(check(W(pc, v0, reply, pc.last_turn)))
            except Exception as e:  # noqa: BLE001
                check_ok, err = False, err or f"check {type(e).__name__}: {e}"
            ok = intent_ok and check_ok and not err
            rows.append(
                {
                    "msg": msg,
                    "ok": ok,
                    "intents": pc.last_turn.get("intents"),
                    "seconds": round(secs, 1),
                    "llm": bool(pc.last_turn.get("llm")),
                    "reply": reply[:900],
                    "error": err,
                }
            )
            print(
                f"{'OK ' if ok else 'BAD'} [{','.join(pc.last_turn.get('intents') or [])}{'+llm' if pc.last_turn.get('llm') else ''}] "
                f"pv{v0}->pv{pc.state['cur']} {secs:5.1f}s  {msg}\n      {reply[:600]}",
                flush=True,
            )
            if not ok:
                print(f"      intent_ok={intent_ok} check_ok={check_ok} {err or ''}")
    finally:
        pc.close()
    n, ok = len(rows), sum(r["ok"] for r in rows)
    usd = planner.cost()[1] if planner is not None else 0.0
    print(f"\n{ok}/{n} turns OK; {sum(r['seconds'] for r in rows):.0f} s of turns; {time.perf_counter() - t_all:.0f} s in all; ${usd:.4f}")
    (OUT / "report.json").write_text(json.dumps({"ok": ok, "turns": n, "rows": rows}, ensure_ascii=False, indent=1), encoding="utf-8")
    return ok, n


if __name__ == "__main__":
    planner = None
    if "--offline" not in sys.argv:
        from ai_pc.llm.planner import ChatPlanner

        planner = ChatPlanner()
    ok, n = run(planner)
    sys.exit(0 if ok == n else 1)
