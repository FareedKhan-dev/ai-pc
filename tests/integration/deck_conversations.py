"""Multi-turn conversations about PowerPoint decks, each turn checked on the deck it produced (PowerPoint renders and
measures every version in the background).

  .venv\\Scripts\\python.exe tests\\integration\\deck_conversations.py [solar] [client] [--offline]
--offline: no model (the turns that need one are skipped). Chats go to out/docs/_tests/deck_chats (wiped at each run).
solar: a deck the agent made (its own shapes, cards, a native chart); client: a deck typed into PowerPoint's default
template (placeholders, inherited theme colours, a table in the table style).
"""

import json
import re
import shutil
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from ai_pc.office import pptx_ops as PO  # noqa: E402
from ai_pc.office.deckchat import DeckChat  # noqa: E402
from ai_pc.office.docx_ops import colour  # noqa: E402

CHATS = ROOT / "out" / "docs" / "_tests" / "deck_chats"
SOLAR = ROOT / "out/docs/going-solar-in-lahore_141351/going-solar-in-lahore.pptx"


def slide_text(slide, notes=False):
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            out += [c.text_frame.text for row in sh.table.rows for c in row.cells]
    if notes and slide.has_notes_slide:
        out.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(out)


class Y:
    def __init__(self, c, v0, v1, reply, turn):
        self.c, self.v0, self.v1, self.reply, self.turn = c, v0, v1, reply, turn
        self.prs = Presentation(str(c.path()))
        self.m = PO.deck_map(self.prs)
        self.m0 = c.map(v0)
        self.prs0 = Presentation(str(c.path(v0)))

    def new(self):
        return self.v1 != self.v0

    def says(self, *words):
        return all(re.search(w, self.reply, re.I) for w in words)

    def titles(self, m=None):
        return [s["title"] for s in (m or self.m)["slides"]]

    def s(self, n, m=None):
        return (m or self.m)["slides"][n - 1]

    def count(self):
        return self.m["count"]

    def words(self, n, m=None):
        return sum(len(t["text"].split()) for t in self.s(n, m)["texts"])

    def text(self, n=None, notes=False, prs=None):
        prs = prs or self.prs
        ss = [prs.slides[n - 1]] if n else list(prs.slides)
        return "\n".join(slide_text(s, notes) for s in ss)

    def runs(self, n, part="all", prs=None):
        return list(PO._part_runs((prs or self.prs).slides[n - 1], part))

    def colours(self, part="titles", ns=None):
        ns = ns or range(1, self.count() + 1)
        return {PO._rgb_of(r.font.color) for n in ns for r in self.runs(n, part)}

    def fonts(self, part="titles"):
        return {r.font.name for n in range(1, self.count() + 1) for r in self.runs(n, part)}

    def sizes(self, n, part, prs=None):
        return [r.font.size.pt if r.font.size else 18.0 for r in self.runs(n, part, prs)]

    def bg(self, n):
        return PO.slide_bg(self.prs.slides[n - 1])

    def bullets(self, n, prs=None):
        sh = PO._body_shape((prs or self.prs).slides[n - 1])
        return len([p for p in sh.text_frame.paragraphs if "".join(r.text for r in p.runs).strip()]) if sh is not None else 0

    def arabic(self, n):
        t = self.text(n)
        letters = sum(1 for ch in t if ch.isalpha())
        return sum(1 for ch in t if "؀" <= ch <= "ۿ") / max(1, letters)

    def notes(self, n, m=None):
        return self.s(n, m)["notes"]

    def readable(self):
        return not PO.unreadable(self.prs, range(1, self.count() + 1))

    def fits(self):
        return not (self.c.version.get("overflow") or [])


def client_deck(path):
    """A deck as a client types it into PowerPoint's default template: placeholders, theme colours, a styled table."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Q3 Sales Review"
    s.placeholders[1].text = "Northwind Traders, Regional Team"

    def bullets(title, items, notes=None):
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = title
        tf = sl.placeholders[1].text_frame
        tf.text = items[0]
        for it in items[1:]:
            tf.add_paragraph().text = it
        if notes:
            sl.notes_slide.notes_text_frame.text = notes
        return sl

    bullets("Agenda", ["Results for Q3", "Regional highlights", "Pipeline for Q4", "Next steps"], "Keep the intro short.")
    bullets(
        "Q3 results", ["Revenue up 12% on Q2", "Gross margin steady at 41%", "Two new key accounts in Lahore"], "Mention the Lahore accounts by name."
    )
    bullets(
        "Regional highlights",
        ["North: strongest growth at 18%", "South: flat, pricing pressure", "Central: new distributor signed", "West: delayed by stock-outs"],
    )
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Sales by region"
    rows = [("Region", "Q2", "Q3"), ("North", "120", "142"), ("South", "98", "97"), ("Central", "76", "88"), ("West", "64", "59")]
    t = s.shapes.add_table(len(rows), 3, Inches(1.0), Inches(1.8), Inches(8.0), Inches(3.0)).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    bullets("Next steps", ["Hire two field reps for the North", "Fix stock-outs in the West before Q4", "Review pricing in the South"])
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Questions?"
    prs.save(str(path))
    return path


DARK_BLUE, NAVY, GOLD, WHITE = colour("dark blue"), colour("navy"), colour("gold"), "FFFFFF"


def _grandparent(x):
    vs = x.c.state["versions"]
    return vs[vs[x.v0]["parent"]]["parent"]


CONVERSATIONS = {
    "solar": (
        SOLAR,
        [
            (
                "how many slides are there and which ones have charts?",
                "question",
                lambda x: x.says(r"\b10 slides", r"charts: \[5\]") and not x.new(),
                False,
            ),
            (
                "delete slide 4, move slide 9 to the start and make the titles dark blue",
                "change",
                lambda x: (
                    x.count() == 9
                    and x.titles()[0].startswith("On-grid")
                    and not any("budget" in t for t in x.titles())
                    and x.colours("titles") == {DARK_BLUE}
                ),
                False,
            ),
            ("swap slides 2 and 3", "change", lambda x: x.titles()[1] == "Agenda" and x.titles()[2].startswith("Going Solar"), False),
            ("actually, undo that", "undo", lambda x: x.titles()[1].startswith("Going Solar"), False),
            (
                "move the title slide back to the start",
                "change",
                lambda x: x.titles()[0].startswith("Going Solar") and x.titles()[1].startswith("On-grid"),
                False,
            ),
            (
                "add a slide about financing options for solar after slide 7",
                "change",
                lambda x: (
                    x.count() == 10 and x.titles()[7] not in x.titles(x.m0) and re.search(r"financ|loan|instal|bank|credit|pay", x.text(8), re.I)
                ),
                True,
            ),
            (
                "make slide 4 shorter and add speaker notes to it",
                "change",
                lambda x: x.words(4) < x.words(4, x.m0) and x.notes(4) and x.notes(4) != x.notes(4, x.m0),
                True,
            ),
            ("translate the last slide into urdu", "change", lambda x: x.arabic(10) > 0.5 and x.arabic(9) < 0.1, True),
            (
                "make the background of the closing slide navy",
                "change",
                lambda x: x.bg(10) == NAVY and x.bg(9) != NAVY and not PO.unreadable(x.prs, [10]),
                True,
            ),
            (
                "we need one more slide comparing a 3 kW and a 5 kW system, right before the thank you slide",
                "change",
                lambda x: x.count() == 11 and re.search(r"3\s?kW", x.text(10)) and x.arabic(11) > 0.5,
                True,
            ),
            (
                "make all the titles Georgia and a bit bigger",
                "change",
                lambda x: x.fonts("titles") == {"Georgia"} and x.sizes(2, "titles")[0] > x.sizes(2, "titles", x.prs0)[0],
                True,
            ),
            ("fix any text that doesn't fit", None, lambda x: (x.new() and x.fits()) or (not x.new() and x.says("fits")), True),
            ("what's on slide 8?", "question", lambda x: x.says(re.escape(x.titles()[7][:12])) and not x.new(), True),
            ("compare v0 with the current version", "compare", lambda x: x.says(r"slides 10 -> 11"), True),
            ("go back to v5", "goto", lambda x: x.c.state["cur"] == 5, True),
            ("export", "export", lambda x: x.says(r"\.pptx", r"\.pdf"), False),
        ],
    ),
    "client": (
        "client",
        [
            ("what are the slide titles?", "question", lambda x: x.says("Agenda", r"Questions\?") and not x.new(), False),
            ("delete slide 40", "ask", lambda x: x.says(r"7 slides") and not x.new(), False),
            (
                "add a bullet to slide 3: 'Cash collection improved to 54 days'",
                "change",
                lambda x: x.bullets(3) == 4 and "54 days" in x.text(3),
                False,
            ),
            ("remove the second bullet from slide 4", "change", lambda x: x.bullets(4) == 3 and "South" not in x.text(4), False),
            ("add a bullet to slide 5: 'West recovers in Q4'", "change", lambda x: not x.new() and x.says("table", "bullets are"), False),
            ("replace Q3 with Q4 everywhere", "change", lambda x: "Q3" not in x.text(notes=True) and "Q4 Sales Review" in x.text(1), False),
            (
                "make the background navy on every slide and the titles gold",
                "change",
                lambda x: (
                    all(x.bg(n) == NAVY for n in range(1, 8)) and x.colours("titles") == {GOLD} and x.colours("body", [3]) == {WHITE} and x.readable()
                ),
                False,
            ),
            (
                "duplicate slide 3 and make its title white",
                "change",
                lambda x: (
                    x.count() == 8 and x.titles()[3] == x.titles()[2] and x.colours("titles", [4]) == {WHITE} and x.colours("titles", [3]) == {GOLD}
                ),
                False,
            ),
            ("make the body text on slide 7 dark blue", "change", lambda x: x.says("hard to read") and x.colours("body", [7]) == {DARK_BLUE}, False),
            ("undo", "undo", lambda x: x.colours("body", [7]) == {WHITE}, False),
            ("rename slide 6's title to 'Revenue by region (PKR m)'", "change", lambda x: x.titles()[5] == "Revenue by region (PKR m)", False),
            (
                "make the table text smaller",
                "change",
                lambda x: max(x.sizes(6, "table")) < 18 and x.sizes(6, "titles") == x.sizes(6, "titles", x.prs0),
                False,
            ),
            (
                "add speaker notes to slides 2-4: 'Pause here for questions'",
                "change",
                lambda x: all(x.notes(n) == "Pause here for questions" for n in (2, 3, 4)) and not x.notes(5),
                False,
            ),
            (
                "write speaker notes for the slides that have none",
                "change",
                lambda x: all(x.notes(n) for n in range(1, 9)) and x.notes(2) == "Pause here for questions",
                True,
            ),
            ("translate slide 3 into urdu", "change", lambda x: x.arabic(3) > 0.5 and x.arabic(4) < 0.1, True),
            ("make slide 5 punchier", "change", lambda x: x.text(5) != x.text(5, prs=x.prs0), True),
            ("undo the last 2 changes", "undo", lambda x: x.c.state["cur"] == _grandparent(x), True),
            ("redo", "redo", lambda x: x.arabic(3) > 0.5, True),
            ("how many slides are there and which have speaker notes?", "question", lambda x: x.says(r"8 slides", r"notes") and not x.new(), False),
            ("export", "export", lambda x: x.says(r"\.pptx", r"\.pdf"), False),
        ],
    ),
}


def run(names, planner):
    shutil.rmtree(CHATS, ignore_errors=True)
    CHATS.mkdir(parents=True, exist_ok=True)
    rows = []
    t_all = time.perf_counter()
    for name in names:
        src, turns = CONVERSATIONS[name]
        if src == "client":
            src = client_deck(CHATS / "q3_sales_review.pptx")
        c = DeckChat.start(src, chats_dir=CHATS, planner=planner)
        print(f"\n=== {name} ({Path(src).name})", flush=True)
        try:
            for msg, want, check, needs_model in turns:
                if needs_model and planner is None:
                    print(f"skip (needs the model)  {msg}")
                    continue
                v0 = c.state["cur"]
                t = time.perf_counter()
                err = None
                try:
                    reply = c.say(msg)
                except Exception as e:  # noqa: BLE001
                    import traceback

                    traceback.print_exc()
                    reply, err = "", f"{type(e).__name__}: {e}"
                    c.last_turn = {"intents": ["crash"], "ops": []}
                secs = time.perf_counter() - t
                intent_ok = want is None or want in (c.last_turn.get("intents") or [])
                try:
                    x = Y(c, v0, c.state["cur"], reply, c.last_turn)
                    check_ok = bool(check(x))
                except Exception as e:  # noqa: BLE001
                    check_ok, err = False, err or f"check {type(e).__name__}: {e}"
                ok = intent_ok and check_ok and not err
                rows.append(
                    {
                        "chat": name,
                        "msg": msg,
                        "ok": ok,
                        "intents": c.last_turn.get("intents"),
                        "seconds": round(secs, 2),
                        "llm": bool(c.last_turn.get("llm")),
                        "reply": reply[:600],
                        "error": err,
                    }
                )
                print(
                    f"{'OK ' if ok else 'BAD'} [{','.join(c.last_turn.get('intents') or [])}{'+llm' if c.last_turn.get('llm') else ''}] "
                    f"v{v0}->v{c.state['cur']} {secs:5.1f}s  {msg}\n      {reply[:300]}",
                    flush=True,
                )
                if not ok:
                    print(f"      intent_ok={intent_ok} check_ok={check_ok} {err or ''}")
        finally:
            c.close()
    n, ok = len(rows), sum(r["ok"] for r in rows)
    usd = planner.cost()[1] if planner is not None else 0.0
    print(
        f"\n{ok}/{n} turns OK; {sum(r['llm'] for r in rows)} used the model for reading the request; "
        f"{sum(r['seconds'] for r in rows) / max(1, n):.1f} s per turn; {time.perf_counter() - t_all:.0f} s in all; ${usd:.4f}"
    )
    (CHATS / "report.json").write_text(json.dumps({"ok": ok, "turns": n, "rows": rows}, ensure_ascii=False, indent=1), encoding="utf-8")
    return ok, n


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    planner = None
    if "--offline" not in sys.argv:
        from ai_pc.llm.planner import ChatPlanner

        planner = ChatPlanner()
    ok, n = run(args or list(CONVERSATIONS), planner)
    sys.exit(0 if ok == n else 1)
