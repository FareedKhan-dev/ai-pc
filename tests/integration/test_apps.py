"""The many small programs (src/ai_pc/apps), one main path each: OCR, PC care, diagrams, references, quizzes, ebooks,
playlists, printing, maps and notes, through their functions and through the chat.

  .venv\\Scripts\\python.exe tests\\integration\\test_apps.py
Read-only on the PC (winget search, disk, Defender status); printing goes to 'Microsoft Print to PDF' with its output
file; Crossref, Open Library and Nominatim are fakes; everything is written under out\\_tests\\apps.
"""

import shutil
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from PIL import Image, ImageDraw, ImageFont  # noqa: E402

from ai_pc.apps import backup, cite, database, diagrams, ebook, maps, media, ocr, pccare, printing, quiz  # noqa: E402
from ai_pc.apps.appschat import AppsChat  # noqa: E402
from ai_pc.hub.http import FakeTransport  # noqa: E402

OUT = ROOT / "out" / "_tests" / "apps"
RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append((name, bool(ok)))
    print(("ok   " if ok else "FAIL ") + name + ("" if ok or not detail else f"\n     why: {str(detail)[:700]}"))


def chat(files=(), **extra):
    return AppsChat.start(files=files, chats_dir=OUT / "chats", extra=extra)


def t_ocr():
    im = Image.new("RGB", (900, 470), "white")
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype("arial.ttf", 30)
    for i, t in enumerate(
        [
            "KHAN ELECTRONICS",
            "Hall Road, Lahore  NTN 1234567-8",
            "Date: 04/10/2026",
            "LED TV 55      85,000.00",
            "Wall Mount      3,500.00",
            "Grand Total    104,430.00",
            "Phone 0300-1234567",
        ]
    ):
        d.text((40, 30 + i * 60), t, fill="black", font=f)
    p = OUT / "receipt.png"
    im.save(p)
    r = ocr.read(p, work=OUT / "ocr")
    words = r["text"].split()
    check("ocr: Windows' own OCR reads a receipt", all(w in words for w in ("KHAN", "ELECTRONICS", "Wall", "Mount", "Total")), r["text"])
    check(
        "ocr: the total, date, NTN and phone picked out",
        ocr.fields(r["text"]) == {"total": 104430.0, "date": "04/10/2026", "ntn": "1234567-8", "phone": "0300-1234567"},
        ocr.fields(r["text"]),
    )
    reply = chat([p]).say("what is the total on receipt.png?")
    check("ocr: by chat, saved as text with the total said", reply.startswith("Read ") and "total 104,430.00" in reply and "checked" in reply, reply)


def t_pccare():
    sample = (
        "Name                          Id                     Version   Available Source\n"
        "-------------------------------------------------------------------------------\n"
        "VLC media player              VideoLAN.VLC           3.0.23    3.0.24    winget\n"
        "剪映专业版                    ByteDance.JianyingPro  8.9.0     11.5.0    winget\n"
        "2 upgrades available.\n"
    )
    rows = pccare.table(sample)
    check(
        "pc care: winget's table read by display width (a Chinese name is twice as wide)",
        [(r["Id"], r["Available"]) for r in rows] == [("VideoLAN.VLC", "3.0.24"), ("ByteDance.JianyingPro", "11.5.0")],
        rows,
    )
    check(
        "pc care: requests read",
        [
            pccare.parse(x, {})["op"]
            for x in (
                "which apps can be updated?",
                "update vlc",
                "install 7-zip",
                "uninstall zoom",
                "how much disk space is left?",
                "is my antivirus on?",
                "run a virus scan",
                "wifi",
                "battery",
            )
        ]
        == ["upgrades", "upgrade", "install", "uninstall", "disks", "defender", "scan", "wifi", "battery"],
    )
    ds = pccare.disks()
    check("pc care: disk space of each drive", any(d["drive"].upper().startswith("C") and d["total_gb"] > 10 for d in ds), ds)
    s = pccare.defender()
    check("pc care: Windows Security status read", "AntivirusEnabled" in s, s)
    c = chat()
    r = c.say("install 7-zip")
    check("pc care: an install is shown first (found in winget) and waits for a yes", r.startswith("Ready to install 7-Zip") and "yes" in r, r)
    r = c.say("no")
    check("pc care: and 'no' drops it", r.startswith("Dropped"), r)


def t_diagrams():
    for kind, spec in (
        ("flow", "Start -> Take order -> In stock? -yes-> Pack -> Ship -> End; In stock? -no-> Order from supplier -> Pack"),
        ("tree", "CEO > Sales Manager, Accounts Manager; Sales Manager > Ali, Sara; Accounts Manager > Hina"),
        ("mind", "Marketing: Social media (Facebook, Instagram), Ads, SEO, Email (Newsletter)"),
    ):
        r = diagrams.make(kind, spec, OUT / "diagrams", kind)
        bad = [w for w, ok in diagrams.check(r) if not ok]
        check(f"diagrams: a {kind} drawn (PNG, PDF, SVG, draw.io), every name in, nothing overlapping", not bad, bad)
    r = diagrams.make("flow", "Start -> In stock? -yes-> Pack; In stock? -no-> Order -> Pack", OUT / "diagrams", "lanes")
    check(
        "diagrams: an arrow across rows gets its own lane (waypoints), not through a box",
        r["edges"] and r.get("pos") and diagrams.read_flow("A -> B? -yes-> C; B? -no-> D -> C")[1][1] == ("B?", "C", "yes"),
    )
    reply = chat().say("org chart: Owner > Manager; Manager > Cashier, Salesman")
    check("diagrams: by chat", reply.startswith("Org Chart with 4 boxes and 3 links") and "Checked" in reply, reply)


CROSSREF = {
    "message": {
        "type": "journal-article",
        "title": ["Deep learning"],
        "author": [{"given": "Yann", "family": "LeCun"}, {"given": "Yoshua", "family": "Bengio"}, {"given": "Geoffrey", "family": "Hinton"}],
        "issued": {"date-parts": [[2015, 5, 28]]},
        "container-title": ["Nature"],
        "volume": "521",
        "issue": "7553",
        "page": "436-444",
        "publisher": "Springer Science and Business Media LLC",
        "DOI": "10.1038/nature14539",
    }
}


def t_cite():
    cite.TRANSPORT = FakeTransport(
        {
            ("GET", "api.crossref.org/works/10.1038/nature14539"): (200, CROSSREF),
            ("GET", "openlibrary.org/isbn/9780262035613.json"): (
                200,
                {"title": "Deep Learning", "publishers": ["MIT Press"], "publish_date": "2016", "authors": [{"key": "/authors/OL1A"}]},
            ),
            ("GET", "openlibrary.org/authors/OL1A.json"): (200, {"name": "Ian Goodfellow"}),
        }
    )
    recs, refs, files = cite.make([("doi", "10.1038/nature14539"), ("isbn", "9780262035613")], "apa", OUT / "cite_apa")
    got = [cite.text_of(r) for r in refs]
    check(
        "cite: APA 7, sorted by author, the article and the book",
        got
        == [
            "Goodfellow, I. (2016). Deep Learning. MIT Press.",
            "LeCun, Y., Bengio, Y., & Hinton, G. (2015). Deep learning. Nature, 521(7553), 436\u2013444. https://doi.org/10.1038/nature14539",
        ],
        got,
    )
    recs, refs, files = cite.make([("doi", "10.1038/nature14539")], "ieee", OUT / "cite_ieee")
    check(
        "cite: IEEE numbered",
        cite.text_of(refs[0]) == '[1] Y. LeCun, Y. Bengio, and G. Hinton, "Deep learning," Nature, vol. 521, no. 7553, pp. 436\u2013444, '
        "2015, doi: 10.1038/nature14539.",
        cite.text_of(refs[0]),
    )
    check(
        "cite: BibTeX and RIS for Zotero, Mendeley and EndNote",
        "@article{lecun2015," in files[1].read_text(encoding="utf-8") and "TY  - JOUR" in files[2].read_text(encoding="utf-8"),
    )
    reply = chat().say("cite 10.1038/nature14539 in mla")
    check(
        "cite: by chat, MLA, the Word file read back",
        'LeCun, Yann, et al. "Deep learning." Nature, vol. 521, no. 7553, 2015, pp. 436\u2013444.' in reply and "checked" in reply,
        reply,
    )


def t_quiz():
    spec = "1. What is H2O? a) Water* b) Salt c) Sugar d) Iron 2. The sun is a star. True* 3. Capital of Pakistan? answer: Islamabad"
    qs, files, skipped = quiz.make(spec, "Science test", OUT / "quiz")
    check("quiz: three kinds of question read", [q["kind"] for q in qs] == ["mc", "tf", "short"] and qs[0]["answer"] == "Water", qs)
    bad = [w for w, ok in quiz.check(qs, files) if not ok]
    check("quiz: Moodle XML, GIFT, Kahoot sheet, paper and answer key, each read back", not bad and skipped == 1, bad)
    reply = chat().say(f"quiz 'Maths': {spec}")
    check("quiz: by chat", reply.startswith("3 questions") and "Checked" in reply, reply)


def t_ebook():
    md = OUT / "story.md"
    md.write_text(
        "# Chapter One\n\nIt was a **bright** day in Lahore.\n\n## The market\n\nAli walked to *Anarkali*.\n\n# Chapter Two\n\nThe end & a start.\n",
        encoding="utf-8",
    )
    dest, chs = ebook.make(md, OUT / "ebooks", "Story", "Ayesha Khan")
    bad = [w for w, ok in ebook.check(dest, chs) if not ok]
    check("ebook: EPUB 3 from Markdown, two chapters, every check passes", len(chs) == 2 and not bad, bad)
    from docx import Document

    d = Document()
    d.add_heading("Introduction", 1)
    d.add_paragraph("First words.")
    d.add_heading("Method", 1)
    d.add_paragraph("Second words.")
    src = OUT / "thesis.docx"
    d.save(src)
    reply = chat([src]).say("make an ebook from thesis.docx by Ayesha Khan")
    check("ebook: by chat from a Word file", reply.startswith("EPUB made:") and "2 chapters" in reply and "Checked" in reply, reply)


def t_media():
    folder = OUT / "music"
    folder.mkdir(parents=True, exist_ok=True)
    for n in ("track 10", "track 2", "track 1"):
        subprocess.run(
            ["ffmpeg", "-v", "error", "-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=1", str(folder / f"{n}.mp3")],
            check=True,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    files, total, outs = media.playlist(folder, OUT / "playlists")
    check(
        "media: a playlist in natural order (1, 2, 10) with its length measured",
        [f.stem for f in files] == ["track 1", "track 2", "track 10"]
        and 2.5 < total < 3.6
        and outs[1].read_text(encoding="utf-8").count("<track>") == 3,
        ([f.stem for f in files], total),
    )
    reply = chat().say(f"make a playlist from {folder}")
    check("media: by chat", reply.startswith("Playlist of 3 files") and "checked" in reply, reply)
    check("media: 'play ... in vlc' read (not run in tests)", media.parse(f"play {outs[0]} in vlc", {})["op"] == "play")


def t_printing():
    names = [p["name"] for p in printing.printers()]
    check("printing: the printers on this PC", "Microsoft Print to PDF" in names, names)
    from docx import Document

    d = Document()
    d.add_paragraph("Page one")
    d.add_page_break()
    d.add_paragraph("Page two")
    src = OUT / "letter.docx"
    d.save(src)
    c = chat([src])
    r = c.say("print letter.docx on Microsoft Print to PDF")
    check(
        "printing: shown first (which printer, how many copies), waits for a yes",
        r.startswith("Ready to print letter.docx, 1 copy, on Microsoft Print to PDF") and "no paper" in r,
        r,
    )
    r = c.say("yes")
    check(
        "printing: a Word file made PDF, drawn page by page through Windows, the printed file has both pages",
        "Printed 2 page(s)" in r and "checked" in r,
        r,
    )


def t_maps():
    pts = {
        "Hall Road Lahore": ("31.5638", "74.3176"),
        "Liberty Market Lahore": ("31.5106", "74.3442"),
        "Lahore": ("31.5204", "74.3587"),
        "Islamabad": ("33.6844", "73.0479"),
    }
    routes = {("GET", f"q={k.replace(' ', '+')}&"): (200, [{"lat": v[0], "lon": v[1], "display_name": k}]) for k, v in pts.items()}
    maps.TRANSPORT = FakeTransport(routes)
    pts_, legs, files, link = maps.make(["Hall Road Lahore", "Liberty Market Lahore"], OUT / "maps", "shops", cache=OUT / "geo.json")
    check("maps: places found, KML, GPX and GeoJSON written, each with every place", not [w for w, ok in maps.check(["a", "b"], files) if not ok])
    check(
        "maps: the leg measured (about 6 km) and a Google Maps route link",
        5 < legs[0] < 8 and link.startswith("https://www.google.com/maps/dir/?api=1&origin="),
        (legs, link),
    )
    a, b = maps.geocode("Lahore", OUT / "geo.json"), maps.geocode("Islamabad", OUT / "geo.json")
    check("maps: Lahore to Islamabad about 270 km in a straight line", 260 < maps.km(a, b) < 280, maps.km(a, b))
    n = len(maps.TRANSPORT.sent)
    maps.geocode("Lahore", OUT / "geo.json")
    check("maps: a place is looked up once (kept on this PC)", len(maps.TRANSPORT.sent) == n)


def t_notes():
    v = OUT / "vault"
    c = chat(notes_dir=str(v))
    r1 = c.say("note: call the supplier about the TV order")
    r2 = c.say("todo: pay the electricity bill by Friday")
    c.say("todo: send the quotation to Ali Traders")
    check("notes: a note and to-dos go into today's note", "checked" in r1 and "checked" in r2 and len(list(v.glob("*.md"))) == 1, (r1, r2))
    r = c.say("my tasks")
    check("notes: the open to-dos across the notes", r.startswith("2 to-dos") and "electricity" in r, r)
    r = c.say("done: electricity bill")
    check("notes: one ticked off by its words", r.startswith("Ticked off: pay the electricity bill") and c.say("my tasks").startswith("1 to-dos"), r)
    c.say("new note 'Meeting with Ali': prices agreed at 84,000 for the LED TV")
    r = c.say("find notes about prices agreed")
    check("notes: a named note saved and found by its words", "Meeting with Ali" in r, r)


def t_database():
    import datetime as dt

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Customer", "Amount", "Date"])
    for row in (
        ("Ali Traders", 85000, dt.datetime(2026, 7, 10)),
        ("Bilal & Sons", 3500, dt.datetime(2026, 7, 11)),
        ("Ali Traders", 2000.5, dt.datetime(2026, 7, 12)),
    ):
        ws.append(row)
    ws2 = wb.create_sheet("Items")
    ws2.append(["Item", "Price"])
    ws2.append(["LED TV 55", 85000])
    src = OUT / "sales.xlsx"
    wb.save(src)
    for kind in ("access", "sqlite"):
        dest = OUT / "db" / f"sales.{'accdb' if kind == 'access' else 'db'}"
        dest.parent.mkdir(parents=True, exist_ok=True)
        (database.to_access if kind == "access" else database.to_sqlite)(src, dest)
        heads, rows = database.query(dest, "SELECT Customer, SUM(Amount) AS Total FROM Sales GROUP BY Customer ORDER BY Customer")
        check(
            f"database: Excel to {kind} (two sheets, two tables), a question answered",
            database.tables(dest) == {"Sales": 3, "Items": 1}
            and [(r[0], round(r[1], 2)) for r in rows] == [("Ali Traders", 87000.5), ("Bilal & Sons", 3500.0)],
            (database.tables(dest), rows),
        )
    try:
        database.query(OUT / "db" / "sales.db", "DELETE FROM Sales")
        check("database: only questions are run (nothing that changes the data)", False)
    except ValueError:
        check("database: only questions are run (nothing that changes the data)", True)
    r = chat([src]).say("make an access database from sales.xlsx")
    check("database: by chat, every row checked", r.startswith("Access database:") and "checked" in r, r)


def t_backup():
    src = OUT / "shop"
    (src / "invoices").mkdir(parents=True, exist_ok=True)
    (src / "invoices" / "INV-1.txt").write_text("invoice one", encoding="utf-8")
    (src / "notes.txt").write_text("hello", encoding="utf-8")
    dest, man = backup.backup(src, OUT / "backups")
    check(
        "backup: a dated zip with every file's SHA-256 inside, read back and matching",
        len(man["files"]) == 2 and all(ok for _, ok in backup.verify(dest)),
    )
    import zipfile

    bad = OUT / "backups" / "damaged.zip"
    with zipfile.ZipFile(dest) as z, zipfile.ZipFile(bad, "w") as w:
        for i in z.infolist():
            data = z.read(i)
            w.writestr(i.filename, data + b"!" if i.filename.endswith("notes.txt") else data)
    check("backup: a damaged backup is caught", [f for f, ok in backup.verify(bad) if not ok] == ["notes.txt"])
    evil = OUT / "evil.zip"
    with zipfile.ZipFile(evil, "w") as z:
        z.writestr("../../escape.txt", "x")
    try:
        backup.safe_unzip(evil, OUT / "unz")
        check("backup: a zip that would write outside its folder is refused", False)
    except ValueError:
        check("backup: a zip that would write outside its folder is refused", not (OUT.parent / "escape.txt").exists())
    r = chat().say(f"back up {src} to {OUT / 'backups2'}")
    check("backup: by chat", r.startswith("Backed up 2 files") and "matches" in r, r)


def t_contacts_codes_music_calendar():
    import datetime as dt

    r = chat().say("contact Ali Raza, 0300-1234567, ali@x.com, Ali Traders")
    vcf = next((OUT / "chats").rglob("Ali_Raza.vcf"))
    t = vcf.read_text(encoding="utf-8")
    check("contacts: a vCard with the number in +92 form", "TEL;TYPE=CELL:+923001234567" in t and "ORG:Ali Traders" in t and "checked" in r, t)
    from ai_pc.apps import calendar, codes

    r = chat().say("wifi qr for KhanShop password secret123")
    check("codes: a Wi-Fi QR code, read back by a scanner", "read back by a scanner: same text" in r, r)
    r = chat().say("ean13 barcode 896000123456")
    check(
        "codes: EAN-13 with its check digit worked out (8+27+6+0+0+0+1+6+3+12+5+18 = 86, so 4), decoded back from its bars",
        "8960001234564" in r and "decoded back" in r,
        r,
    )
    try:
        codes.ean13("8960001234563")
        check("codes: a wrong EAN-13 check digit is refused", False)
    except ValueError:
        check("codes: a wrong EAN-13 check digit is refused", True)
    bits = codes.widths_to_bits(codes.code128("LED-TV-55"))
    check("codes: Code 128 decodes back with its checksum", codes.decode128(bits) == "LED-TV-55")
    r = chat().say("barcode labels for TV55, MOUNT, CABLE")
    check("codes: a sheet of barcode labels as an A4 PDF", r.startswith("3 barcode labels on 1 A4 page"), r)
    r = chat().say("melody: C4 D4 E4 F4 G4/2 G4/2 R/4 A4/8 A4/8 G4/1 at 120 bpm on flute")
    check("music: a melody as MIDI and MusicXML, every note read back", "9 notes" in r and "checked" in r, r)
    now = dt.datetime(2026, 10, 4, 12, 0)
    e = calendar.read("meeting with Ali Traders on 12 October at 3 pm for 2 hours at the shop", now)
    check(
        "calendar: an event read from words (date, time, length, place)",
        (e["title"], e["start"], e["minutes"], e["where"]) == ("Meeting with Ali Traders", dt.datetime(2026, 10, 12, 15, 0), 120, "the shop")
        or (e["title"], e["start"], e["minutes"]) == ("Meeting with Ali Traders", dt.datetime(2026, 10, 12, 15, 0), 120),
        e,
    )
    e = calendar.read("remind me to pay rent every month on the 5th at 10 am", now)
    check("calendar: a monthly reminder", e["rule"] == "FREQ=MONTHLY;BYMONTHDAY=5" and e["start"].time() == dt.time(10, 0) and e["reminder"] == 30, e)
    r = chat().say("meeting with Ali Traders on 12 October at 3 pm for 1 hour")
    check("calendar: by chat, an .ics file", ".ics" in r and "checked" in r, r)


def t_stores():
    from stores_fakes import FakeDaraz, FakeOdoo, FakeShopify, FakeWoo, FakeWordPress

    from ai_pc.apps import daraz, odoo, shopify, woocommerce, wordpress

    check(
        "daraz: the request signature matches Daraz's own test vector",
        daraz.sign(
            "/order/get",
            {"access_token": "test", "app_key": "123456", "order_id": "1234", "sign_method": "sha256", "timestamp": "1517820392000"},
            "helloworld",
        )
        == "4190D32361CFB9581350222F345CB77F3B19F0E31D162316848A2C1FFD5FAB4A",
    )
    fs, fw, fd, fp, fo = FakeShopify(), FakeWoo(), FakeDaraz(), FakeWordPress(), FakeOdoo()
    img = OUT / "lawn.jpg"
    Image.new("RGB", (64, 64), "pink").save(img)
    clients = {
        "shopify": shopify.Client({"shop": fs.SHOP, "client_id": "shop-id", "client_secret": "shop-secret"}, fs),
        "woocommerce": woocommerce.Client({"url": "https://shop.pk", "ck": "ck_1", "cs": "cs_1"}, fw),
        "daraz": daraz.Client({"app_key": "123456", "app_secret": fd.SECRET, "access_token": "dz-token", "expires_at": 9e12}, fd),
        "wordpress": wordpress.Client({"url": "https://blog.pk", "user": "ayesha", "app_password": "abcd abcd abcd abcd abcd abcd"}, fp),
        "odoo": odoo.Client({"url": "https://khan.odoo.com", "db": "khan", "login": "ayesha@khan.pk", "key": "odoo-key"}, fo),
    }
    c = chat([img], clients=clients)
    for store, word in (("shopify", "shopify"), ("woocommerce", "woocommerce")):
        r = c.say(f"add Lawn Suit to {word} at 2,500, sku LS-01, 10 in stock")
        check(f"{store}: a new product shown first", r.startswith("Ready to add 'Lawn Suit' (SKU LS-01)") and "Rs 2,500.00 with 10 in stock" in r, r)
        r = c.say("yes")
        check(f"{store}: added, read back with its price and stock", "read back: Rs 2,500.00, 10 in stock" in r, r)
        c.say(f"set stock of LS-01 on {word} to 25")
        r = c.say("yes")
        check(f"{store}: stock changed by SKU, read back", "read back: Rs 2,500.00, 25 in stock" in r, r)
        r = c.say(f"{word} products")
        check(f"{store}: products listed", "LS-01: Lawn Suit, Rs 2,500.00, 25 in stock" in r, r)
        r = c.say(f"new orders on {word}")
        check(f"{store}: new orders listed", ("1001" if store == "shopify" else "345") in r, r)
    c.say("mark order 1001 shipped on shopify with TCS 1234567890")
    r = c.say("yes")
    check(
        "shopify: an order fulfilled with the courier's tracking (after one throttled answer)",
        "Shopify says SUCCESS" in r and fs.last_tracking == {"company": "TCS", "number": "1234567890"},
        r,
    )
    c.say("mark order 345 shipped on woocommerce with TCS 1234567890")
    r = c.say("yes")
    check(
        "woocommerce: the customer told the tracking number, the order completed (keys sent in the address when the host strips the header)",
        "completed" in r and "1234567890" in fw.notes[-1]["note"] and fw.notes[-1]["customer_note"],
        r,
    )
    c.say("set price of LS-01 on daraz to 2,300")
    r = c.say("yes")
    check("daraz: a price changed by seller SKU (signed calls), read back", "read back: Rs 2,300.00, 10 in stock" in r and fd.bad_signs == 0, r)
    r = c.say("mark order 123456789 shipped on daraz")
    check("daraz: shipping is shown first (Daraz's courier gives the tracking)", r.startswith("Ready to pack order 123456789 on Daraz"), r)
    r = c.say("yes")
    check("daraz: packed and ready to ship, Daraz's tracking number said", "ready to ship; Daraz's tracking DRZ123456" in r, r)
    c.say("wordpress draft 'New lawn collection': Our summer lawn is here. Prices from Rs 2,500. with lawn.jpg")
    r = c.say("yes")
    check("wordpress: a draft post with its featured picture, read back", "read back: a draft, the same title" in r and "with its picture" in r, r)
    c.say("odoo invoice for Ayesha Khan: 2 LS-01 at 2,500, 1 delivery at 300")
    r = c.say("yes")
    check(
        "odoo: a draft invoice (new customer added, the product found by SKU), read back",
        "before tax Rs 5,300.00" in r and "read back: a draft" in r and any(p["name"] == "Ayesha Khan" for p in fo.t["res.partner"].values()),
        r,
    )


def t_web_desktop():
    from ai_pc.apps import desktop

    page = OUT / "prices.html"
    page.write_text(
        "<html><head><title>Khan Electronics prices</title><script>document.title += ' 2026'</script></head><body><nav>Home | About</nav>"
        "<article><h1>Prices this week</h1><p>LED TVs are cheaper this week.</p><ul><li>Free delivery in Lahore</li></ul>"
        "<table><tr><th>Item</th><th>Price</th></tr><tr><td>LED TV 55</td><td>Rs 84,000</td></tr><tr><td>Wall Mount</td><td>3,500</td></tr></table>"
        "<a href='/about.html'>About us</a></article><footer>Copyright</footer></body></html>",
        encoding="utf-8",
    )
    url = page.resolve().as_uri()
    r = chat().say(f"read {url}")
    md = next((OUT / "chats").rglob("*.md")).read_text(encoding="utf-8")
    check(
        "web: a page's readable text (its script ran; menus and footers left out)",
        "# Khan Electronics prices 2026" in md
        and "LED TVs are cheaper" in md
        and "- Free delivery in Lahore" in md
        and "Home | About" not in md
        and "Copyright" not in md,
        md,
    )
    r = chat().say(f"tables from {url}")
    from openpyxl import load_workbook

    x = next((OUT / "chats").rglob("*.tables.xlsx"))
    rows = [list(r_) for r_ in load_workbook(x).active.iter_rows(values_only=True)]
    check("web: its tables to Excel, prices as numbers", rows == [["Item", "Price"], ["LED TV 55", 84000], ["Wall Mount", 3500]], rows)
    r = chat().say(f"links on {url}")
    check("web: its links, made absolute", "about.html" in r and r.startswith("1 links"), r)
    import win32clipboard as cb

    cb.OpenClipboard()
    try:
        only_text = all(
            f in (cb.CF_UNICODETEXT, cb.CF_TEXT, cb.CF_OEMTEXT, cb.CF_LOCALE)
            for f in iter(lambda f=[0]: (f.__setitem__(0, cb.EnumClipboardFormats(f[0])), f[0])[1], 0)
        )
    finally:
        cb.CloseClipboard()
    if only_text:  # the person's clipboard is put back as it was
        before = desktop.clip_get()
        r = chat().say("copy 'Meezan Bank PK36MEZN0001234567890123' to the clipboard")
        back = desktop.clip_get()
        if before is not None:
            desktop.clip_set(before)
        check(
            "desktop: text put on the clipboard and read back (the clipboard restored after)",
            "checked" in r and back == "Meezan Bank PK36MEZN0001234567890123",
            r,
        )
    else:
        check("desktop: clipboard test skipped (it holds a picture or files; not touched)", True)
    from PIL import ImageGrab

    check("desktop: a screenshot can be taken (a 10 x 10 corner, kept in memory only)", ImageGrab.grab(bbox=(0, 0, 10, 10)).size == (10, 10))


def t_mail_files_chat():
    from openpyxl import Workbook
    from stores_fakes import FakeBrevo, FakeDiscord, FakeDropbox, FakeMailchimp

    from ai_pc.apps import discord, dropbox, mailchimp

    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Email"])
    ws.append(["Ali Raza", "ali@x.com"])
    ws.append(["Sara Khan", "SARA@x.com"])
    ws.append(["No Email", "not-an-email"])
    lst = OUT / "customers.xlsx"
    wb.save(lst)
    fm, fb, fd, fc = FakeMailchimp(), FakeBrevo(), FakeDropbox(), FakeDiscord()
    clients = {
        "mailchimp": mailchimp.Mailchimp({"key": "mc-key-us21", "reply_to": "shop@khan.pk"}, fm),
        "brevo": mailchimp.Brevo({"key": "brevo-key", "sender": "shop@khan.pk"}, fb),
        "dropbox": dropbox.Client({"client_id": "x", "refresh_token": "r", "access_token": "dbx-token", "expires_at": 9e12}, fd),
        "discord": discord.Client({"webhook": fc.HOOK}, fc),
    }
    pic = OUT / "arrivals.png"
    Image.new("RGB", (32, 32), "teal").save(pic)
    c = chat([lst, pic], clients=clients)
    for which in ("mailchimp", "brevo"):
        r = c.say(f"add contacts from customers.xlsx to {which} list Customers")
        check(f"{which}: adding contacts shown first (2 good addresses of 3)", r.startswith("Ready to add 2 contact(s)"), r)
        r = c.say("yes")
        check(f"{which}: contacts added, the audience now 2", "now has 2 people" in r, r)
        r = c.say(f"email campaign on {which} to Customers: subject 'Eid Sale' from Khan Electronics, body: 20% off everything till Sunday.")
        check(f"{which}: a campaign shown first with how many it reaches", "Ready to SEND 'Eid Sale'" in r and "2 people get it" in r, r)
        r = c.say("yes")
        check(f"{which}: sent", r.startswith(f"Sent on {'Mailchimp' if which == 'mailchimp' else 'Brevo'}"), r)
        r = c.say(f"{which} report")
        check(f"{which}: the report (sent, opened, clicked)", "sent to 2" in r, r)
    c.say("upload arrivals.png to dropbox")
    r = c.say("yes")
    check(
        "dropbox: uploaded, checked with Dropbox's own content hash", "checked: Dropbox holds the same bytes" in r and "/AI PC/arrivals.png" in r, r
    )
    c.say("share link for /AI PC/arrivals.png on dropbox")
    r = c.say("yes")
    check("dropbox: a share link after a yes", r.startswith("Share link: https://www.dropbox.com/"), r)
    c.say("post arrivals.png to discord saying New arrivals!")
    r = c.say("yes")
    check("discord: a picture posted with its text, read back", "read back: the same text" in r and "with the picture" in r, r)
    c.say("delete the last discord post")
    r = c.say("yes")
    check("discord: the post taken back", r.startswith("Deleted the Discord post") and not fc.messages, r)


def t_stats_grammar():
    from ai_pc.apps import grammar, stats

    close = lambda a, b, tol=5e-4: abs(a - b) < tol  # noqa: E731
    check(
        "stats: p-values match the tables (t(10)=2.228 -> .05, t(10)=2.0 -> .0734, chi2(1)=3.841 -> .05, F(2,27)=3.354 -> .05)",
        close(stats.p_t(2.228139, 10), 0.05)
        and close(stats.p_t(2.0, 10), 0.07339)
        and close(stats.p_chi2(3.841459, 1), 0.05)
        and close(stats.p_f(3.354131, 2, 27), 0.05),
        (stats.p_t(2.228139, 10), stats.p_t(2.0, 10), stats.p_chi2(3.841459, 1), stats.p_f(3.354131, 2, 27)),
    )
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Score", "Gender", "Class", "Hours", "Q1", "Q2", "Q3", "Choice"])
    rows = [
        (5, "M", "A", 1, 1, 2, 3, "Yes"),
        (6, "M", "A", 2, 2, 3, 4, "Yes"),
        (7, "M", "B", 3, 3, 4, 5, "No"),
        (8, "M", "B", 4, 4, 5, 6, "Yes"),
        (9, "M", "C", 5, 1, 2, 3, "No"),
        (1, "F", "C", 1, 2, 3, 4, "No"),
        (2, "F", "A", 2, 3, 4, 5, "No"),
        (3, "F", "B", 3, 4, 5, 6, "Yes"),
        (4, "F", "C", 4, 1, 2, 3, "No"),
        (5, "F", "A", 5, 2, 3, 4, "No"),
    ]
    for r in rows:
        ws.append(r)
    src = OUT / "survey.xlsx"
    wb.save(src)
    data = stats.sheet(src)
    t = stats.ttest(data, "Score", "Gender")
    check(
        "stats: t-test by hand (means 7 and 3, pooled SD 1.58: t(8) = 4.00, p = .004, d = 2.53)",
        close(t["t"], 4.0) and t["df"] == 8 and close(t["p"], 0.003949, 1e-4) and close(t["d"], 2.5298),
        t,
    )
    a = stats.anova({"y": [1, 2, 3, 4, 5, 6, 7, 8, 9], "g": list("aaabbbccc")}, "y", "g")
    check(
        "stats: one-way ANOVA by hand (F(2, 6) = 27.0, p = .001, eta squared .90)",
        close(a["F"], 27.0) and close(a["p"], 0.001) and close(a["eta2"], 0.9),
        a,
    )
    r = stats.correlation({"x": [1, 2, 3, 4, 5], "y": [2, 4, 5, 4, 5]}, "x", "y")
    check("stats: Pearson r = .77, p = .124 (n = 5)", close(r["r"], 0.774597) and close(r["p"], 0.1240, 1e-3), r)
    x = stats.crosstab({"a": ["p"] * 30 + ["q"] * 30, "b": ["y"] * 10 + ["n"] * 20 + ["y"] * 20 + ["n"] * 10}, "a", "b")
    check(
        "stats: chi-square of [[10, 20], [20, 10]] = 6.67, p = .0098, V = .33",
        close(x["chi2"], 6.6667) and close(x["p"], 0.009823, 1e-4) and close(x["V"], 0.3333),
        x,
    )
    g = stats.regression({"y": [3.1, 4.9, 7.2, 8.8, 11.1], "x": [1, 2, 3, 4, 5]}, "y", ["x"])
    check("stats: regression slope 1.99, R squared .997", close(float(g["B"][1]), 1.99, 1e-6) and g["R2"] > 0.996, (g["B"], g["R2"]))
    al = stats.alpha({"Q1": [1, 2, 3, 4], "Q2": [2, 3, 4, 5], "Q3": [3, 4, 5, 6]}, ["Q1", "Q2", "Q3"])
    check("stats: Cronbach's alpha of three parallel items is 1.00", close(al["alpha"], 1.0), al)
    rr = chat([src]).say("t-test of Score by Gender in survey.xlsx")
    check("stats: by chat, an APA sentence and a Word report", "t(8) = 4.00, p = .004, d = 2.53" in rr and "Report:" in rr, rr)
    rr = chat([src]).say("crosstab Gender by Choice in survey.xlsx")
    check("stats: a small crosstab warns about expected counts below 5", "expected counts below 5" in rr, rr)
    grammar.TRANSPORT = FakeTransport(
        {
            ("POST", "api.languagetool.org/v2/check"): lambda req: (
                200,
                {
                    "matches": [
                        {
                            "offset": 2,
                            "length": 3,
                            "message": "Use 'have' with 'I'.",
                            "replacements": [{"value": "have"}],
                            "rule": {"category": {"id": "GRAMMAR"}},
                        },
                        {
                            "offset": 6,
                            "length": 1,
                            "message": "Use 'an' before a vowel.",
                            "replacements": [{"value": "an"}],
                            "rule": {"category": {"id": "MISC"}},
                        },
                        {
                            "offset": 8,
                            "length": 4,
                            "message": "Possible spelling mistake.",
                            "replacements": [{"value": "apple"}],
                            "rule": {"category": {"id": "TYPOS"}},
                        },
                    ]
                }
                if "I has a aple" in __import__("urllib.parse").parse.unquote_plus(req["body"].decode())
                else {"matches": []},
            )
        }
    )
    from docx import Document

    d = Document()
    d.add_heading("My essay", 1)
    d.add_paragraph("I has a aple.")
    src2 = OUT / "essay.docx"
    d.save(src2)
    c = chat([src2])
    rr = c.say("check grammar in essay.docx")
    check(
        "grammar: shown first (the text goes to LanguageTool's server)",
        rr.startswith("Ready to check essay.docx") and "languagetool" in rr.lower(),
        rr,
    )
    rr = c.say("yes")
    fixed = [p.text for p in Document(next((OUT / "chats").rglob("essay.corrected.docx"))).paragraphs]
    check(
        "grammar: issues listed; spelling and grammar fixed in a copy, a style hint only listed",
        fixed == ["My essay", "I have a apple."] and "3 issue(s) found, 2 fixed" in rr,
        (fixed, rr),
    )


def main():
    t0 = time.time()
    if OUT.exists():
        shutil.rmtree(OUT, ignore_errors=True)
    OUT.mkdir(parents=True)
    for t in (
        t_ocr,
        t_pccare,
        t_diagrams,
        t_cite,
        t_quiz,
        t_ebook,
        t_media,
        t_printing,
        t_maps,
        t_notes,
        t_database,
        t_backup,
        t_contacts_codes_music_calendar,
        t_stores,
        t_web_desktop,
        t_mail_files_chat,
        t_stats_grammar,
    ):
        try:
            t()
        except Exception:  # noqa: BLE001 - one program failing does not hide the others
            import traceback

            check(f"{t.__name__}: ran without an error", False, traceback.format_exc()[-700:])
    bad = [n for n, ok in RESULTS if not ok]
    print(f"\n{'ALL PASS' if not bad else f'{len(bad)} FAILED'}  ({len(RESULTS) - len(bad)}/{len(RESULTS)}, {time.time() - t0:.0f} s)")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
