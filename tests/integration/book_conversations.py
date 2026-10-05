"""Multi-turn conversations about Excel workbooks, each turn checked on the workbook it produced (Excel does every edit in
the background; the checks read the saved file again with openpyxl and compare with Python's own arithmetic).

  .venv\\Scripts\\python.exe tests\\integration\\book_conversations.py [sales] [marks] [--offline]
--offline: no model (the turns that need one are skipped). Chats go to out/docs/_tests/book_chats (wiped at each run).
sales: a messy client workbook (dates typed three ways, 'Rs 42,000' prices, four spellings of Lahore, a repeated order,
an empty row, a price list on another sheet); marks: the class marks workbook the agent made (formulas, a summary row,
a summary sheet with a chart).
"""

import json
import re
import shutil
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
for _s in (sys.stdout, sys.stderr):
    _s.reconfigure(encoding="utf-8", errors="replace")

import warnings  # noqa: E402

warnings.simplefilter("ignore")
import openpyxl  # noqa: E402
from messy_books import PRODUCTS, sales_book  # noqa: E402

from ai_pc.office import xlsx_map as XM  # noqa: E402
from ai_pc.office.bookchat import BookChat  # noqa: E402

CHATS = ROOT / "out" / "docs" / "_tests" / "book_chats"
MARKS = ROOT / "out/docs/student-marks-report_144504/student-marks-report.xlsx"


class Z:
    def __init__(self, c, v0, v1, reply, turn):
        self.c, self.v0, self.v1, self.reply, self.turn = c, v0, v1, reply, turn
        self.m = c.map()
        self.m0 = c.map(v0)

    def new(self):
        return self.v1 != self.v0

    def says(self, *words):
        return all(re.search(w, self.reply, re.I) for w in words)

    def sheet(self, name=None, m=None):
        return XM.sheet(m or self.m, name)

    def names(self, sheet=None, m=None):
        s = self.sheet(sheet, m)
        return [c["name"] for c in s["table"]["cols"]] if s and s.get("table") else []

    def col(self, name, sheet=None, m=None):
        s = self.sheet(sheet, m)
        c = XM.col(None, s, name)
        return XM.values(s, c) if c else None

    def recs(self, sheet=None, m=None):
        return [r for r in XM.records(self.sheet(sheet, m)) if any(v not in (None, "") for v in r.values())]

    def rows(self, sheet=None, m=None):
        return len(self.recs(sheet, m))

    def total(self, name, sheet=None):
        s = self.sheet(sheet)
        t = s["table"]
        if not t["total"]:
            return None
        wb = openpyxl.load_workbook(str(self.c.path()), data_only=True)
        c = XM.col(None, s, name)
        return wb[s["name"]].cell(t["total"], c["idx"]).value

    def cf(self, sheet=None, m=None):
        return self.sheet(sheet, m)["cf"]

    def charts(self, m=None):
        return sum(s["charts"] for s in (m or self.m)["sheets"])

    def kinds(self, m=None):
        return [s["kind"] for s in (m or self.m)["sheets"]]

    def checks_ok(self):
        v = self.c.version
        return bool(v.get("checks")) and all(ch["ok"] for ch in v["checks"])

    def ws(self, sheet=None):
        wb = openpyxl.load_workbook(str(self.c.path()))
        return wb[self.sheet(sheet)["name"]]


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _money(x):
    return f"{round(x):,}"


def _pending(x):
    rs = [r for r in x.recs() if r.get("Status") == "Pending"]
    return len(rs), sum(_num(r["Amount"]) or 0 for r in rs)


def _avg_pct(x):
    vals = [_num(r["Percentage"]) for r in x.recs("Marks") if _num(r.get("Percentage")) is not None]
    return f"{sum(vals) / len(vals):.1%}"


CAT = {p: c for p, c, _ in PRODUCTS}

CONVERSATIONS = {
    "sales": (
        "sales",
        [
            ("what sheets are there?", "question", lambda x: x.says("Orders", "Products") and not x.new(), False),
            (
                "clean up the data",
                "change",
                lambda x: (
                    all(isinstance(v, str) and re.match(r"\d{4}-\d{2}-\d{2}", v) for v in x.col("Date"))
                    and all(_num(v) for v in x.col("Unit Price"))
                    and set(x.col("City")) == {"Lahore", "Karachi", "Islamabad"}
                    and x.rows() == 16
                    and x.checks_ok()
                ),
                False,
            ),
            (
                "add a column Amount = Qty times Unit Price after unit price, and add a total row",
                "change",
                lambda x: (
                    x.names().index("Amount") == x.names().index("Unit Price") + 1
                    and all(abs(_num(r["Amount"]) - _num(r["Qty"]) * _num(r["Unit Price"])) < 1e-6 for r in x.recs())
                    and abs((x.total("Amount") or 0) - sum(_num(r["Amount"]) for r in x.recs())) < 1e-6
                    and x.total("Unit Price") is None
                    and x.checks_ok()
                ),
                False,
            ),
            (
                "bring the category from the products sheet",
                "change",
                lambda x: "Category" in x.names() and all(r["Category"] == CAT[r["Product"]] for r in x.recs()) and x.checks_ok(),
                False,
            ),
            ("which city sold the most?", "question", lambda x: x.says(r"^Lahore") and not x.new(), False),
            (
                "how many orders are pending and what is their total amount?",
                "question",
                lambda x: x.says(rf"\b{_pending(x)[0]} row", _money(_pending(x)[1])) and not x.new(),
                False,
            ),
            (
                "highlight orders over 5 lakh in green. also highlight the overdue ones in red",
                "change",
                lambda x: x.cf() >= x.cf(m=x.m0) + 2 and x.checks_ok(),
                False,
            ),
            ("sort by date, newest first", "change", lambda x: x.col("Date") == sorted(x.col("Date"), reverse=True) and x.checks_ok(), False),
            ("make a pivot of amount by city and month", "change", lambda x: "pivot" in x.kinds() and x.checks_ok(), False),
            (
                "chart the amount by month",
                "change",
                lambda x: x.charts() == x.charts(x.m0) + 1 and len(x.m["sheets"]) == len(x.m0["sheets"]) + 1,
                False,
            ),
            (
                "add an order for Noor Electric: 2 Inverter 6kW at 183,000 in Lahore, pending, dated 2 April 2026",
                "change",
                lambda x: (
                    x.rows() == x.rows(m=x.m0) + 1
                    and any(
                        r["Customer"] == "Noor Electric" and _num(r["Qty"]) == 2 and _num(r["Unit Price"]) == 183000 and _num(r["Amount"]) == 366000
                        for r in x.recs()
                    )
                ),
                True,
            ),
            ("undo", "undo", lambda x: x.rows() == 16, False),
            (
                "rename the Orders sheet to Sales Q1 and freeze the header",
                "change",
                lambda x: "Sales Q1" in x.m["names"] and x.ws("Sales Q1").freeze_panes == "A4",
                False,
            ),
            ("delete the unit price column", "change", lambda x: not x.new() and x.says("anyway"), False),
            ("compare v0 with the current version", "compare", lambda x: x.says("columns added", "Amount"), False),
            ("how many rows does the products sheet have?", "question", lambda x: x.says(r"\b5 row"), False),
            (
                "add a column Days Left that shows how many days until the due date",
                "change",
                lambda x: "Days Left" in x.names() and x.checks_ok(),
                True,
            ),
            ("export", "export", lambda x: x.says(r"\.xlsx", r"\.pdf"), False),
        ],
    ),
    "marks": (
        MARKS,
        [
            ("who scored the highest percentage?", "question", lambda x: x.says("Fatima Noor") and not x.new(), False),
            (
                "add a column Average = average of Maths, English, Science and Urdu after Urdu, and round it to 1 decimal",
                "change",
                lambda x: (
                    x.names("Marks").index("Average") == x.names("Marks").index("Urdu") + 1
                    and all(
                        abs(_num(r["Average"]) - round((r["Maths"] + r["English"] + r["Science"] + r["Urdu"]) / 4 + 1e-9, 1)) < 1e-9
                        for r in x.recs("Marks")
                    )
                    and x.checks_ok()
                ),
                False,
            ),
            ("highlight students who failed in red", "change", lambda x: x.cf("Marks") == x.cf("Marks", x.m0) + 1 and x.checks_ok(), False),
            (
                "sort by total, highest first",
                "change",
                lambda x: x.col("Total", "Marks") == sorted(x.col("Total", "Marks"), reverse=True) and x.checks_ok(),
                False,
            ),
            (
                "add a student Zara Iqbal with Maths 81, English 77, Science 85 and Urdu 90",
                "change",
                lambda x: (
                    any(r["Student"] == "Zara Iqbal" and r["Total"] == 333 and abs(_num(r["Average"]) - 83.3) < 1e-9 for r in x.recs("Marks"))
                    and x.rows("Marks") == x.rows("Marks", x.m0) + 1
                    and x.checks_ok()
                ),
                False,
            ),
            ("what is the class average percentage now?", "question", lambda x: x.says(re.escape(_avg_pct(x))), False),
            ("delete maths", "change", lambda x: not x.new() and x.says("anyway"), False),
            ("delete maths anyway", "change", lambda x: "Maths" not in x.names("Marks") and x.checks_ok(), False),
            ("undo", "undo", lambda x: "Maths" in x.names("Marks"), False),
            ("make a bar chart of total by student", "change", lambda x: x.charts() == x.charts(x.m0) + 1 and x.checks_ok(), False),
            ("add a dropdown to result with Pass and Fail", "change", lambda x: x.sheet("Marks")["validations"] >= 1 and x.checks_ok(), False),
            ("make the header dark green with white text", "change", lambda x: str(x.ws("Marks")["A1"].fill.fgColor.rgb).endswith("1B5E20"), False),
            ("what did you change?", "history", lambda x: x.says("header") and not x.new(), False),
            ("go back to v2", "goto", lambda x: x.c.state["cur"] == 2, False),
            ("export", "export", lambda x: x.says(r"\.xlsx", r"\.pdf"), False),
        ],
    ),
}


def run(names, planner):
    shutil.rmtree(CHATS, ignore_errors=True)
    CHATS.mkdir(parents=True, exist_ok=True)
    rows = []
    t_all = time.perf_counter()
    for name in names:
        src, turns = CONVERSATIONS[name]
        if src == "sales":
            src = sales_book(CHATS / "sales_q1.xlsx")
        c = BookChat.start(src, chats_dir=CHATS, planner=planner)
        print(f"\n=== {name} ({Path(src).name})", flush=True)
        try:
            for msg, want, check, needs_model in turns:
                if needs_model and planner is None:
                    print(f"skip (needs the model)  {msg}")
                    continue
                v0 = c.state["cur"]
                t = time.perf_counter()
                err = None
                try:
                    reply = c.say(msg)
                except Exception as e:  # noqa: BLE001
                    import traceback

                    traceback.print_exc()
                    reply, err = "", f"{type(e).__name__}: {e}"
                    c.last_turn = {"intents": ["crash"], "ops": []}
                secs = time.perf_counter() - t
                intent_ok = want is None or want in (c.last_turn.get("intents") or [])
                try:
                    x = Z(c, v0, c.state["cur"], reply, c.last_turn)
                    check_ok = bool(check(x))
                except Exception as e:  # noqa: BLE001
                    check_ok, err = False, err or f"check {type(e).__name__}: {e}"
                ok = intent_ok and check_ok and not err
                rows.append(
                    {
                        "chat": name,
                        "msg": msg,
                        "ok": ok,
                        "intents": c.last_turn.get("intents"),
                        "seconds": round(secs, 2),
                        "llm": bool(c.last_turn.get("llm")),
                        "reply": reply[:700],
                        "error": err,
                    }
                )
                print(
                    f"{'OK ' if ok else 'BAD'} [{','.join(c.last_turn.get('intents') or [])}{'+llm' if c.last_turn.get('llm') else ''}] "
                    f"v{v0}->v{c.state['cur']} {secs:5.1f}s  {msg}\n      {reply[:420]}",
                    flush=True,
                )
                if not ok:
                    print(f"      intent_ok={intent_ok} check_ok={check_ok} {err or ''}")
        finally:
            c.close()
    n, ok = len(rows), sum(r["ok"] for r in rows)
    usd = planner.cost()[1] if planner is not None else 0.0
    print(
        f"\n{ok}/{n} turns OK; {sum(r['llm'] for r in rows)} used the model for reading the request; "
        f"{sum(r['seconds'] for r in rows) / max(1, n):.1f} s per turn; {time.perf_counter() - t_all:.0f} s in all; ${usd:.4f}"
    )
    (CHATS / "report.json").write_text(json.dumps({"ok": ok, "turns": n, "rows": rows}, ensure_ascii=False, indent=1), encoding="utf-8")
    return ok, n


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    planner = None
    if "--offline" not in sys.argv:
        from ai_pc.llm.planner import ChatPlanner

        planner = ChatPlanner()
    ok, n = run(args or list(CONVERSATIONS), planner)
    sys.exit(0 if ok == n else 1)
