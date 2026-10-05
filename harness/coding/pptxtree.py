"""A PowerPoint file (Canva's export of a design, or any .pptx) read into the same shape as Figma's description of a
frame, so design2code turns Canva designs into web pages too. Each slide is a frame; text boxes become texts (fonts,
sizes, weights, colours, alignment, line spacing, letter spacing, capitals, bullets, links, insets and anchoring kept);
pictures become picture fills (crops applied); rectangles, rounded rectangles, ovals and lines become shapes; freeforms
become vector outlines; groups keep their children. What has no web equivalent (charts, tables, SmartArt, unusual shapes)
is cut from the page's own picture when one is given, so it still looks right.

  tree, images, notes = load("design.pptx", pages=[1], pictures={1: "page1.png"}, work=folder)
  tree: {"document": {"type": "DOCUMENT", "children": [{"type": "CANVAS", "children": [frame]}]}}
"""
import hashlib
import io
import math
import re
from pathlib import Path

from .design2code import family_weight

EMU_PX = 9525.0  # 914400 EMU an inch, 96 px an inch
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
THEME_KEYS = {1: "dk1", 2: "lt1", 3: "dk2", 4: "lt2", 5: "accent1", 6: "accent2", 7: "accent3", 8: "accent4", 9: "accent5", 10: "accent6",
              11: "hlink", 12: "folHlink", 13: "dk1", 14: "lt1", 15: "dk2", 16: "lt2"}


def _px(v):
    return float(v or 0) / EMU_PX


def _hex(h, alpha=1.0):
    h = (h or "000000").lstrip("#")
    return {"r": int(h[0:2], 16) / 255.0, "g": int(h[2:4], 16) / 255.0, "b": int(h[4:6], 16) / 255.0, "a": float(alpha)}


class Reader:
    def __init__(self, prs, work, pictures=None):
        self.prs, self.work, self.pictures = prs, Path(work), pictures or {}
        self.images, self.notes, self.n = {}, [], 0
        self.theme, self.theme_fonts = self._theme()
        self.page_img = {}

    # ---------------------------------------------------------------- colours
    def _theme(self):
        out, fonts = {}, {}
        try:
            from lxml import etree
            from pptx.opc.constants import RELATIONSHIP_TYPE as RT
            root = etree.fromstring(self.prs.slide_master.part.part_related_by(RT.THEME).blob)
            for el in root.iter(A + "clrScheme"):
                for c in el:
                    key = c.tag.replace(A, "")
                    v = c.find(A + "srgbClr")
                    s = c.find(A + "sysClr")
                    out[key] = v.get("val") if v is not None else (s.get("lastClr") if s is not None else None)
                break
            for kind, key in (("majorFont", "+mj"), ("minorFont", "+mn")):
                el = root.find(f".//{A}{kind}/{A}latin")
                if el is not None and el.get("typeface"):
                    fonts[key] = el.get("typeface")
        except Exception:  # noqa: BLE001  (no theme: plain black and white, Arial)
            pass
        return out, fonts

    def color_el(self, el):
        """An a:solidFill (or any element holding a colour) -> Figma colour dict, or None."""
        if el is None:
            return None
        for tag in ("srgbClr", "schemeClr", "sysClr", "prstClr"):
            c = el.find(A + tag)
            if c is None:
                continue
            if tag == "srgbClr":
                h = c.get("val")
            elif tag == "schemeClr":
                v = c.get("val")
                h = self.theme.get({"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}.get(v, v)) or "000000"
            elif tag == "sysClr":
                h = c.get("lastClr") or "000000"
            else:
                h = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "blue": "0000FF", "green": "008000"}.get(c.get("val"), "000000")
            alpha = c.find(A + "alpha")
            col = _hex(h, int(alpha.get("val")) / 100000.0 if alpha is not None else 1.0)
            lm, lo = c.find(A + "lumMod"), c.find(A + "lumOff")
            if lm is not None or lo is not None:  # theme tints and shades
                mod = int(lm.get("val")) / 100000.0 if lm is not None else 1.0
                off = int(lo.get("val")) / 100000.0 if lo is not None else 0.0
                for k in ("r", "g", "b"):
                    col[k] = max(0.0, min(1.0, col[k] * mod + off))
            return col
        return None

    def fills_of(self, sppr, w, h, shape=None):
        """Fills from a shape's properties (solid, gradient, picture) -> Figma paints; [] when it has none."""
        if sppr is None:
            return []
        if sppr.find(A + "noFill") is not None:
            return []
        sf = sppr.find(A + "solidFill")
        if sf is not None:
            c = self.color_el(sf)
            return [{"type": "SOLID", "color": c}] if c else []
        gf = sppr.find(A + "gradFill")
        if gf is not None:
            stops = []
            for gs in gf.iter(A + "gs"):
                c = self.color_el(gs)
                if c:
                    stops.append({"color": c, "position": int(gs.get("pos", "0")) / 100000.0})
            lin = gf.find(A + "lin")
            ang = int(lin.get("ang", "0")) / 60000.0 if lin is not None else 90.0  # clockwise from left-to-right
            dx, dy = math.cos(math.radians(ang)), math.sin(math.radians(ang))
            if stops:
                return [{"type": "GRADIENT_LINEAR", "gradientStops": sorted(stops, key=lambda s: s["position"]),
                         "gradientHandlePositions": [{"x": 0.5 - dx / 2, "y": 0.5 - dy / 2}, {"x": 0.5 + dx / 2, "y": 0.5 + dy / 2},
                                                     {"x": 0.5 - dy / 2, "y": 0.5 + dx / 2}]}]
            return []
        bf = sppr.find(A + "blipFill")
        if bf is not None and shape is not None:
            ref = self.blip(bf, shape.part if hasattr(shape, "part") else shape)
            return [{"type": "IMAGE", "imageRef": ref, "scaleMode": "FILL"}] if ref else []
        return []

    def blip(self, blipfill, part, crop=None):
        b = blipfill.find(A + "blip")
        if b is None:
            return None
        rid = b.get(R + "embed")
        try:
            blob = part.related_part(rid).blob
        except Exception:  # noqa: BLE001
            return None
        src = blipfill.find(A + "srcRect")
        crop = crop or ({k: int(src.get(k, "0")) / 100000.0 for k in ("l", "t", "r", "b")} if src is not None else None)
        if crop and any(abs(v) > 1e-4 for v in crop.values()):
            blob = self._crop(blob, crop)
        ref = hashlib.sha1(blob).hexdigest()[:16]
        if ref not in self.images:
            ext = "png" if blob[:4] == b"\x89PNG" else "jpg" if blob[:2] == b"\xff\xd8" else "gif" if blob[:4] == b"GIF8" else "png"
            if ext == "png" and blob[:4] != b"\x89PNG":  # an EMF, WMF or other: re-encode what Pillow can read
                blob = self._repng(blob)
                if blob is None:
                    self.note("a picture in a format the browser cannot show was left out")
                    return None
            p = self.work / f"{ref}.{ext}"
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_bytes(blob)
            self.images[ref] = str(p)
        return ref

    @staticmethod
    def _crop(blob, crop):
        from PIL import Image
        im = Image.open(io.BytesIO(blob))
        w, h = im.size
        box = (int(round(w * max(0, crop["l"]))), int(round(h * max(0, crop["t"]))), int(round(w * (1 - max(0, crop["r"])))),
               int(round(h * (1 - max(0, crop["b"])))))
        if box[2] - box[0] < 2 or box[3] - box[1] < 2:
            return blob
        out = io.BytesIO()
        fmt = "PNG" if im.mode in ("RGBA", "LA", "P") else "JPEG"
        im.crop(box).save(out, fmt, quality=92)
        return out.getvalue()

    @staticmethod
    def _repng(blob):
        from PIL import Image
        try:
            im = Image.open(io.BytesIO(blob))
            out = io.BytesIO()
            im.save(out, "PNG")
            return out.getvalue()
        except Exception:  # noqa: BLE001
            return None

    def note(self, s):
        if s not in self.notes:
            self.notes.append(s)

    def nid(self):
        self.n += 1
        return f"{self.page}:{self.n}"

    # ---------------------------------------------------------------- shapes
    def xfrm(self, el):
        sppr = el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
        if sppr is None:
            sppr = el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}grpSpPr")
        x = sppr.find(A + "xfrm") if sppr is not None else None
        return sppr, x

    def node_box(self, x, y, w, h, rot=0.0, flip_h=False, flip_v=False):
        """Figma geometry for a box at (x, y, w, h) px turned rot degrees clockwise (PowerPoint's way)."""
        n = {"size": {"x": w, "y": h}}
        th = math.radians(rot)
        a, b, c, d = math.cos(th), math.sin(th), -math.sin(th), math.cos(th)  # clockwise on screen (y down)
        if flip_h:
            a, b = -a, -b
        if flip_v:
            c, d = -c, -d
        cx, cy = x + w / 2, y + h / 2
        xs = [a * u + c * v for u, v in ((-w / 2, -h / 2), (w / 2, -h / 2), (-w / 2, h / 2), (w / 2, h / 2))]
        ys = [b * u + d * v for u, v in ((-w / 2, -h / 2), (w / 2, -h / 2), (-w / 2, h / 2), (w / 2, h / 2))]
        n["absoluteBoundingBox"] = {"x": cx + min(xs), "y": cy + min(ys), "width": max(xs) - min(xs), "height": max(ys) - min(ys)}
        tx, ty = cx - (a * w / 2 + c * h / 2), cy - (b * w / 2 + d * h / 2)
        n["relativeTransform"] = [[a, c, tx], [b, d, ty]]
        return n

    def shapes(self, shapes, tf=lambda x, y, w, h: (x, y, w, h)):
        out = []
        for sh in shapes:
            try:
                out += self.shape(sh, tf)
            except Exception as e:  # noqa: BLE001  (one odd shape never spoils the page)
                self.note(f"'{getattr(sh, 'name', '?')}' could not be read ({type(e).__name__})")
        return out

    def shape(self, sh, tf):
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        el = sh._element
        sppr, xf = self.xfrm(el)
        if sh.left is None or sh.width is None:
            return []
        x, y, w, h = tf(_px(sh.left), _px(sh.top), _px(sh.width), _px(sh.height))
        rot = float(getattr(sh, "rotation", 0.0) or 0.0)
        fh = xf is not None and xf.get("flipH") == "1"
        fv = xf is not None and xf.get("flipV") == "1"
        st = sh.shape_type
        name = sh.name or ""
        if st == MSO_SHAPE_TYPE.GROUP:
            return [self.group(sh, xf, tf, name)]
        base = {"id": self.nid(), "name": name, "visible": True, **self.node_box(x, y, w, h, rot, fh, fv)}
        if st == MSO_SHAPE_TYPE.PICTURE:
            bf = el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}blipFill")
            ref = self.blip(bf, sh.part) if bf is not None else None
            if not ref:
                return self.art(base, name)
            node = dict(base, type="RECTANGLE", fills=[{"type": "IMAGE", "imageRef": ref, "scaleMode": "STRETCH"}])
            geom = sppr.find(A + "prstGeom") if sppr is not None else None
            if geom is not None and geom.get("prst") == "ellipse":
                node["type"] = "ELLIPSE"
            elif geom is not None and geom.get("prst") == "roundRect":
                node["cornerRadius"] = self._round(geom, w, h)
            return [node]
        if st in (MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.DIAGRAM) or \
                el.tag.endswith("graphicFrame"):
            return self.art(base, name)
        out = []
        geom = sppr.find(A + "prstGeom") if sppr is not None else None
        cust = sppr.find(A + "custGeom") if sppr is not None else None
        prst = geom.get("prst") if geom is not None else ("custom" if cust is not None else "rect")
        fills = self.fills_of(sppr, w, h, sh)
        strokes, weight = self.line(sppr)
        style = el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:  # the theme's look, for a shape that does not set its own
            if not fills and (sppr is None or sppr.find(A + "noFill") is None) and sppr is not None and sppr.find(A + "solidFill") is None \
                    and sppr.find(A + "gradFill") is None and sppr.find(A + "blipFill") is None:
                fr = style.find(A + "fillRef")
                if fr is not None and fr.get("idx", "0") != "0":
                    c = self.color_el(fr)
                    fills = [{"type": "SOLID", "color": c}] if c else []
            ln = sppr.find(A + "ln") if sppr is not None else None
            if not strokes and (ln is None or (ln.find(A + "noFill") is None and ln.find(A + "solidFill") is None)):
                lr = style.find(A + "lnRef")
                if lr is not None and lr.get("idx", "0") != "0":
                    c = self.color_el(lr)
                    if c:
                        strokes, weight = [{"type": "SOLID", "color": c}], _px(ln.get("w")) if ln is not None and ln.get("w") else 0.75 * 96 / 72
        if el.tag.endswith("}cxnSp") or prst in ("line", "straightConnector1"):
            if strokes:  # a line is drawn as a thin box along it
                lw = max(weight, 1.0)
                length = math.hypot(w, h)
                ang = rot + math.degrees(math.atan2(h, w)) * (-1 if fh != fv else 1)
                cx, cy = x + w / 2, y + h / 2
                out.append(dict(base, type="RECTANGLE", fills=strokes, **self.node_box(cx - length / 2, cy - lw / 2, length, lw, ang)))
            return out
        if fills or strokes:
            if prst in ("rect", "roundRect", "ellipse", "snip1Rect", "flowChartProcess", "flowChartAlternateProcess"):
                node = dict(base, type="ELLIPSE" if prst == "ellipse" else "RECTANGLE", fills=fills, strokes=strokes, strokeWeight=weight,
                            strokeAlign="CENTER")
                if prst in ("roundRect", "flowChartAlternateProcess"):
                    node["cornerRadius"] = self._round(geom, w, h)
                out.append(node)
            elif prst == "custom":
                out.append(self.freeform(base, cust, fills, strokes, weight, w, h))
            else:
                out += self.art(base, name, f"its shape ({prst})")
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = self.text(sh, base, x, y, w, h, rot, fh, fv)
            if t:
                out.append(t)
        return out

    @staticmethod
    def _round(geom, w, h):
        adj = 16667
        for gd in geom.iter(A + "gd"):
            m = re.search(r"val\s+(\d+)", gd.get("fmla", ""))
            if m:
                adj = int(m.group(1))
        return min(w, h) * adj / 100000.0

    def line(self, sppr):
        ln = sppr.find(A + "ln") if sppr is not None else None
        if ln is None or ln.find(A + "noFill") is not None:
            return [], 0.0
        sf = ln.find(A + "solidFill")
        c = self.color_el(sf) if sf is not None else None
        if not c:
            return [], 0.0
        return [{"type": "SOLID", "color": c}], _px(ln.get("w", "12700"))

    def freeform(self, base, cust, fills, strokes, weight, w, h):
        paths = []
        for path in cust.iter(A + "path"):
            pw, ph = float(path.get("w") or w * EMU_PX or 1), float(path.get("h") or h * EMU_PX or 1)
            sx, sy = w / pw if pw else 1, h / ph if ph else 1
            d, cur = [], (0.0, 0.0)
            for cmd in path:
                tag = cmd.tag.replace(A, "")
                pts = [(float(p.get("x")) * sx, float(p.get("y")) * sy) for p in cmd.iter(A + "pt")]
                if tag == "moveTo" and pts:
                    d.append(f"M{pts[0][0]:.2f} {pts[0][1]:.2f}")
                    cur = pts[0]
                elif tag == "lnTo" and pts:
                    d.append(f"L{pts[0][0]:.2f} {pts[0][1]:.2f}")
                    cur = pts[0]
                elif tag == "cubicBezTo" and len(pts) == 3:
                    d.append("C" + " ".join(f"{a:.2f} {b:.2f}" for a, b in pts))
                    cur = pts[2]
                elif tag == "quadBezTo" and len(pts) == 2:
                    d.append("Q" + " ".join(f"{a:.2f} {b:.2f}" for a, b in pts))
                    cur = pts[1]
                elif tag == "arcTo":
                    wr, hr = float(cmd.get("wR")) * sx, float(cmd.get("hR")) * sy
                    st, sw = math.radians(int(cmd.get("stAng")) / 60000.0), math.radians(int(cmd.get("swAng")) / 60000.0)
                    cx, cy = cur[0] - wr * math.cos(st), cur[1] - hr * math.sin(st)
                    end = (cx + wr * math.cos(st + sw), cy + hr * math.sin(st + sw))
                    d.append(f"A{wr:.2f} {hr:.2f} 0 {int(abs(sw) > math.pi)} {int(sw > 0)} {end[0]:.2f} {end[1]:.2f}")
                    cur = end
                elif tag == "close":
                    d.append("Z")
            if d:
                paths.append({"path": " ".join(d), "windingRule": "NONZERO"})
        node = dict(base, type="VECTOR", fills=fills, fillGeometry=paths, strokes=[], strokeGeometry=[])
        if strokes and weight:
            node["strokes"], node["strokeWeight"] = strokes, weight
            node["_svg_stroke"] = True
            self.note("outlines of freeform shapes are drawn as fills only")
        return node

    def group(self, sh, xf, tf, name):
        off, ext = xf.find(A + "off"), xf.find(A + "ext")
        choff, chext = xf.find(A + "chOff"), xf.find(A + "chExt")
        gx, gy, gw, gh = _px(off.get("x")), _px(off.get("y")), _px(ext.get("cx")), _px(ext.get("cy"))
        cx0, cy0 = (_px(choff.get("x")), _px(choff.get("y"))) if choff is not None else (gx, gy)
        cw, ch = (_px(chext.get("cx")), _px(chext.get("cy"))) if chext is not None else (gw, gh)
        kx, ky = (gw / cw if cw else 1.0), (gh / ch if ch else 1.0)

        def inner(x, y, w, h):  # child coordinates -> the group's place on the slide, then the parents' mapping
            return tf(gx + (x - cx0) * kx, gy + (y - cy0) * ky, w * kx, h * ky)
        kids = self.shapes(sh.shapes, inner)
        X, Y, W, H = tf(gx, gy, gw, gh)
        if float(getattr(sh, "rotation", 0) or 0):
            self.note(f"'{name}': a turned group is shown unturned")
        return {"id": self.nid(), "name": name, "type": "GROUP", "visible": True, "children": kids,
                "absoluteBoundingBox": {"x": X, "y": Y, "width": W, "height": H}}

    def art(self, base, name, why="it"):
        """Something with no web equivalent: cut from the page's picture, so it still looks right."""
        img = self.page_picture()
        if img is None:
            self.note(f"'{name}': {why} has no web equivalent and no page picture was given; it is left out")
            return []
        b = base["absoluteBoundingBox"]
        W, H = self.page_size
        sx, sy = img.size[0] / W, img.size[1] / H
        box = (max(0, int(b["x"] * sx)), max(0, int(b["y"] * sy)), min(img.size[0], int(math.ceil((b["x"] + b["width"]) * sx))),
               min(img.size[1], int(math.ceil((b["y"] + b["height"]) * sy))))
        if box[2] - box[0] < 2 or box[3] - box[1] < 2:
            return []
        out = io.BytesIO()
        img.crop(box).save(out, "PNG")
        blob = out.getvalue()
        ref = "art-" + hashlib.sha1(blob).hexdigest()[:12]
        p = self.work / f"{ref}.png"
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(blob)
        self.images[ref] = str(p)
        self.note(f"'{name}': {why} has no web equivalent, so it is a picture cut from the design")
        X0, Y0 = box[0] / sx, box[1] / sy
        return [{"id": base["id"], "name": name or "artwork", "type": "RECTANGLE", "visible": True,
                 "absoluteBoundingBox": {"x": X0, "y": Y0, "width": (box[2] - box[0]) / sx, "height": (box[3] - box[1]) / sy},
                 "fills": [{"type": "IMAGE", "imageRef": ref, "scaleMode": "STRETCH"}]}]

    def page_picture(self):
        if self.page not in self.page_img:
            p = self.pictures.get(self.page)
            if p and Path(p).exists():
                from PIL import Image
                self.page_img[self.page] = Image.open(p).convert("RGBA")
            else:
                self.page_img[self.page] = None
        return self.page_img[self.page]

    # ---------------------------------------------------------------- text
    def text(self, sh, base, x, y, w, h, rot, fh, fv):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        tf = sh.text_frame
        body = tf._txBody.find(A + "bodyPr")
        ins = {k: _px(body.get(k)) if body is not None and body.get(k) is not None else d
               for k, d in (("lIns", 9.6), ("tIns", 4.8), ("rIns", 9.6), ("bIns", 4.8))}
        scale = 1.0
        if body is not None:
            na = body.find(A + "normAutofit")
            if na is not None and na.get("fontScale"):
                scale = int(na.get("fontScale")) / 100000.0
        chars, ov, table, types, styles = [], [], {}, [], []
        first = None
        for pi, p in enumerate(tf.paragraphs):
            if pi:
                chars.append("\n")
                ov.append(0)
            ppr = p._p.find(A + "pPr")
            bullet = "NONE"
            if ppr is not None:
                if ppr.find(A + "buChar") is not None:
                    bullet = "UNORDERED"
                elif ppr.find(A + "buAutoNum") is not None:
                    bullet = "ORDERED"
            types.append(bullet)
            for r in p.runs:
                st = self.run_style(r, p, scale)
                if first is None:
                    first = dict(st, align=p.alignment, spacing=p.line_spacing)
                key = tuple(sorted((k, str(v)) for k, v in st.items()))
                if key not in styles:
                    styles.append(key)
                sid = styles.index(key)
                table[sid] = st
                for ch in r.text:
                    chars.append(ch)
                    ov.append(sid)
                    if ord(ch) > 0xFFFF:
                        ov.append(sid)
        if first is None:
            return {}
        text = "".join(chars)
        size = first["fontSize"]
        sp = first.get("spacing")
        if isinstance(sp, float):
            lh = size * sp * 1.2
        elif sp is not None:
            lh = _px(sp)
        else:
            lh = size * 1.2
        align = {PP_ALIGN.CENTER: "CENTER", PP_ALIGN.RIGHT: "RIGHT", PP_ALIGN.JUSTIFY: "JUSTIFIED"}.get(first.get("align"), "LEFT")
        anchor = {MSO_ANCHOR.MIDDLE: "CENTER", MSO_ANCHOR.BOTTOM: "BOTTOM"}.get(tf.vertical_anchor, "TOP")
        if tf.vertical_anchor is None and body is not None:
            anchor = {"ctr": "CENTER", "b": "BOTTOM"}.get(body.get("anchor"), "TOP")
        wrap = not (body is not None and body.get("wrap") == "none")
        bx, by = x + ins["lIns"], y + ins["tIns"]
        bw, bh = max(1.0, w - ins["lIns"] - ins["rIns"]), max(1.0, h - ins["tIns"] - ins["bIns"])
        base0 = table[0] if 0 in table else first
        style = {"fontFamily": base0["fontFamily"], "fontWeight": base0["fontWeight"], "italic": base0["italic"], "fontSize": base0["fontSize"],
                 "lineHeightPx": lh, "letterSpacing": base0.get("letterSpacing", 0), "textAlignHorizontal": align, "textAlignVertical": anchor,
                 "textAutoResize": "NONE" if wrap else "WIDTH_AND_HEIGHT"}
        if base0.get("textCase"):
            style["textCase"] = base0["textCase"]
        if base0.get("textDecoration"):
            style["textDecoration"] = base0["textDecoration"]
        over = {}
        for sid, st in table.items():
            if sid == 0:
                continue
            d = {k: v for k, v in st.items() if k != "color" and st.get(k) != base0.get(k)}
            if st.get("color") != base0.get("color"):
                d["fills"] = [{"type": "SOLID", "color": st["color"]}]
            if st.get("hyperlink"):
                d["hyperlink"] = {"type": "URL", "url": st["hyperlink"]}
            over[sid] = d
        node = {"id": self.nid(), "name": text.strip().split("\n")[0][:40] or sh.name or "Text", "type": "TEXT", "visible": True, "characters": text,
                "style": style,
                "fills": [{"type": "SOLID", "color": base0["color"]}], "characterStyleOverrides": ov if over else [],
                "styleOverrideTable": over, "lineTypes": types, **self.node_box(bx, by, bw, bh, rot, fh, fv)}
        if base0.get("hyperlink"):
            node["style"]["hyperlink"] = {"type": "URL", "url": base0["hyperlink"]}
        return node

    def run_style(self, r, p, scale):
        f = r.font
        rpr = r._r.find(A + "rPr")
        fam = f.name or (p.font.name if p.font is not None else None) or self._default_font()
        if fam.startswith("+"):  # the theme's heading or body font
            fam = self.theme_fonts.get(fam[:3], "Arial")
        bold = f.bold if f.bold is not None else (p.font.bold if p.font is not None else None)
        fam, wt = family_weight(fam, 700 if bold else None)
        size_pt = f.size.pt if f.size is not None else (p.font.size.pt if p.font is not None and p.font.size is not None else 18.0)
        col = None
        if rpr is not None:
            col = self.color_el(rpr.find(A + "solidFill"))
        st = {"fontFamily": fam, "fontWeight": wt, "italic": bool(f.italic), "fontSize": round(size_pt * 96 / 72 * scale, 2),
              "color": col or {"r": 0, "g": 0, "b": 0, "a": 1}}
        if rpr is not None and rpr.get("spc"):
            st["letterSpacing"] = int(rpr.get("spc")) / 100.0 * 96 / 72
        if rpr is not None and rpr.get("cap") == "all":
            st["textCase"] = "UPPER"
        if f.underline:
            st["textDecoration"] = "UNDERLINE"
        if rpr is not None and rpr.get("strike") not in (None, "noStrike"):
            st["textDecoration"] = "STRIKETHROUGH"
        try:
            if r.hyperlink.address:
                st["hyperlink"] = r.hyperlink.address
        except Exception:  # noqa: BLE001
            pass
        return st

    def _default_font(self):
        try:
            latin = self.prs.slide_master.part._element.find(".//" + A + "latin")
            if latin is not None and latin.get("typeface"):
                t = latin.get("typeface")
                return self.theme_fonts.get(t[:3], "Arial") if t.startswith("+") else t
        except Exception:  # noqa: BLE001
            pass
        return "Arial"

    # ---------------------------------------------------------------- slides
    def background(self, slide, W, H):
        P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
        for src in (slide, slide.slide_layout, slide.slide_layout.slide_master):
            bg = src._element.find(P + "cSld/" + P + "bg")
            if bg is None:
                continue
            bgpr = bg.find(P + "bgPr")
            if bgpr is not None:
                return self.fills_of(bgpr, W, H, src)
            ref = bg.find(P + "bgRef")
            if ref is not None:
                c = self.color_el(ref)
                return [{"type": "SOLID", "color": c}] if c else []
        return [{"type": "SOLID", "color": _hex("FFFFFF")}]

    def slide(self, slide, no, title):
        self.page, self.n = no, 0
        W, H = _px(self.prs.slide_width), _px(self.prs.slide_height)
        self.page_size = (W, H)
        kids = self.shapes(slide.shapes)
        return {"id": f"{no}:0", "name": title, "type": "FRAME", "visible": True, "clipsContent": True,
                "absoluteBoundingBox": {"x": 0.0, "y": 0.0, "width": W, "height": H}, "fills": self.background(slide, W, H), "children": kids}


def load(path, pages=None, pictures=None, work=None, title=None):
    """-> (tree in Figma's shape, images {ref: file}, notes). pages: [1, 2] (1-based) or None for all; several pages are
    stacked top to bottom as sections of one page."""
    from pptx import Presentation
    prs = Presentation(str(path))
    work = Path(work or Path(path).with_suffix("")).resolve()
    rd = Reader(prs, work, pictures)
    slides = list(prs.slides)
    want = [p for p in (pages or range(1, len(slides) + 1)) if 1 <= p <= len(slides)]
    if not want:
        raise ValueError("the file has no such page")
    name = title or Path(path).stem
    frames = [rd.slide(slides[i - 1], i, f"{name} page {i}" if len(want) > 1 else name) for i in want]
    if len(frames) == 1:
        root = frames[0]
    else:
        W = max(f["absoluteBoundingBox"]["width"] for f in frames)
        y = 0.0
        for f in frames:
            _shift(f, 0.0, y)
            f["name"] = f"Section {f['id'].split(':')[0]}"
            y += f["absoluteBoundingBox"]["height"]
        root = {"id": "0:0", "name": name, "type": "FRAME", "visible": True, "clipsContent": True, "children": frames,
                "absoluteBoundingBox": {"x": 0.0, "y": 0.0, "width": W, "height": y}, "fills": [{"type": "SOLID", "color": _hex("FFFFFF")}]}
    doc = {"document": {"id": "0:doc", "type": "DOCUMENT", "children": [{"id": "0:page", "type": "CANVAS", "name": name, "children": [root]}]}}
    return doc, rd.images, rd.notes


def _shift(n, dx, dy):
    b = n.get("absoluteBoundingBox")
    if b:
        b["x"] += dx
        b["y"] += dy
    t = n.get("relativeTransform")
    if t:
        t[0][2] += dx
        t[1][2] += dy
    for k in n.get("children") or []:
        _shift(k, dx, dy)
