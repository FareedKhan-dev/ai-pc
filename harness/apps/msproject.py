"""Microsoft Project (and the free ProjectLibre) by project files: a plan from words ('foundation 10 days; walls 3 weeks
after foundation by Mason team; handover milestone after finishing') or from an Excel sheet (Task, Duration,
Predecessors, Resource), scheduled on working days with the critical path found, written as Project's XML (MSPDI: File >
Open in Project or ProjectLibre) and as an Excel Gantt chart. Checked by reading the XML back and scheduling it again:
the same dates, every task starts after the tasks it waits for, no loops, and the critical tasks have no slack.

  "project plan 'House build' starting 2 november: foundation 10 days; walls 15 days after foundation; roof 7 days after walls"
  'ms project plan from tasks.xlsx, 6 day week'
"""
import datetime as dt
import re
from pathlib import Path
from xml.etree import ElementTree as ET

NAME, LABEL = "msproject", "Microsoft Project: plans with a critical path as Project XML + Excel Gantt"
EXAMPLES = ["project plan 'House build' starting 2 november: foundation 10 days; walls 15 days after foundation; roof 7 days after walls",
            "ms project plan from tasks.xlsx, 6 day week"]
NS = "http://schemas.microsoft.com/project"
HOURS = 8
# the order Project's XML schema wants (Microsoft's MSPDI schema reference; Active and Manual come after Name since Project 2010)
ORDER = {
    "Project": ("SaveVersion UID Name Title Subject Category Company Manager Author CreationDate Revision LastSaved ScheduleFromStart StartDate FinishDate "
                "FYStartDate CriticalSlackLimit CurrencyDigits CurrencySymbol CurrencyCode CurrencySymbolPosition CalendarUID DefaultStartTime DefaultFinishTime "
                "MinutesPerDay MinutesPerWeek DaysPerMonth DefaultTaskType DefaultFixedCostAccrual DefaultStandardRate DefaultOvertimeRate DurationFormat WorkFormat "
                "EditableActualCosts HonorConstraints EarnedValueMethod InsertedProjectsLikeSummary MultipleCriticalPaths NewTasksEffortDriven NewTasksEstimated "
                "SplitsInProgressTasks SpreadActualCost SpreadPercentComplete TaskUpdatesResource FiscalYearStart WeekStartDay MoveCompletedEndsBack "
                "MoveRemainingStartsBack MoveRemainingStartsForward MoveCompletedEndsForward BaselineForEarnedValue AutoAddNewResourcesAndTasks StatusDate CurrentDate "
                "MicrosoftProjectServerURL Autolink NewTaskStartDate DefaultTaskEVMethod ProjectExternallyEdited ExtendedCreationDate ActualsInSync RemoveFileProperties "
                "AdminProject OutlineCodes WBSMasks ExtendedAttributes Calendars Tasks Resources Assignments"),
    "Task": ("UID GUID ID Name Active Manual Type IsNull CreateDate Contact WBS WBSLevel OutlineNumber OutlineLevel Priority Start Finish Duration DurationFormat "
             "Work Stop Resume ResumeValid EffortDriven Recurring OverAllocated Estimated Milestone Summary Critical IsSubproject IsSubprojectReadOnly "
             "SubprojectName ExternalTask ExternalTaskProject EarlyStart EarlyFinish LateStart LateFinish StartVariance FinishVariance WorkVariance FreeSlack "
             "TotalSlack FixedCost FixedCostAccrual PercentComplete PercentWorkComplete Cost OvertimeCost OvertimeWork ActualStart ActualFinish ActualDuration "
             "ActualCost ActualOvertimeCost ActualWork ActualOvertimeWork RegularWork RemainingDuration RemainingCost RemainingWork RemainingOvertimeCost "
             "RemainingOvertimeWork ACWP CV ConstraintType CalendarUID ConstraintDate Deadline LevelAssignments LevelingCanSplit LevelingDelay "
             "LevelingDelayFormat PreLeveledStart PreLeveledFinish Hyperlink HyperlinkAddress HyperlinkSubAddress IgnoreResourceCalendar Notes HideBar Rollup "
             "BCWS BCWP PhysicalPercentComplete EarnedValueMethod PredecessorLink ActualWorkProtected ActualOvertimeWorkProtected ExtendedAttribute Baseline "
             "OutlineCode IsPublished StatusManager CommitmentStart CommitmentFinish CommitmentType TimephasedData"),
    "PredecessorLink": "PredecessorUID Type CrossProject CrossProjectName LinkLag LagFormat",
    "Resource": ("UID ID Name Type IsNull Initials Phonetics NTAccount MaterialLabel Code Group WorkGroup EmailAddress Hyperlink HyperlinkAddress "
                 "HyperlinkSubAddress MaxUnits PeakUnits OverAllocated AvailableFrom AvailableTo Start Finish CanLevel AccrueAt"),
    "Assignment": ("UID TaskUID ResourceUID PercentWorkComplete ActualCost ActualFinish ActualOvertimeCost ActualOvertimeWork ActualStart ActualWork ACWP "
                   "Confirmed Cost CostRateTable CostVariance CV Delay Finish FinishVariance Hyperlink HyperlinkAddress HyperlinkSubAddress WorkVariance "
                   "HasFixedRateUnits FixedMaterial LevelingDelay LevelingDelayFormat LinkedFields Milestone Notes Overallocated OvertimeCost OvertimeWork "
                   "PeakUnits RegularWork RemainingCost RemainingOvertimeCost RemainingOvertimeWork RemainingWork ResponsePending Start Stop Resume "
                   "StartVariance Summary SV Units UpdateNeeded VAC Work WorkContour"),
    "Calendar": "UID Name IsBaseCalendar BaseCalendarUID WeekDays",
    "WeekDay": "DayType DayWorking TimePeriod WorkingTimes",
}


def in_schema_order(root):
    """Every element's children in the order the schema lists them (unknown names count as wrong)."""
    for e in root.iter():
        tag = e.tag.split("}")[-1]
        if tag not in ORDER:
            continue
        seq = ORDER[tag].split()
        pos = [seq.index(c.tag.split("}")[-1]) if c.tag.split("}")[-1] in seq else -1 for c in e]
        if -1 in pos or pos != sorted(pos):
            return False
    return True


class PlanError(ValueError):
    pass


def tasks_from_text(body):
    out = []
    for item in [x.strip() for x in re.split(r"\s*[;\n]\s*", body) if x.strip()]:
        m = re.match(r"^(?:milestone\s+)?(?P<name>.+?)(?:\s+milestone)?(?:\s+after\s+(?P<after>.+?))?(?:\s+by\s+(?P<res>.+?))?\.?$", item, re.I) \
            if re.search(r"\bmilestone\b", item, re.I) else \
            re.match(r"^(?P<name>.+?)\s*[:,-]?\s+(?P<n>\d+(?:\.\d+)?)\s*(?P<u>working\s+days?|days?|d|weeks?|wks?|w)\b(?:\s+after\s+(?P<after>.+?))?(?:\s+by\s+(?P<res>.+?))?\.?$",
                     item, re.I)
        if not m:
            raise PlanError(f"can't read '{item}': say it like 'walls 15 days after foundation'")
        g = m.groupdict()
        days = 0 if g.get("n") is None else float(g["n"]) * (5 if g["u"].lower().startswith("w") else 1)
        out.append({"name": g["name"].strip(), "days": round(days), "after": [a.strip() for a in re.split(r"\s*,\s*|\s+and\s+", g["after"] or "") if a.strip()],
                    "res": (g.get("res") or "").strip() or None})
    return out


def tasks_from_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = [["" if c is None else str(c).strip() for c in r] for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    head = [h.lower() for h in rows[0]]
    col = lambda *names: next((i for i, h in enumerate(head) if any(n in h for n in names)), None)  # noqa: E731
    ci, cd, cp, cr = col("task", "name", "activity"), col("duration", "days"), col("predecessor", "after", "depends"), col("resource", "who", "by")
    if ci is None or cd is None:
        raise PlanError("the sheet needs a Task column and a Duration column (Predecessors and Resource are optional)")
    body = [r for r in rows[1:] if len(r) > ci and r[ci]]
    out = []
    for r in body:
        n = re.match(r"(\d+(?:\.\d+)?)\s*(w)?", (r[cd] or "0").lower())
        after = [a.strip() for a in re.split(r"[,;]", r[cp]) if a.strip()] if cp is not None and cp < len(r) and r[cp] else []
        after = [body[int(a) - 1][ci] if a.isdigit() and 0 < int(a) <= len(body) else a for a in after]  # row numbers or names
        out.append({"name": r[ci], "days": round(float(n.group(1)) * (5 if n.group(2) else 1)) if n else 0, "after": after,
                    "res": (r[cr] if cr is not None and cr < len(r) else "") or None})
    return out


def schedule(tasks, start, workdays):
    """Critical path on working days: early/late start in working-day numbers, slack, and the calendar dates."""
    names = {t["name"].lower(): i for i, t in enumerate(tasks)}
    for t in tasks:
        t["pred"] = []
        for a in t["after"]:
            k = names.get(a.lower()) if a.lower() in names else next((i for n, i in names.items() if n.startswith(a.lower())), None)
            if k is None:
                raise PlanError(f"'{t['name']}' waits for '{a}', which is not a task in the plan")
            t["pred"].append(k)
    order, state = [], {}

    def visit(i, path):
        if state.get(i) == 1:
            raise PlanError("the tasks wait for each other in a loop: " + " -> ".join(tasks[j]["name"] for j in path + [i]))
        if state.get(i) == 2:
            return
        state[i] = 1
        for p in tasks[i]["pred"]:
            visit(p, path + [i])
        state[i] = 2
        order.append(i)
    for i in range(len(tasks)):
        visit(i, [])
    for i in order:
        t = tasks[i]
        t["es"] = max([tasks[p]["ef"] for p in t["pred"]], default=0)
        t["ef"] = t["es"] + t["days"]
    end = max(t["ef"] for t in tasks)
    for i in reversed(order):
        t = tasks[i]
        succ = [s for s in tasks if i in s["pred"]]
        t["lf"] = min([s["ls"] for s in succ], default=end)
        t["ls"] = t["lf"] - t["days"]
        t["slack"] = t["ls"] - t["es"]
        t["critical"] = t["slack"] == 0
    day = start
    while day.weekday() not in workdays:
        day += dt.timedelta(days=1)
    cal = [day]
    while len(cal) <= end + 1:
        day += dt.timedelta(days=1)
        if day.weekday() in workdays:
            cal.append(day)
    for t in tasks:
        if t["days"]:
            t["start"] = dt.datetime.combine(cal[t["es"]], dt.time(8))
            t["finish"] = dt.datetime.combine(cal[t["ef"] - 1], dt.time(17))
        else:  # a milestone sits at the end of what it waits for
            t["start"] = t["finish"] = dt.datetime.combine(cal[t["es"] - 1], dt.time(17)) if t["es"] else dt.datetime.combine(cal[0], dt.time(8))
    return tasks, end


def mspdi(title, tasks, workdays):
    iso = lambda d: d.strftime("%Y-%m-%dT%H:%M:%S")  # noqa: E731
    dur = lambda days: f"PT{days * HOURS}H0M0S"  # noqa: E731
    start, finish = min(t["start"] for t in tasks), max(t["finish"] for t in tasks)
    x = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>', f'<Project xmlns="{NS}">', "<SaveVersion>14</SaveVersion>",
         f"<Name>{_e(title)}.xml</Name>", f"<Title>{_e(title)}</Title>", "<ScheduleFromStart>1</ScheduleFromStart>", f"<StartDate>{iso(start)}</StartDate>",
         f"<FinishDate>{iso(finish)}</FinishDate>", "<CalendarUID>1</CalendarUID>", "<DefaultStartTime>08:00:00</DefaultStartTime>",
         "<DefaultFinishTime>17:00:00</DefaultFinishTime>", f"<MinutesPerDay>{HOURS * 60}</MinutesPerDay>", f"<MinutesPerWeek>{HOURS * 60 * len(workdays)}</MinutesPerWeek>",
         "<DaysPerMonth>20</DaysPerMonth>", "<DurationFormat>7</DurationFormat>", "<WorkFormat>2</WorkFormat>", "<WeekStartDay>1</WeekStartDay>",
         "<Calendars><Calendar><UID>1</UID><Name>Standard</Name><IsBaseCalendar>1</IsBaseCalendar><BaseCalendarUID>-1</BaseCalendarUID><WeekDays>"]
    for daytype in range(1, 8):  # 1 = Sunday
        working = ((daytype - 2) % 7) in workdays
        x.append(f"<WeekDay><DayType>{daytype}</DayType><DayWorking>{int(working)}</DayWorking>" + (
            "<WorkingTimes><WorkingTime><FromTime>08:00:00</FromTime><ToTime>12:00:00</ToTime></WorkingTime>"
            "<WorkingTime><FromTime>13:00:00</FromTime><ToTime>17:00:00</ToTime></WorkingTime></WorkingTimes>" if working else "") + "</WeekDay>")
    x.append("</WeekDays></Calendar></Calendars><Tasks>")
    total = max(t["ef"] for t in tasks)
    x.append(f"<Task><UID>0</UID><ID>0</ID><Name>{_e(title)}</Name><Manual>0</Manual><Type>1</Type><IsNull>0</IsNull><OutlineNumber>0</OutlineNumber>"
             f"<OutlineLevel>0</OutlineLevel><Start>{iso(start)}</Start><Finish>{iso(finish)}</Finish><Duration>{dur(total)}</Duration>"
             "<DurationFormat>7</DurationFormat><Milestone>0</Milestone><Summary>1</Summary></Task>")
    for i, t in enumerate(tasks, 1):
        links = "".join(f"<PredecessorLink><PredecessorUID>{p + 1}</PredecessorUID><Type>1</Type><CrossProject>0</CrossProject><LinkLag>0</LinkLag>"
                        "<LagFormat>7</LagFormat></PredecessorLink>" for p in t["pred"])
        x.append(f"<Task><UID>{i}</UID><ID>{i}</ID><Name>{_e(t['name'])}</Name><Manual>0</Manual><Type>0</Type><IsNull>0</IsNull><WBS>{i}</WBS>"
                 f"<OutlineNumber>{i}</OutlineNumber><OutlineLevel>1</OutlineLevel><Start>{iso(t['start'])}</Start><Finish>{iso(t['finish'])}</Finish>"
                 f"<Duration>{dur(t['days'])}</Duration><DurationFormat>7</DurationFormat><Milestone>{int(t['days'] == 0)}</Milestone><Summary>0</Summary>"
                 f"<Critical>{int(t['critical'])}</Critical><TotalSlack>{t['slack'] * HOURS * 600}</TotalSlack><PercentComplete>0</PercentComplete>"
                 f"<ConstraintType>0</ConstraintType>{links}</Task>")
    x.append("</Tasks><Resources><Resource><UID>0</UID><ID>0</ID><IsNull>0</IsNull></Resource>")
    res = list(dict.fromkeys(t["res"] for t in tasks if t["res"]))
    for i, r in enumerate(res, 1):
        x.append(f"<Resource><UID>{i}</UID><ID>{i}</ID><Name>{_e(r)}</Name><Type>1</Type><IsNull>0</IsNull><MaxUnits>1.00</MaxUnits></Resource>")
    x.append("</Resources><Assignments>")
    for i, t in enumerate([t for t in tasks if t["res"]], 1):
        x.append(f"<Assignment><UID>{i}</UID><TaskUID>{tasks.index(t) + 1}</TaskUID><ResourceUID>{res.index(t['res']) + 1}</ResourceUID>"
                 f"<Finish>{iso(t['finish'])}</Finish><Start>{iso(t['start'])}</Start><Units>1</Units><Work>{dur(t['days'])}</Work></Assignment>")
    x.append("</Assignments></Project>")
    return "\n".join(x)


def _e(s):
    import html
    return html.escape(str(s), quote=False)


def gantt(path, title, tasks, end):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    days = sorted({t["start"].date() for t in tasks} | {t["finish"].date() for t in tasks})
    first, last = days[0], days[-1]
    cal = [first + dt.timedelta(days=k) for k in range((last - first).days + 1)]
    ws.append(["Task", "Days", "Start", "Finish", "Slack", "Waits for", "Resource"] + [d.strftime("%d %b") for d in cal])
    red, blue, grey = PatternFill("solid", fgColor="E53935"), PatternFill("solid", fgColor="1E88E5"), PatternFill("solid", fgColor="EEEEEE")
    for t in tasks:
        ws.append([t["name"], t["days"], t["start"].strftime("%d %b %Y"), t["finish"].strftime("%d %b %Y"), t["slack"],
                   ", ".join(tasks[p]["name"] for p in t["pred"]), t["res"] or ""] + [""] * len(cal))
        row = ws.max_row
        for k, d in enumerate(cal):
            cell = ws.cell(row, 8 + k)
            if t["start"].date() <= d <= t["finish"].date():
                cell.fill = red if t["critical"] else blue
                if not t["days"]:
                    cell.value = "◆"
            elif d.weekday() not in t.get("_workdays", range(7)):
                cell.fill = grey
    ws.column_dimensions["A"].width = 28
    for k in range(len(cal)):
        ws.column_dimensions[ws.cell(2, 8 + k).column_letter].width = 4.5
        ws.cell(2, 8 + k).alignment = Alignment(text_rotation=90)
    ws.append([])
    ws.append(["Red: critical path (any delay moves the end date). Blue: has slack."])
    wb.save(path)


def check(path, title, tasks, workdays):
    root = ET.parse(path).getroot()
    q = lambda e, tag: e.find(f"{{{NS}}}{tag}")  # noqa: E731
    back = []
    for e in root.find(f"{{{NS}}}Tasks"):
        if q(e, "UID").text == "0":
            continue
        hours = int(re.match(r"PT(\d+)H", q(e, "Duration").text).group(1))
        back.append({"name": q(e, "Name").text, "days": hours // HOURS, "res": None,
                     "after": [], "_pred": [int(p.find(f"{{{NS}}}PredecessorUID").text) - 1 for p in e.findall(f"{{{NS}}}PredecessorLink")],
                     "xs": q(e, "Start").text, "xf": q(e, "Finish").text})
    for t in back:
        t["after"] = [back[p]["name"] for p in t["_pred"]]
    first = dt.datetime.fromisoformat(back[0]["xs"]).date() if back else None
    again, _ = schedule([dict(t) for t in back], min(dt.datetime.fromisoformat(t["xs"]).date() for t in back) if back else first, workdays)
    same = all(a["start"].strftime("%Y-%m-%dT%H:%M:%S") == b["xs"] and a["finish"].strftime("%Y-%m-%dT%H:%M:%S") == b["xf"] for a, b in zip(again, back))
    order_ok = all(tasks[p]["finish"] <= t["start"] for t in tasks for p in t["pred"])
    crit = [t for t in tasks if t["critical"]]
    return [("every part in the order Project's XML schema wants", in_schema_order(root)),
            ("Project's XML read back: every task, the same dates when scheduled again", len(back) == len(tasks) and same),
            ("every task starts after the tasks it waits for", order_ok),
            ("the critical path has no slack and ends on the project's last day", bool(crit) and all(t["slack"] == 0 for t in crit)
             and max(t["finish"] for t in crit) == max(t["finish"] for t in tasks))]


def parse(text, ctx):
    from .appschat import find_file
    c = text.lower()
    if not re.search(r"\b(?:ms|microsoft)\s+project\b|\bprojectlibre\b|\bproject\s+(?:plan|schedule)\b|\bgantt\b", c):
        return None
    m = re.search(r"['\"]([^'\"]+)['\"]", text)
    workdays = set(range(6)) if re.search(r"\b(?:6|six)[- ]day\b|\bsaturdays?\s+(?:working|too|included)\b|\bmon(?:day)?\s*-\s*sat", c) else set(range(5))
    head, _, body = text.partition(":")
    f = find_file(text, ctx, {".xlsx"})
    import datetime as _dt
    from ..hub.hubparse import when
    start = when(head.lower(), ctx.get("now") or _dt.datetime.now())[0] if re.search(r"\b(?:start|starting|from|begin)", head.lower()) else None
    if f and not body.strip():
        return {"op": "plan", "title": m.group(1) if m else Path(f).stem.replace("_", " ").title(), "file": f, "body": None, "start": start, "workdays": sorted(workdays)}
    if not body.strip():
        return None
    return {"op": "plan", "title": m.group(1) if m else "Project Plan", "file": None, "body": body, "start": start, "workdays": sorted(workdays)}


def run(op, ctx):
    workdays = set(op["workdays"])
    try:
        tasks = tasks_from_xlsx(op["file"]) if op.get("file") else tasks_from_text(op["body"])
        if not tasks:
            raise PlanError("no tasks found")
        now = (ctx.get("now") or dt.datetime.now()).date()
        start = op.get("start") or now + dt.timedelta(days=(7 - now.weekday()) % 7 or 7)  # next Monday
        tasks, end = schedule(tasks, start, workdays)
    except PlanError as e:
        return f"Couldn't plan it: {e}."
    out = Path(ctx["out"]) / "msproject"
    out.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^\w-]+", "_", op["title"]).strip("_")
    xml = out / f"{stem}.xml"
    xml.write_text(mspdi(op["title"], tasks, workdays), encoding="utf-8")
    for t in tasks:
        t["_workdays"] = workdays
    gantt(out / f"{stem}_gantt.xlsx", op["title"], tasks, end)
    checks = check(xml, op["title"], tasks, workdays)
    bad = [w for w, ok in checks if not ok]
    crit = " -> ".join(t["name"] for t in sorted([t for t in tasks if t["critical"]], key=lambda t: (t["es"], t["ef"])))
    finish = max(t["finish"] for t in tasks)
    slack = [f"{t['name']} {t['slack']} day{'s' if t['slack'] != 1 else ''}" for t in tasks if t["slack"]]
    return (f"Project plan '{op['title']}': {len(tasks)} tasks, {end} working days ({'Mon-Sat' if len(workdays) == 6 else 'Mon-Fri'}), "
            f"{min(t['start'] for t in tasks):%a %d %b %Y} to {finish:%a %d %b %Y}. Critical path: {crit}." +
            (f" Slack: {', '.join(slack)}." if slack else "") +
            f" Files: {xml} (open in Microsoft Project or ProjectLibre: File > Open) and {stem}_gantt.xlsx. " +
            ("Checked: " + "; ".join(w for w, _ in checks) + "." if not bad else "NOT right: " + "; ".join(bad) + "."))
