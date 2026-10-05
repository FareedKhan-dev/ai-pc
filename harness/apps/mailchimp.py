"""Email campaigns with Mailchimp (Marketing API 3.0) or Brevo (API v3): audiences listed, contacts added from an Excel or
CSV list (only people who agreed to hear from you), a campaign written and sent to an audience after a yes (shown
first with how many people it reaches), and the last campaign's report (sent, opened, clicked).

  'mailchimp lists'   'add contacts from customers.xlsx to mailchimp list Customers'
  "email campaign on mailchimp to Customers: subject 'Eid Sale' from Khan Electronics, body: 20% off everything till Sunday."   'mailchimp report'
  (the same with 'brevo')
"""
import base64
import hashlib
import html
import re
from pathlib import Path

from ..hub import vault
from ..hub.http import Api, HubError

NAME, LABEL = "mailchimp", "Email campaigns: Mailchimp and Brevo (audiences, contacts, send, reports)"
EXAMPLES = ["mailchimp lists", "add contacts from customers.xlsx to mailchimp list Customers",
            "email campaign on brevo to Customers: subject 'Eid Sale' from Khan Electronics, body: 20% off till Sunday."]
OUTWARD = {"add", "send"}
APP = {"label": "Mailchimp",
       "fields": [("key", "Mailchimp API key (ends in -us21 or similar)", True), ("reply_to", "Your reply-to email address", False)],
       "steps": ["Mailchimp: Profile > Extras > API keys > Create A Key; copy it (shown once). Run 'apps.py connect mailchimp'.",
                 "Brevo instead: SMTP & API > API Keys > Generate; run 'apps.py connect brevo'. A verified sender email is needed in Brevo (Senders & IP)."],
       "notes": "Send only to people who agreed to get your emails; both services stop accounts that send to bought or scraped lists."}
BREVO_APP = {"label": "Brevo", "fields": [("key", "Brevo API key", True), ("sender", "Your verified sender email", False)],
             "steps": APP["steps"][1:], "notes": APP["notes"]}


def people(path):
    """[{email, first, last}] from an Excel or CSV list (columns found by their headings)."""
    p = Path(path)
    if p.suffix.lower() == ".csv":
        import csv
        with p.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
    else:
        from openpyxl import load_workbook
        rows = [list(r) for r in load_workbook(p, data_only=True, read_only=True).active.iter_rows(values_only=True)]
    head = [str(h or "").lower() for h in rows[0]]
    ei = next((i for i, h in enumerate(head) if "mail" in h), None)
    ni = next((i for i, h in enumerate(head) if "name" in h), None)
    out = []
    for r in rows[1:]:
        e = str(r[ei] or "").strip() if ei is not None else ""
        if re.fullmatch(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", e):
            first, _, last = str(r[ni] or "").strip().partition(" ") if ni is not None else ("", "", "")
            out.append({"email": e.lower(), "first": first, "last": last})
    return out


class Mailchimp:
    label = "Mailchimp"

    def __init__(self, creds, transport=None):
        dc = creds["key"].rsplit("-", 1)[-1]
        auth = base64.b64encode(f"aipc:{creds['key']}".encode()).decode()
        self.c, self.api = creds, Api(f"https://{dc}.api.mailchimp.com/3.0", headers={"Authorization": f"Basic {auth}"}, service="mailchimp", transport=transport)

    def call(self, method, path, body=None, **kw):
        try:
            return self.api.request(method, path, json_body=body, retries=0 if method == "POST" else 3, **kw)
        except HubError as e:
            b = e.body if isinstance(e.body, dict) else {}
            raise RuntimeError(f"Mailchimp: {b.get('detail') or b.get('title') or e}") from e

    def lists(self):
        return [{"id": x["id"], "name": x["name"], "count": (x.get("stats") or {}).get("member_count", 0)}
                for x in self.call("GET", "lists", params={"count": 50}).get("lists") or []]

    def add(self, list_id, ps):
        for p in ps:
            h = hashlib.md5(p["email"].lower().encode()).hexdigest()  # noqa: S324 - Mailchimp's own id for a member
            self.call("PUT", f"lists/{list_id}/members/{h}", {"email_address": p["email"], "status_if_new": "subscribed",
                                                              "merge_fields": {"FNAME": p["first"], "LNAME": p["last"]}})
        return len(ps)

    def send(self, list_id, subject, from_name, body_html):
        cid = self.call("POST", "campaigns", {"type": "regular", "recipients": {"list_id": list_id}, "settings": {
            "subject_line": subject, "from_name": from_name, "reply_to": self.c.get("reply_to") or "", "title": subject}})["id"]
        self.call("PUT", f"campaigns/{cid}/content", {"html": body_html})
        ready = self.call("GET", f"campaigns/{cid}/send-checklist")
        if not ready.get("is_ready"):
            raise RuntimeError("Mailchimp says the campaign is not ready: " + "; ".join(i.get("details", "") for i in ready.get("items") or [] if i.get("type") == "error"))
        self.call("POST", f"campaigns/{cid}/actions/send")
        return cid

    def report(self, cid):
        r = self.call("GET", f"reports/{cid}")
        return {"sent": r.get("emails_sent", 0), "opened": (r.get("opens") or {}).get("unique_opens", 0), "clicked": (r.get("clicks") or {}).get("unique_clicks", 0)}


class Brevo:
    label = "Brevo"

    def __init__(self, creds, transport=None):
        self.c, self.api = creds, Api("https://api.brevo.com/v3", headers={"api-key": creds["key"]}, service="brevo", transport=transport)

    def call(self, method, path, body=None, **kw):
        try:
            return self.api.request(method, path, json_body=body, retries=0 if method == "POST" else 3, **kw)
        except HubError as e:
            b = e.body if isinstance(e.body, dict) else {}
            raise RuntimeError(f"Brevo: {b.get('message') or e}") from e

    def lists(self):
        return [{"id": x["id"], "name": x["name"], "count": x.get("totalSubscribers", x.get("uniqueSubscribers", 0))}
                for x in self.call("GET", "contacts/lists", params={"limit": 50}).get("lists") or []]

    def add(self, list_id, ps):
        for p in ps:
            self.call("POST", "contacts", {"email": p["email"], "attributes": {"FIRSTNAME": p["first"], "LASTNAME": p["last"]}, "listIds": [int(list_id)],
                                           "updateEnabled": True})
        return len(ps)

    def send(self, list_id, subject, from_name, body_html):
        cid = self.call("POST", "emailCampaigns", {"name": subject, "subject": subject, "sender": {"name": from_name, "email": self.c.get("sender") or ""},
                                                   "htmlContent": body_html, "recipients": {"listIds": [int(list_id)]}})["id"]
        self.call("POST", f"emailCampaigns/{cid}/sendNow")
        return cid

    def report(self, cid):
        s = ((self.call("GET", f"emailCampaigns/{cid}", params={"statistics": "globalStats"}).get("statistics") or {}).get("globalStats") or {})
        return {"sent": s.get("sent", 0), "opened": s.get("uniqueViews", s.get("uniqueOpens", 0)), "clicked": s.get("uniqueClicks", 0)}


def client(ctx, which):
    inj = (ctx.get("clients") or {}).get(which)
    if inj:
        return inj
    c = vault.get(which)
    if not c:
        raise RuntimeError(f"{which.title()} is not connected: 'apps.py steps mailchimp' shows how")
    return (Mailchimp if which == "mailchimp" else Brevo)(c)


def connect(values, transport=None, store=None, which="mailchimp"):
    c = (Mailchimp if which == "mailchimp" else Brevo)(dict(values), transport)
    n = len(c.lists())
    (store or (lambda v: vault.put(which, v)))(dict(values))
    return {"who": f"{n} audience(s)", "where": which.title()}


def parse(text, ctx):
    from .appschat import find_file
    c = text.lower()
    which = "brevo" if re.search(r"\bbrevo\b|\bsendinblue\b", c) else "mailchimp" if re.search(r"\bmailchimp\b", c) else None
    if not which:
        return None
    if re.search(r"\b(?:lists?|audiences?)\b", c) and not re.search(r"\badd\b|\bcampaign\b", c):
        return {"op": "lists", "which": which}
    m = re.search(r"\badd\s+contacts?\s+from\s+(\S+)\s+to\s+\S+\s+(?:list|audience)\s+(.+?)\s*$", text, re.I)
    if m:
        f = find_file(m.group(1), ctx, {".xlsx", ".csv"})
        return {"op": "add", "which": which, "file": f, "list": m.group(2).strip(" '\"")} if f else None
    m = re.search(r"\bcampaign\s+on\s+\S+\s+to\s+(.+?):\s*subject\s+(?:'([^']+)'|\"([^\"]+)\")\s+from\s+(.+?),\s*body\s*:\s*(.+)$", text, re.I | re.S)
    if m:
        return {"op": "send", "which": which, "list": m.group(1).strip(" '\""), "subject": m.group(2) or m.group(3), "from": m.group(4).strip(), "body": m.group(5).strip()}
    if re.search(r"\breport\b|\bhow did\b|\bopens?\b", c):
        return {"op": "report", "which": which}
    return None


def _list(cl, name):
    hit = next((x for x in cl.lists() if x["name"].lower() == name.lower()), None)
    if not hit:
        raise RuntimeError(f"no audience called '{name}' (say '{cl.label.lower()} lists')")
    return hit


def preview(op, ctx):
    cl = client(ctx, op["which"])
    lst = _list(cl, op["list"])
    op["list_id"] = lst["id"]
    if op["op"] == "add":
        n = len(people(op["file"]))
        return f"Ready to add {n} contact(s) from {Path(op['file']).name} to {cl.label} audience '{lst['name']}' (only people who agreed to hear from you)."
    return f"Ready to SEND '{op['subject']}' from {op['from']} on {cl.label} to '{lst['name']}': {lst['count']} people get it at once."


def run(op, ctx):
    cl = client(ctx, op["which"])
    if op["op"] == "lists":
        ls = cl.lists()
        return f"{cl.label} audiences: " + "; ".join(f"{x['name']} ({x['count']} people)" for x in ls) if ls else f"No audiences on {cl.label} yet."
    if op["op"] == "report":
        cid = ctx["memo"].get(f"{op['which']}_campaign") or (vault.get(f"{op['which']}_last") or {}).get("id")
        if not cid:
            return "No campaign sent from here yet."
        r = cl.report(cid)
        return f"Campaign {cid}: sent to {r['sent']}, opened by {r['opened']}, clicked by {r['clicked']}."
    if not op.get("confirmed"):
        return preview(op, ctx)
    if op["op"] == "add":
        n = cl.add(op["list_id"], people(op["file"]))
        after = _list(cl, op["list"])
        return f"{n} contact(s) added or updated on {cl.label}; the audience now has {after['count']} people."
    body = "".join(f"<p>{html.escape(p.strip())}</p>" for p in re.split(r"\n\s*\n", op["body"]) if p.strip())
    cid = cl.send(op["list_id"], op["subject"], op["from"], body)
    if ctx.get("clients") is None:
        vault.put(f"{op['which']}_last", {"id": cid})
    ctx["memo"][f"{op['which']}_campaign"] = cid
    return f"Sent on {cl.label} (campaign {cid}). Say '{op['which']} report' later for opens and clicks."
