"""Anki (the flashcard app students everywhere use) by deck files: cards from words ('France = Paris; Japan = Tokyo'), from
a notes file (lines like 'term - meaning', 'Q: ... A: ...'), or from a CSV / Excel sheet's first two columns; cloze cards
from sentences with the hidden part in {braces}; 'both ways' adds the reverse card. Written as an .apkg with genanki
(double-click it and Anki imports the deck; the same cards made again update, not duplicate). Checked by opening the
package's own database: the deck is there with every note, the right number of cards, and the first card's text.

  "anki deck 'Capitals': France = Paris; Japan = Tokyo; Pakistan = Islamabad, both ways"   'anki deck from biology.xlsx'
  'anki cloze: The {mitochondria} is the powerhouse of the cell; Water boils at {100} degrees'
"""
import html
import json
import re
import sqlite3
import tempfile
import zipfile
import zlib
from pathlib import Path

NAME, LABEL = "anki", "Anki: flashcard decks (.apkg) from lists, notes, CSV or Excel"
EXAMPLES = ["anki deck 'Capitals': France = Paris; Japan = Tokyo; Pakistan = Islamabad, both ways", "anki deck from biology.xlsx",
            "anki cloze: The {mitochondria} is the powerhouse of the cell"]
SOURCES = {".txt", ".md", ".csv", ".tsv", ".xlsx"}
SEPS = r"\s*::\s*|\s+=\s+|\s+[-–—]\s+|\t|:\s+"
HEADER = re.compile(r"^(?:front|question|term|word|q)$", re.I)


def pairs_from_text(text):
    """Cards from lines or ';'-separated items: 'front = back', 'front - back', 'front :: back', or Q:/A: blocks."""
    qa = re.findall(r"(?:^|\n)\s*Q[:.)]\s*(.+?)\s*\n\s*A[:.)]\s*(.+?)(?=\n\s*Q[:.)]|\Z)", text, re.S | re.I)
    if qa:
        return [(q.strip(), a.strip()) for q, a in qa]
    out = []
    for item in re.split(r"\s*[;\n]\s*", text):
        parts = re.split(SEPS, item.strip(), maxsplit=1)
        if len(parts) == 2 and parts[0] and parts[1]:
            out.append((parts[0].strip(), parts[1].strip().rstrip(".")))
    return out


def pairs_from_file(path):
    p = Path(path)
    if p.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        rows = [[("" if c is None else str(c)).strip() for c in r[:2]] for r in wb.active.iter_rows(values_only=True)]
        wb.close()
    elif p.suffix.lower() in (".csv", ".tsv"):
        import csv
        text = p.read_text(encoding="utf-8-sig")
        rows = [[c.strip() for c in r[:2]] for r in csv.reader(text.splitlines(), delimiter="\t" if p.suffix.lower() == ".tsv" or "\t" in text[:500] else ",")]
    else:
        return pairs_from_text(p.read_text(encoding="utf-8-sig"))
    rows = [r for r in rows if len(r) == 2 and r[0] and r[1]]
    if rows and HEADER.match(rows[0][0]):
        rows = rows[1:]
    return [tuple(r) for r in rows]


def clozes(text):
    """'The {mitochondria} is ...' -> Anki cloze text 'The {{c1::mitochondria}} is ...' (each brace its own card)."""
    out = []
    for item in re.split(r"\s*[;\n]\s*", text):
        n = [0]

        def sub(m):
            n[0] += 1
            return "{{c%d::%s}}" % (n[0], html.escape(m.group(1)))
        body = re.sub(r"\{\{?(?:c\d+::)?([^{}]+)\}?\}", sub, html.escape(item.strip(), quote=False).replace("&#x27;", "'"))
        if n[0]:
            out.append(body)
    return out


def _id(name, salt):
    return 1 << 30 | zlib.crc32(f"{salt}:{name}".encode()) & 0x3FFFFFFF  # the same deck name gives the same id, so Anki updates it


def write(path, deck_name, pairs=(), cloze=(), both=False):
    import genanki
    deck = genanki.Deck(_id(deck_name, "deck"), deck_name)
    model = genanki.BASIC_AND_REVERSED_CARD_MODEL if both else genanki.BASIC_MODEL
    for f, b in pairs:
        deck.add_note(genanki.Note(model=model, fields=[html.escape(f), html.escape(b)]))
    for t in cloze:
        deck.add_note(genanki.Note(model=genanki.CLOZE_MODEL, fields=[t, ""]))
    genanki.Package(deck).write_to_file(str(path))
    return deck


def check(path, deck_name, notes, cards, first):
    with zipfile.ZipFile(path) as z, tempfile.TemporaryDirectory() as tmp:
        names = z.namelist()
        z.extract("collection.anki2", tmp)
        db = sqlite3.connect(Path(tmp) / "collection.anki2")
        try:
            decks = json.loads(db.execute("select decks from col").fetchone()[0])
            did = next((int(k) for k, d in decks.items() if d["name"] == deck_name), None)
            n_notes = db.execute("select count(*) from notes").fetchone()[0]
            n_cards = db.execute("select count(*) from cards where did = ?", (did,)).fetchone()[0]
            flds = db.execute("select flds from notes order by id limit 1").fetchone()[0].split("\x1f")[0]
        finally:
            db.close()
    return [("the package holds Anki's collection and media list", "collection.anki2" in names and "media" in names),
            (f"the deck '{deck_name}' is in it with all {notes} notes", did is not None and n_notes == notes),
            (f"{cards} cards to study", n_cards == cards),
            ("the first card reads as given", html.unescape(flds) == html.unescape(first))]


def parse(text, ctx):
    from .appschat import find_file
    c = text.lower()
    if not re.search(r"\banki\b", c):
        return None
    m = re.search(r"\b(?:deck|called|named)\s+['\"]([^'\"]+)['\"]", text, re.I)
    both = bool(re.search(r"\b(?:both\s+ways|both\s+directions|reversed?|two[- ]way)\b", c))
    body = text.split(":", 1)[1] if ":" in text else ""
    body = re.sub(r",?\s*\b(?:both\s+ways|both\s+directions|reversed?|two[- ]way)\b\.?\s*$", "", body, flags=re.I)
    src = find_file(text, ctx, SOURCES)
    if re.search(r"\bcloze\b", c) or re.search(r"\{[^{}]+\}", body):
        items = clozes(body if body.strip() else Path(src).read_text(encoding="utf-8-sig") if src else "")
        return {"op": "deck", "name": m.group(1) if m else "Cloze Cards", "pairs": [], "cloze": items, "both": False} if items else None
    if src and not body.strip():
        return {"op": "deck", "name": m.group(1) if m else Path(src).stem.replace("_", " ").title(), "file": src, "pairs": None, "cloze": [], "both": both}
    pairs = pairs_from_text(body)
    return {"op": "deck", "name": m.group(1) if m else "AI PC Cards", "pairs": pairs, "cloze": [], "both": both} if pairs else None


def run(op, ctx):
    pairs = op["pairs"] if op.get("pairs") is not None else pairs_from_file(op["file"])
    if not pairs and not op["cloze"]:
        return "No cards found: give lines like 'term - meaning', 'Q: ... A: ...', or a sheet with the front in column A and the back in column B."
    out = Path(ctx["out"]) / "anki"
    out.mkdir(parents=True, exist_ok=True)
    path = out / (re.sub(r"[^\w-]+", "_", op["name"]).strip("_") + ".apkg")
    write(path, op["name"], pairs, op["cloze"], op["both"])
    notes = len(pairs) + len(op["cloze"])
    cards = len(pairs) * (2 if op["both"] else 1) + sum(len(set(re.findall(r"\{\{c(\d+)::", t))) for t in op["cloze"])
    first = pairs[0][0] if pairs else op["cloze"][0]
    checks = check(path, op["name"], notes, cards, html.escape(first) if pairs else first)
    bad = [w for w, ok in checks if not ok]
    sample = "; ".join(f"{f} -> {b}" for f, b in pairs[:3]) if pairs else "; ".join(re.sub(r"\{\{c\d+::([^}]+)\}\}", r"[\1]", html.unescape(t)) for t in op["cloze"][:2])
    return (f"Anki deck '{op['name']}': {path} ({notes} notes, {cards} cards{', both ways' if op['both'] else ''}; e.g. {sample}). "
            "Double-click it (or Anki: File > Import) to add it to Anki; making it again updates the same cards. " +
            ("Checked: " + "; ".join(w for w, _ in checks) + "." if not bad else "NOT right: " + "; ".join(bad) + "."))
