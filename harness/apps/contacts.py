"""Contact cards: people and businesses as vCards (.vcf, version 3.0) that phones, Gmail and Outlook import, from words
or from an Excel/CSV list (name, phone, email, company columns found by their headings); Pakistani mobile numbers
written in the +92 form; a QR code of a card for a visiting card or shop counter. Checks: the file reads back with
every contact.

  'contact Ali Raza, 0300-1234567, ali@x.com, Ali Traders'   'contacts from customers.xlsx'   'qr of my contact card'
"""
import re
from pathlib import Path

NAME, LABEL = "contacts", "Contact cards (vCard .vcf) from words or Excel, with QR codes"
EXAMPLES = ["contact Ali Raza, 0300-1234567, ali@x.com, Ali Traders", "contacts from customers.xlsx"]


def phone(s):
    d = re.sub(r"[^\d+]", "", str(s or ""))
    if re.fullmatch(r"03\d{9}", d):
        return "+92" + d[1:]
    if re.fullmatch(r"923\d{9}", d):
        return "+" + d
    return d or None


def card(c):
    def esc(s):
        return str(s).replace("\\", "\\\\").replace(",", "\\,").replace(";", "\\;")
    first, _, last = c["name"].partition(" ")
    lines = ["BEGIN:VCARD", "VERSION:3.0", f"N:{esc(last)};{esc(first)};;;", f"FN:{esc(c['name'])}"]
    if c.get("company"):
        lines.append(f"ORG:{esc(c['company'])}")
    if c.get("phone"):
        lines.append(f"TEL;TYPE=CELL:{c['phone']}")
    if c.get("email"):
        lines.append(f"EMAIL;TYPE=INTERNET:{c['email']}")
    if c.get("address"):
        lines.append(f"ADR;TYPE=WORK:;;{esc(c['address'])};;;;Pakistan")
    return "\r\n".join(lines + ["END:VCARD"]) + "\r\n"


def from_words(s):
    parts = [p.strip() for p in s.split(",") if p.strip()]
    c = {"name": parts[0] if parts else ""}
    for p in parts[1:]:
        if re.fullmatch(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", p):
            c["email"] = p
        elif re.fullmatch(r"[+\d][\d\s-]{7,16}", p):
            c["phone"] = phone(p)
        elif not c.get("company"):
            c["company"] = p
        else:
            c["address"] = p
    return c


def from_sheet(path):
    p = Path(path)
    if p.suffix.lower() == ".csv":
        import csv
        with p.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
    else:
        from openpyxl import load_workbook
        rows = [list(r) for r in load_workbook(p, data_only=True, read_only=True).active.iter_rows(values_only=True)]
    head = [str(h or "").lower() for h in rows[0]]

    def col(*words):
        return next((i for i, h in enumerate(head) if any(w in h for w in words)), None)
    ni, pi, ei, ci, ai = col("name"), col("phone", "mobile", "cell", "whatsapp", "number"), col("email", "e-mail"), col("company", "business", "shop"), col("address", "city")
    out = []
    for r in rows[1:]:
        if ni is None or not r[ni]:
            continue
        out.append({"name": str(r[ni]).strip(), "phone": phone(r[pi]) if pi is not None else None, "email": str(r[ei]).strip() if ei is not None and r[ei] else None,
                    "company": str(r[ci]).strip() if ci is not None and r[ci] else None, "address": str(r[ai]).strip() if ai is not None and r[ai] else None})
    return out


def parse(text, ctx):
    from .appschat import find_file
    m = re.match(r"^\s*(?:add\s+)?(?:a\s+)?contact(?:\s+card)?\s*:?\s+(.+)$", text, re.I)
    if m and not re.match(r"^\s*(?:cards?\s+)?from\b", m.group(1), re.I):
        return {"op": "card", "contacts": [from_words(m.group(1))], "qr": bool(re.search(r"\bqr\b", text, re.I))}
    if re.search(r"\b(?:mailchimp|brevo|sendinblue|audience|campaign)\b", text, re.I):
        return None  # contacts for an email service are that service's
    if re.search(r"\bcontacts?\b.*\bfrom\b|\bvcards?\b|\.vcf\b", text, re.I):
        f = find_file(text, ctx, {".xlsx", ".csv"})
        return {"op": "sheet", "file": f} if f else None
    return None


def run(op, ctx):
    out = Path(ctx["out"]) / "contacts"
    out.mkdir(parents=True, exist_ok=True)
    cs = op["contacts"] if op["op"] == "card" else from_sheet(op["file"])
    if not cs or not cs[0].get("name"):
        return "No contacts found (a 'Name' column, or 'contact <name>, <phone>, <email>')."
    dest = out / (re.sub(r"[^\w-]+", "_", cs[0]["name"]) + ".vcf" if len(cs) == 1 else "contacts.vcf")
    dest.write_text("".join(card(c) for c in cs), encoding="utf-8", newline="")
    back = dest.read_text(encoding="utf-8")
    ok = back.count("BEGIN:VCARD") == len(cs) and all(c["name"] in back for c in cs)
    reply = f"{len(cs)} contact card(s): {dest} ({'checked' if ok else 'NOT right'})."
    if op.get("qr") and len(cs) == 1:
        import segno
        q = out / (dest.stem + "_qr.png")
        segno.make(card(cs[0]), error="m").save(str(q), scale=8, border=2)
        reply += f" QR code of the card: {q}."
    return reply + " Phones import .vcf files from Contacts > Import; Gmail from Google Contacts > Import."
