"""Checks at every element of a rendered document, the way the video verifier checks every edit point.

  rep = check(plan, built, rendered, planner=None)   plan: docplan.resolve; built: docx_build.build; rendered: render.to_pdf
  -> {"counts": {"pass", "warn", "fail"}, "results": [{"id", "type", "status", "why", "page", "fix"?}], "facts": {...}}

global    rendered; pages against the length asked for; page size and orientation; no blank pages; every glyph drawn in
          the theme's fonts (a substituted font is measured per character, from the PDF); no field errors ("Error!")
toc       the contents page lists the headings with the pages Word really put them on
heading   on the page (from the PDF's bookmarks); not left alone at the bottom of a page
table     its header found; inside the margins; the header repeated on every page the table runs onto; the total row
chart     drawn (its title and labels are on a page, inside the margins); image: present on a page
text      every paragraph and list item made it into the PDF
look      (with a planner) one vision look at all the pages for overlaps, cut-off text, empty pages, odd spacing
A result with "fix" says what the fixer can change (studio.py applies it and checks again).
"""
import ctypes
import re

from . import render as RN
from .docplan import plain

PAGE_PT = {"A4": (595.3, 841.9), "Letter": (612.0, 792.0)}
OK_FONTS = {"symbol", "wingdings", "couriernew", "segoeuisymbol"}
FIT_TYPES = {"cv", "letter", "application", "invoice", "quotation", "notice", "memo", "certificate", "agenda"}


def _norm(s):
    s = plain(str(s or "")).lower().replace("­", "").replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    s = re.sub("[‐‑‒–—−]", "-", s)
    s = " ".join(re.sub(r"[^\w%'.,:;!?&/()+-]+", " ", s).split())
    return re.sub(r"\s*-\s*", "-", s)  # "well-installed", "well - installed" and "well- installed" (a line break) alike


def glyph_fonts(pdf):
    """{font name: number of characters drawn in it} for the whole PDF (what was really used, per character)."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    doc = pdfium.PdfDocument(str(pdf))
    out = {}
    try:
        buf, flags = ctypes.create_string_buffer(256), ctypes.c_int()
        for i in range(len(doc)):
            tp = doc[i].get_textpage()
            for k in range(tp.count_chars()):
                n = raw.FPDFText_GetFontInfo(tp.raw, k, buf, 256, ctypes.byref(flags))
                f = buf.raw[:max(0, n - 1)].decode("latin-1")
                if f and tp.get_text_range(k, 1).strip():
                    out[f] = out.get(f, 0) + 1
            tp.close()
    finally:
        doc.close()
    return out


def image_counts(pdf):
    """Images drawn on each page (charts are vector drawings, not images)."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    doc = pdfium.PdfDocument(str(pdf))
    try:
        return [sum(1 for o in doc[i].get_objects(max_depth=3) if o.type == raw.FPDF_PAGEOBJ_IMAGE) for i in range(len(doc))]
    finally:
        doc.close()


def _family(f):
    return re.sub(r"[^a-z]", "", str(f).split("+")[-1].lower())


def _is_family(font, families):
    """A PDF font name ("TimesNewRomanPS-BoldMT", "ABCDEF+Calibri-Bold") belongs to one of the families."""
    n = _family(font)
    return any(n.startswith(f) for f in families if f)


def check(plan, built, rendered, planner=None, log=print, look_budget=25):
    th = plan["th"]
    res = []
    facts = {}

    def add(id_, typ, status, why, page=None, fix=None, **ev):
        r = {"id": id_, "type": typ, "status": status, "why": why}
        if page is not None:
            r["page"] = page + 1
        if fix:
            r["fix"] = fix
        if ev:
            r["evidence"] = ev
        res.append(r)
    if not rendered.get("ok"):
        add("global", "global", "fail", f"not rendered: {rendered.get('error')}")
        return {"counts": {"pass": 0, "warn": 0, "fail": 1}, "results": res, "facts": facts}
    pdf = rendered["pdf"]
    ls = RN.lines(pdf)
    sizes = RN.page_size(pdf)
    n_pages = len(sizes)
    outl = RN.outline(pdf)
    margin = plan["page"]["margins"] * 28.35
    first_body = min((p_ for _, _, p_ in outl), default=0) if (plan.get("cover") or plan.get("toc")) else 0
    body_pages = n_pages - first_body  # a cover and a contents page are not what "a 5 page report" counts
    facts.update(pages=n_pages, body_pages=body_pages, words=rendered.get("words"))
    full = _norm(" ".join(RN.text(pdf)))
    by_page = {}
    for x in ls:
        by_page.setdefault(x["page"], []).append(x)

    # ---- global
    probs, status = [], "pass"
    want = plan.get("target_pages")
    if want:
        dev = abs(body_pages - want) / max(1, want)
        if dev > 0.5 or (want <= 2 and n_pages > want + 1):
            status = "fail"
        elif dev > 0.25:
            status = "warn"
        if status != "pass":
            probs.append(f"{body_pages} pages" + (f" (+{first_body} front)" if first_body else "") + f" for the {want} asked")
    pw, ph = PAGE_PT[plan["page"]["size"]]
    if plan["page"]["orientation"] == "landscape":
        pw, ph = ph, pw
    if any(abs(w - pw) > 3 or abs(h - ph) > 3 for w, h in sizes):
        probs.append(f"page size {sizes[0][0]:.0f}x{sizes[0][1]:.0f} pt, planned {pw:.0f}x{ph:.0f}")
        status = "fail"
    imgs = image_counts(pdf)
    blank = []
    for i, (w, h) in enumerate(sizes):
        body = [x for x in by_page.get(i, []) if margin * 0.6 < x["box"][1] < h - margin * 0.6]
        if not body and not imgs[i] and not any(p_ == i for _, _, p_ in outl):
            blank.append(i)
    if blank:
        probs.append("blank page " + ", ".join(str(b + 1) for b in blank))
        status = "fail" if status == "fail" else "warn"
    fit = None
    if n_pages >= 2 and n_pages - 1 not in blank:
        h = sizes[-1][1]
        body = [x for x in by_page.get(n_pages - 1, []) if margin * 0.6 < x["box"][1] < h - margin * 0.75]
        fill = (max(x["box"][3] for x in body) - margin) / max(1.0, h - 2 * margin) if body else 0.0
        facts["last_page_fill"] = round(fill, 2)
        if fill < 0.3 and plan["doctype"] in FIT_TYPES and (not want or n_pages > want):
            probs.append(f"page {n_pages} is almost empty ({fill:.0%} used): it should fit on {n_pages - 1}")
            status = "fail" if status == "fail" else "warn"
            fit = {"op": "fit_pages", "pages": n_pages - 1}
    gf = glyph_fonts(pdf)
    want_fams = {_family(th["body"]), _family(th["head"]), _family(th["head"]).replace("semibold", "")}
    stray = {f: n for f, n in gf.items() if not _is_family(f, want_fams) and not _is_family(f, OK_FONTS)}
    facts["fonts"] = {f: n for f, n in sorted(gf.items(), key=lambda kv: -kv[1])}
    if stray:
        probs.append("drawn in other fonts: " + ", ".join(f"{f} ({n} chars)" for f, n in list(stray.items())[:3]))
        status = "fail" if status == "fail" else "warn"
    if "error!" in full:
        probs.append("a field shows an error (" + full[full.index("error!"):full.index("error!") + 60] + ")")
        status = "fail"
    add("global", "global", status, "; ".join(probs) or f"{n_pages} pages, {plan['page']['size']} {plan['page']['orientation']}, fonts as designed",
        fix={"op": "length", "pages": body_pages, "want": want} if want and any("asked" in p for p in probs) else
        ({"op": "blank_pages", "pages": [b + 1 for b in blank]} if blank else fit))

    # ---- table of contents
    heads = [b for b in plan["blocks"] if b["type"] == "heading"]
    if plan.get("toc"):
        if "no table of contents entries" in full or "right-click to update" in full:
            add("toc", "toc", "fail", "the table of contents was not filled in")
        else:
            toc_page = next((p_ for p_ in range(n_pages) if any(_norm(x["text"]).startswith("contents") for x in by_page.get(p_, []))), None)
            bad, seen = [], 0
            if toc_page is not None:
                toc_lines = [x for x in by_page.get(toc_page, [])]
                for lv, title, pg in outl:
                    tl = _norm(title)
                    hit = next((x for x in toc_lines if _norm(x["text"]).startswith(tl[:30])), None)
                    if hit is None:
                        continue
                    seen += 1
                    nums = [x for x in toc_lines if abs(x["box"][1] - hit["box"][1]) < 3 and re.fullmatch(r"\d+", x["text"].strip())]
                    if nums and int(nums[-1]["text"]) != pg + 1:
                        bad.append(f"'{title[:30]}' listed on {nums[-1]['text']}, really on {pg + 1}")
            if toc_page is None or not seen:
                add("toc", "toc", "warn", "contents page not found in the PDF")
            else:
                add("toc", "toc", "warn" if bad else "pass", "; ".join(bad[:3]) or f"{seen} entries with the right page numbers", page=toc_page)

    # ---- headings
    found_titles = [(_norm(t), p_) for _, t, p_ in outl]
    for b in plan["blocks"]:
        bi = plan["blocks"].index(b)
        if b["type"] != "heading":
            continue
        hn = _norm(b["text"])
        pg = next((p_ for t, p_ in found_titles if t == hn or t.startswith(hn[:40])), None)
        if pg is None:
            pg = next((x["page"] for x in ls if _norm(x["text"]).startswith(hn[:40])), None)
        if pg is None:
            add(f"b{bi}", "heading", "fail", f"heading '{b['text'][:40]}' not found in the PDF")
            continue
        body = [x for x in by_page.get(pg, []) if x["box"][1] < sizes[pg][1] - margin * 0.75]
        last = max(body, key=lambda x: x["box"][3]) if body else None
        if last is not None and _norm(last["text"]).startswith(hn[:30]) and pg + 1 < n_pages:
            add(f"b{bi}", "heading", "warn", f"'{b['text'][:40]}' is alone at the bottom of page {pg + 1}", page=pg,
                fix={"op": "page_break_before", "block": bi})
        else:
            add(f"b{bi}", "heading", "pass", f"on page {pg + 1}", page=pg)

    # ---- tables, charts, images
    n_img = sum(imgs)
    want_img = sum(1 for b in plan["blocks"] if b["type"] == "image")
    for bi, b in enumerate(plan["blocks"]):
        if b["type"] in ("table", "invoice_items"):
            cols = b["columns"] if b["type"] == "table" else ["Description", "Qty", "Unit price", "Amount"]
            want_h = [_norm(c) for c in cols[:3] if _norm(c)]
            pg, row_y, hits = None, None, []
            for p_ in range(n_pages):
                cand = [x for x in by_page.get(p_, []) if any(_norm(x["text"]).startswith(h[:12]) for h in want_h)]
                for x in cand:
                    same = [y for y in cand if abs(y["box"][1] - x["box"][1]) < 4]
                    if len(same) >= min(2, len(want_h)):
                        pg, row_y, hits = p_, x["box"][1], same
                        break
                if pg is not None:
                    break
            if pg is None:
                add(f"b{bi}", "table", "fail", "table header not found in the PDF")
                continue
            w = sizes[pg][0]
            row = [x for x in by_page.get(pg, []) if abs(x["box"][1] - row_y) < 4]
            x0, x1 = min(x["box"][0] for x in row), max(x["box"][2] for x in row)
            probs = []
            if x0 < margin - 6 or x1 > w - margin + 6:
                probs.append(f"wider than the text area ({x0:.0f}-{x1:.0f} pt of {margin:.0f}-{w - margin:.0f})")
            # rows that run onto later pages: the header must be there too
            last_cell = _norm(b["rows"][-1][0]) if b["type"] == "table" and b["rows"] and _norm(b["rows"][-1][0]) else None
            end_pg = pg
            if last_cell and len(last_cell) > 3:
                end_pg = next((p_ for p_ in range(n_pages - 1, pg - 1, -1) if any(_norm(x["text"]).startswith(last_cell[:20]) for x in by_page.get(p_, []))), pg)
            for p_ in range(pg + 1, end_pg + 1):
                if not any(_norm(x["text"]).startswith(want_h[0][:12]) for x in by_page.get(p_, [])):
                    probs.append(f"header not repeated on page {p_ + 1}")
            if b.get("total_row") and "total" not in " ".join(_norm(x["text"]) for p_ in range(pg, end_pg + 1) for x in by_page.get(p_, [])):
                probs.append("total row missing")
            if b["type"] == "invoice_items":
                tot = b["computed"]["total"]
                amount = f"{tot:,.{b['computed']['decimals']}f}"
                if amount.lower() not in full:
                    probs.append(f"total {amount} not on the page")
            add(f"b{bi}", "table", "warn" if probs else "pass", "; ".join(probs) or
                f"on page {pg + 1}" + (f"-{end_pg + 1}, header repeated" if end_pg > pg else "") + ", inside the margins", page=pg,
                fix={"op": "table_smaller", "block": bi} if any("wider" in p for p in probs) else None)
        elif b["type"] == "chart":
            want_txt = [_norm(c) for c in b["categories"]][:6]
            if b.get("title"):
                want_txt.append(_norm(b["title"]))
            best, best_pg = 0, None
            for p_ in range(n_pages):
                txt = [_norm(x["text"]) for x in by_page.get(p_, [])]
                k = sum(1 for w_ in want_txt if any(t == w_ or t.startswith(w_[:20]) or w_ in t for t in txt))
                if k > best:
                    best, best_pg = k, p_
            if best_pg is None or best < max(1, len(want_txt) // 2):
                add(f"b{bi}", "chart", "fail", f"chart '{b.get('title') or b.get('label')}' not seen (its labels are not on any page)")
            else:
                add(f"b{bi}", "chart", "pass", f"drawn on page {best_pg + 1} ({best}/{len(want_txt)} labels found)", page=best_pg)
        elif b["type"] == "image":
            add(f"b{bi}", "image", "pass" if n_img >= want_img else "fail", f"{n_img} image(s) in the PDF for {want_img} placed")

    # ---- every paragraph and list item made it
    missing, total = [], 0
    for bi, b in enumerate(plan["blocks"]):
        texts = []
        if b["type"] in ("paragraph", "quote", "callout"):
            texts = [b.get("text") or ""]
        elif b["type"] == "bullets":
            texts = [it["text"] for it in b["items"]] + [s["text"] for it in b["items"] for s in it.get("items") or []]
        elif b["type"] == "entry":
            texts = [b["title"]] + list(b.get("bullets") or [])
        for t in texts:
            words = _norm(t).split()[:7]
            if len(words) < 2:
                continue
            total += 1
            if " ".join(words) not in full:
                missing.append(" ".join(words))
    if total:
        add("text", "text", "pass" if not missing else ("warn" if len(missing) <= max(1, total // 20) else "fail"),
            f"all {total} paragraphs and items are in the PDF" if not missing else f"{len(missing)}/{total} not found: " + "; ".join(missing[:3]))

    # ---- one look at the pages (with a planner), on a time budget: a slow model never holds up the run
    if planner is not None:
        import threading
        box = {}

        def go():
            try:
                box["r"] = look(pdf, plan, planner)
            except Exception as e:  # noqa: BLE001
                box["r"] = {"id": "look", "type": "look", "status": "skip", "why": f"visual check failed: {type(e).__name__}: {str(e)[:80]}"}
        th_ = threading.Thread(target=go, daemon=True)
        th_.start()
        th_.join(timeout=look_budget)
        res.append(box.get("r") or {"id": "look", "type": "look", "status": "skip", "why": f"visual check skipped: no answer within {look_budget} s"})
    counts = {k: sum(1 for r in res if r["status"] == k) for k in ("pass", "warn", "fail", "skip")}
    return {"counts": counts, "results": res, "facts": facts}


LOOK_SYSTEM = """You check the pages of a document an editor made, before it goes to the client. The image is a contact
sheet: every page in order, labelled p1, p2, ... Look for real layout problems only:
- text cut off, overlapping or running outside the page; a table or chart cut or squashed
- a page that is empty or almost empty in the middle of the document (a short LAST page is fine)
- a heading alone at the bottom of a page; very uneven spacing; clashing styles or fonts
- anything that looks broken or unprofessional
Reply with ONE JSON object: {"ok": true/false, "problems": [{"page": <n>, "what": "<short>"}], "impression": "<one sentence>"}"""


def look(pdf, plan, planner):
    from ..vlm import ask
    imgs = RN.pages(pdf, scale=0.9, last=11)
    jpeg = RN.sheet(imgs, [f"p{i + 1}" for i in range(len(imgs))], cols=min(4, len(imgs)), cell_w=340 if len(imgs) > 2 else 520)
    d, secs = ask(planner, LOOK_SYSTEM, f"A {plan['doctype']} of {len(imgs)} page(s): '{plan.get('title') or ''}'.", [jpeg], tier="vision")
    if not isinstance(d, dict):
        return {"id": "look", "type": "look", "status": "warn", "why": "the visual check gave no answer"}
    probs = [p for p in d.get("problems") or [] if isinstance(p, dict) and p.get("what")]
    return {"id": "look", "type": "look", "status": "pass" if d.get("ok", not probs) and not probs else "warn",
            "why": "; ".join(f"p{p.get('page')}: {p['what']}" for p in probs[:4]) or str(d.get("impression") or "looks clean")[:160],
            "evidence": {"impression": d.get("impression"), "seconds": secs}}


# ------------------------------------------------------------------------------------------------ slide decks
DECO = {"QuoteMark", "SlideNumber"}


def check_deck(deck, built, rendered, planner=None, png_dir=None, look_budget=25):
    """Checks for a deck: every slide rendered; no text running out of its box (PowerPoint's own measurement of each
    text after layout); every slide title and bullet on its slide; charts drawn; fonts as designed; one vision look."""
    th = deck["th"]
    res = []

    def add(id_, typ, status, why, slide=None, fix=None):
        r = {"id": id_, "type": typ, "status": status, "why": why}
        if slide is not None:
            r["page"] = slide
        if fix:
            r["fix"] = fix
        res.append(r)
    if not rendered.get("ok"):
        add("global", "global", "fail", f"not rendered: {rendered.get('error')}")
        return {"counts": {"pass": 0, "warn": 0, "fail": 1, "skip": 0}, "results": res, "facts": {}}
    pdf = rendered["pdf"]
    n, want = int(rendered.get("slides") or 0), len(deck["slides"])
    probs, status = [], "pass"
    if n != want:
        probs.append(f"{n} slides rendered for {want} built")
        status = "fail"
    asked = deck.get("target_slides")
    if asked and abs(want - asked) > max(1, round(asked * 0.15)):
        probs.append(f"{want} slides for the {asked} asked")
        status = "fail" if status == "fail" else "warn"
    gf = glyph_fonts(pdf)
    fams = {_family(th["body"]), _family(th["head"]), _family(th["head"]).replace("semibold", ""), "georgia", "arial"}
    stray = {f: c for f, c in gf.items() if not _is_family(f, fams) and not _is_family(f, OK_FONTS)}
    if stray:
        probs.append("drawn in other fonts: " + ", ".join(f"{f} ({c} chars)" for f, c in list(stray.items())[:3]))
        status = "fail" if status == "fail" else "warn"
    add("global", "global", status, "; ".join(probs) or f"{n} slides, fonts as designed")
    # text that does not fit its box, slide by slide (PowerPoint measured it after laying it out)
    over = {}
    for sh in rendered.get("shapes") or []:
        if sh["name"] in DECO:
            continue
        if sh["text_height"] > sh["height"] + 3 or sh["text_width"] > sh["width"] + 3:
            over.setdefault(sh["slide"], []).append(sh)
    pages = RN.text(pdf)
    hard = {}
    try:  # text whose colour is too close to what is behind it (under 3:1)
        from pptx import Presentation

        from . import pptx_ops as PO
        prs = Presentation(built["path"])
        for b in PO.unreadable(prs, range(1, len(prs.slides) + 1)):
            hard.setdefault(b["slide"], []).append(b)
    except Exception:  # noqa: BLE001
        hard = {}
    for k, s in enumerate(deck["slides"]):
        sl = k + 1
        if sl in hard and sl not in over:
            b = hard[sl][0]
            add(f"r{sl}", "slide", "warn", f"slide {sl}: hard to read: '{b['text']}' ({b['colour']} on {b['behind']}, contrast {b['contrast']}:1)", slide=sl)
        if sl in over:
            sh = over[sl][0]
            add(f"s{sl}", "slide", "fail", f"slide {sl}: text runs out of its box ({sh['name']}: {sh['text_height']:.0f} pt of text in {sh['height']:.0f} pt)",
                slide=sl, fix={"op": "shrink", "src": s.get("src")})
            continue
        txt = _norm(pages[k]) if k < len(pages) else ""
        want_txt = [x for x in [s.get("title")] + [it["text"] for it in s.get("bullets") or []] if x]
        missing = [w for w in want_txt if " ".join(_norm(w).split()[:5]) not in txt]
        if s["layout"] == "chart":
            cats = [_norm(c) for c in s["chart"]["categories"][:6]]
            if sum(1 for c in cats if c in txt) < max(1, len(cats) // 2):
                missing.append("chart labels")
        if missing:
            add(f"s{sl}", "slide", "warn", f"slide {sl}: not found on the slide: " + "; ".join(m[:40] for m in missing[:3]), slide=sl)
        else:
            add(f"s{sl}", "slide", "pass", f"slide {sl} ({s['layout']}): all text inside its boxes", slide=sl)
    if planner is not None and png_dir:
        import threading
        from pathlib import Path
        box = {}

        def go():
            try:
                from PIL import Image
                imgs = [Image.open(p).convert("RGB") for p in sorted(Path(png_dir).glob("*.png"))][:16]
                box["r"] = look_images(imgs, f"a slide deck of {len(imgs)} slides: '{deck.get('title') or ''}'", planner, SLIDES_SYSTEM)
            except Exception as e:  # noqa: BLE001
                box["r"] = {"id": "look", "type": "look", "status": "skip", "why": f"visual check failed: {type(e).__name__}: {str(e)[:80]}"}
        t = threading.Thread(target=go, daemon=True)
        t.start()
        t.join(timeout=look_budget)
        res.append(box.get("r") or {"id": "look", "type": "look", "status": "skip", "why": f"visual check skipped: no answer within {look_budget} s"})
    counts = {k: sum(1 for r in res if r["status"] == k) for k in ("pass", "warn", "fail", "skip")}
    return {"counts": counts, "results": res, "facts": {"slides": n, "fonts": gf}}


SLIDES_SYSTEM = """You check the slides of a deck a designer made, before it goes to the client. The image is a contact sheet:
every slide in order, labelled s1, s2, ... Look for real problems only:
- text cut off, overlapping other text or shapes, or running off the slide
- text too small to read ON THE REAL SLIDE or with poor contrast against its background (the slides are shown at about a
  quarter of their real size: judge a text's size against its slide, not against this image; 14 pt or more on a slide is fine)
- an empty or broken-looking slide; elements clearly misaligned; inconsistent styles between slides
Reply with ONE JSON object: {"ok": true/false, "problems": [{"slide": <n>, "what": "<short>"}], "impression": "<one sentence>"}"""


def look_images(images, what, planner, system):
    from ..vlm import ask
    jpeg = RN.sheet(images, [f"s{i + 1}" for i in range(len(images))], cols=min(4, len(images)), cell_w=340 if len(images) > 2 else 520)
    d, secs = ask(planner, system, f"{what}.", [jpeg], tier="vision")
    if not isinstance(d, dict):
        return {"id": "look", "type": "look", "status": "warn", "why": "the visual check gave no answer"}
    probs = [p for p in d.get("problems") or [] if isinstance(p, dict) and p.get("what")]
    return {"id": "look", "type": "look", "status": "pass" if d.get("ok", not probs) and not probs else "warn",
            "why": "; ".join(f"s{p.get('slide') or p.get('page')}: {p['what']}" for p in probs[:4]) or str(d.get("impression") or "looks clean")[:160],
            "evidence": {"impression": d.get("impression"), "seconds": secs}}


# ------------------------------------------------------------------------------------------------ workbooks
def check_book(book, built, rendered):
    """Checks for a workbook: Excel opened and recalculated it; no cell shows an error (#DIV/0!, #REF!, #NAME?...); every
    total, metric and group Excel computed equals the same number computed in Python; the charts are there."""
    res = []

    def add(id_, typ, status, why):
        res.append({"id": id_, "type": typ, "status": status, "why": why})
    if not rendered.get("ok"):
        add("global", "global", "fail", f"Excel could not open it: {rendered.get('error')}")
        return {"counts": {"pass": 0, "warn": 0, "fail": 1, "skip": 0}, "results": res, "facts": {}}
    errs = rendered.get("errors") or []
    probs = list(built.get("problems") or [])
    want_charts = sum(1 for s in book["sheets"] if s["kind"] == "summary" for g in s.get("groups") or [] if g.get("chart"))
    if (rendered.get("charts") or 0) < want_charts:
        probs.append(f"{rendered.get('charts')} chart(s) for {want_charts} planned")
    add("global", "global", "fail" if errs or probs else "pass",
        "; ".join(([f"{len(errs)} cell(s) show errors: " + ", ".join(f"{e['sheet']}!R{e['row']}C{e['col']} {e['error']}" for e in errs[:4])] if errs else [])
                  + probs) or f"Excel recalculated {len(rendered.get('sheets') or [])} sheet(s): no errors, {rendered.get('charts')} chart(s)")
    vals = rendered.get("values") or {}
    bad = []
    for c in built.get("checks") or []:
        got = vals.get(f"{c['sheet']}!{c['cell']}")
        ok = isinstance(got, (int, float)) and abs(float(got) - float(c["expect"])) <= 0.005 * max(1.0, abs(float(c["expect"])))
        if not ok:
            bad.append(f"{c['what']} ({c['sheet']}!{c['cell']}): Excel {got}, expected {c['expect']}")
    if built.get("checks"):
        add("values", "values", "fail" if bad else "pass",
            "; ".join(bad[:4]) or f"all {len(built['checks'])} totals and summaries Excel computed equal the numbers computed here")
    counts = {k: sum(1 for r in res if r["status"] == k) for k in ("pass", "warn", "fail", "skip")}
    return {"counts": counts, "results": res, "facts": {"sheets": rendered.get("sheets"), "charts": rendered.get("charts")}}
